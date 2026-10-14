import inspect
import uuid
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.common.contracts import AppError, PageParams
from app.core.database import SessionLocal
from app.infrastructure.adapters import LocalFileStorage
from app.modules.bid.api import router as bid_router_module
from app.modules.bid.api.router import download_bid_project_file
from app.modules.bid.application.service import MAX_PROJECT_ITEMS, BidProjectService
from app.modules.bid.domain.lifecycle import (
    BidImportStatus,
    BidItemStatus,
    BidProjectStatus,
    BidRecommendationType,
    ensure_transition,
)
from app.modules.bid.infrastructure.models import BidProject, BidProjectItem, BidTemplate
from app.modules.bid.infrastructure.repository import BidProjectRepository
from app.modules.bid.schemas import BidInputMappingUpdateRequest, BidOutputMappingUpdateRequest
from app.modules.catalog.application.product_export_columns import PRODUCT_EXPORT_COLUMN_KEYS
from app.modules.matching.application.service import MatchingService
from app.modules.matching.application.type1_agent import (
    Type1AgentContractError,
    Type1TemplateMappingSuggestion,
)
from app.modules.matching.application.type2_agent import Type2TemplateMappingSuggestion


def _workbook_bytes() -> tuple[bytes, list[str]]:
    headers = [
        "采购项编码",
        "商品名称",
        "品牌",
        "型号",
        "规格",
        "类目",
        "数量",
        "单位",
        "限价",
        "报价",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "需求明细"
    sheet.append(headers)
    sheet.append(["BUY-1", "测试商品", "测试品牌", "M-1", "规格A", "数码", 2, "台", 100, None])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), headers


@pytest.mark.asyncio
async def test_project_creation_parses_frozen_template(tmp_path: Path) -> None:
    file_bytes, headers = _workbook_bytes()
    actor_id = uuid.uuid4()
    storage = LocalFileStorage(tmp_path)
    async with SessionLocal() as session:
        service = BidProjectService(session, storage)
        template = BidTemplate(
            template_code=f"test-{uuid.uuid4()}",
            template_name="测试模板",
            version=1,
            sheet_name="需求明细",
            header_row=1,
            data_start_row=2,
            import_mapping={
                "buyer_item_code": "采购项编码",
                "product_name": "商品名称",
                "brand": "品牌",
                "model": "型号",
                "specification": "规格",
                "category_text": "类目",
                "quantity": "数量",
                "unit": "单位",
                "max_price": "限价",
            },
            export_mapping={"selected_unit_price": "报价"},
            fingerprint=service._fingerprint("需求明细", 1, headers),
        )
        session.add(template)
        await session.flush()
        response = await service.create(
            project_name="测试投标项目",
            buyer_name="测试需求商",
            start_at=None,
            deadline_at=None,
            remark=None,
            filename="需求.xlsx",
            file_bytes=file_bytes,
            actor_id=actor_id,
        )
        assert response.import_status == BidImportStatus.PARSED
        assert response.total_item_count == 1
        rows = await service.items(response.id, page_params=PageParams(), status=None)
        assert rows.total == 1
        assert rows.items[0].product_name == "测试商品"
        await session.rollback()


def test_type1_type2_row_parser_skips_formula_summary_rows() -> None:
    headers = ["商品编码", "商品名称", "数量合计", "参考价格"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    sheet.append(["SKU-001", "正常商品", 2, 10])
    sheet.append(["SKU-TOTAL", "合计", 1, 20])
    sheet.append([None, "合计", "=SUM(C2:C3)", None])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    template = BidTemplate(
        template_code="summary-row-test",
        template_name="汇总行测试",
        version=1,
        sheet_name="Sheet1",
        header_row=1,
        data_start_row=2,
        import_mapping={
            "buyer_item_code": "商品编码",
            "product_name": "商品名称",
            "quantity": "数量合计",
            "max_price": "参考价格",
        },
        export_mapping={},
        fingerprint="summary-row-test",
    )

    service = BidProjectService.__new__(BidProjectService)
    rows = service._parse_rows(output.getvalue(), template)

    assert [row["product_name"] for row in rows] == ["正常商品", "合计"]
    assert [row["quantity"] for row in rows] == [Decimal("2"), Decimal("1")]


@pytest.mark.asyncio
async def test_type1_ai_mapping_failure_still_allows_manual_confirmation(
    tmp_path: Path,
) -> None:
    file_bytes, _headers = _workbook_bytes()
    actor_id = uuid.uuid4()

    class FailedType1MappingAgent:
        async def suggest_template_mapping(
            self, _file_bytes: bytes, *, project_remark: str | None
        ) -> Type1TemplateMappingSuggestion:
            raise Type1AgentContractError("invalid mapping")

    async with SessionLocal() as session:
        service = BidProjectService(
            session,
            LocalFileStorage(tmp_path),
            type1_agent=FailedType1MappingAgent(),  # type: ignore[arg-type]
        )
        created = await service.create(
            project_name="类型1人工兜底",
            buyer_name="测试需求商",
            start_at=None,
            deadline_at=None,
            remark=None,
            filename="需求.xlsx",
            file_bytes=file_bytes,
            recommendation_type=BidRecommendationType.TYPE_1_SPECIFICATION,
            actor_id=actor_id,
        )

        assert created.import_status == BidImportStatus.MAPPING_REQUIRED
        mapping = await service.input_mapping(created.id)
        assert mapping.mapping == {}
        assert mapping.sheet_name == "需求明细"
        assert mapping.header_row == 1
        assert "商品名称" in mapping.headers
        await session.rollback()


@pytest.mark.asyncio
async def test_type1_output_mapping_is_project_scoped_and_requires_confirmation(
    tmp_path: Path,
) -> None:
    file_bytes, headers = _workbook_bytes()
    actor_id = uuid.uuid4()
    suggestion = Type1TemplateMappingSuggestion.model_validate(
        {
            "sheet_name": "需求明细",
            "header_row": 1,
            "data_start_row": 2,
            "input_mapping": {"product_name": "商品名称", "specification": "规格"},
            "output_mapping": {"agreement_price": "报价"},
        }
    )

    class StubType1MappingAgent:
        async def suggest_template_mapping(
            self, _file_bytes: bytes, *, project_remark: str | None
        ) -> Type1TemplateMappingSuggestion:
            assert project_remark is None
            return suggestion

    async with SessionLocal() as session:
        service = BidProjectService(
            session,
            LocalFileStorage(tmp_path),
            type1_agent=StubType1MappingAgent(),  # type: ignore[arg-type]
        )
        template = BidTemplate(
            template_code=f"type1-mapping-{uuid.uuid4()}",
            template_name="类型1回填映射测试",
            version=1,
            sheet_name="需求明细",
            header_row=1,
            data_start_row=2,
            import_mapping={"product_name": "商品名称", "specification": "规格"},
            export_mapping={"selected_unit_price": "历史写死报价列"},
            fingerprint=service._fingerprint("需求明细", 1, headers),
        )
        session.add(template)
        await session.flush()

        created = await service.create(
            project_name="类型1回填映射",
            buyer_name="测试需求商",
            start_at=None,
            deadline_at=None,
            remark=None,
            filename="需求.xlsx",
            file_bytes=file_bytes,
            recommendation_type=BidRecommendationType.TYPE_1_SPECIFICATION,
            actor_id=actor_id,
        )
        project = await session.get(BidProject, created.id)
        assert project is not None
        assert project.template_id != template.id
        assert created.import_status == BidImportStatus.MAPPING_REQUIRED
        assert created.total_item_count == 0
        assert project.output_mapping == {"agreement_price": "报价"}
        assert project.output_mapping_confirmed_at is None

        input_mapping = await service.input_mapping(created.id)
        assert input_mapping.confirmed is False
        assert input_mapping.mapping == {
            "product_name": "商品名称",
            "specification": "规格",
        }
        assert input_mapping.sheet_names == ["需求明细"]

        with pytest.raises(AppError, match="请先人工确认客户需求字段映射"):
            await MatchingService(session).start(created.id, actor_id)

        confirmed_input = await service.confirm_input_mapping(
            created.id,
            BidInputMappingUpdateRequest(
                sheet_name="需求明细",
                header_row=1,
                data_start_row=2,
                mapping={"product_name": "商品名称", "specification": "规格"},
            ),
            actor_id,
        )
        assert confirmed_input.confirmed is True
        assert project.import_status == BidImportStatus.PARSED.value
        assert project.total_item_count == 1

        initial = await service.output_mapping(created.id)
        assert initial.confirmed is False
        assert initial.mapping == {"agreement_price": "报价"}
        assert "报价" in initial.headers
        assert initial.suggested_headers == ["报价"]
        field_keys = {field.key for field in initial.fields}
        assert set(PRODUCT_EXPORT_COLUMN_KEYS).issubset(field_keys)
        assert "agreement_price" in field_keys
        assert "selected_unit_price" not in field_keys

        with pytest.raises(AppError, match="请先确认客户原表的回填字段映射"):
            await MatchingService(session).start(created.id, actor_id)

        confirmed = await service.update_output_mapping(
            created.id,
            BidOutputMappingUpdateRequest(
                mapping={"agreement_price": "报价", "note": "单位"}
            ),
            actor_id,
        )

        assert confirmed.confirmed is True
        assert confirmed.mapping == {"agreement_price": "报价", "note": "单位"}
        assert template.export_mapping == {"selected_unit_price": "历史写死报价列"}
        await session.rollback()


@pytest.mark.asyncio
async def test_output_mapping_returns_headers_resolved_against_the_current_workbook(
    tmp_path: Path,
) -> None:
    sheet_name = f"变体-{uuid.uuid4().hex[:8]}"
    discount_header = "是否满足平台承诺折扣率（暂不做要求）"
    validity_header = "价格有效期\n至少到12.31"
    headers = ["品牌", "型号", "商品报价（单价）", discount_header, validity_header]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    sheet.append(["联想", "LT1821", None, None, None])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    file_bytes = output.getvalue()

    actor_id = uuid.uuid4()
    suggestion = Type2TemplateMappingSuggestion.model_validate(
        {
            "sheet_name": sheet_name,
            "header_row": 1,
            "data_start_row": 2,
            "input_mapping": {"brand": "品牌", "model": "型号"},
            "output_mapping": {},
        }
    )

    class StubType2MappingAgent:
        async def suggest_template_mapping(
            self, _file_bytes: bytes, *, project_remark: str | None
        ) -> Type2TemplateMappingSuggestion:
            assert project_remark is None
            return suggestion

    async with SessionLocal() as session:
        service = BidProjectService(
            session,
            LocalFileStorage(tmp_path),
            type2_agent=StubType2MappingAgent(),  # type: ignore[arg-type]
        )
        template = BidTemplate(
            template_code=f"type2-header-variant-{uuid.uuid4()}",
            template_name="类型2结果列表头变体",
            version=1,
            sheet_name=sheet_name,
            header_row=1,
            data_start_row=2,
            import_mapping={"brand": "品牌", "model": "型号"},
            export_mapping={
                "selected_unit_price": "商品报价（单价）",
                "discount_compliance": "是否满足平台承诺折扣率（88折）",
                "price_valid_until": "价格有效期 至少到12.31",
            },
            fingerprint=service._fingerprint(sheet_name, 1, headers),
        )
        session.add(template)
        await session.flush()

        created = await service.create(
            project_name="类型2表头变体",
            buyer_name="测试需求商",
            start_at=None,
            deadline_at=None,
            remark=None,
            filename="需求.xlsx",
            file_bytes=file_bytes,
            recommendation_type=BidRecommendationType.TYPE_2_IDENTIFIED_PRODUCT,
            actor_id=actor_id,
        )
        project = await session.get(BidProject, created.id)
        assert project is not None
        assert project.template_id != template.id
        assert created.import_status == BidImportStatus.MAPPING_REQUIRED
        await service.confirm_input_mapping(
            created.id,
            BidInputMappingUpdateRequest(
                sheet_name=sheet_name,
                header_row=1,
                data_start_row=2,
                mapping={"brand": "品牌", "model": "型号"},
            ),
            actor_id,
        )
        assert project.import_status == BidImportStatus.PARSED.value
        assert project.total_item_count == 1
        assert project.output_mapping == {}
        project.output_mapping = {
            "selected_unit_price": "商品报价（单价）",
            "discount_compliance": "是否满足平台承诺折扣率（88折）",
            "price_valid_until": "价格有效期 至少到12.31",
        }
        await session.flush()
        mapping = await service.output_mapping(created.id)

        assert mapping.mapping["discount_compliance"] == discount_header
        assert mapping.mapping["price_valid_until"] == validity_header
        assert mapping.suggested_headers == [
            "商品报价（单价）",
            discount_header,
            validity_header,
        ]
        assert "是否满足平台承诺折扣率（88折）" not in mapping.headers
        assert "价格有效期 至少到12.31" not in mapping.headers
        await session.rollback()


def test_output_mapping_suggests_only_the_result_block_after_quote_column() -> None:
    headers = [
        "物料代码",
        "物料名称",
        "限价调整",
        "2025年采购量",
        "需报价物料",
        "供应商报价",
        "商品编码（sku）",
        "商品型号",
        "报价单位",
        "备注",
        "供应商",
        "采销",
    ]

    suggested = BidProjectService._suggested_output_headers(
        headers,
        {
            "agreement_price": "供应商报价",
            "sku": "商品编码（sku）",
            "model": "商品型号",
            "unit": "报价单位",
            "note": "备注",
            "supplier_name": "供应商",
        },
        {
            "buyer_item_code": "物料代码",
            "product_name": "物料名称",
            "max_price": "限价调整",
            "quantity": "2025年采购量",
        },
    )

    assert suggested == headers[5:]


def test_output_mapping_infers_unique_obvious_supplier_quote_column() -> None:
    mapping = BidProjectService._with_obvious_price_mapping(
        ["物料名称", "限价调整", "供应商报价", "商品编码（sku）", "备注"],
        {},
    )

    assert mapping == {"agreement_price": "供应商报价"}


@pytest.mark.asyncio
async def test_quote_export_loads_database_timestamp_before_building_response(
    tmp_path: Path,
) -> None:
    file_bytes, headers = _workbook_bytes()
    actor_id = uuid.uuid4()
    storage = LocalFileStorage(tmp_path)
    async with SessionLocal() as session:
        service = BidProjectService(session, storage)
        template = BidTemplate(
            template_code=f"export-{uuid.uuid4()}",
            template_name="导出响应测试模板",
            version=1,
            sheet_name="需求明细",
            header_row=1,
            data_start_row=2,
            import_mapping={"product_name": "商品名称"},
            export_mapping={"selected_unit_price": "报价"},
            fingerprint=service._fingerprint("需求明细", 1, headers),
        )
        session.add(template)
        await session.flush()
        created = await service.create(
            project_name="导出响应测试",
            buyer_name="测试需求商",
            start_at=None,
            deadline_at=None,
            remark=None,
            filename="需求.xlsx",
            file_bytes=file_bytes,
            actor_id=actor_id,
        )
        project = await session.get(BidProject, created.id)
        item = await session.scalar(
            select(BidProjectItem).where(BidProjectItem.project_id == created.id)
        )
        assert project is not None
        assert item is not None
        project.status = BidProjectStatus.READY.value
        item.status = BidItemStatus.NO_QUOTE.value
        await session.flush()

        exported = await service.export(created.id, actor_id)

        assert exported.file_type == "QUOTED_EXPORT"
        assert exported.created_at is not None
        assert exported.original_filename.endswith(".xlsx")
        await session.rollback()


@pytest.mark.asyncio
async def test_type2_export_only_fills_blank_result_cells(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价明细"
    sheet.append(["型号", "商品报价", "是否满足88折", "备注", "税收编码"])
    sheet.append(["LT1821", None, None, None, None])
    sheet.append(["TN9999", "客户原值", None, "客户已有备注", None])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    template = BidTemplate(
        template_code="type2-export-test",
        template_name="类型2导出测试",
        version=1,
        sheet_name="报价明细",
        header_row=1,
        data_start_row=2,
        import_mapping={"model": "型号"},
        export_mapping={
            "selected_unit_price": "商品报价",
            "discount_compliance": "是否满足88折",
            "note": "备注",
            "tax_code": "税收编码",
        },
        fingerprint="test",
    )
    selections = [
        (
            SimpleNamespace(
                source_row_number=2,
                unit="支",
                buyer_item_code=None,
                brand=None,
                model="LT1821",
            ),
            SimpleNamespace(
                selected_unit_price=Decimal("88.0000"),
                product_snapshot={
                    "type2_match_type": "TYPE2_AI_EXACT",
                    "tax_code": "1090512990000000000",
                },
                supplier_snapshot={},
                price_snapshot={"discount_compliance": "是"},
                note="类型2 AI 判定完全同品",
            ),
        ),
        (
            SimpleNamespace(
                source_row_number=3,
                unit="支",
                buyer_item_code=None,
                brand=None,
                model="TN9999",
            ),
            SimpleNamespace(
                selected_unit_price=Decimal("66.0000"),
                product_snapshot={
                    "type2_match_type": "TYPE2_AI_SUBSTITUTE",
                    "tax_code": "1070101010000000000",
                },
                supplier_snapshot={},
                price_snapshot={"discount_compliance": "是"},
                note="同类型商品替代：型号不同",
            ),
        ),
    ]

    async with SessionLocal() as session:
        service = BidProjectService(session, LocalFileStorage(tmp_path))
        result = service._build_export(source.getvalue(), template, selections)

    exported = load_workbook(BytesIO(result), data_only=False)
    try:
        result_sheet = exported["报价明细"]
        assert result_sheet["B2"].value == 88
        assert result_sheet["C2"].value == "是"
        assert result_sheet["D2"].value == "类型2 AI 判定完全同品"
        assert result_sheet["E2"].value == "1090512990000000000"
        assert result_sheet["B3"].value == "客户原值"
        assert result_sheet["C3"].value == "是"
        assert result_sheet["D3"].value == "客户已有备注"
        assert result_sheet["E3"].value == "1070101010000000000"
    finally:
        exported.close()


@pytest.mark.asyncio
async def test_type1_export_uses_original_workbook_and_only_fills_blank_result_cells(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客户物料表"
    sheet.append(["物料名称", "规格参数", "供应商报价", "商品编码（sku）", "备注"])
    sheet.append(["硒鼓", "黑色，打印量3000页", None, None, None])
    sheet.append(["复印纸", "A4 80g", "客户原值", None, "客户已有备注"])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    template = BidTemplate(
        template_code="type1-export-test",
        template_name="类型1导出测试",
        version=1,
        sheet_name="客户物料表",
        header_row=1,
        data_start_row=2,
        import_mapping={
            "product_name": "物料名称",
            "specification": "规格参数",
        },
        export_mapping={
            "selected_unit_price": "供应商报价",
            "sku": "商品编码（sku）",
            "note": "备注",
        },
        fingerprint="test",
    )
    selections = [
        (
            SimpleNamespace(
                source_row_number=2,
                unit="支",
                buyer_item_code=None,
                brand=None,
                model=None,
            ),
            SimpleNamespace(
                selected_unit_price=Decimal("88.0000"),
                product_snapshot={"sku": "SKU-001"},
                supplier_snapshot={},
                price_snapshot={},
                note="规格参数匹配",
            ),
        ),
        (
            SimpleNamespace(
                source_row_number=2,
                unit="支",
                buyer_item_code=None,
                brand=None,
                model=None,
            ),
            SimpleNamespace(
                selected_unit_price=Decimal("89.0000"),
                product_snapshot={"sku": "SKU-001-B"},
                supplier_snapshot={},
                price_snapshot={},
                note="第二个规格匹配商品",
            ),
        ),
        (
            SimpleNamespace(
                source_row_number=3,
                unit="包",
                buyer_item_code=None,
                brand=None,
                model=None,
            ),
            SimpleNamespace(
                selected_unit_price=Decimal("66.0000"),
                product_snapshot={"sku": "SKU-002"},
                supplier_snapshot={},
                price_snapshot={},
                note="规格参数匹配",
            ),
        ),
    ]

    async with SessionLocal() as session:
        result = BidProjectService(
            session, LocalFileStorage(tmp_path)
        )._build_export(
            source.getvalue(), template, selections, expand_multiple_rows=True
        )

    exported = load_workbook(BytesIO(result), data_only=False)
    try:
        result_sheet = exported["客户物料表"]
        assert result_sheet["A2"].value == "硒鼓"
        assert result_sheet["B2"].value == "黑色，打印量3000页"
        assert result_sheet["C2"].value == 88
        assert result_sheet["D2"].value == "SKU-001"
        assert result_sheet["E2"].value == "规格参数匹配"
        assert result_sheet["A3"].value == "硒鼓"
        assert result_sheet["B3"].value == "黑色，打印量3000页"
        assert result_sheet["C3"].value == 89
        assert result_sheet["D3"].value == "SKU-001-B"
        assert result_sheet["E3"].value == "第二个规格匹配商品"
        assert result_sheet["C4"].value == "客户原值"
        assert result_sheet["D4"].value == "SKU-002"
        assert result_sheet["E4"].value == "客户已有备注"
    finally:
        exported.close()


@pytest.mark.asyncio
async def test_type2_export_resolves_layout_and_instruction_variants_in_result_headers(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["说明"])
    sheet.append(
        [
            "型号",
            "商品报价（单价）",
            "是否满足平台承诺折扣率（暂不做要求）",
            "价格有效期\n至少到12.31",
            "备注",
        ]
    )
    sheet.append(["LT1821", None, None, None, None])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    template = BidTemplate(
        template_code="type2-old-beijing-export-mapping",
        template_name="类型2旧北京烟草模板映射",
        version=1,
        sheet_name="Sheet1",
        header_row=2,
        data_start_row=3,
        import_mapping={"model": "型号"},
        export_mapping={
            "selected_unit_price": "商品报价（单价）",
            "discount_compliance": "是否满足平台承诺折扣率（88折）",
            "price_valid_until": "价格有效期 至少到12.31",
            "note": "备注",
        },
        fingerprint="test",
    )
    selections = [
        (
            SimpleNamespace(
                source_row_number=3,
                unit="支",
                buyer_item_code=None,
                brand="联想",
                model="LT1821",
            ),
            SimpleNamespace(
                selected_unit_price=Decimal("88.0000"),
                product_snapshot={"type2_match_type": "TYPE2_AI_EXACT"},
                supplier_snapshot={},
                price_snapshot={
                    "discount_compliance": "是",
                    "price_valid_until": None,
                },
                note="类型2 AI 判定完全同品",
            ),
        )
    ]

    async with SessionLocal() as session:
        result = BidProjectService(
            session, LocalFileStorage(tmp_path)
        )._build_export(source.getvalue(), template, selections)

    exported = load_workbook(BytesIO(result), data_only=False)
    try:
        result_sheet = exported["Sheet1"]
        assert result_sheet["B3"].value == 88
        assert result_sheet["C3"].value == "是"
        assert result_sheet["D3"].value is None
        assert result_sheet["E3"].value == "类型2 AI 判定完全同品"
    finally:
        exported.close()


@pytest.mark.asyncio
async def test_project_start_time_accepts_equal_or_earlier_deadline_and_rejects_later(
    tmp_path: Path,
) -> None:
    service = BidProjectService(SessionLocal(), LocalFileStorage(tmp_path))
    start = datetime(2026, 9, 16, 9, 0)
    assert start <= start
    with pytest.raises(AppError, match="项目开始时间不能晚于截止时间"):
        await service.create(
            project_name="时间校验",
            buyer_name="需求商",
            start_at=start + timedelta(minutes=1),
            deadline_at=start,
            remark=None,
            filename="request.xlsx",
            file_bytes=b"",
            actor_id=uuid.uuid4(),
        )


def test_project_lifecycle_rejects_illegal_transition() -> None:
    with pytest.raises(AppError) as exc_info:
        ensure_transition(BidProjectStatus.IMPORTED.value, BidProjectStatus.SUBMITTED)
    assert exc_info.value.status_code == 409


def test_project_item_limit_is_100k_and_rejects_only_above_boundary() -> None:
    assert MAX_PROJECT_ITEMS == 100_000
    assert not 100_000 > MAX_PROJECT_ITEMS
    assert 100_001 > MAX_PROJECT_ITEMS


@pytest.mark.asyncio
async def test_bid_file_download_uses_utf8_content_disposition(monkeypatch) -> None:
    class StubBidProjectService:
        def __init__(self, _session) -> None:
            pass

        async def download(self, _project_id, _file_id):
            return SimpleNamespace(original_filename="BID-测试-报价文件-V1.xlsx"), b"xlsx"

    monkeypatch.setattr(bid_router_module, "BidProjectService", StubBidProjectService)
    response = await download_bid_project_file(
        uuid.uuid4(), uuid.uuid4(), SimpleNamespace(), SimpleNamespace()
    )

    assert response.headers["content-disposition"] == (
        "attachment; filename*=UTF-8''BID-%E6%B5%8B%E8%AF%95-"
        "%E6%8A%A5%E4%BB%B7%E6%96%87%E4%BB%B6-V1.xlsx"
    )


def test_item_page_contract_uses_sql_keyword_filters_and_outer_join_snapshot_read_model() -> None:
    source = inspect.getsource(BidProjectRepository.item_page)
    assert "outerjoin" in source
    assert "BidProjectItem.current_selection_id == BidItemSelection.id" in source
    for field in ("product_name", "brand", "model", "buyer_item_code"):
        assert f"BidProjectItem.{field}.contains(keyword)" in source
    assert "matched_first" in source
    assert 'BidProjectItem.status.in_(("SELECTED", "UNIQUE_MATCH", "MULTIPLE_MATCH"))' in source
    assert "for item" not in source


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "recommendation_type",
    [
        BidRecommendationType.TYPE_1_SPECIFICATION.value,
        BidRecommendationType.TYPE_2_IDENTIFIED_PRODUCT.value,
    ],
)
async def test_type1_and_type2_item_pages_request_matched_rows_first(
    tmp_path: Path,
    recommendation_type: str,
) -> None:
    project_id = uuid.uuid4()

    class StubRepository:
        matched_first: bool | None = None

        async def project_by_id(self, _project_id: uuid.UUID, *, lock: bool = False):
            assert _project_id == project_id
            assert lock is False
            return SimpleNamespace(recommendation_type=recommendation_type)

        async def item_page(self, _project_id: uuid.UUID, **kwargs):
            assert _project_id == project_id
            self.matched_first = kwargs["matched_first"]
            return [], 0

    async with SessionLocal() as session:
        service = BidProjectService(session, LocalFileStorage(tmp_path))
        repository = StubRepository()
        service.repository = repository  # type: ignore[assignment]

        result = await service.items(
            project_id,
            PageParams(page=1, page_size=20),
            status=None,
        )

    assert result.total == 0
    assert repository.matched_first is True
