"""Real MySQL regressions for durability, lifecycle and independent queues.

Fixtures are UUID-scoped and cleaned in dependency order, including on failure.
Remote providers are replaced; MySQL, workbook parsing and HTTP are real.
"""

from __future__ import annotations

import asyncio
import hashlib
import os
import socket
import sys
import threading
import uuid
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from httpx import ASGITransport, AsyncClient, HTTPError
from openpyxl import Workbook
from sqlalchemy import delete, func, select

from app.common.async_work import run_blocking
from app.common.contracts import AppError
from app.core.database import SessionLocal
from app.infrastructure.adapters import get_object_storage
from app.jobs import recommendation_agent, task_worker
from app.main import app
from app.modules.auth.dependencies import get_current_user, get_current_user_before_response
from app.modules.auth.schemas import CurrentUser
from app.modules.bid.application.service import BidProjectService
from app.modules.bid.infrastructure.models import (
    BidItemSelection,
    BidProject,
    BidProjectEvent,
    BidProjectFile,
    BidProjectItem,
    BidTemplate,
    MatchCandidate,
    MatchTask,
)
from app.modules.catalog.infrastructure.models import Product
from app.modules.matching.application.service import MatchingService
from app.modules.recommendation.application.agent_schemas import RequirementAnalysis
from app.modules.recommendation.infrastructure.models import (
    PptGenerationTask,
    PptRecommendationConfig,
    RecommendationCandidate,
    RecommendationCategoryChoice,
    RecommendationConfirmation,
    RecommendationRun,
)
from app.modules.recommendation.template.models import RecommendationTemplateMapping
from app.modules.supplier.infrastructure.models import Supplier
from app.modules.system.mapping_analysis_service import MappingAnalysisService
from app.modules.system.models import AsyncTask, MappingAnalysisCache
from app.modules.system.task_schemas import AsyncTaskCreate, AsyncTaskKind, AsyncTaskResponse
from app.modules.system.task_service import AsyncTaskService
from app.modules.type3_recommendation.application.service import Type3RecommendationService
from app.modules.type3_recommendation.infrastructure.models import (
    Type3Candidate,
    Type3Export,
    Type3Run,
    Type3SheetConfig,
)
from app.modules.type3_recommendation.schemas import Type3RemarkAnalysis


@asynccontextmanager
async def fixture(kind: AsyncTaskKind):
    project_id, actor_id, file_id, template_id, supplier_id, product_id = (
        uuid.uuid4() for _ in range(6)
    )
    token = uuid.uuid4().hex
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active.append(["商品名称", "协议价"])
    workbook.active.append(["1.1.1 测试类目", None])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    key = f"queue-tests/{token}.xlsx"
    storage = get_object_storage()
    key = await storage.save(key, content)
    project_type = (
        "FILTER_RECOMMENDATION"
        if kind in {AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH}
        else "PPT_SOLUTION"
        if kind
        in {AsyncTaskKind.TYPE5_RECOMMENDATION_GENERATION, AsyncTaskKind.TYPE5_PPT_GENERATION}
        else "FREE_RECOMMENDATION"
    )
    rec_type = {
        AsyncTaskKind.TYPE1_MATCH: "TYPE_1_SPECIFICATION",
        AsyncTaskKind.TYPE2_MATCH: "TYPE_2_IDENTIFIED_PRODUCT",
        AsyncTaskKind.TYPE3_CANDIDATE_GENERATION: "TYPE_3_CATEGORY",
    }.get(kind, "TYPE_4_FREE" if project_type == "FREE_RECOMMENDATION" else "TYPE_5_PPT")
    now = datetime.now(UTC).replace(tzinfo=None)
    async with SessionLocal.begin() as session:
        session.add(
            Supplier(
                id=supplier_id,
                supplier_code=f"QT{token[:12]}",
                supplier_name=token,
                main_brands="queue-test",
                advantage="queue-test",
                archive_status="ARCHIVED",
                cooperation_status="NORMAL",
            )
        )
        session.add(
            BidTemplate(
                id=template_id,
                template_code=f"QT{token[:12]}",
                template_name=token,
                version=1,
                sheet_name="Sheet1",
                header_row=1,
                data_start_row=2,
                import_mapping={},
                export_mapping={"agreement_price": "协议价"},
                fingerprint=token,
            )
        )
        await session.flush()
        session.add(
            BidProject(
                id=project_id,
                project_code=f"QT{token[:12]}",
                project_name=token,
                buyer_name="queue test",
                status="IMPORTED",
                import_status="PARSED"
                if project_type == "FILTER_RECOMMENDATION"
                else "NOT_REQUIRED",
                project_type=project_type,
                recommendation_type=rec_type,
                template_id=template_id,
                template_version=1,
                remark=f"这是只用于队列集成测试的二十字以上需求说明。必须品牌{token}",
                created_by=actor_id,
            )
        )
        await session.flush()
        session.add(
            Product(
                id=product_id,
                source_supplier_id=supplier_id,
                sku=token,
                product_name=token,
                brand=token,
                category_level1_name="测试一级",
                category_level2_name="测试二级",
                category_level3_name="测试类目",
                agreement_price=Decimal("987654321.1234"),
                jd_price=Decimal("150"),
                gross_margin=Decimal("0.2"),
                status="ACTIVE",
            )
        )
        session.add(
            BidProjectFile(
                id=file_id,
                project_id=project_id,
                file_type="ORIGINAL"
                if project_type == "FILTER_RECOMMENDATION"
                else "RECOMMENDATION_TEMPLATE",
                version_no=1,
                original_filename="test.xlsx",
                storage_key=key,
                file_size=len(content),
                sha256=hashlib.sha256(content).hexdigest(),
                created_by=actor_id,
            )
        )
        await session.flush()
        if project_type == "FREE_RECOMMENDATION":
            session.add(
                RecommendationTemplateMapping(
                    project_id=project_id,
                    template_file_id=file_id,
                    template_sha256=hashlib.sha256(content).hexdigest(),
                    sheet_name="Sheet1",
                    header_row=1,
                    data_start_row=2,
                    mapping_json={"product_name": "商品名称"},
                    confirmed_by=actor_id,
                    confirmed_at=now,
                )
            )
        if project_type == "PPT_SOLUTION":
            session.add(
                PptRecommendationConfig(
                    project_id=project_id,
                    recommendation_mode="SINGLE",
                    price_bands=[{"max_price": "200", "item_count": 1}],
                    candidate_count_per_band=1,
                    plan_count_per_band=1,
                    created_by=actor_id,
                    updated_by=actor_id,
                )
            )
    current = CurrentUser(
        user_id=actor_id,
        username=token,
        roles=[],
        role_names={},
        permissions=[
            "bid:match",
            "bid:detail",
            "bid:update",
            "bid:void",
            "recommendation:create",
            "recommendation:detail",
            "recommendation:run",
            "recommendation:review",
            "recommendation:export",
        ],
        permission_names={},
    )
    app.dependency_overrides[get_current_user] = lambda: current
    app.dependency_overrides[get_current_user_before_response] = lambda: current
    try:
        yield project_id, actor_id, file_id, token, current
    finally:
        app.dependency_overrides.pop(get_current_user, None)
        app.dependency_overrides.pop(get_current_user_before_response, None)
        async with SessionLocal.begin() as session:
            run_ids = select(RecommendationRun.id).where(RecommendationRun.project_id == project_id)
            cids = select(RecommendationCandidate.id).where(
                RecommendationCandidate.run_id.in_(run_ids)
            )
            await session.execute(
                delete(RecommendationConfirmation).where(
                    RecommendationConfirmation.candidate_id.in_(cids)
                )
            )
            await session.execute(
                delete(RecommendationCategoryChoice).where(
                    RecommendationCategoryChoice.run_id.in_(run_ids)
                )
            )
            await session.execute(
                delete(RecommendationCandidate).where(RecommendationCandidate.run_id.in_(run_ids))
            )
            await session.execute(
                delete(PptGenerationTask).where(PptGenerationTask.project_id == project_id)
            )
            await session.execute(
                delete(RecommendationRun).where(RecommendationRun.project_id == project_id)
            )
            await session.execute(
                delete(PptRecommendationConfig).where(
                    PptRecommendationConfig.project_id == project_id
                )
            )
            t3ids = select(Type3Run.id).where(Type3Run.project_id == project_id)
            await session.execute(delete(Type3Export).where(Type3Export.run_id.in_(t3ids)))
            await session.execute(delete(Type3Candidate).where(Type3Candidate.run_id.in_(t3ids)))
            await session.execute(delete(Type3Run).where(Type3Run.project_id == project_id))
            await session.execute(
                delete(Type3SheetConfig).where(Type3SheetConfig.project_id == project_id)
            )
            mids = select(MatchTask.id).where(MatchTask.project_id == project_id)
            await session.execute(
                delete(BidItemSelection).where(
                    BidItemSelection.project_item_id.in_(
                        select(BidProjectItem.id).where(BidProjectItem.project_id == project_id)
                    )
                )
            )
            await session.execute(
                delete(MatchCandidate).where(MatchCandidate.match_task_id.in_(mids))
            )
            await session.execute(delete(MatchTask).where(MatchTask.project_id == project_id))
            await session.execute(
                delete(MappingAnalysisCache).where(MappingAnalysisCache.project_id == project_id)
            )
            await session.execute(delete(AsyncTask).where(AsyncTask.project_id == project_id))
            await session.execute(
                delete(RecommendationTemplateMapping).where(
                    RecommendationTemplateMapping.project_id == project_id
                )
            )
            await session.execute(
                delete(BidProjectFile).where(BidProjectFile.project_id == project_id)
            )
            await session.execute(
                delete(BidProjectItem).where(BidProjectItem.project_id == project_id)
            )
            await session.execute(
                delete(BidProjectEvent).where(BidProjectEvent.project_id == project_id)
            )
            await session.execute(delete(BidProject).where(BidProject.id == project_id))
            await session.execute(delete(BidTemplate).where(BidTemplate.id == template_id))
            await session.execute(delete(Product).where(Product.id == product_id))
            await session.execute(delete(Supplier).where(Supplier.id == supplier_id))
        await storage.delete(key)


async def warm_type3(project_id: uuid.UUID, actor_id: uuid.UUID, file_id: uuid.UUID):
    async with SessionLocal.begin() as session:
        task = await MappingAnalysisService(session).cached_or_enqueue(
            project_id=project_id,
            file_id=file_id,
            analysis_key="TYPE3_TEMPLATE_PREVIEW",
            payload_json={"analysis_type": "TYPE3_TEMPLATE_PREVIEW"},
            actor_id=actor_id,
        )
    assert isinstance(task, AsyncTaskResponse)
    assert await task_worker.run_once({AsyncTaskKind.MAPPING_STRUCTURE_ANALYSIS}, task.id)
    async with SessionLocal() as session:
        result = await MappingAnalysisService(session).require_cached(
            project_id, file_id, "TYPE3_TEMPLATE_PREVIEW"
        )
        return {"sheets": result["sheets"]}


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "kind",
    [
        AsyncTaskKind.TYPE1_MATCH,
        AsyncTaskKind.TYPE2_MATCH,
        AsyncTaskKind.TYPE3_CANDIDATE_GENERATION,
        AsyncTaskKind.TYPE4_RECOMMENDATION_AGENT,
        AsyncTaskKind.TYPE5_RECOMMENDATION_GENERATION,
    ],
)
async def test_http_concurrent_dispatch_cancel_and_state_recovery(kind):
    async with fixture(kind) as (project_id, actor_id, file_id, _token, _current):
        payload = (
            await warm_type3(project_id, actor_id, file_id)
            if kind == AsyncTaskKind.TYPE3_CANDIDATE_GENERATION
            else None
        )
        path = (
            f"/api/v1/bid-projects/{project_id}/commands/start-matching"
            if kind in {AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH}
            else f"/api/v1/type3-recommendation-projects/{project_id}/runs"
            if (kind == AsyncTaskKind.TYPE3_CANDIDATE_GENERATION)
            else f"/api/v1/recommendation-projects/{project_id}/runs"
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            responses = await asyncio.gather(
                client.post(path, json=payload), client.post(path, json=payload)
            )
            assert sorted(r.status_code for r in responses) == [202, 409], [
                r.text for r in responses
            ]
            data = next(r for r in responses if r.status_code == 202).json()["data"]
            async with SessionLocal() as session:
                run_model = (
                    Type3Run
                    if kind == AsyncTaskKind.TYPE3_CANDIDATE_GENERATION
                    else RecommendationRun
                )
                if kind not in {AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH}:
                    assert (
                        await session.scalar(
                            select(func.count())
                            .select_from(run_model)
                            .where(run_model.project_id == project_id)
                        )
                        == 1
                    )
            cancelled = await client.post(
                f"/api/v1/bid-projects/async-tasks/{data['id']}/commands/cancel"
            )
            assert cancelled.status_code == 200, cancelled.text
            async with SessionLocal() as session:
                project = await session.get(BidProject, project_id)
                assert project.status == "IMPORTED"
                task = await session.get(AsyncTask, uuid.UUID(data["id"]))
                assert task.status == "CANCELLED"


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", list(AsyncTaskKind))
async def test_running_stop_is_acknowledged_before_new_work(monkeypatch, kind):
    async with fixture(kind) as (pid, actor, fid, _token, current):
        rid = uuid.uuid4()
        payload = {}
        domain = None
        async with SessionLocal.begin() as session:
            if kind == AsyncTaskKind.TYPE3_CANDIDATE_GENERATION:
                domain = Type3Run(
                    id=rid,
                    project_id=pid,
                    status="ANALYZING",
                    raw_remark_snapshot="test",
                    sheet_snapshot=[],
                    created_by=actor,
                )
            elif kind in {AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH}:
                domain = MatchTask(id=rid, project_id=pid, status="RUNNING", created_by=actor)
            elif kind != AsyncTaskKind.MAPPING_STRUCTURE_ANALYSIS:
                run = RecommendationRun(
                    id=rid,
                    project_id=pid,
                    status="RANKING",
                    raw_requirement_snapshot="test",
                    created_by=actor,
                )
                session.add(run)
                await session.flush()
                domain = run
                if kind == AsyncTaskKind.TYPE5_PPT_GENERATION:
                    domain = PptGenerationTask(
                        project_id=pid,
                        run_id=rid,
                        status="RUNNING",
                        template_code="SYSTEM_DEFAULT",
                        template_version="1",
                        provider="mock",
                        model="mock",
                        prompt_version="mock",
                        created_by=actor,
                    )
                    session.add(domain)
                    await session.flush()
                    payload["ppt_generation_task_id"] = str(domain.id)
            if domain is not None:
                session.add(domain)
            task = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    run_id=rid,
                    task_kind=kind,
                    created_by=actor,
                    payload_json=payload,
                )
            )
            if kind == AsyncTaskKind.MAPPING_STRUCTURE_ANALYSIS:
                session.add(
                    MappingAnalysisCache(
                        project_id=pid,
                        file_id=fid,
                        async_task_id=task.id,
                        analysis_key="stop-test",
                        status="RUNNING",
                    )
                )
            project = await session.get(BidProject, pid)
            project.status = "READY" if kind == AsyncTaskKind.TYPE5_PPT_GENERATION else "MATCHING"
        entered, stopped = asyncio.Event(), asyncio.Event()

        async def blocked(_task):
            entered.set()
            try:
                await asyncio.Event().wait()
            finally:
                stopped.set()

        monkeypatch.setattr(task_worker, "execute", blocked)
        worker = asyncio.create_task(task_worker.run_once({kind}, task.id))
        try:
            await asyncio.wait_for(entered.wait(), 5)
            async with AsyncClient(
                transport=ASGITransport(app=app), base_url="http://test"
            ) as client:
                path = f"/api/v1/bid-projects/async-tasks/{task.id}"
                first = await client.post(f"{path}/commands/cancel")
                assert first.status_code == 200, first.text
                assert first.json()["data"]["status"] == "RUNNING"
                assert first.json()["data"]["cancel_requested_at"]
                duplicate = await client.post(f"{path}/commands/cancel")
                assert duplicate.status_code == 200
                if kind != AsyncTaskKind.MAPPING_STRUCTURE_ANALYSIS:
                    async with SessionLocal() as session:
                        with pytest.raises(AppError) as error:
                            await AsyncTaskService(session).enqueue(
                                AsyncTaskCreate(
                                    project_id=pid,
                                    task_kind=kind,
                                    created_by=actor,
                                )
                            )
                        assert error.value.code == "ASYNC_TASK_ALREADY_ACTIVE"
                assert await asyncio.wait_for(worker, 5)
                assert stopped.is_set()
                final = await client.get(path)
                assert final.json()["data"]["status"] == "CANCELLED"
                async with SessionLocal() as session:
                    await AsyncTaskService(session).succeed(task.id, attempt=1)
                saved_permissions = current.permissions
                current.permissions = ["bid:detail"]
                assert (await client.post(f"{path}/commands/cancel")).status_code == 403
                current.permissions = saved_permissions
                assert (await client.get(path)).status_code == 200
                assert (await client.post(f"{path}/commands/retry")).status_code == 409
                listed = await client.get(f"/api/v1/bid-projects/{pid}/async-tasks")
                assert listed.json()["data"]["total"] == 1
            async with SessionLocal() as session:
                row = await session.get(AsyncTask, task.id)
                assert row.status == "CANCELLED"
                assert await session.get(BidProject, pid) is not None
                assert (await AsyncTaskService(session).latest_for_projects([pid]))[
                    pid
                ].id == task.id
                assert (
                    await session.scalar(
                        select(func.count())
                        .select_from(BidProjectEvent)
                        .where(
                            BidProjectEvent.project_id == pid,
                            BidProjectEvent.event_type == "ASYNC_TASK_CANCEL_REQUESTED",
                        )
                    )
                ) == 1
                if domain is not None:
                    row = await session.get(type(domain), domain.id)
                    assert row.status in {"FAILED", "CANCELLED"}
        finally:
            if not worker.done():
                worker.cancel()
                await asyncio.gather(worker, return_exceptions=True)


@pytest.mark.asyncio
async def test_stopped_mapping_remains_in_task_history():
    async with fixture(AsyncTaskKind.TYPE3_CANDIDATE_GENERATION) as (pid, actor, fid, *_):
        args = dict(
            project_id=pid,
            file_id=fid,
            actor_id=actor,
            analysis_key="TYPE3_TEMPLATE_PREVIEW",
            payload_json={"analysis_type": "TYPE3_TEMPLATE_PREVIEW"},
        )
        async with SessionLocal() as session:
            task = await MappingAnalysisService(session).cached_or_enqueue(**args)
            await AsyncTaskService(session).cancel(task.id, actor)
            await session.commit()
        async with SessionLocal() as session:
            cached = await MappingAnalysisService(session).cached_or_enqueue(**args)
            assert cached.id == task.id
            assert cached.status == "CANCELLED"
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(AsyncTask)
                    .where(
                        AsyncTask.project_id == pid,
                    )
                )
            ) == 1


@pytest.mark.asyncio
async def test_project_purge_physically_removes_database_rows_and_rejects_active_work():
    async with fixture(AsyncTaskKind.TYPE1_MATCH) as (pid, actor, _fid, _token, _current):
        async with SessionLocal() as session:
            task = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid, task_kind=AsyncTaskKind.TYPE1_MATCH, created_by=actor
                )
            )
            project = await session.get(BidProject, pid)
            project_code = project.project_code
            with pytest.raises(AppError) as active_error:
                await BidProjectService(session).purge(pid, project_code, actor)
            assert active_error.value.code == "BID_PROJECT_PURGE_ACTIVE_TASK"
            await AsyncTaskService(session).cancel(task.id, actor)
            await session.commit()
        async with SessionLocal.begin() as session:
            with pytest.raises(AppError) as confirm_error:
                await BidProjectService(session).purge(pid, "BID00000000", actor)
            assert confirm_error.value.code == "BID_PROJECT_PURGE_CONFIRM_MISMATCH"
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.request(
                "DELETE",
                f"/api/v1/bid-projects/{pid}",
                json={"confirm_project_code": project_code},
            )
        assert response.status_code == 200
        assert response.json()["data"] == {"id": str(pid), "status": "PURGED"}
        async with SessionLocal() as session:
            assert await session.get(BidProject, pid) is None
            assert await session.get(AsyncTask, task.id) is None
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(BidProjectFile)
                    .where(BidProjectFile.project_id == pid)
                )
                == 0
            )


@pytest.mark.asyncio
async def test_expired_stop_request_converges_and_ignores_wrong_attempt():
    async with fixture(AsyncTaskKind.TYPE1_MATCH) as (pid, actor, *_):
        async with SessionLocal() as session:
            task = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    task_kind=AsyncTaskKind.TYPE1_MATCH,
                    created_by=actor,
                )
            )
        async with SessionLocal() as session:
            await AsyncTaskService(session).claim_next({AsyncTaskKind.TYPE1_MATCH}, task.id)
        async with SessionLocal() as session:
            await AsyncTaskService(session).cancel(task.id, actor)
            await session.commit()
        async with SessionLocal() as session:
            await AsyncTaskService(session).acknowledge_cancel(task.id, attempt=2)
        async with SessionLocal.begin() as session:
            row = await session.get(AsyncTask, task.id)
            assert row.status == "RUNNING"
            row.lease_expires_at = datetime.now(UTC).replace(tzinfo=None) - timedelta(seconds=1)
        async with SessionLocal() as session:
            await AsyncTaskService(session).recover_stale()
        async with SessionLocal() as session:
            assert (await session.get(AsyncTask, task.id)).status == "CANCELLED"


@pytest.mark.asyncio
async def test_blocking_work_drains_before_acknowledging_cancellation():
    entered, released = asyncio.Event(), threading.Event()
    finished = threading.Event()
    loop = asyncio.get_running_loop()

    def block():
        loop.call_soon_threadsafe(entered.set)
        released.wait(timeout=5)
        finished.set()
        return "late output"

    work = asyncio.create_task(run_blocking(block))
    await asyncio.wait_for(entered.wait(), 2)
    work.cancel()
    await asyncio.sleep(0.02)
    assert not work.done()
    released.set()
    with pytest.raises(asyncio.CancelledError):
        await work
    assert finished.is_set()


@pytest.mark.asyncio
async def test_real_mapping_stop_waits_for_parser_and_discards_its_output(monkeypatch):
    async with fixture(AsyncTaskKind.TYPE3_CANDIDATE_GENERATION) as (pid, actor, fid, *_):
        entered, released = asyncio.Event(), threading.Event()
        loop = asyncio.get_running_loop()
        inspect = Type3RecommendationService._inspect

        def block(content):
            loop.call_soon_threadsafe(entered.set)
            released.wait(timeout=5)
            return inspect(content)

        monkeypatch.setattr(Type3RecommendationService, "_inspect", staticmethod(block))
        async with SessionLocal() as session:
            task = await MappingAnalysisService(session).cached_or_enqueue(
                project_id=pid,
                file_id=fid,
                actor_id=actor,
                analysis_key="TYPE3_TEMPLATE_PREVIEW",
                payload_json={"analysis_type": "TYPE3_TEMPLATE_PREVIEW"},
            )
        worker = asyncio.create_task(
            task_worker.run_once({AsyncTaskKind.MAPPING_STRUCTURE_ANALYSIS}, task.id)
        )
        try:
            await asyncio.wait_for(entered.wait(), 3)
            async with SessionLocal() as session:
                await AsyncTaskService(session).cancel(task.id, actor)
                await session.commit()
            await asyncio.sleep(1.2)
            assert not worker.done()
            async with SessionLocal() as session:
                assert (await session.get(AsyncTask, task.id)).status == "RUNNING"
            released.set()
            assert await asyncio.wait_for(worker, 5)
            async with SessionLocal() as session:
                assert (await session.get(AsyncTask, task.id)).status == "CANCELLED"
                cache = await session.scalar(
                    select(MappingAnalysisCache).where(
                        MappingAnalysisCache.async_task_id == task.id
                    )
                )
                assert cache.status == "FAILED" and cache.result_json is None
        finally:
            released.set()
            if not worker.done():
                worker.cancel()
            await asyncio.gather(worker, return_exceptions=True)


@pytest.mark.asyncio
async def test_real_type4_stop_cancels_provider_and_restores_project(monkeypatch):
    async with fixture(AsyncTaskKind.TYPE4_RECOMMENDATION_AGENT) as (pid, _actor, *_):
        entered, stopped = asyncio.Event(), asyncio.Event()

        class BlockedProvider:
            provider = "mock"
            model = "mock"
            prompt_version = "mock"

            async def structured_completion(self, **kwargs):
                entered.set()
                try:
                    await asyncio.Event().wait()
                finally:
                    stopped.set()

        monkeypatch.setattr(recommendation_agent, "DeepSeekClient", BlockedProvider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(f"/api/v1/recommendation-projects/{pid}/runs")
            assert created.status_code == 202, created.text
            task_id = uuid.UUID(created.json()["data"]["id"])
            worker = asyncio.create_task(
                task_worker.run_once({AsyncTaskKind.TYPE4_RECOMMENDATION_AGENT}, task_id)
            )
            try:
                await asyncio.wait_for(entered.wait(), 3)
                requested = await client.post(
                    f"/api/v1/bid-projects/async-tasks/{task_id}/commands/cancel"
                )
                assert requested.status_code == 200
                assert await asyncio.wait_for(worker, 5)
                assert stopped.is_set()
                async with SessionLocal() as session:
                    task = await session.get(AsyncTask, task_id)
                    assert task.status == "CANCELLED"
                    assert (await session.get(RecommendationRun, task.run_id)).status == "CANCELLED"
                    assert (await session.get(BidProject, pid)).status == "IMPORTED"
            finally:
                if not worker.done():
                    worker.cancel()
                await asyncio.gather(worker, return_exceptions=True)


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", [AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH])
async def test_input_confirmation_and_queue_dispatch_keep_each_others_gates(kind):
    async with fixture(kind) as (pid, actor, _fid, _token, _current):
        async with SessionLocal.begin() as session:
            project = await session.get(BidProject, pid)
            project.import_status = "MAPPING_REQUIRED"
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            rejected = await client.post(f"/api/v1/bid-projects/{pid}/commands/start-matching")
            assert rejected.status_code == 409
            assert rejected.json()["code"] == "BID_PROJECT_NOT_PARSED"
            async with SessionLocal.begin() as session:
                project = await session.get(BidProject, pid)
                project.import_status = "PARSED"
            async with SessionLocal() as session:
                await AsyncTaskService(session).enqueue(
                    AsyncTaskCreate(project_id=pid, task_kind=kind, created_by=actor)
                )
            rejected = await client.put(
                f"/api/v1/bid-projects/{pid}/input-mapping",
                json={
                    "sheet_name": "Sheet1",
                    "header_row": 1,
                    "data_start_row": 2,
                    "mapping": {"product_name": "商品名称"},
                },
            )
            assert rejected.status_code == 409
            assert rejected.json()["code"] == "ASYNC_TASK_ALREADY_ACTIVE"


@pytest.mark.asyncio
async def test_output_cache_and_worker_use_the_frozen_input_sheet_and_header():
    async with fixture(AsyncTaskKind.TYPE1_MATCH) as (pid, _actor, fid, _token, _current):
        workbook = Workbook()
        workbook.active.title = "Sheet1"
        workbook.active.append(["商品名称", "协议价"])
        workbook.active.append(["旧需求", None])
        sheet = workbook.create_sheet("Sheet2")
        sheet.append(["说明"])
        sheet.append(["商品名称", "新协议价"])
        sheet.append(["新需求", None])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        content = output.getvalue()
        storage = get_object_storage()
        key = await storage.save(f"queue-tests/{uuid.uuid4().hex}.xlsx", content)
        try:
            async with SessionLocal.begin() as session:
                file = await session.get(BidProjectFile, fid)
                file.storage_key = key
                file.sha256 = hashlib.sha256(content).hexdigest()
                file.file_size = len(content)
            async with AsyncClient(
                transport=ASGITransport(app=app), base_url="http://test"
            ) as client:
                path = f"/api/v1/bid-projects/{pid}/output-mapping"
                first = await client.get(path)
                assert first.status_code == 202, first.text
                old_task_id = uuid.UUID(first.json()["data"]["id"])
                confirmed = await client.put(
                    f"/api/v1/bid-projects/{pid}/input-mapping",
                    json={
                        "sheet_name": "Sheet2",
                        "header_row": 2,
                        "data_start_row": 3,
                        "mapping": {"product_name": "商品名称"},
                    },
                )
                assert confirmed.status_code == 200, confirmed.text
                assert await task_worker.run_once(
                    {AsyncTaskKind.MAPPING_STRUCTURE_ANALYSIS}, old_task_id
                )
                async with SessionLocal() as session:
                    old_cache = await session.scalar(
                        select(MappingAnalysisCache).where(
                            MappingAnalysisCache.async_task_id == old_task_id
                        )
                    )
                    assert old_cache.result_json["headers"] == ["商品名称", "协议价"]
                second = await client.get(path)
                assert second.status_code == 202, second.text
                new_task_id = uuid.UUID(second.json()["data"]["id"])
                assert new_task_id != old_task_id
                assert await task_worker.run_once(
                    {AsyncTaskKind.MAPPING_STRUCTURE_ANALYSIS}, new_task_id
                )
                cached = await client.get(path)
                assert cached.status_code == 200, cached.text
                assert cached.json()["data"]["headers"] == ["商品名称", "新协议价"]
        finally:
            await storage.delete(key)


@pytest.mark.asyncio
async def test_mapping_cache_commits_deduplicates_and_does_not_wait_for_heavy():
    async with fixture(AsyncTaskKind.TYPE3_CANDIDATE_GENERATION) as (
        pid,
        actor,
        fid,
        _token,
        _current,
    ):
        async with SessionLocal() as session:
            heavy = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    task_kind=AsyncTaskKind.TYPE3_CANDIDATE_GENERATION,
                    created_by=actor,
                )
            )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            path = f"/api/v1/type3-recommendation-projects/{pid}/template-preview"
            responses = await asyncio.gather(*(client.get(path) for _ in range(8)))
            assert all(r.status_code == 202 for r in responses), [r.text for r in responses]
            assert len({r.json()["data"]["id"] for r in responses}) == 1
            task_id = uuid.UUID(responses[0].json()["data"]["id"])
            permissions = list(_current.permissions)
            _current.permissions = ["bid:detail", "bid:update"]
            denied = await client.post(
                f"/api/v1/bid-projects/async-tasks/{task_id}/commands/cancel"
            )
            assert denied.status_code == 403
            _current.permissions = permissions
            assert await task_worker.run_once({AsyncTaskKind.MAPPING_STRUCTURE_ANALYSIS}, task_id)
            cached = await client.get(path)
            assert cached.status_code == 200, cached.text
            assert cached.json()["data"]["template_file_id"] == str(fid)
            async with SessionLocal() as session:
                assert (await session.get(AsyncTask, heavy.id)).status == "PENDING"
                assert (await session.get(AsyncTask, task_id)).status == "SUCCEEDED"


@pytest.mark.asyncio
async def test_stale_type3_recovers_atomically_and_retry_uses_new_run():
    async with fixture(AsyncTaskKind.TYPE3_CANDIDATE_GENERATION) as (
        pid,
        actor,
        fid,
        _token,
        _current,
    ):
        payload = await warm_type3(pid, actor, fid)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.post(
                f"/api/v1/type3-recommendation-projects/{pid}/runs", json=payload
            )
            data = response.json()["data"]
        task_id, old_run_id = uuid.UUID(data["id"]), uuid.UUID(data["run_id"])
        async with SessionLocal() as session:
            task = await AsyncTaskService(session).claim_next(
                {AsyncTaskKind.TYPE3_CANDIDATE_GENERATION}, task_id
            )
            assert task.attempt_count == 1
        async with SessionLocal.begin() as session:
            task = await session.get(AsyncTask, task_id)
            task.lease_expires_at = datetime.now(UTC).replace(tzinfo=None) - timedelta(seconds=1)
        async with SessionLocal() as session:
            assert len(await AsyncTaskService(session).recover_stale()) >= 1
        async with SessionLocal() as session:
            assert (await session.get(AsyncTask, task_id)).status == "FAILED"
            assert (await session.get(Type3Run, old_run_id)).status == "FAILED"
            assert (await session.get(BidProject, pid)).status == "IMPORTED"
        async with SessionLocal() as session:
            retried = await AsyncTaskService(session).retry(task_id)
            assert retried.run_id != old_run_id
        async with SessionLocal() as session:
            claimed = await AsyncTaskService(session).claim_next(
                {AsyncTaskKind.TYPE3_CANDIDATE_GENERATION}, task_id
            )
            assert claimed.attempt_count == 2
        async with SessionLocal() as session:
            await AsyncTaskService(session).succeed(task_id, attempt=1)
        async with SessionLocal() as session:
            assert (await session.get(AsyncTask, task_id)).status == "RUNNING"


@pytest.mark.asyncio
async def test_failed_retry_rejects_another_active_heavy_task():
    async with fixture(AsyncTaskKind.TYPE1_MATCH) as (pid, actor, _fid, _token, _current):
        async with SessionLocal() as session:
            first = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    task_kind=AsyncTaskKind.TYPE1_MATCH,
                    created_by=actor,
                )
            )
        async with SessionLocal() as session:
            await AsyncTaskService(session).fail(first.id, code="TEST", message="test")
        async with SessionLocal() as session:
            await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    task_kind=AsyncTaskKind.TYPE1_MATCH,
                    created_by=actor,
                )
            )
        async with SessionLocal() as session:
            with pytest.raises(AppError, match="项目已有进行中的任务"):
                await AsyncTaskService(session).retry(first.id)


@pytest.mark.asyncio
async def test_worker_type4_result_survives_session_close(monkeypatch):
    async with fixture(AsyncTaskKind.TYPE4_RECOMMENDATION_AGENT) as (
        pid,
        _actor,
        _fid,
        token,
        current,
    ):

        class Provider:
            provider, model, prompt_version = "mock", "mock", "mock"

            async def structured_completion(self, **kwargs):
                return RequirementAnalysis(
                    summary="test",
                    keywords=["test"],
                    required_brands=[token],
                    agreement_price_min=Decimal("987654321.1234"),
                    agreement_price_max=Decimal("987654321.1234"),
                )

        monkeypatch.setattr(recommendation_agent, "DeepSeekClient", Provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.post(f"/api/v1/recommendation-projects/{pid}/runs")
            assert response.status_code == 202, response.text
            data = response.json()["data"]
            # A user with only recommendation permissions must be able to query tasks.
            current.permissions = ["recommendation:detail", "recommendation:run"]
            detail = await client.get(f"/api/v1/bid-projects/async-tasks/{data['id']}")
            assert detail.status_code == 200
        assert await task_worker.run_once(task_worker._HEAVY_KINDS, uuid.UUID(data["id"]))
        async with SessionLocal() as session:
            assert (await session.get(AsyncTask, uuid.UUID(data["id"]))).status == "SUCCEEDED"
            assert (
                await session.get(RecommendationRun, uuid.UUID(data["run_id"]))
            ).status == "WAITING_CONFIRMATION"
            assert (await session.get(BidProject, pid)).status == "SELECTING"
            assert (
                await session.scalar(
                    select(func.count())
                    .select_from(RecommendationCandidate)
                    .where(RecommendationCandidate.run_id == uuid.UUID(data["run_id"]))
                )
                == 1
            )


@pytest.mark.asyncio
async def test_type3_success_enables_confirmation():
    async with fixture(AsyncTaskKind.TYPE3_CANDIDATE_GENERATION) as (
        pid,
        actor,
        fid,
        token,
        _current,
    ):
        payload = await warm_type3(pid, actor, fid)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                f"/api/v1/type3-recommendation-projects/{pid}/runs", json=payload
            )
            data = created.json()["data"]
            rid = uuid.UUID(data["run_id"])
            analysis = Type3RemarkAnalysis(required_brands=[token])
            async with SessionLocal() as session:
                frozen = await Type3RecommendationService(session).prepare_for_category_mapping(
                    rid,
                    analysis,
                    provider="mock",
                    model="mock",
                    prompt_version="mock",
                )
            sheets = frozen["sheets"]
            catalog = frozen["catalog"]
            async with SessionLocal() as session:
                await Type3RecommendationService(session).execute(
                    rid,
                    analysis,
                    provider="mock",
                    model="mock",
                    prompt_version="mock",
                    category_matches={
                        node["node_id"]: [catalog[0]["category_key"]]
                        for sheet in sheets
                        for node in sheet["selected_nodes"]
                    },
                )
            async with SessionLocal() as session:
                assert (await session.get(Type3Run, rid)).status == "WAITING_CONFIRMATION"
                assert (await session.get(BidProject, pid)).status == "SELECTING"
                candidate_id = await session.scalar(
                    select(Type3Candidate.id).where(Type3Candidate.run_id == rid)
                )
            confirmed = await client.post(
                f"/api/v1/type3-recommendation-projects/runs/{rid}/candidate-confirmations",
                json={"candidate_ids": [str(candidate_id)]},
            )
            assert confirmed.status_code == 200, confirmed.text


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", [AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH])
async def test_matching_worker_links_exact_domain_task(kind):
    async with fixture(kind) as (pid, actor, _fid, _token, _current):
        async with SessionLocal() as session:
            task = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    task_kind=kind,
                    created_by=actor,
                )
            )
        assert await task_worker.run_once({kind}, task.id)
        async with SessionLocal() as session:
            assert (await session.get(AsyncTask, task.id)).status == "SUCCEEDED"
            assert (await session.get(MatchTask, task.run_id)).status == "COMPLETED"
            assert (await session.get(BidProject, pid)).status == "READY"


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", list(task_worker._HEAVY_KINDS))
async def test_stale_recovery_finishes_every_domain(kind):
    async with fixture(kind) as (pid, actor, _fid, _token, _current):
        rid = uuid.uuid4()
        payload = {}
        async with SessionLocal.begin() as session:
            if kind == AsyncTaskKind.TYPE3_CANDIDATE_GENERATION:
                session.add(
                    Type3Run(
                        id=rid,
                        project_id=pid,
                        status="ANALYZING",
                        raw_remark_snapshot="test",
                        sheet_snapshot=[],
                        created_by=actor,
                    )
                )
            elif kind not in {AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH}:
                session.add(
                    RecommendationRun(
                        id=rid,
                        project_id=pid,
                        status="RANKING",
                        raw_requirement_snapshot="test",
                        created_by=actor,
                    )
                )
                await session.flush()
                if kind == AsyncTaskKind.TYPE5_PPT_GENERATION:
                    generation = PptGenerationTask(
                        project_id=pid,
                        run_id=rid,
                        status="RUNNING",
                        template_code="SYSTEM_DEFAULT",
                        template_version="1",
                        provider="mock",
                        model="mock",
                        prompt_version="mock",
                        created_by=actor,
                    )
                    session.add(generation)
                    await session.flush()
                    payload["ppt_generation_task_id"] = str(generation.id)
            task = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    run_id=rid,
                    task_kind=kind,
                    created_by=actor,
                    payload_json=payload,
                )
            )
            if kind in {AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH}:
                session.add(MatchTask(id=rid, project_id=pid, status="RUNNING", created_by=actor))
            project = await session.get(BidProject, pid)
            project.status = "MATCHING" if kind != AsyncTaskKind.TYPE5_PPT_GENERATION else "READY"
        async with SessionLocal() as session:
            await AsyncTaskService(session).claim_next({kind}, task.id)
        async with SessionLocal.begin() as session:
            row = await session.get(AsyncTask, task.id)
            row.lease_expires_at = datetime.now(UTC).replace(tzinfo=None) - timedelta(seconds=1)
        async with SessionLocal() as session:
            await AsyncTaskService(session).recover_stale()
        async with SessionLocal() as session:
            assert (await session.get(AsyncTask, task.id)).status == "FAILED"
            project = await session.get(BidProject, pid)
            assert project.status != "MATCHING"
            if kind in {AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH}:
                assert (await session.get(MatchTask, rid)).status == "FAILED"
            elif kind == AsyncTaskKind.TYPE3_CANDIDATE_GENERATION:
                assert (await session.get(Type3Run, rid)).status == "FAILED"
            elif kind == AsyncTaskKind.TYPE5_PPT_GENERATION:
                assert (await session.get(PptGenerationTask, generation.id)).status == "FAILED"
            else:
                assert (await session.get(RecommendationRun, rid)).status == "FAILED"


@pytest.mark.asyncio
async def test_skip_locked_claims_once_retry_limit_and_failure_visibility(monkeypatch):
    async with fixture(AsyncTaskKind.TYPE1_MATCH) as (pid, actor, _fid, _token, _current):
        async with SessionLocal() as session:
            task = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    task_kind=AsyncTaskKind.TYPE1_MATCH,
                    created_by=actor,
                    max_attempts=1,
                    payload_json={"prompt": "secret", "nested": {"token": "secret"}},
                )
            )
        async with SessionLocal.begin() as locked:
            await locked.get(AsyncTask, task.id, with_for_update=True)
            async with SessionLocal() as session:
                assert (
                    await AsyncTaskService(session).claim_next({AsyncTaskKind.TYPE1_MATCH}, task.id)
                    is None
                )
        monkeypatch.setattr(task_worker, "execute", AsyncMock(side_effect=RuntimeError("secret")))
        assert await task_worker.run_once({AsyncTaskKind.TYPE1_MATCH}, task.id)
        async with SessionLocal() as session:
            row = await session.get(AsyncTask, task.id)
            assert row.status == "FAILED" and "secret" not in str(row.payload_json)
            assert "secret" not in row.error_message
        async with SessionLocal() as session:
            with pytest.raises(AppError) as error:
                await AsyncTaskService(session).retry(task.id)
            assert error.value.code == "ASYNC_TASK_RETRY_EXHAUSTED"


@pytest.mark.asyncio
async def test_workers_execute_in_separate_processes():
    async with fixture(AsyncTaskKind.TYPE1_MATCH) as (pid, actor, _fid, _token, _current):
        async with SessionLocal() as session:
            heavy = await AsyncTaskService(session).enqueue(
                AsyncTaskCreate(
                    project_id=pid,
                    task_kind=AsyncTaskKind.TYPE1_MATCH,
                    created_by=actor,
                )
            )
        async with fixture(AsyncTaskKind.TYPE3_CANDIDATE_GENERATION) as (
            mpid,
            mactor,
            mfid,
            _mtoken,
            _mcurrent,
        ):
            async with SessionLocal() as session:
                mapping = await MappingAnalysisService(session).cached_or_enqueue(
                    project_id=mpid,
                    file_id=mfid,
                    actor_id=mactor,
                    analysis_key="TYPE3_TEMPLATE_PREVIEW",
                    payload_json={"analysis_type": "TYPE3_TEMPLATE_PREVIEW"},
                )
            processes = []
            process_logs: list[tuple[bytes, bytes]] = []
            try:
                for queue, task_id in [("heavy", heavy.id), ("mapping", mapping.id)]:
                    processes.append(
                        await asyncio.create_subprocess_exec(
                            sys.executable,
                            "-m",
                            "app.jobs.task_worker",
                            "--queue",
                            queue,
                            "--once",
                            "--task-id",
                            str(task_id),
                            stdout=asyncio.subprocess.PIPE,
                            stderr=asyncio.subprocess.PIPE,
                        )
                    )
                for process in processes:
                    out, err = await asyncio.wait_for(process.communicate(), 45)
                    process_logs.append((out, err))
                    assert process.returncode == 0, err.decode(errors="replace")
                async with SessionLocal() as session:
                    assert (await session.get(AsyncTask, heavy.id)).status == "SUCCEEDED", (
                        "\n".join(
                            f"stdout:\n{out.decode(errors='replace')}\nstderr:\n{err.decode(errors='replace')}"
                            for out, err in process_logs
                        )
                    )
                    assert (await session.get(AsyncTask, mapping.id)).status == "SUCCEEDED"
                    cache = await MappingAnalysisService(session).require_cached(
                        mpid, mfid, "TYPE3_TEMPLATE_PREVIEW"
                    )
                    assert cache["sheets"]
            finally:
                for process in processes:
                    if process.returncode is None:
                        process.kill()
                        await process.wait()


@pytest.mark.asyncio
async def test_two_uvicorn_workers_serve_health_concurrently():
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        "-m",
        "uvicorn",
        "app.main:app",
        "--host",
        "127.0.0.1",
        "--port",
        str(port),
        "--workers",
        "2",
        "--log-level",
        "warning",
        stdout=asyncio.subprocess.DEVNULL,
        stderr=asyncio.subprocess.DEVNULL,
    )
    try:
        async with AsyncClient(base_url=f"http://127.0.0.1:{port}", timeout=3) as client:
            for _ in range(40):
                try:
                    response = await client.get("/health/ready")
                    if response.status_code == 200:
                        break
                except HTTPError:
                    pass
                assert process.returncode is None
                await asyncio.sleep(0.25)
            else:
                pytest.fail("two-worker API did not become ready")
            responses = await asyncio.gather(*(client.get("/health/live") for _ in range(16)))
            assert all(response.status_code == 200 for response in responses)
    finally:
        if process.returncode is None:
            if os.name == "nt":
                killer = await asyncio.create_subprocess_exec(
                    "taskkill",
                    "/PID",
                    str(process.pid),
                    "/T",
                    "/F",
                    stdout=asyncio.subprocess.DEVNULL,
                    stderr=asyncio.subprocess.DEVNULL,
                )
                await killer.wait()
            else:
                process.terminate()
            await asyncio.wait_for(process.wait(), 10)


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", [AsyncTaskKind.TYPE1_MATCH, AsyncTaskKind.TYPE2_MATCH])
async def test_late_failed_matching_handler_does_not_reset_new_task(kind):
    async with fixture(kind) as (pid, actor, _fid, _token, _current):
        old_id = uuid.uuid4()
        async with SessionLocal.begin() as session:
            session.add(MatchTask(id=old_id, project_id=pid, status="FAILED", created_by=actor))
            project = await session.get(BidProject, pid)
            project.status = "MATCHING"
        async with SessionLocal() as session:
            service = MatchingService(session)
            callback = (
                service._fail_type1_ai_matching
                if kind == AsyncTaskKind.TYPE1_MATCH
                else service._fail_type2_ai_matching
            )
            if kind == AsyncTaskKind.TYPE1_MATCH:
                await callback(old_id, pid, actor, "late callback", "IMPORTED")
            else:
                await callback(old_id, pid, actor, "late callback")
        async with SessionLocal() as session:
            assert (await session.get(BidProject, pid)).status == "MATCHING"
