from __future__ import annotations

import uuid
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select

from app.core.database import SessionLocal
from app.modules.bid.domain.lifecycle import (
    BidFileType,
    BidImportStatus,
    BidProjectStatus,
    BidProjectType,
)
from app.modules.bid.infrastructure.models import BidProject, BidProjectEvent, BidProjectFile
from app.modules.catalog.domain.lifecycle import ProductStatus
from app.modules.catalog.infrastructure.models import Product
from app.modules.recommendation.api.recommendation_router import router as recommendation_router
from app.modules.recommendation.application.export_service import RecommendationExportService
from app.modules.recommendation.application.service import RecommendationService
from app.modules.recommendation.infrastructure.models import (
    RecommendationCandidate,
    RecommendationConfirmation,
    RecommendationExport,
    RecommendationRun,
)
from app.modules.recommendation.schemas import ConfirmationUpdateRequest
from app.modules.recommendation.template.models import RecommendationTemplateMapping
from app.modules.recommendation.template.schemas import RecommendationRunStatus
from app.modules.supplier.domain.rules import ArchiveStatus, CooperationStatus
from app.modules.supplier.infrastructure.models import Supplier


class MemoryStorage:
    def __init__(self) -> None:
        self.files: dict[str, bytes] = {}

    async def save(self, name: str, content: bytes) -> str:
        self.files[name] = content
        return name

    async def read(self, key: str) -> bytes:
        return self.files[key]

    async def delete(self, key: str) -> None:
        self.files.pop(key, None)


def test_export_transaction_commits_before_response_is_sent() -> None:
    route = next(
        route
        for route in recommendation_router.routes
        if getattr(route, "path", "").endswith("/{project_id}/runs/{run_id}/exports")
        and "POST" in getattr(route, "methods", set())
    )
    session_dependency = next(
        dependency for dependency in route.dependant.dependencies if dependency.name == "session"
    )

    assert session_dependency.scope == "function"


def _template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "推荐清单"
    sheet["A1"] = "自由推品确认结果"
    headers = [
        "品牌",
        "名称",
        "京东价",
        "协议价",
        "毛利率",
        "是否厂直",
        "所属公司",
        "活动价",
        "履约说明",
        "依据与备注",
        "校验",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(2, column).value = header
    sheet["A3"].font = Font(bold=True)
    sheet["E3"].number_format = "0.00%"
    sheet["A3"], sheet["B3"] = "示例品牌A", "示例商品A"
    sheet["A4"], sheet["B4"] = "示例品牌B", "示例商品B"
    sheet["A5"], sheet["B5"] = "示例品牌C", "示例商品C"
    sheet["K3"] = '=IF(B3<>"","OK","")'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.asyncio
async def test_export_confirmed_candidates_preserves_template_and_versions() -> None:
    actor_id = uuid.uuid4()
    token = uuid.uuid4().hex
    storage = MemoryStorage()
    template_key = f"test/recommendation-template-{token}.xlsx"
    storage.files[template_key] = _template()

    async with SessionLocal() as session:
        supplier = Supplier(
            supplier_code=f"RE{token[:12]}",
            supplier_name=f"导出测试供应商-{token}",
            main_brands="导出品牌",
            advantage="导出测试",
            archive_status=ArchiveStatus.ARCHIVED.value,
            cooperation_status=CooperationStatus.NORMAL.value,
        )
        project = BidProject(
            project_code=f"EX{token[:10]}",
            project_name="自由推品导出测试",
            buyer_name="测试客户",
            status=BidProjectStatus.IMPORTED.value,
            import_status=BidImportStatus.NOT_REQUIRED.value,
            project_type=BidProjectType.FREE_RECOMMENDATION.value,
            remark="导出已确认自由推品候选。",
            created_by=actor_id,
        )
        session.add_all([supplier, project])
        await session.flush()
        template_file = BidProjectFile(
            project_id=project.id,
            file_type=BidFileType.RECOMMENDATION_TEMPLATE.value,
            version_no=1,
            original_filename="推荐模板.xlsx",
            storage_key=template_key,
            file_size=len(storage.files[template_key]),
            sha256="a" * 64,
            created_by=actor_id,
        )
        product = Product(
            source_supplier_id=supplier.id,
            sku=f"SKU-{token[:8]}",
            product_name="确认导出商品",
            brand="导出品牌",
            category_level1_name="食品饮料",
            category_level2_name="休闲食品",
            category_level3_name="坚果",
            jd_price=Decimal("100"),
            agreement_price=Decimal("88"),
            gross_margin=Decimal("0.12"),
            status=ProductStatus.ACTIVE.value,
        )
        session.add_all([template_file, product])
        await session.flush()
        mapping = RecommendationTemplateMapping(
            project_id=project.id,
            template_file_id=template_file.id,
            template_sha256=template_file.sha256,
            sheet_name="推荐清单",
            header_row=2,
            data_start_row=3,
            mapping_json={
                "brand": "品牌",
                "product_name": "名称",
                "jd_price": "京东价",
                "agreement_price": "协议价",
                "gross_margin": "毛利率",
                "factory_direct": "是否厂直",
                "company_name": "所属公司",
                "campaign_price": "活动价",
                "fulfillment_cycle": "履约说明",
                "evidence": "依据与备注",
            },
            confirmed_by=actor_id,
            confirmed_at=datetime.now(UTC).replace(tzinfo=None),
        )
        run = RecommendationRun(
            project_id=project.id,
            status=RecommendationRunStatus.CONFIRMED.value,
            raw_requirement_snapshot="导出已确认自由推品候选。",
            parsed_requirement={},
            created_by=actor_id,
        )
        session.add_all([mapping, run])
        await session.flush()
        candidate = RecommendationCandidate(
            run_id=run.id,
            product_id=product.id,
            rank=1,
            score=Decimal("99"),
            product_snapshot={
                "brand": "导出品牌",
                "product_name": "确认导出商品",
                "company_name": "导出所属公司",
            },
            supplier_snapshot={},
            price_snapshot={
                "jd_price": "100",
                "agreement_price": "88",
                "gross_margin": "0.12",
            },
        )
        session.add(candidate)
        await session.flush()
        session.add(
            RecommendationConfirmation(
                candidate_id=candidate.id,
                campaign_price=Decimal("77"),
                factory_direct="YES",
                fulfillment_cycle="48小时",
                evidence="供应商确认",
                confirmed_by=actor_id,
                confirmed_at=datetime.now(UTC).replace(tzinfo=None),
            )
        )
        await session.flush()

        service = RecommendationExportService(session, storage)  # type: ignore[arg-type]
        first = await service.export(project.id, run.id, actor_id)
        second = await service.export(project.id, run.id, actor_id)

        assert first.file_type == BidFileType.RECOMMENDATION_EXPORT
        assert first.version_no == 1
        assert second.version_no == 2
        assert run.status == RecommendationRunStatus.EXPORTED.value
        downloaded_file, exported = await service.download_export(run.id, first.id)
        assert downloaded_file.id == first.id
        workbook = load_workbook(BytesIO(exported), data_only=False)
        sheet = workbook["推荐清单"]
        assert sheet["A1"].value == "自由推品确认结果"
        assert sheet["A3"].value == "导出品牌"
        assert sheet["B3"].value == "确认导出商品"
        assert sheet["C3"].value == 100
        assert sheet["D3"].value == 88
        assert Decimal(str(sheet["E3"].value)) == Decimal("0.12")
        assert sheet["E3"].number_format == "0.00%"
        assert sheet["F3"].value == "是"
        assert sheet["G3"].value == "导出所属公司"
        assert sheet["H3"].value == 77
        assert sheet["I3"].value == "48小时"
        assert sheet["J3"].value == "供应商确认"
        assert sheet["K3"].value == '=IF(B3<>"","OK","")'
        assert sheet["A4"].value is None
        assert sheet["B4"].value is None
        assert sheet["A5"].value is None
        assert sheet["B5"].value is None
        assert sheet["A3"].font.bold is True
        workbook.close()

        records = list(
            (
                await session.scalars(
                    select(RecommendationExport)
                    .where(RecommendationExport.run_id == run.id)
                    .order_by(RecommendationExport.version_no)
                )
            ).all()
        )
        assert [record.version_no for record in records] == [1, 2]
        assert records[0].mapping_snapshot["mapping_json"]["factory_direct"] == "是否厂直"
        events = list(
            (
                await session.scalars(
                    select(BidProjectEvent)
                    .where(BidProjectEvent.project_id == project.id)
                    .order_by(BidProjectEvent.occurred_at)
                )
            ).all()
        )
        assert [event.event_type for event in events] == [
            "RECOMMENDATION_EXPORTED",
            "RECOMMENDATION_EXPORTED",
        ]
        await RecommendationService(session).confirm_candidate(
            candidate.id,
            ConfirmationUpdateRequest(campaign_price=Decimal("66")),
            actor_id,
        )
        assert run.status == RecommendationRunStatus.CONFIRMED.value
        third = await service.export(project.id, run.id, actor_id)
        assert third.version_no == 3
        assert run.status == RecommendationRunStatus.EXPORTED.value
        await session.rollback()
