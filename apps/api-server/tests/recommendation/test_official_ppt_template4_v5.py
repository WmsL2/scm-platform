"""Fictional structural tests for the independent template 4 v5 renderer."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Pt

from app.common.contracts import AppError
from app.modules.recommendation.application import ppt_official_template4_v5 as template4
from app.modules.recommendation.application.ppt_renderer import PptRenderer
from app.modules.recommendation.application.ppt_template_manifest import get_pptx_manifest
from app.modules.recommendation.application.ppt_template_registry import get_ppt_template


def _png(colour: tuple[int, int, int]) -> bytes:
    output = BytesIO()
    Image.new("RGB", (120, 80), colour).save(output, format="PNG")
    return output.getvalue()


def _source_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = 12_192_000, 6_858_000
    prototype = deck.slides.add_slide(deck.slide_layouts[6])
    prototype.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 5_721_719, 765_621, 5_721_719, 6_399_904)
    body = prototype.shapes.add_textbox(294_640, 1_668_145, 5_331_460, 4_731_385)
    body.text = (
        "sku：（自动填入）\n商品配置：\n（自动填入）\n京东价：（自动填入）\n协议价：（自动填入）"
    )
    title = prototype.shapes.add_textbox(150_495, 235_585, 9_812_020, 755_650)
    title.text = "商品名称：（自动填入）"
    slot = prototype.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 6_572_250, 1_419_225, 4_680_000, 4_680_000
    )
    slot.text = "图片自动填入"
    sample = deck.slides.add_slide(deck.slide_layouts[6])
    sample.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 5_721_719, 765_621, 5_721_719, 6_399_904)
    sample.shapes.add_textbox(294_333, 1_122_251, 4_747_260, 5_277_653).text = "虚构示例"
    sample.shapes.add_textbox(150_495, 235_774, 10_249_001, 755_650).text = "虚构示例商品"
    sample.shapes.add_picture(
        BytesIO(_png((180, 20, 20))), 6_143_625, 991_235, 5_400_000, 5_400_000
    )
    deck.save(path)


def _runtime(tmp_path: Path) -> Path:
    source = tmp_path / "fictional-template4.pptx"
    spec = tmp_path / "fictional-template4.txt"
    runtime = tmp_path / "UNION_QUOTE_WHITE_OFFICIAL_CLEAN_v5.pptx"
    _source_deck(source)
    spec.write_text("虚构说明", encoding="utf-8")
    template4.standardize_union_quote_white_v5(
        source,
        spec,
        runtime,
        expected_source_sha256=None,
        expected_specification_sha256=None,
        expected_source_fingerprint=template4._structure_fingerprint(Presentation(source)),
    )
    template4.verify_no_sample_media(source, runtime)
    return runtime


def _single_source() -> dict[str, Any]:
    return {
        "single_products": [
            {
                "product": {
                    "id": "p1",
                    "sku": "SKU-001",
                    "product_name": "虚构办公礼包",
                    "product_specification": "蓝色，标准装",
                    "reference_url": "https://example.invalid/p1",
                    "price_validity": "2026-12-31",
                    "shipping_courier": "京东物流",
                    "remark": "虚构商品备注",
                    "purchasing_agent": "虚构采销",
                },
                "prices": {"jd_price": "229", "agreement_price": "199"},
                "manual": {"evidence": "人工复核通过"},
            }
        ]
    }


def _render(
    runtime: Path,
    source: dict[str, Any],
    monkeypatch: pytest.MonkeyPatch,
    images: dict[str, bytes] | None = None,
) -> Presentation:
    monkeypatch.setattr(template4, "resolve_ppt_template_asset", lambda _template: runtime)
    content = PptRenderer().render(
        source,
        template_code="UNION_QUOTE_WHITE",
        template_version="5",
        product_images=images or {},
    )
    return Presentation(BytesIO(content))


def _shapes(deck: Presentation, slide_index: int = 0) -> dict[str, Any]:
    return {shape.name: shape for shape in deck.slides[slide_index].shapes}


def test_registry_adds_exact_v5_without_changing_v3_default(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "app.modules.recommendation.application.ppt_template_registry.is_local_v5_default_enabled",
        lambda: False,
    )
    assert get_ppt_template("UNION_QUOTE_WHITE").version == "3"
    template = get_ppt_template("UNION_QUOTE_WHITE", "5")
    assert template.version == "5" and template.single_product_only is False
    manifest = get_pptx_manifest("UNION_QUOTE_WHITE", "5")
    assert manifest is not None
    assert manifest.clean_filename == "UNION_QUOTE_WHITE_OFFICIAL_CLEAN_v5.pptx"
    assert manifest.source_text_slots == {"SCM_FIELD_UNION_BODY": 1, "SCM_FIELD_TITLE": 2}


def test_builder_keeps_only_clean_prototype_and_removes_sample_media(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    assert len(deck.slides) == 1
    slots = _shapes(deck)
    assert set(template4._GEOMETRY) <= set(slots)
    assert slots["SCM_FIELD_TITLE"].text == ""
    assert slots["SCM_FIELD_UNION_BODY"].text == ""
    assert slots["SCM_SLOT_IMAGE_01"].shape_type == MSO_SHAPE_TYPE.PICTURE
    assert slots["SCM_SLOT_IMAGE_01"].image.size == (1, 1)


def test_single_product_uses_exact_fields_styles_and_one_image(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = _render(_runtime(tmp_path), _single_source(), monkeypatch, {"p1": _png((20, 120, 200))})
    assert len(deck.slides) == 1
    slots = _shapes(deck)
    title = slots["SCM_FIELD_TITLE"].text_frame.paragraphs[0].runs[0]
    assert title.text == "虚构办公礼包"
    assert title.font.name == "Microsoft YaHei"
    assert title.font.size == Pt(24) and title.font.bold and title.font.underline
    body = slots["SCM_FIELD_UNION_BODY"]
    assert body.text_frame.paragraphs[2].text == "蓝色，标准装"
    assert "京东价：¥ 229.00" in body.text
    assert "协议价：¥ 199.00" in body.text
    assert "参考链接：https://example.invalid/p1" in body.text
    assert "备注：虚构商品备注；人工复核通过" in body.text
    assert all(paragraph.line_spacing == 1.8 for paragraph in body.text_frame.paragraphs)
    agreement = next(p for p in body.text_frame.paragraphs if p.text.startswith("协议价："))
    assert agreement.runs[0].font.color.rgb == RGBColor(255, 0, 0)
    pictures = [
        shape for shape in deck.slides[0].shapes if shape.name.startswith("SCM_RENDERED_IMAGE_")
    ]
    assert len(pictures) == 1 and pictures[0].image.size == (120, 80)
    assert not any(shape.is_placeholder for shape in deck.slides[0].shapes)


def test_combination_lists_names_sums_exact_prices_and_uses_one_image_per_product(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    items = []
    images: dict[str, bytes] = {}
    for index, (quantity, jd_price, agreement_price) in enumerate(
        ((2, 10, 8), (1, 20, 15), (3, 5, 4)), start=1
    ):
        product_id = f"p{index}"
        images[product_id] = _png((index * 40, 80, 140))
        items.append(
            {
                "quantity": quantity,
                "product_snapshot": {
                    "id": product_id,
                    "sku": f"S{index}",
                    "product_name": f"虚构商品{index}",
                    "shipping_courier": "京东物流",
                    "purchasing_agent": "虚构采销",
                },
                "price_snapshot": {"jd_price": jd_price, "agreement_price": agreement_price},
            }
        )
    source = {
        "packages": [{"id": "pkg1", "name": "虚构组合方案", "reason": "组合说明", "items": items}]
    }
    deck = _render(_runtime(tmp_path), source, monkeypatch, images)
    body = _shapes(deck)["SCM_FIELD_UNION_BODY"]
    assert "1. 虚构商品1 × 2" in body.text
    assert "2. 虚构商品2 × 1" in body.text
    assert "3. 虚构商品3 × 3" in body.text
    assert "京东价：¥ 55.00" in body.text
    assert "协议价：¥ 43.00" in body.text
    assert "备注：组合说明" in body.text
    pictures = [
        shape for shape in deck.slides[0].shapes if shape.name.startswith("SCM_RENDERED_IMAGE_")
    ]
    assert len(pictures) == 3


def test_missing_values_keep_labels_without_placeholder_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = {"single_products": [{"product": {"id": "p1", "product_name": "虚构空字段"}}]}
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    text = _shapes(deck)["SCM_FIELD_UNION_BODY"].text
    assert "京东价：" in text and "协议价：" in text and "价格有效期：" in text
    assert "自动填入" not in text and "{{" not in text


def test_long_title_and_body_fail_instead_of_shrinking_or_dropping(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _single_source()
    source["single_products"][0]["product"]["product_name"] = "超长" * 30
    with pytest.raises(AppError) as error:
        _render(_runtime(tmp_path), source, monkeypatch)
    assert error.value.code == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"


def test_multiline_specification_and_tracking_url_fit_without_losing_content(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _single_source()
    product = source["single_products"][0]["product"]
    product["sku"] = "100269697775"
    product["product_specification"] = (
        "\u9632\u6c34\u7b49\u7ea7\nIPX5\n\u5145\u7535\u63a5\u53e3\nType-C\n\u9632\u5c18\u7b49\u7ea7\nIP5X\n"
        "\u5305\u88c5\u5f62\u5f0f\n\u666e\u901a\u88c5\n\u97f3\u9891\u8ba4\u8bc1\n"
        "Hi-Res Audio\n\u4f69\u6234\u65b9\u5f0f\n"
        "\u5165\u8033\u5f0f\n\u5305\u88c5\u6e05\u5355\n\u8033\u673a*1\uff0c\u5145\u7535\u7ebf*1"
    )
    product["reference_url"] = (
        "https://item.jd.com/10219427574784.html?pcdk=tracking-value&spmTag=more#switch-sku"
    )
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    body = _shapes(deck)["SCM_FIELD_UNION_BODY"].text
    assert "\u9632\u6c34\u7b49\u7ea7\uff1bIPX5\uff1b\u5145\u7535\u63a5\u53e3\uff1bType-C" in body
    assert "\u8033\u673a*1\uff0c\u5145\u7535\u7ebf*1" in body
    assert "https://item.jd.com/10219427574784.html" in body
    assert "pcdk=" not in body and "#switch-sku" not in body
