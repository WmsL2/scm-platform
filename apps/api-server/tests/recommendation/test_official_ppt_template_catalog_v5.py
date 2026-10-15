"""Fictional structural and rendering coverage for CATALOG_MINIMAL v5."""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path
from shutil import copyfile
from tempfile import NamedTemporaryFile
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from lxml import etree  # type: ignore[import-untyped]
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from app.common.contracts import AppError
from app.core.config import PROJECT_ROOT
from app.modules.recommendation.application import ppt_official_template_catalog_v5 as catalog
from app.modules.recommendation.application import ppt_template_registry
from app.modules.recommendation.application.ppt_renderer import PptRenderer
from app.modules.recommendation.application.ppt_template_manifest import get_pptx_manifest
from app.modules.recommendation.application.ppt_template_registry import (
    get_ppt_template,
    is_ppt_template_available,
)
from scripts.build_private_official_ppt_template_v5 import (
    build_catalog_minimal_v5_runtime_asset,
)


def _png(colour: tuple[int, int, int], size: tuple[int, int] = (160, 100)) -> bytes:
    output = BytesIO()
    Image.new("RGB", size, colour).save(output, format="PNG")
    return output.getvalue()


def _style(run: Any, size: int, colour: tuple[int, int, int], *, bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*colour)


def _set_runs(
    paragraph: Any,
    values: tuple[tuple[str, tuple[int, int, int], bool], ...],
    size: int,
) -> None:
    paragraph.clear()
    paragraph.line_spacing = 1.5
    for text, colour, bold in values:
        run = paragraph.add_run()
        run.text = text
        _style(run, size, colour, bold=bold)


def _source_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    empty = deck.slides.add_slide(deck.slide_layouts[6])
    image_instruction = empty.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.556), Inches(0.678), Inches(5.446), Inches(6.153)
    )
    image_instruction.text = "图片"
    content = empty.shapes.add_textbox(
        Inches(6.365), Inches(0.692), Inches(5.366), Inches(6.600)
    )
    frame = content.text_frame
    frame.word_wrap = True
    while len(frame.paragraphs) < 21:
        frame.add_paragraph()
    _set_runs(frame.paragraphs[0], (("（型号，自动填入）", (175, 149, 138), True),), 18)
    _set_runs(frame.paragraphs[1], (("（商品名称，自动填入）", (89, 89, 89), False),), 18)
    _set_runs(frame.paragraphs[3], (("规格参数：", (89, 89, 89), False),), 9)
    _set_runs(frame.paragraphs[4], (("（自动填入）", (89, 89, 89), False),), 9)
    _set_runs(
        frame.paragraphs[9],
        (("零  售  价：", (89, 89, 89), True), ("（自动填入）", (175, 149, 138), True)),
        9,
    )
    _set_runs(
        frame.paragraphs[10],
        (("团购价：", (255, 0, 0), True), ("（自动填入）", (255, 0, 0), True)),
        9,
    )
    _set_runs(frame.paragraphs[12], (("产品特点：", (89, 89, 89), True),), 9)
    _set_runs(frame.paragraphs[13], (("（自动填入卖点）", (89, 89, 89), False),), 9)
    _set_runs(
        frame.paragraphs[19],
        (("京东链接：", (89, 89, 89), True), ("（自动填入）", (89, 89, 89), False)),
        9,
    )
    empty.shapes.add_picture(
        BytesIO(_png((157, 144, 139), (262, 270))),
        Inches(11.508), Inches(0.789), Inches(1.191), Inches(1.226),
    )
    empty.shapes.add_picture(
        BytesIO(_png((167, 139, 127), (262, 287))),
        Inches(11.508), Inches(2.971), Inches(1.191), Inches(1.303),
    )

    sample = deck.slides.add_slide(deck.slide_layouts[6])
    sample.shapes.add_textbox(
        Inches(6.365), Inches(0.692), Inches(5.251), Inches(6.600)
    ).text = "样例业务文本"
    sample.shapes.add_textbox(
        Inches(11.009), Inches(2.333), Inches(1.774), Inches(0.416)
    ).text = "样例分类"
    sample.shapes.add_picture(
        BytesIO(_png((157, 144, 139), (262, 270))),
        Inches(11.508), Inches(0.789), Inches(1.191), Inches(1.226),
    )
    sample.shapes.add_picture(
        BytesIO(_png((167, 139, 127), (262, 287))),
        Inches(11.508), Inches(2.971), Inches(1.191), Inches(1.303),
    )
    sample.shapes.add_picture(
        BytesIO(_png((12, 34, 56), (398, 446))),
        Inches(0.533), Inches(0.692), Inches(5.493), Inches(6.156),
    )
    deck.save(path)


def _inject_fictional_package_metadata(path: Path) -> None:
    """Add synthetic identity/history parts without using customer metadata."""
    with ZipFile(path) as archive:
        infos = archive.infolist()
        payloads = {info.filename: archive.read(info.filename) for info in infos}
    presentation_rels = "ppt/_rels/presentation.xml.rels"
    relationships = etree.fromstring(payloads[presentation_rels])
    namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    for relation_id, relation_type, target in (
        (
            "rIdFictionalAuthors",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors",
            "commentAuthors.xml",
        ),
        (
            "rIdFictionalTags",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags",
            "tags/tag-fictional.xml",
        ),
    ):
        relationship = etree.SubElement(relationships, f"{{{namespace}}}Relationship")
        relationship.set("Id", relation_id)
        relationship.set("Type", relation_type)
        relationship.set("Target", target)
    payloads[presentation_rels] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    root_rels = etree.fromstring(payloads["_rels/.rels"])
    custom_relationship = etree.SubElement(root_rels, f"{{{namespace}}}Relationship")
    custom_relationship.set("Id", "rIdFictionalCustom")
    custom_relationship.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties",
    )
    custom_relationship.set("Target", "docProps/custom.xml")
    payloads["_rels/.rels"] = etree.tostring(
        root_rels, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    core = etree.fromstring(payloads["docProps/core.xml"])
    namespace_map = {
        "dc": "http://purl.org/dc/elements/1.1/",
        "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    }
    creator = core.find("dc:creator", namespaces=namespace_map)
    modified = core.find("cp:lastModifiedBy", namespaces=namespace_map)
    if creator is None:
        creator = etree.SubElement(core, f"{{{namespace_map['dc']}}}creator")
    if modified is None:
        modified = etree.SubElement(core, f"{{{namespace_map['cp']}}}lastModifiedBy")
    creator.text = "fictional-reviewer@example.invalid"
    modified.text = "Fictional Reviewer"
    payloads["docProps/core.xml"] = etree.tostring(
        core, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    app = etree.fromstring(payloads["docProps/app.xml"])
    app_namespace = {
        "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    }
    app.find("ep:Slides", namespaces=app_namespace).text = "2"  # type: ignore[union-attr]
    app.find("ep:Notes", namespaces=app_namespace).text = "1"  # type: ignore[union-attr]
    payloads["docProps/app.xml"] = etree.tostring(
        app, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    content_types = etree.fromstring(payloads["[Content_Types].xml"])
    content_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    for part_name, content_type in (
        (
            "/ppt/commentAuthors.xml",
            "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml",
        ),
        (
            "/ppt/tags/tag-fictional.xml",
            "application/vnd.openxmlformats-officedocument.presentationml.tags+xml",
        ),
        (
            "/docProps/custom.xml",
            "application/vnd.openxmlformats-officedocument.custom-properties+xml",
        ),
    ):
        override = etree.SubElement(content_types, f"{{{content_namespace}}}Override")
        override.set("PartName", part_name)
        override.set("ContentType", content_type)
    payloads["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    additions = {
        "ppt/commentAuthors.xml": (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<p:cmAuthorLst xmlns:p="http://schemas.openxmlformats.org/'
            b'presentationml/2006/main"><p:cmAuthor id="0" name="fictional-reviewer@'
            b'example.invalid" initials="FR" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>'
        ),
        "ppt/tags/tag-fictional.xml": (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<p:tagLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            b'<p:tag name="fictional-device" val="fictional-value"/></p:tagLst>'
        ),
        "docProps/custom.xml": (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/'
            b'2006/custom-properties"/>'
        ),
    }
    with NamedTemporaryFile(dir=path.parent, suffix=".pptx", delete=False) as handle:
        candidate = Path(handle.name)
    try:
        with ZipFile(candidate, "w", compression=ZIP_DEFLATED) as archive:
            for info in infos:
                archive.writestr(info, payloads[info.filename])
            for name, payload in additions.items():
                archive.writestr(name, payload)
        candidate.replace(path)
    finally:
        candidate.unlink(missing_ok=True)


def _runtime(tmp_path: Path) -> Path:
    source = tmp_path / "fictional-template3.pptx"
    runtime = tmp_path / "CATALOG_MINIMAL_OFFICIAL_CLEAN_v5.pptx"
    _source_deck(source)
    deck = Presentation(source)
    catalog.standardize_catalog_minimal_v5(
        source,
        runtime,
        expected_source_fingerprint=catalog._structure_fingerprint(deck),
        expected_source_sha256=None,
    )
    catalog.verify_no_sample_media(source, runtime)
    return runtime


def _item(index: int = 1) -> dict[str, object]:
    return {
        "product": {
            "id": f"fictional-{index}",
            "model": f"MODEL-{index:03d}",
            "product_name": f"虚构极简商品 {index}",
            "product_specification": "颜色：虚构蓝\n尺寸：100×80×60mm",
            "selling_points": "虚构卖点一；虚构卖点二；虚构卖点三",
            "jd_same_product_url": f"https://example.invalid/jd/{index}",
        },
        "prices": {
            "market_price": "999.00",
            "jd_price": "299.00",
            "agreement_price": "239.00",
        },
    }


def _render(
    runtime: Path,
    source: dict[str, object],
    monkeypatch: pytest.MonkeyPatch,
    images: dict[str, bytes] | None = None,
) -> Presentation:
    monkeypatch.setattr(catalog, "resolve_ppt_template_asset", lambda _template: runtime)
    content = PptRenderer(layout_preview=False).render(
        source,
        product_images=images or {},
        template_code="CATALOG_MINIMAL",
        template_version="5",
    )
    return Presentation(BytesIO(content))


def _content(slide: Any) -> Any:
    return next(shape for shape in slide.shapes if shape.name == "SCM_FIELD_CATALOG_CONTENT")


def test_standardizer_removes_sample_page_media_and_notes(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    assert len(deck.slides) == 1
    assert {shape.name for shape in deck.slides[0].shapes} == {
        "SCM_SLOT_IMAGE_MAIN",
        "SCM_FIELD_CATALOG_CONTENT",
        "SCM_STATIC_DECORATION_01",
        "SCM_STATIC_DECORATION_02",
    }
    assert "样例" not in "\n".join(
        shape.text for shape in deck.slides[0].shapes if shape.has_text_frame
    )
    assert all("notesSlide" not in rel.reltype for rel in deck.slides[0].part.rels.values())


def test_standardizer_removes_identity_history_and_package_relationships(
    tmp_path: Path,
) -> None:
    source = tmp_path / "fictional-template3-with-metadata.pptx"
    runtime = tmp_path / "CATALOG_MINIMAL_OFFICIAL_CLEAN_v5.pptx"
    _source_deck(source)
    _inject_fictional_package_metadata(source)
    deck = Presentation(source)
    catalog.standardize_catalog_minimal_v5(
        source,
        runtime,
        expected_source_fingerprint=catalog._structure_fingerprint(deck),
        expected_source_sha256=None,
    )
    catalog.validate_catalog_minimal_v5_runtime(
        runtime, get_pptx_manifest("CATALOG_MINIMAL", "5")  # type: ignore[arg-type]
    )
    with ZipFile(runtime) as archive:
        names = set(archive.namelist())
        assert "ppt/commentAuthors.xml" not in names
        assert "docProps/custom.xml" not in names
        assert not any(
            name.startswith(("ppt/notesSlides/", "ppt/notesMasters/", "ppt/tags/"))
            for name in names
        )
        assert not any(
            b"fictional-reviewer" in archive.read(name)
            for name in names
            if name.endswith((".xml", ".rels"))
        )
        for rels_name in (name for name in names if name.endswith(".rels")):
            source_part = catalog._relationship_source(rels_name)
            relationships = etree.fromstring(archive.read(rels_name))
            for relationship in relationships:
                relation_name = relationship.get("Type", "").rsplit("/", 1)[-1].lower()
                assert relation_name not in catalog._FORBIDDEN_RELATIONSHIP_NAMES
                target = relationship.get("Target")
                assert target
                assert catalog._relationship_target(source_part, target) in names
        app = etree.fromstring(archive.read("docProps/app.xml"))
        values = {etree.QName(child).localname: child.text for child in app}
        assert values["Slides"] == "1"
        assert values["Notes"] == values["HiddenSlides"] == "0"
        core = etree.fromstring(archive.read("docProps/core.xml"))
        assert {etree.QName(child).localname for child in core} == {
            "identifier",
            "keywords",
        }


@pytest.mark.parametrize(
    ("rels_name", "expected_source"),
    (
        ("_rels/.rels", ""),
        ("ppt/_rels/presentation.xml.rels", "ppt/presentation.xml"),
        ("ppt/slides/_rels/slide1.xml.rels", "ppt/slides/slide1.xml"),
        (
            "ppt/slideLayouts/_rels/slideLayout1.xml.rels",
            "ppt/slideLayouts/slideLayout1.xml",
        ),
        (
            "ppt/slideMasters/_rels/slideMaster1.xml.rels",
            "ppt/slideMasters/slideMaster1.xml",
        ),
    ),
)
def test_relationship_source_uses_ooxml_posix_paths(
    rels_name: str, expected_source: str
) -> None:
    assert catalog._relationship_source(rels_name) == expected_source


@pytest.mark.parametrize(
    ("source", "target", "expected_part"),
    (
        ("", "ppt/presentation.xml", "ppt/presentation.xml"),
        ("ppt/presentation.xml", "slides/slide1.xml", "ppt/slides/slide1.xml"),
        (
            "ppt/slides/slide1.xml",
            "../slideLayouts/slideLayout1.xml",
            "ppt/slideLayouts/slideLayout1.xml",
        ),
        (
            "ppt/slideLayouts/slideLayout1.xml",
            "../slideMasters/slideMaster1.xml",
            "ppt/slideMasters/slideMaster1.xml",
        ),
        (
            "ppt/slideMasters/slideMaster1.xml",
            "../theme/theme1.xml",
            "ppt/theme/theme1.xml",
        ),
        ("ppt/presentation.xml", "/docProps/app.xml", "docProps/app.xml"),
    ),
)
def test_relationship_target_resolves_relative_ooxml_paths(
    source: str, target: str, expected_part: str
) -> None:
    assert catalog._relationship_target(source, target) == expected_part


def test_runtime_validator_rejects_forbidden_metadata_parts(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    _inject_fictional_package_metadata(runtime)
    with pytest.raises(AppError) as raised:
        catalog.validate_catalog_minimal_v5_runtime(
            runtime, get_pptx_manifest("CATALOG_MINIMAL", "5")  # type: ignore[arg-type]
        )
    assert raised.value.code == "PPT_TEMPLATE_ASSET_STRUCTURE_INVALID"


def test_private_builder_publishes_only_after_full_validation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "fictional-template3.pptx"
    runtime = tmp_path / "runtime.pptx"
    _source_deck(source)
    monkeypatch.setattr(catalog, "_SOURCE_SHA256", sha256(source.read_bytes()).hexdigest())
    monkeypatch.setattr(
        catalog,
        "_SOURCE_STRUCTURE_FINGERPRINT",
        catalog._structure_fingerprint(Presentation(source)),
    )
    # The public helper's default arguments are bound at import time; use the
    # generic builder contract through an accepted source copy instead.
    monkeypatch.setattr(
        "scripts.build_private_official_ppt_template_v5.standardize_catalog_minimal_v5",
        lambda original, destination: catalog.standardize_catalog_minimal_v5(
            original,
            destination,
            expected_source_fingerprint=catalog._structure_fingerprint(Presentation(original)),
            expected_source_sha256=None,
        ),
    )
    build_catalog_minimal_v5_runtime_asset(source, runtime)
    catalog.validate_catalog_minimal_v5_runtime(runtime, get_pptx_manifest("CATALOG_MINIMAL", "5"))  # type: ignore[arg-type]


def test_txt_priority_styles_are_applied_at_run_level(tmp_path: Path) -> None:
    content = _content(Presentation(_runtime(tmp_path)).slides[0])
    paragraphs = content.text_frame.paragraphs
    model = paragraphs[0].runs[0]
    name = paragraphs[1].runs[0]
    specification = paragraphs[4].runs[0]
    retail_value = paragraphs[9].runs[1]
    feature_label, feature_value = paragraphs[12].runs[0], paragraphs[13].runs[0]
    link_label, link_value = paragraphs[19].runs
    assert model.font.color.rgb == RGBColor(157, 144, 139)
    assert name.font.color.rgb == specification.font.color.rgb == RGBColor(0, 0, 0)
    assert retail_value.font.color.rgb == RGBColor(157, 144, 139)
    assert feature_label.font.color.rgb == feature_value.font.color.rgb == RGBColor(0, 0, 0)
    assert link_label.font.color.rgb == link_value.font.color.rgb == RGBColor(0, 0, 0)
    assert link_label.font.bold is True and link_value.font.bold is True
    assert model.font.size == name.font.size == Pt(18)
    assert specification.font.size == retail_value.font.size == link_value.font.size == Pt(9)


def test_complete_mapping_uses_frozen_jd_and_agreement_prices_and_real_hyperlink(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = _render(_runtime(tmp_path), {"single_products": [_item()]}, monkeypatch)
    content = _content(deck.slides[0])
    text = content.text
    assert "MODEL-001" in text and "虚构极简商品 1" in text
    assert "299.00元" in text and "239.00元" in text and "999.00" not in text
    link = content.text_frame.paragraphs[-1].runs[1]
    assert link.text == "https://example.invalid/jd/1"
    assert link.hyperlink.address == "https://example.invalid/jd/1"
    assert link.font.color.rgb == RGBColor(0, 0, 0) and link.font.bold is True
    assert link.font.underline is False


def test_empty_fields_do_not_invent_units_links_or_values(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = {"single_products": [{"product": {"id": "fictional-empty"}, "prices": {}}]}
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    content = _content(deck.slides[0])
    assert "{{" not in content.text and "自动填入" not in content.text
    price_paragraphs = [p for p in content.text_frame.paragraphs if "价：" in p.text]
    assert all("元" not in paragraph.text for paragraph in price_paragraphs)
    assert content.text_frame.paragraphs[-1].runs[1].hyperlink.address is None


def test_price_mapping_preserves_zero_missing_values_and_product_isolation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    zero = _item(1)
    zero["prices"] = {
        "market_price": "999.00",
        "jd_price": 0,
        "agreement_price": "0.00",
    }
    missing = _item(2)
    missing["prices"] = {"market_price": "888.00"}
    distinct = _item(3)
    distinct["prices"] = {
        "market_price": "777.00",
        "jd_price": "301.00",
        "agreement_price": "241.00",
    }
    deck = _render(
        _runtime(tmp_path),
        {"single_products": [zero, missing, distinct]},
        monkeypatch,
    )
    assert "0.00元" in _content(deck.slides[0]).text
    missing_prices = [
        paragraph.text
        for paragraph in _content(deck.slides[1]).text_frame.paragraphs
        if "价：" in paragraph.text
    ]
    assert missing_prices == ["零  售  价：", "团购价："]
    assert "301.00元" in _content(deck.slides[2]).text
    assert "241.00元" in _content(deck.slides[2]).text
    rendered_text = "\n".join(_content(slide).text for slide in deck.slides)
    assert all(value not in rendered_text for value in ("999.00", "888.00", "777.00"))


@pytest.mark.parametrize(
    "url",
    (
        "https://example.invalid/jd/short",
        "https://example.invalid/jd/" + "a" * 300,
    ),
)
def test_link_chunks_preserve_url_and_real_hyperlink_relationship(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, url: str
) -> None:
    item = _item()
    item["product"]["jd_same_product_url"] = url  # type: ignore[index]
    deck = _render(_runtime(tmp_path), {"single_products": [item]}, monkeypatch)
    content = _content(deck.slides[0])
    link_paragraphs = content.text_frame.paragraphs[
        -catalog._line_count(url, 9, first_line_prefix=catalog._JD_LINK_LABEL) :
    ]
    assert "".join(paragraph.runs[1].text for paragraph in link_paragraphs) == url
    assert link_paragraphs[0].runs[0].text == "京东链接："
    assert all(paragraph.runs[0].text == "" for paragraph in link_paragraphs[1:])
    assert all(paragraph.runs[1].hyperlink.address == url for paragraph in link_paragraphs)
    assert all(
        paragraph.runs[1].font.color.rgb == RGBColor(0, 0, 0)
        and paragraph.runs[1].font.bold is True
        and paragraph.runs[1].font.underline is False
        for paragraph in link_paragraphs
    )
    targets = [
        relation.target_ref
        for relation in deck.slides[0].part.rels.values()
        if relation.is_external and relation.reltype.endswith("/hyperlink")
    ]
    assert targets and set(targets) == {url}


def test_link_first_line_reserves_fixed_label_width(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    base = "https://example.invalid/jd/"
    suffix = ""
    while (
        catalog._estimated_width(catalog._JD_LINK_LABEL + base + suffix + "a", 9)
        <= catalog._SAFE_TEXT_WIDTH_PT
    ):
        suffix += "a"
    url = base + suffix + "a"
    chunks = catalog._wrapped_chunks(url, 9, first_line_prefix=catalog._JD_LINK_LABEL)
    assert len(chunks) == 2
    assert "".join(chunks) == url
    assert (
        catalog._estimated_width(catalog._JD_LINK_LABEL + chunks[0], 9)
        <= catalog._SAFE_TEXT_WIDTH_PT
    )
    assert (
        catalog._estimated_width(catalog._JD_LINK_LABEL + chunks[0] + chunks[1][0], 9)
        > catalog._SAFE_TEXT_WIDTH_PT
    )

    item = _item()
    item["product"]["jd_same_product_url"] = url  # type: ignore[index]
    deck = _render(_runtime(tmp_path), {"single_products": [item]}, monkeypatch)
    content = _content(deck.slides[0])
    link_paragraphs = content.text_frame.paragraphs[-len(chunks) :]
    assert link_paragraphs[0].runs[0].text == catalog._JD_LINK_LABEL
    assert link_paragraphs[0].runs[1].text == chunks[0]
    assert (
        catalog._estimated_width(link_paragraphs[0].text, 9)
        <= catalog._SAFE_TEXT_WIDTH_PT
    )
    assert "".join(paragraph.runs[1].text for paragraph in link_paragraphs) == url


@pytest.mark.parametrize(
    "invalid_url",
    (
        "https://example.invalid/jd/with\nnewline",
        "https://example.invalid/jd/with\ttab",
        "https://example.invalid/jd/with space",
        " https://example.invalid/jd/leading-space",
        "https://example.invalid/jd/trailing-space ",
        "ftp://example.invalid/file",
    ),
)
def test_link_validation_rejects_whitespace_controls_and_protocols(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, invalid_url: str
) -> None:
    item = _item()
    item["product"]["jd_same_product_url"] = invalid_url  # type: ignore[index]
    with pytest.raises(AppError) as raised:
        _render(_runtime(tmp_path), {"single_products": [item]}, monkeypatch)
    assert raised.value.code == "PPT_TEMPLATE_LINK_INVALID"


def test_empty_and_invalid_links_do_not_create_invalid_relationships(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    empty = _item()
    empty["product"]["jd_same_product_url"] = ""  # type: ignore[index]
    deck = _render(_runtime(tmp_path), {"single_products": [empty]}, monkeypatch)
    assert not [
        relation
        for relation in deck.slides[0].part.rels.values()
        if relation.is_external and relation.reltype.endswith("/hyperlink")
    ]


@pytest.mark.parametrize("size", ((320, 80), (80, 320), (160, 160), (800, 20)))
def test_main_image_uses_contain_without_crop_or_stretch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    size: tuple[int, int],
) -> None:
    deck = _render(
        _runtime(tmp_path),
        {"single_products": [_item()]},
        monkeypatch,
        {"fictional-1": _png((20, 90, 160), size)},
    )
    picture = next(shape for shape in deck.slides[0].shapes if shape.name == "SCM_SLOT_IMAGE_MAIN")
    assert picture.shape_type == MSO_SHAPE_TYPE.PICTURE
    assert picture.crop_left == picture.crop_right == picture.crop_top == picture.crop_bottom == 0
    assert picture.width / picture.height == pytest.approx(size[0] / size[1], rel=0.01)
    assert picture.left >= Inches(0.556) and picture.top >= Inches(0.678)
    assert picture.left + picture.width <= Inches(0.556 + 5.446)
    assert picture.top + picture.height <= Inches(0.678 + 6.153)


def test_missing_image_keeps_private_transparent_slot(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = _render(_runtime(tmp_path), {"single_products": [_item()]}, monkeypatch)
    picture = next(shape for shape in deck.slides[0].shapes if shape.name == "SCM_SLOT_IMAGE_MAIN")
    assert Image.open(BytesIO(picture.image.blob)).size == (1, 1)


def test_long_fields_add_paragraphs_without_overlap_or_font_shrink(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    item = _item()
    item["product"].update(  # type: ignore[index]
        model="MODEL-LONG-2026-ABCDEFGHIJKLMN-EXTRA-SEGMENT",
        product_name="虚构多行商品名称用于验证自然换行与固定字号",
        product_specification="\n".join(
            f"虚构属性{index}：" + "参数" * 12 for index in range(1, 8)
        ),
        selling_points="\n".join(f"虚构卖点{index}：" + "说明" * 14 for index in range(1, 7)),
        jd_same_product_url="https://example.invalid/jd/" + "continuous-ascii-" * 18,
    )
    deck = _render(_runtime(tmp_path), {"single_products": [item]}, monkeypatch)
    content = _content(deck.slides[0])
    assert content.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert len(content.text_frame.paragraphs) > 21
    assert content.text_frame.paragraphs[0].runs[0].font.size == Pt(18)
    assert all(
        paragraph.runs[0].font.size == Pt(9)
        for paragraph in content.text_frame.paragraphs[4:]
        if paragraph.runs
    )
    assert content.top + content.height <= deck.slide_height
    model = "MODEL-LONG-2026-ABCDEFGHIJKLMN-EXTRA-SEGMENT"
    chunks = catalog._wrapped_chunks(model, 18)
    assert len(chunks) > 1
    assert model == "".join(
        paragraph.text for paragraph in content.text_frame.paragraphs[: len(chunks)]
    )
    values = catalog._item_values(item)
    high_lines = len(catalog._wrapped_chunks(values["model"], 18)) + len(
        catalog._wrapped_chunks(values["product_name"], 18)
    ) + 1
    specification_lines = sum(
        len(catalog._wrapped_chunks(value, 9))
        for value in catalog._logical_lines(values["specification"])
    )
    selling_lines = sum(
        len(catalog._wrapped_chunks(value, 9))
        for value in catalog._logical_lines(values["selling_points"], semicolon_separator=True)
    )
    low_lines = (
        1
        + max(catalog._MIN_SPEC_LINES, specification_lines)
        + 2
        + 1
        + 1
        + max(catalog._MIN_SELLING_LINES, selling_lines)
        + catalog._line_count(
            values["jd_link"], 9, first_line_prefix=catalog._JD_LINK_LABEL
        )
    )
    assert catalog._visible_text_bottom(
        content, high_lines=high_lines, low_lines=low_lines
    ) <= deck.slide_height


def test_genuine_text_overflow_is_rejected_without_truncation_or_shrink(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    item = _item()
    item["product"].update(  # type: ignore[index]
        product_specification="\n".join("超长规格" * 35 for _ in range(20)),
        selling_points="\n".join("超长卖点" * 35 for _ in range(20)),
    )
    with pytest.raises(AppError) as raised:
        _render(_runtime(tmp_path), {"single_products": [item]}, monkeypatch)
    assert raised.value.code == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"
    assert "fictional-1" in raised.value.message and "超出" in raised.value.message


def test_five_products_are_isolated_and_keep_distinct_images(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = {"single_products": [_item(index) for index in range(1, 6)]}
    images = {
        f"fictional-{index}": _png((30 * index, 90, 170), (120 + index, 90))
        for index in range(1, 6)
    }
    deck = _render(_runtime(tmp_path), source, monkeypatch, images)
    assert len(deck.slides) == 5
    hashes: list[str] = []
    for index, slide in enumerate(deck.slides, start=1):
        assert f"MODEL-{index:03d}" in _content(slide).text
        assert all(
            f"MODEL-{other:03d}" not in _content(slide).text
            for other in range(1, 6)
            if other != index
        )
        picture = next(shape for shape in slide.shapes if shape.name == "SCM_SLOT_IMAGE_MAIN")
        hashes.append(sha256(picture.image.blob).hexdigest())
        link = _content(slide).text_frame.paragraphs[-1].runs[1]
        assert link.hyperlink.address == f"https://example.invalid/jd/{index}"
    assert len(set(hashes)) == 5


def test_combination_invalid_link_and_corrupt_runtime_fail_safely(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    with pytest.raises(AppError) as combination:
        _render(runtime, {"packages": [{"items": []}]}, monkeypatch)
    assert combination.value.code == "PPT_TEMPLATE_COMBINATION_UNSUPPORTED"
    invalid = _item()
    invalid["product"]["jd_same_product_url"] = "javascript:alert(1)"  # type: ignore[index]
    with pytest.raises(AppError) as link:
        _render(runtime, {"single_products": [invalid]}, monkeypatch)
    assert link.value.code == "PPT_TEMPLATE_LINK_INVALID"
    broken = Presentation(runtime)
    _content(broken.slides[0]).name = "BROKEN"
    broken.save(runtime)
    with pytest.raises(AppError) as structure:
        _render(runtime, {"single_products": [_item()]}, monkeypatch)
    assert structure.value.code == "PPT_TEMPLATE_ASSET_STRUCTURE_INVALID"


def test_manifest_registry_and_explicit_v5_dispatch_are_versioned() -> None:
    manifest = get_pptx_manifest("CATALOG_MINIMAL", "5")
    template = get_ppt_template("CATALOG_MINIMAL", "5")
    assert manifest is not None
    assert manifest.clean_filename == "CATALOG_MINIMAL_OFFICIAL_CLEAN_v5.pptx"
    assert manifest.image_fit_modes == ("contain",)
    assert template.version == "5" and template.single_product_only is True
    assert get_ppt_template("CATALOG_MINIMAL", "3").version == "3"
    assert get_ppt_template("CATALOG_MINIMAL", "4").version == "4"


def test_v5_renderer_entry_resolves_private_manifest_asset_and_dispatches(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    private_root = tmp_path / "private-runtime"
    private_root.mkdir()
    manifest = get_pptx_manifest("CATALOG_MINIMAL", "5")
    assert manifest is not None
    installed = private_root / manifest.clean_filename
    copyfile(runtime, installed)
    monkeypatch.setattr(
        ppt_template_registry,
        "get_settings",
        lambda: type(
            "Settings",
            (),
            {
                "ppt_template_asset_dir": private_root,
                "ppt_v5_preview_enabled": True,
                "app_env": "development",
                "ppt_local_v5_default_enabled": False,
            },
        )(),
    )
    assert ppt_template_registry.resolve_ppt_template_asset(
        get_ppt_template("CATALOG_MINIMAL", "5")
    ) == installed

    invoked: list[Path] = []
    original = catalog.render_catalog_minimal_v5

    def recording_renderer(*args: Any, **kwargs: Any) -> bytes:
        invoked.append(ppt_template_registry.resolve_ppt_template_asset(kwargs["template"]))
        return original(*args, **kwargs)

    monkeypatch.setattr(catalog, "render_catalog_minimal_v5", recording_renderer)
    content = PptRenderer(layout_preview=False).render(
        {"single_products": [_item()]},
        template_code="CATALOG_MINIMAL",
        template_version="5",
    )
    assert invoked == [installed]
    assert len(Presentation(BytesIO(content)).slides) == 1


def test_public_v5_asset_uses_default_root_and_official_renderer(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A fresh checkout resolves the reviewed public v5 asset without private files."""
    manifest = get_pptx_manifest("CATALOG_MINIMAL", "5")
    assert manifest is not None
    installed = PROJECT_ROOT / "assets" / "ppt-templates" / manifest.clean_filename
    assert installed.is_file()
    assert sha256(installed.read_bytes()).hexdigest() == (
        "953743f5947b7012170eeffa5dbf4a16c710dcbf8dd26c7b7f63eb263931c735"
    )
    monkeypatch.setattr(
        ppt_template_registry,
        "get_settings",
        lambda: type(
            "Settings",
            (),
            {
                "ppt_template_asset_dir": None,
                "ppt_v5_preview_enabled": True,
                "app_env": "development",
                "ppt_local_v5_default_enabled": False,
            },
        )(),
    )
    assert ppt_template_registry.resolve_ppt_template_asset(
        get_ppt_template("CATALOG_MINIMAL", "5")
    ) == installed

    invoked: list[Path] = []
    original = catalog.render_catalog_minimal_v5

    def recording_renderer(*args: Any, **kwargs: Any) -> bytes:
        invoked.append(ppt_template_registry.resolve_ppt_template_asset(kwargs["template"]))
        return original(*args, **kwargs)

    monkeypatch.setattr(catalog, "render_catalog_minimal_v5", recording_renderer)
    content = PptRenderer(layout_preview=False).render(
        {"single_products": [_item()]},
        template_code="CATALOG_MINIMAL",
        template_version="5",
    )
    assert invoked == [installed]
    assert len(Presentation(BytesIO(content)).slides) == 1


def test_public_v5_asset_has_complete_cross_platform_ooxml_relationships() -> None:
    manifest = get_pptx_manifest("CATALOG_MINIMAL", "5")
    assert manifest is not None
    installed = PROJECT_ROOT / "assets" / "ppt-templates" / manifest.clean_filename

    catalog.validate_catalog_minimal_v5_runtime(installed, manifest)

    with ZipFile(installed) as archive:
        names = set(archive.namelist())
        sources: set[str] = set()
        for rels_name in sorted(name for name in names if name.endswith(".rels")):
            source = catalog._relationship_source(rels_name)
            sources.add(source)
            root = etree.fromstring(archive.read(rels_name))
            for relationship in root:
                assert relationship.get("TargetMode") != "External"
                target = relationship.get("Target")
                assert target
                assert catalog._relationship_target(source, target) in names

    assert {
        "",
        "ppt/presentation.xml",
        "ppt/slides/slide1.xml",
        "ppt/slideLayouts/slideLayout1.xml",
        "ppt/slideMasters/slideMaster1.xml",
    } <= sources


def test_private_v5_asset_is_unavailable_without_controlled_runtime(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        ppt_template_registry,
        "get_settings",
        lambda: type(
            "Settings",
            (),
            {"ppt_template_asset_dir": tmp_path, "ppt_v5_preview_enabled": True},
        )(),
    )
    assert is_ppt_template_available(get_ppt_template("CATALOG_MINIMAL", "5")) is False
