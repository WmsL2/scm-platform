"""Public CI uses a wholly fictional two-slide structural fixture."""

from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app.common.contracts import AppError
from app.modules.recommendation.application import ppt_official_template_festive_v5 as festive
from app.modules.recommendation.application import ppt_template_registry
from app.modules.recommendation.application.ppt_renderer import PptRenderer
from app.modules.recommendation.application.ppt_template_manifest import get_pptx_manifest
from app.modules.recommendation.application.ppt_template_registry import get_ppt_template
from scripts import build_private_official_ppt_template_v5 as builder


def _png(color: tuple[int, int, int], size: tuple[int, int] = (160, 80)) -> bytes:
    output = BytesIO()
    Image.new("RGB", size, color).save(output, format="PNG")
    return output.getvalue()


def _run(paragraph: Any, value: str, size: int = 14, bold: bool = False) -> None:
    run = paragraph.add_run()
    run.text = value
    run.font.name = "京东朗正体 玲珑"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(255, 255, 255)


def _set_theme_colour(run: Any) -> None:
    """Model the approved customer's theme-colour Run, not only RGB fixtures."""
    properties = run._r.get_or_add_rPr()
    for child in list(properties):
        if child.tag.endswith("}solidFill"):
            properties.remove(child)
    fill = OxmlElement("a:solidFill")
    colour = OxmlElement("a:schemeClr")
    colour.set("val", "lt1")
    fill.append(colour)
    properties.append(fill)


def _source_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.335), Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    cover.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "虚构固定封面"
    product = deck.slides.add_slide(deck.slide_layouts[6])
    product.shapes.add_picture(
        BytesIO(_png((120, 20, 20))), 0, 0, deck.slide_width, deck.slide_height
    )
    geometries = (
        (0.357, 0.995, 5.343, 5.320),
        (5.863, 2.702, 7.029, 1.816),
        (5.863, 1.620, 7.028, 1.160),
        (5.820, 0.587, 7.161, 1.110),
        (0.795, 0.310, 1.714, 0.517),
        (6.001, 4.380, 7.029, 0.605),
        (5.881, 4.931, 6.015, 1.885),
        (6.053, 5.010, 1.811, 1.733),
        (9.912, 5.010, 1.811, 1.733),
        (7.983, 5.010, 1.811, 1.733),
    )
    for geometry, label in zip(
        geometries,
        (
            "主图1说明",
            "",
            "",
            "",
            "方案编号",
            "组套后或内外包装图片：",
            "",
            "图片1",
            "图片3",
            "图片2",
        ),
        strict=True,
    ):
        shape = product.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Inches(v) for v in geometry))
        shape.text = label
    title = product.shapes[4].text_frame.paragraphs[0]
    _run(title, "自动填入商品名称说明", 20, True)
    product.shapes[4].text_frame.add_paragraph()
    selling = product.shapes[2].text_frame
    _run(selling.paragraphs[0], "商品卖点：", 20, True)
    selling.paragraphs[0].line_spacing = 1.5
    for _ in range(4):
        paragraph = selling.add_paragraph()
        paragraph.line_spacing = 1.5
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", "342900")
        properties.set("indent", "-342900")
        font = OxmlElement("a:buFont")
        font.set("typeface", "Wingdings")
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "u")
        properties.append(font)
        properties.append(bullet)
        _run(paragraph, "自动填入卖点说明", 12)
    reference = product.shapes[3].text_frame
    first = reference.paragraphs[0]
    first.line_spacing = 1.5
    for text in ("市场参考价：", " " * 13, "京东大客户价：", " " * 11, "折扣率：", " ", "  "):
        _run(first, text)
    first.add_line_break()
    for text in ("SKU", "：", " " * 20, "市场参考价不要虚高，客户会查！"):
        _run(first, text, bold=text.startswith("市场参考"))
    second = reference.add_paragraph()
    second.line_spacing = 1.5
    for text in (
        "货期：",
        " " * 15,
        "    发货物流",
        "/",
        "快递",
        "：",
        " " * 5,
        " " * 5,
        "质保",
        "：",
        " ",
    ):
        _run(second, text, bold=True)
    deck.save(path)


def _source() -> dict[str, Any]:
    return {
        "single_products": [
            {
                "product": {
                    "id": "fictional-1",
                    "product_name": "虚构节庆商品",
                    "selling_points": "卖点一\n卖点二",
                    "sku": "FEST001",
                    "shipping_courier": "京东",
                    "warranty_period": "一年",
                },
                "prices": {
                    "market_price": "199.5",
                    "agreement_price": "169.5",
                    "jd_price": "888",
                    "discount_rate": "0.85",
                },
                "manual": {"fulfillment_cycle": "3天"},
            }
        ]
    }


def _runtime(tmp_path: Path) -> Path:
    original, runtime = tmp_path / "fictional-source.pptx", tmp_path / "runtime.pptx"
    _source_deck(original)
    festive.standardize_jd_festive_red_v5(
        original,
        runtime,
        expected_source_fingerprint=festive._structure_fingerprint(Presentation(original)),
        expected_source_sha256=None,
    )
    festive.verify_no_sample_media(original, runtime)
    return runtime


def _render(
    runtime: Path,
    source: dict[str, Any],
    monkeypatch: pytest.MonkeyPatch,
    images: dict[str, bytes] | None = None,
    *,
    layout_preview: bool = False,
) -> Presentation:
    monkeypatch.setattr(festive, "resolve_ppt_template_asset", lambda _template: runtime)
    result = PptRenderer(layout_preview=layout_preview).render(
        source, template_code="JD_FESTIVE_RED", template_version="5", product_images=images or {}
    )
    return Presentation(BytesIO(result))


def _slots(deck: Any, slide_index: int = 1) -> dict[str, Any]:
    return {shape.name: shape for shape in deck.slides[slide_index].shapes}


def _reference_run(slots: dict[str, Any], field: str) -> Any:
    paragraph_index, run_index, _width = festive._VALUE_RUNS[field]
    return slots["SCM_FIELD_REFERENCE"].text_frame.paragraphs[paragraph_index].runs[
        run_index
    ]


def _reference_value(slots: dict[str, Any], field: str) -> str:
    return _reference_run(slots, field).text.rstrip("\t\u3000")


def _paragraph_properties_without_tabs(paragraph: Any) -> str:
    properties = deepcopy(paragraph._p.get_or_add_pPr())
    for child in list(properties):
        if child.tag.endswith("}tabLst"):
            properties.remove(child)
    return properties.xml


def _tab_positions(paragraph: Any) -> tuple[int, ...]:
    tab_list = paragraph._p.pPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tabLst"
    )
    assert tab_list is not None
    return tuple(int(tab.get("pos")) for tab in tab_list)


def test_fixed_reference_xml_cover_and_static_geometry_are_preserved(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    original = Presentation(tmp_path / "fictional-source.pptx")
    deck = _render(runtime, _source(), monkeypatch)
    assert deck.slides[0]._element.xml == original.slides[0]._element.xml
    source_reference = original.slides[1].shapes[3]
    rendered_reference = _slots(deck)["SCM_FIELD_REFERENCE"]
    dynamic_runs = {(paragraph, run) for paragraph, run, _width in festive._VALUE_RUNS.values()}
    for paragraph_index, (source_p, rendered_p) in enumerate(
        zip(
            source_reference.text_frame.paragraphs,
            rendered_reference.text_frame.paragraphs,
            strict=True,
        )
    ):
        assert _paragraph_properties_without_tabs(source_p) == _paragraph_properties_without_tabs(
            rendered_p
        )
        for run_index, (source_run, rendered_run) in enumerate(
            zip(source_p.runs, rendered_p.runs, strict=True)
        ):
            assert source_run._r.rPr.xml == rendered_run._r.rPr.xml
            if (paragraph_index, run_index) not in dynamic_runs | {(1, 2), (1, 7)}:
                assert source_run.text == rendered_run.text
    assert rendered_reference.text_frame.paragraphs[1].runs[2].text == "发货物流"
    assert rendered_reference.text_frame.paragraphs[1].runs[7].text == ""
    assert not any(shape.name.startswith("SCM_VALUE_") for shape in deck.slides[1].shapes)
    for index, name in (
        (2, "SCM_FIELD_SELLING_POINTS"),
        (3, "SCM_FIELD_REFERENCE"),
        (4, "SCM_FIELD_TITLE"),
        (5, "SCM_STATIC_PROJECT_NUMBER"),
        (6, "SCM_STATIC_IMAGE_CAPTION"),
    ):
        source, rendered = original.slides[1].shapes[index], _slots(deck)[name]
        assert (source.left, source.top, source.width, source.height) == (
            rendered.left,
            rendered.top,
            rendered.width,
            rendered.height,
        )


@pytest.mark.parametrize(
    ("market", "agreement", "expected_market", "expected_agreement"),
    (
        ("199.5", "169.5", "199.50元", "169.50元"),
        ("99", None, "99.00元", ""),
        (None, "89", "", "89.00元"),
        (None, None, "", ""),
        (0, 0, "0.00元", "0.00元"),
    ),
)
def test_frozen_price_mapping_never_uses_jd_price_or_campaign_price(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    market: object,
    agreement: object,
    expected_market: str,
    expected_agreement: str,
) -> None:
    source = _source()
    source["single_products"][0]["prices"].update(market_price=market, agreement_price=agreement)
    source["single_products"][0]["manual"]["campaign_price"] = "777"
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    slots = _slots(deck)
    assert _reference_value(slots, "market_price") == expected_market
    assert _reference_value(slots, "jd_customer_price") == expected_agreement
    assert "888" not in slots["SCM_FIELD_REFERENCE"].text
    assert "777" not in slots["SCM_FIELD_REFERENCE"].text


def test_values_use_exact_original_gap_runs_and_requested_styles(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    runtime_slots = _slots(Presentation(runtime))
    deck = _render(runtime, _source(), monkeypatch)
    slots = _slots(deck)
    assert not any(name.startswith("SCM_VALUE_") for name in slots)
    for field, (paragraph_index, run_index, _width) in festive._VALUE_RUNS.items():
        target = _reference_run(slots, field)
        source_run = _reference_run(runtime_slots, field)
        assert target.font.name == "京东朗正体 玲珑"
        assert target.font.size == Pt(14) and target.font.color.rgb == RGBColor(255, 255, 255)
        assert target.font.bold is (paragraph_index == 1)
        assert target._r.rPr.xml == source_run._r.rPr.xml
    assert all(
        paragraph.line_spacing == 1.5
        for paragraph in slots["SCM_FIELD_REFERENCE"].text_frame.paragraphs
    )
    assert "市场参考价：199.50元" in slots["SCM_FIELD_REFERENCE"].text
    assert "京东大客户价：169.50元" in slots["SCM_FIELD_REFERENCE"].text
    assert "折扣率：85%" in slots["SCM_FIELD_REFERENCE"].text
    assert "a:tabLst" in slots["SCM_FIELD_REFERENCE"].text_frame.paragraphs[0]._p.pPr.xml
    assert "a:tabLst" in slots["SCM_FIELD_REFERENCE"].text_frame.paragraphs[1]._p.pPr.xml
    assert _tab_positions(slots["SCM_FIELD_REFERENCE"].text_frame.paragraphs[0]) == tuple(
        int(point * festive._EMU_PER_POINT) for point in (90.0, 157.0, 333.0)
    )
    assert _tab_positions(slots["SCM_FIELD_REFERENCE"].text_frame.paragraphs[1]) == tuple(
        int(point * festive._EMU_PER_POINT) for point in (118.0, 311.0)
    )


def test_original_gap_run_preserves_theme_colour_without_accessing_rgb(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Exercise the actual renderer path used by the theme-colour customer asset."""
    source = tmp_path / "theme-colour-source.pptx"
    runtime = tmp_path / "theme-colour-runtime.pptx"
    _source_deck(source)
    deck = Presentation(source)
    _set_theme_colour(deck.slides[1].shapes[3].text_frame.paragraphs[0].runs[1])
    deck.save(source)
    festive.standardize_jd_festive_red_v5(
        source,
        runtime,
        expected_source_fingerprint=festive._structure_fingerprint(Presentation(source)),
        expected_source_sha256=None,
    )

    rendered = _render(runtime, _source(), monkeypatch)
    value_run = _reference_run(_slots(rendered), "market_price")
    assert "schemeClr" in value_run._r.rPr.xml
    assert value_run.text == "199.50元\t"


def test_selling_bullets_paragraph_positions_and_run_styles_survive(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    original = Presentation(tmp_path / "fictional-source.pptx").slides[1].shapes[2]
    deck = _render(runtime, _source(), monkeypatch)
    selling = _slots(deck)["SCM_FIELD_SELLING_POINTS"]
    assert [p.text for p in selling.text_frame.paragraphs] == [
        "商品卖点：",
        "卖点一",
        "卖点二",
        "",
        "",
    ]
    for source_p, target_p in zip(
        original.text_frame.paragraphs, selling.text_frame.paragraphs, strict=True
    ):
        assert source_p._p.pPr.xml == target_p._p.pPr.xml
        assert source_p.runs[0]._r.rPr.xml == target_p.runs[0]._r.rPr.xml
    assert selling.text_frame.paragraphs[0].runs[0].font.size == Pt(20)


def test_missing_fields_clear_only_values_and_all_picture_instructions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = {
        "single_products": [
            {
                "product": {
                    "product_name": "虚构空字段",
                    "brand": "NEVER-RENDER-BRAND",
                    "model": "NEVER-RENDER-MODEL",
                    "product_specification": "NEVER-RENDER-SPECIFICATION",
                }
            }
        ]
    }
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    slots = _slots(deck)
    assert all(_reference_value(slots, field) == "" for field in festive._VALUE_RUNS)
    assert not any(name.startswith("SCM_VALUE_") for name in slots)
    text = "\n".join(s.text for s in deck.slides[1].shapes if s.has_text_frame)
    assert all(
        word not in text
        for word in (
            "自动填入",
            "图片1",
            "主图1",
            "{{",
            "品牌：",
            "型号：",
            "京东价：",
            "协议价：",
            "NEVER-RENDER",
        )
    )
    assert all(slots[name].image.size == (1, 1) for name in festive._IMAGE_SLOT_GEOMETRY)


@pytest.mark.parametrize(
    ("field", "value"),
    (
        ("product_name", "宽" * 55),
        ("sku", "W" * 18),
        ("shipping_courier", "宽" * 12),
        ("selling_points", "宽" * 70),
        ("sku", "ABC\nDEF"),
    ),
)
def test_long_text_wraps_without_dropping_value(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, field: str, value: str
) -> None:
    source = _source()
    source["single_products"][0]["product"][field] = value
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    slots = _slots(deck)
    if field == "selling_points":
        assert slots["SCM_FIELD_SELLING_POINTS"].text_frame.paragraphs[1].text == value
    elif field == "product_name":
        assert slots["SCM_FIELD_TITLE"].text_frame.paragraphs[0].text == value
    else:
        assert _reference_value(slots, field) == value


def test_title_and_sku_wraps_accumulate_downstream_offsets(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    source = _source()
    product = source["single_products"][0]["product"]
    product.update(product_name="宽" * 80, sku="W" * 18)
    baseline = _slots(Presentation(runtime))
    slots = _slots(_render(runtime, source, monkeypatch))
    assert slots["SCM_FIELD_TITLE"].height > baseline["SCM_FIELD_TITLE"].height
    assert slots["SCM_FIELD_REFERENCE"].top > baseline["SCM_FIELD_REFERENCE"].top
    assert slots["SCM_FIELD_SELLING_POINTS"].top > baseline["SCM_FIELD_SELLING_POINTS"].top
    assert slots["SCM_SLOT_IMAGE_SUPPLEMENT_01"].top > baseline["SCM_SLOT_IMAGE_SUPPLEMENT_01"].top
    assert _reference_value(slots, "sku") == "W" * 18


def test_wrapped_sku_expands_original_reference_and_moves_downstream_shapes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    baseline_source = _source()
    baseline = _slots(_render(runtime, baseline_source, monkeypatch))
    source = _source()
    source["single_products"][0]["product"].update(  # type: ignore[index]
        sku="W" * 36,
    )
    slots = _slots(_render(runtime, source, monkeypatch))
    reference = slots["SCM_FIELD_REFERENCE"]
    second_row = reference.text_frame.paragraphs[1]
    assert second_row.text.startswith("货期：")
    assert (
        second_row.runs[0]._r.rPr.xml
        == baseline["SCM_FIELD_REFERENCE"].text_frame.paragraphs[1].runs[0]._r.rPr.xml
    )
    assert reference.height > baseline["SCM_FIELD_REFERENCE"].height
    assert _reference_value(slots, "sku") == "W" * 36
    assert not any(name.startswith("SCM_VALUE_") for name in slots)
    assert slots["SCM_FIELD_SELLING_POINTS"].top > baseline["SCM_FIELD_SELLING_POINTS"].top


def test_page_overflow_is_a_controlled_error_after_dynamic_layout(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source()
    source["single_products"][0]["product"]["product_name"] = "宽" * 300
    with pytest.raises(AppError, match="product_name") as error:
        _render(_runtime(tmp_path), source, monkeypatch)
    assert error.value.code == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"


@pytest.mark.parametrize("count", (1, 4, 5, 6))
def test_selling_points_use_actual_page_capacity_instead_of_four_item_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, count: int
) -> None:
    source = _source()
    points = [f"虚构卖点 {index}" for index in range(1, count + 1)]
    source["single_products"][0]["product"]["selling_points"] = "\n".join(  # type: ignore[index]
        points
    )
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    selling = _slots(deck)["SCM_FIELD_SELLING_POINTS"]
    assert [paragraph.text for paragraph in selling.text_frame.paragraphs[1 : count + 1]] == (
        points
    )
    assert len(deck.slides) == 2


def test_seven_short_selling_points_still_fail_when_page_really_overflows(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source()
    source["single_products"][0]["product"]["selling_points"] = "\n".join(  # type: ignore[index]
        f"虚构卖点 {index}" for index in range(1, 8)
    )
    with pytest.raises(AppError, match="selling_points") as error:
        _render(_runtime(tmp_path), source, monkeypatch)
    assert error.value.code == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"


def test_selling_point_parser_preserves_semicolons_inside_newline_items() -> None:
    value = "第一条含中文分号；但仍是一条\n\n第二条含 ASCII; 仍是一条\n3. 第三条"
    assert festive._selling_point_items(value) == [
        "第一条含中文分号；但仍是一条",
        "第二条含 ASCII; 仍是一条",
        "3. 第三条",
    ]


def test_single_line_semicolon_list_drops_only_empty_items() -> None:
    assert festive._selling_point_items("第一条；；第二条; ;第三条") == [
        "第一条",
        "第二条",
        "第三条",
    ]


def test_long_single_selling_point_wraps_but_remains_one_bullet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source()
    value = "虚构长卖点" + "宽" * 70
    source["single_products"][0]["product"]["selling_points"] = value  # type: ignore[index]
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    selling = _slots(deck)["SCM_FIELD_SELLING_POINTS"]
    assert selling.text_frame.paragraphs[1].text == value
    assert all(not paragraph.text for paragraph in selling.text_frame.paragraphs[2:5])


def test_five_points_with_other_long_fields_keep_all_content_and_move_lower_shapes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    baseline = _slots(Presentation(runtime))
    source = _source()
    product = source["single_products"][0]["product"]  # type: ignore[index]
    points = [f"虚构组合卖点 {index}" for index in range(1, 6)]
    product.update(
        product_name="虚构中英文 Mixed 2026 商品名称" + "宽" * 12,
        sku="SKU-2026-LONG-001",
        selling_points="\n".join(points),
    )
    deck = _render(runtime, source, monkeypatch)
    slots = _slots(deck)
    assert [p.text for p in slots["SCM_FIELD_SELLING_POINTS"].text_frame.paragraphs[1:6]] == (
        points
    )
    assert slots["SCM_STATIC_IMAGE_CAPTION"].top > baseline["SCM_STATIC_IMAGE_CAPTION"].top
    assert slots["SCM_SLOT_IMAGE_SUPPLEMENT_01"].top > baseline[
        "SCM_SLOT_IMAGE_SUPPLEMENT_01"
    ].top
    assert max(shape.top + shape.height for shape in slots.values()) <= deck.slide_height


def test_layout_preview_keeps_extra_selling_points_and_multiline_sku(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    source = _source()
    product = source["single_products"][0]["product"]  # type: ignore[index]
    selling_points = [f"虚构完整卖点 {index}-" + "宽" * 35 for index in range(1, 7)]
    product.update(sku="SKU-LONG-" + "W" * 36, selling_points="\n".join(selling_points))
    deck = _render(
        _runtime(tmp_path),
        source,
        monkeypatch,
        layout_preview=True,
    )
    output = BytesIO()
    deck.save(output)
    reopened = Presentation(BytesIO(output.getvalue()))
    main_selling = _slots(reopened)["SCM_FIELD_SELLING_POINTS"]
    assert [paragraph.text for paragraph in main_selling.text_frame.paragraphs[1:7]] == (
        selling_points
    )
    diagnostic_text = "\n".join(
        shape.text
        for slide in list(reopened.slides)[2:]
        for shape in slide.shapes
        if shape.has_text_frame
    )
    normalized = diagnostic_text.replace("\n", "")
    assert "仅用于排版调试" in diagnostic_text
    assert all(point in normalized for point in selling_points)
    assert "PPT v5 layout preview bypassed capacity" in caplog.text


def test_common_hyphenated_sku_uses_original_gap_capacity(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source()
    source["single_products"][0]["product"]["sku"] = "SKU-2026-001"  # type: ignore[index]
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    assert _reference_value(_slots(deck), "sku") == "SKU-2026-001"


def test_two_line_normal_title_uses_existing_title_area(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source()
    source["single_products"][0]["product"]["product_name"] = "宽" * 35  # type: ignore[index]
    deck = _render(_runtime(tmp_path), source, monkeypatch)
    assert _slots(deck)["SCM_FIELD_TITLE"].text_frame.paragraphs[0].text == "宽" * 35


def test_four_images_follow_original_slot_order_and_are_not_cropped(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source()
    source["single_products"][0]["renderer_supplementary_image_ids"] = ["one", "two", "three"]
    images = {
        identifier: _png(color, size)
        for identifier, color, size in (
            ("fictional-1", (200, 20, 20), (300, 80)),
            ("one", (20, 200, 20), (80, 300)),
            ("two", (20, 20, 200), (100, 100)),
            ("three", (200, 200, 20), (140, 100)),
        )
    }
    runtime = _runtime(tmp_path)
    baseline = _slots(Presentation(runtime))
    deck = _render(runtime, source, monkeypatch, images)
    for name, identifier in zip(festive._IMAGE_SLOT_GEOMETRY, images, strict=True):
        original, picture = baseline[name], _slots(deck)[name]
        assert sha256(picture.image.blob).digest() == sha256(images[identifier]).digest()
        assert original.left <= picture.left and original.top <= picture.top
        assert picture.left + picture.width <= original.left + original.width
        assert picture.top + picture.height <= original.top + original.height
        assert (
            abs(picture.width / picture.height - picture.image.size[0] / picture.image.size[1])
            < 0.001
        )
        assert (
            picture.crop_left == picture.crop_right == picture.crop_top == picture.crop_bottom == 0
        )


def test_missing_middle_picture_does_not_shift_or_duplicate_other_images(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source()
    source["single_products"][0]["renderer_supplementary_image_ids"] = ["one", "missing", "three"]
    images = {"one": _png((10, 200, 20)), "three": _png((20, 10, 200))}
    slots = _slots(_render(_runtime(tmp_path), source, monkeypatch, images))
    assert slots["SCM_SLOT_IMAGE_MAIN"].image.size == (1, 1)
    assert slots["SCM_SLOT_IMAGE_SUPPLEMENT_02"].image.size == (1, 1)
    assert slots["SCM_SLOT_IMAGE_SUPPLEMENT_01"].image.blob == images["one"]
    assert slots["SCM_SLOT_IMAGE_SUPPLEMENT_03"].image.blob == images["three"]


def test_unconfirmed_combination_price_contract_is_controlled(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    with pytest.raises(AppError) as error:
        _render(
            _runtime(tmp_path),
            {
                "packages": [
                    {"total_price": "999", "items": [{"price_snapshot": {"agreement_price": "99"}}]}
                ]
            },
            monkeypatch,
        )
    assert error.value.code == "PPT_TEMPLATE_COMBINATION_UNSUPPORTED"


def test_five_products_keep_fields_and_media_isolated(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _source()
    first = source["single_products"][0]
    source["single_products"] = [deepcopy(first) for _ in range(5)]
    for index, item in enumerate(source["single_products"]):
        item["product"].update(id=f"fictional-{index}", product_name=f"虚构商品{index}")
        item["prices"]["agreement_price"] = str(100 + index)
    images = {f"fictional-{i}": _png((30 * i, 100, 20)) for i in range(5)}
    deck = _render(_runtime(tmp_path), source, monkeypatch, images)
    assert len(deck.slides) == 6
    for index in range(5):
        slots = _slots(deck, index + 1)
        assert slots["SCM_FIELD_TITLE"].text == f"虚构商品{index}\n"
        assert _reference_value(slots, "jd_customer_price") == f"{100 + index}.00元"
        assert slots["SCM_SLOT_IMAGE_MAIN"].image.blob == images[f"fictional-{index}"]


@pytest.mark.parametrize(
    "corruption", ("token", "duplicate", "geometry", "style", "label", "bytes")
)
def test_corrupt_runtime_is_a_controlled_app_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, corruption: str
) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    title = _slots(deck)["SCM_FIELD_TITLE"]
    if corruption == "token":
        title.text_frame.paragraphs[0].runs[0].text = "bad"
    elif corruption == "duplicate":
        title.name = "SCM_FIELD_REFERENCE"
    elif corruption == "geometry":
        title.left = Inches(4)
    elif corruption == "style":
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    elif corruption == "label":
        _slots(deck)["SCM_FIELD_REFERENCE"].text_frame.paragraphs[0].runs[0].text = "市场价："
    deck.save(runtime)
    if corruption == "bytes":
        runtime.write_bytes(b"not a zip")
    with pytest.raises(AppError) as error:
        _render(runtime, _source(), monkeypatch)
    assert error.value.code == "PPT_TEMPLATE_ASSET_STRUCTURE_INVALID"


def test_source_sha_and_two_page_contract_prevent_old_asset_build(tmp_path: Path) -> None:
    original = tmp_path / "source.pptx"
    _source_deck(original)
    with pytest.raises(AppError):
        festive.standardize_jd_festive_red_v5(original, tmp_path / "runtime.pptx")
    deck = Presentation(original)
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(original)
    with pytest.raises(AppError):
        festive.standardize_jd_festive_red_v5(
            original,
            tmp_path / "runtime.pptx",
            expected_source_fingerprint=festive._structure_fingerprint(deck),
            expected_source_sha256=None,
        )


def test_atomic_build_does_not_publish_failed_candidate(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source, runtime = tmp_path / "source.pptx", tmp_path / "runtime.pptx"
    _source_deck(source)
    runtime.write_bytes(b"previous-approved-asset")

    def fail(_source: Path, destination: Path) -> None:
        destination.write_bytes(b"unreviewed-candidate")
        raise AppError("PPT_TEMPLATE_ASSET_STRUCTURE_INVALID", "invalid", 409)

    monkeypatch.setattr(builder, "standardize_jd_festive_red_v5", fail)
    with pytest.raises(AppError):
        builder.build_jd_festive_red_v5_runtime_asset(source, runtime)
    assert runtime.read_bytes() == b"previous-approved-asset"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["runtime.pptx", "source.pptx"]


def test_default_stays_v3_and_explicit_dispatch_is_v5(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Developer machines may intentionally replace the default with v5.  The
    # production-safe baseline remains v3 when that local-only flag is absent.
    monkeypatch.setattr(ppt_template_registry, "is_local_v5_default_enabled", lambda: False)
    assert get_ppt_template("JD_FESTIVE_RED").version == "3"
    assert (
        get_pptx_manifest("JD_FESTIVE_RED", "5").clean_filename
        == "JD_FESTIVE_RED_OFFICIAL_CLEAN_v5.pptx"
    )
    deck = _render(_runtime(tmp_path), _source(), monkeypatch)
    assert deck.core_properties.identifier.startswith("JD_FESTIVE_RED:5:two-page:")
