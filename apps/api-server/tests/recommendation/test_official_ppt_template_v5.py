from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.common.contracts import AppError
from app.modules.recommendation.application import ppt_official_template_v5
from app.modules.recommendation.application.ppt_official_template_v5 import (
    render_jd_detail_red_v5,
    standardize_jd_detail_red_v5,
    verify_no_sample_media,
)
from app.modules.recommendation.application.ppt_renderer import PptRenderer
from app.modules.recommendation.application.ppt_template_manifest import get_pptx_manifest
from app.modules.recommendation.application.ppt_template_registry import get_ppt_template
from scripts import build_private_official_ppt_template_v5
from scripts.build_private_official_ppt_template_v5 import build_runtime_asset


def _png(color: tuple[int, int, int], size: tuple[int, int] = (160, 80)) -> bytes:
    output = BytesIO()
    Image.new("RGB", size, color).save(output, format="PNG")
    return output.getvalue()


def _source_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    cover.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "固定封面"
    page = deck.slides.add_slide(deck.slide_layouts[6])
    # The standardiser deliberately reads audited source indices once only.
    dynamic_geometry = {
        0: (6.379, 3.290, 6.299, 0.503),
        1: (6.379, 3.757, 6.299, 0.503),
        2: (6.389, 4.995, 6.299, 1.715),
        3: (6.389, 4.236, 6.299, 0.301),
        4: (6.465, 0.854, 6.268, 0.335),
        # Static divider below the brand row and fixed customization row below
        # delivery model the linked objects in the audited customer layout.
        5: (6.379, 3.082, 6.299, 0.001),
        6: (6.379, 2.615, 6.299, 0.301),
        7: (6.360, 1.186, 3.457, 1.688),
        8: (9.757, 1.165, 3.322, 1.816),
        10: (6.389, 4.488, 6.299, 0.301),
    }
    for index in range(15):
        left, top, width, height = dynamic_geometry.get(
            index,
            (0.3 + (index % 5) * 2.5, 0.3 + (index // 5) * 2.2, 2.2, 1.5),
        )
        shape = page.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        shape.text = f"填写说明 {index}"
    sample = deck.slides.add_slide(deck.slide_layouts[6])
    sample.shapes.add_picture(
        BytesIO(_png((220, 10, 10))), Inches(1), Inches(1), Inches(2), Inches(2)
    )
    deck.save(path)


def _source() -> dict[str, object]:
    return {
        "single_products": [
            {
                "product": {
                    "id": "fictional-1",
                    "product_name": "虚构单品名称",
                    "brand": "虚构品牌",
                    "product_specification": "虚构规格 100mm",
                    "packaging_list": "虚构包装清单",
                    "selling_points": "虚构卖点 A\n虚构卖点 B",
                    "shipping_courier": "虚构快递",
                    "sku": "FICT-001",
                },
                "prices": {
                    "market_price": "199.5",
                    "agreement_price": "159.25",
                    "discount_rate": "0.85",
                },
                "manual": {"fulfillment_cycle": "3 个工作日"},
            }
        ]
    }


def _runtime(tmp_path: Path) -> Path:
    source = tmp_path / "template1.pptx"
    runtime = tmp_path / "JD_DETAIL_RED_OFFICIAL_CLEAN_v5.pptx"
    _source_deck(source)
    source_deck = Presentation(source)
    standardize_jd_detail_red_v5(
        source,
        runtime,
        expected_source_fingerprint=ppt_official_template_v5._structure_fingerprint(source_deck),
    )
    verify_no_sample_media(source, runtime)
    return runtime


def test_standardizer_removes_sample_slide_and_media(tmp_path: Path) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    assert len(deck.slides) == 2
    assert {shape.name for shape in deck.slides[1].shapes} >= {
        "SCM_FIELD_TITLE",
        "SCM_FIELD_REFERENCE",
        "SCM_SLOT_IMAGE_MAIN",
    }
    assert "填写说明 0" not in "\n".join(
        shape.text for shape in deck.slides[1].shapes if shape.has_text_frame
    )


def test_standardizer_rejects_source_that_does_not_match_formal_contract(tmp_path: Path) -> None:
    source = tmp_path / "template1.pptx"
    _source_deck(source)
    with pytest.raises(AppError) as error:
        standardize_jd_detail_red_v5(source, tmp_path / "runtime.pptx")
    assert error.value.code == "PPT_TEMPLATE_ASSET_STRUCTURE_INVALID"


def test_v5_standardizer_restores_sample_brand_and_supplement_layout(tmp_path: Path) -> None:
    deck = Presentation(_runtime(tmp_path))
    page = deck.slides[1]
    brand = next(shape for shape in page.shapes if shape.name == "SCM_FIELD_BRAND")
    brand_geometry = (brand.left.inches, brand.top.inches, brand.width.inches, brand.height.inches)
    assert brand_geometry == pytest.approx((6.379, 2.615, 6.299, 0.301), abs=0.001)
    slots = [
        next(shape for shape in page.shapes if shape.name == name)
        for name in (
            "SCM_SLOT_IMAGE_SUPPLEMENT_01",
            "SCM_SLOT_IMAGE_SUPPLEMENT_02",
            "SCM_SLOT_IMAGE_SUPPLEMENT_03",
        )
    ]
    expected_geometry = (
        (0.448, 5.404, 1.367, 1.969),
        (2.258, 5.404, 1.354, 1.969),
        (4.069, 5.329, 1.313, 1.969),
    )
    for slot, expected in zip(slots, expected_geometry, strict=True):
        geometry = (slot.left.inches, slot.top.inches, slot.width.inches, slot.height.inches)
        assert geometry == pytest.approx(expected, abs=0.001)
    assert [slot.left for slot in slots] == sorted(slot.left for slot in slots)


def test_v5_standardizer_keeps_customer_price_label_black_and_value_red(tmp_path: Path) -> None:
    deck = Presentation(_runtime(tmp_path))
    prices = next(shape for shape in deck.slides[1].shapes if shape.name == "SCM_FIELD_PRICES")
    label, value = prices.text_frame.paragraphs[0].runs
    assert label.text == "京东大客户价：" and label.font.color.rgb == RGBColor(0, 0, 0)
    assert value.text == "{{jd_customer_price}}" and value.font.color.rgb == RGBColor(255, 0, 0)

    # This is deliberately a renderer-only fictional value; it does not add a
    # product, quotation, API, or frozen business-field source.
    ppt_official_template_v5._replace_token(prices, "{{jd_customer_price}}", "321.00")
    label, value = prices.text_frame.paragraphs[0].runs
    assert label.text == "京东大客户价：" and label.font.color.rgb == RGBColor(0, 0, 0)
    assert value.text == "321.00" and value.font.color.rgb == RGBColor(255, 0, 0)


def test_v5_standardizer_declares_jd_font_for_latin_and_east_asian_text(
    tmp_path: Path,
) -> None:
    page = Presentation(_runtime(tmp_path)).slides[1]
    for name in ("SCM_FIELD_REFERENCE", "SCM_FIELD_PRICES"):
        shape = next(shape for shape in page.shapes if shape.name == name)
        for paragraph in shape.text_frame.paragraphs:
            assert paragraph.line_spacing == 1.5
            for run in paragraph.runs:
                properties = run._r.get_or_add_rPr()
                assert run.font.size is not None and run.font.size.pt == 14
                assert properties.find(qn("a:latin")).get("typeface") == "京东朗正体 玲珑"
                assert properties.find(qn("a:ea")).get("typeface") == "京东朗正体 玲珑"


@pytest.mark.parametrize("agreement_price", ("321.50", None))
def test_v5_renderer_maps_agreement_price_and_preserves_customer_price_style(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    agreement_price: str | None,
) -> None:
    runtime = _runtime(tmp_path)
    # Simulate the already-built local runtime, which originally had only
    # a:latin, and verify the renderer repairs the cloned output slide.
    deck = Presentation(runtime)
    runtime_prices = next(
        shape for shape in deck.slides[1].shapes if shape.name == "SCM_FIELD_PRICES"
    )
    for paragraph in runtime_prices.text_frame.paragraphs:
        for run in paragraph.runs:
            properties = run._r.get_or_add_rPr()
            east_asian = properties.find(qn("a:ea"))
            if east_asian is not None:
                properties.remove(east_asian)
    deck.save(runtime)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["prices"]["agreement_price"] = agreement_price  # type: ignore[index]
    result = PptRenderer().render(
        source,
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    reopened = Presentation(BytesIO(result))
    prices = next(
        shape for shape in reopened.slides[1].shapes if shape.name == "SCM_FIELD_PRICES"
    )
    label, value = prices.text_frame.paragraphs[0].runs
    assert label.text == "京东大客户价："
    assert label.font.color.rgb == RGBColor(0, 0, 0)
    assert value.text == ("321.50" if agreement_price else "")
    assert value.font.color.rgb == RGBColor(255, 0, 0)
    assert value.font.size is not None and value.font.size.pt == 14
    assert prices.text_frame.paragraphs[0].line_spacing == 1.5
    for run in (label, value):
        properties = run._r.get_or_add_rPr()
        assert properties.find(qn("a:latin")).get("typeface") == "京东朗正体 玲珑"
        assert properties.find(qn("a:ea")).get("typeface") == "京东朗正体 玲珑"


def test_v5_standardizer_uses_single_line_spacing_only_for_requested_fields(
    tmp_path: Path,
) -> None:
    page = Presentation(_runtime(tmp_path)).slides[1]
    for name in (
        "SCM_FIELD_SPECIFICATION",
        "SCM_FIELD_PACKAGING",
        "SCM_FIELD_DELIVERY",
        "SCM_FIELD_SELLING_POINTS",
    ):
        assert all(paragraph.line_spacing == 1.0 for paragraph in next(
            shape for shape in page.shapes if shape.name == name
        ).text_frame.paragraphs)
    assert next(
        shape for shape in page.shapes if shape.name == "SCM_FIELD_BRAND"
    ).text_frame.paragraphs[0].line_spacing == 1.5


def test_v5_renderer_numbers_only_nonempty_selling_points(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["selling_points"] = "1、第一项\n\n第二项\n3.第三项"  # type: ignore[index]
    page = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    ).slides[1]
    selling = next(shape for shape in page.shapes if shape.name == "SCM_FIELD_SELLING_POINTS")
    assert [paragraph.text for paragraph in selling.text_frame.paragraphs] == [
        "产品卖点（介绍详细）",
        "1、第一项",
        "2、第二项",
        "3、第三项",
    ]
    assert all(paragraph.line_spacing == 1.0 for paragraph in selling.text_frame.paragraphs)


def test_v5_renderer_keeps_empty_selling_points_as_header_only(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["selling_points"] = "\n\n"  # type: ignore[index]
    page = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    ).slides[1]
    selling = next(shape for shape in page.shapes if shape.name == "SCM_FIELD_SELLING_POINTS")
    assert [paragraph.text for paragraph in selling.text_frame.paragraphs] == [
        "产品卖点（介绍详细）"
    ]


def test_v5_renderer_grows_selling_points_for_more_than_six_rows(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["selling_points"] = "\n".join(  # type: ignore[index]
        f"卖点 {index}" for index in range(1, 8)
    )
    page = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    ).slides[1]
    selling = next(shape for shape in page.shapes if shape.name == "SCM_FIELD_SELLING_POINTS")
    assert [paragraph.text for paragraph in selling.text_frame.paragraphs[1:]] == [
        f"{index}、卖点 {index}" for index in range(1, 8)
    ]


def test_v5_renderer_grows_for_wrapped_selling_points(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["selling_points"] = "\n".join(  # type: ignore[index]
        ["宽" * 40] * 5
    )
    baseline = Presentation(runtime).slides[1]
    page = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
                layout_preview=True,
            )
        )
    ).slides[1]
    baseline_selling = next(
        shape for shape in baseline.shapes if shape.name == "SCM_FIELD_SELLING_POINTS"
    )
    selling = next(shape for shape in page.shapes if shape.name == "SCM_FIELD_SELLING_POINTS")
    assert selling.height > baseline_selling.height
    assert selling.top + selling.height <= Inches(7.5)


def test_v5_renderer_keeps_six_safe_selling_points_complete_and_readable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["selling_points"] = "\n".join(  # type: ignore[index]
        f"虚构卖点 {index}" for index in range(1, 7)
    )
    page = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    ).slides[1]
    selling = next(shape for shape in page.shapes if shape.name == "SCM_FIELD_SELLING_POINTS")
    assert [paragraph.text for paragraph in selling.text_frame.paragraphs[1:]] == [
        f"{index}、虚构卖点 {index}" for index in range(1, 7)
    ]


@pytest.mark.parametrize(
    ("field", "value"),
    (
        ("product_name", "长标题" * 20),
        ("sku", "SKU" * 20),
        ("product_specification", "长规格" * 20),
    ),
)
def test_v5_renderer_wraps_and_grows_long_title_sku_and_specification(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, field: str, value: str
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"][field] = value  # type: ignore[index]
    assert Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    ).slides[1]


@pytest.mark.parametrize(
    ("raw", "expected"),
    (
        (
            "产品材质：双层牛津布，铁管\n产品尺寸：24*28*28",
            "产品材质：双层牛津布，铁管；产品尺寸：24*28*28",
        ),
        (
            "产品材质：双层牛津布；\n产品尺寸：24*28*28",
            "产品材质：双层牛津布；产品尺寸：24*28*28",
        ),
        (
            "产品材质：双层牛津布；\n；产品尺寸：24*28*28",
            "产品材质：双层牛津布；产品尺寸：24*28*28",
        ),
        (
            "产品材质：双层牛津布\n；产品尺寸：24*28*28",
            "产品材质：双层牛津布；产品尺寸：24*28*28",
        ),
        (
            "产品材质：双层牛津布，\n产品尺寸：24*28*28",
            "产品材质：双层牛津布，产品尺寸：24*28*28",
        ),
        ("产品尺寸：24*28*28", "产品尺寸：24*28*28"),
    ),
)
def test_v5_specification_formatter_joins_rows_without_duplicate_separator(
    raw: str, expected: str
) -> None:
    assert ppt_official_template_v5._format_specification(raw) == expected


def test_v5_renderer_joins_specification_rows_and_uses_natural_wrapping(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    raw = (
        "产品材质：双层牛津布，铁管\r\n\r\n"
        "产品尺寸：24*28*28\n"
        "包装尺寸：31.5*12*12CM\r"
        "产品箱规：26*33*64CM/8公斤/10个"
    )
    expected = (
        "产品材质：双层牛津布，铁管；产品尺寸：24*28*28；"
        "包装尺寸：31.5*12*12CM；产品箱规：26*33*64CM/8公斤/10个"
    )
    source = _source()
    source["single_products"][0]["product"]["product_specification"] = raw  # type: ignore[index]
    result = PptRenderer().render(
        source,
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    reopened = Presentation(BytesIO(result))
    specification = next(
        shape
        for shape in reopened.slides[1].shapes
        if shape.name == "SCM_FIELD_SPECIFICATION"
    )
    assert len(specification.text_frame.paragraphs) == 1
    assert specification.text == "商品规格：" + expected
    assert "\n" not in specification.text and "\r" not in specification.text
    assert specification.text_frame.word_wrap is True
    assert specification.text_frame.paragraphs[0].line_spacing == 1.0
    assert all(
        run.font.name == "Microsoft YaHei"
        and run.font.size is not None
        and run.font.size.pt == 12
        for run in specification.text_frame.paragraphs[0].runs
    )


def test_v5_specification_natural_wrap_grows_only_by_required_height(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    normal_result = PptRenderer().render(
        _source(),
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    source = _source()
    raw_items = [f"规格属性{index}：" + "宽" * 20 for index in range(1, 7)]
    source["single_products"][0]["product"]["product_specification"] = "\n".join(  # type: ignore[index]
        raw_items
    )
    wrapped_result = PptRenderer().render(
        source,
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    normal = {
        shape.name: shape for shape in Presentation(BytesIO(normal_result)).slides[1].shapes
    }
    wrapped_deck = Presentation(BytesIO(wrapped_result))
    wrapped = {shape.name: shape for shape in wrapped_deck.slides[1].shapes}
    specification = wrapped["SCM_FIELD_SPECIFICATION"]
    growth = specification.height - normal["SCM_FIELD_SPECIFICATION"].height
    assert growth > 0
    assert "；".join(raw_items) in specification.text
    assert "\n" not in specification.text
    assert ppt_official_template_v5._shape_line_count(specification) > 1
    assert ppt_official_template_v5._required_text_height(specification) <= specification.height
    assert wrapped["SCM_FIELD_PACKAGING"].top - normal["SCM_FIELD_PACKAGING"].top == growth
    assert wrapped["SCM_FIELD_DELIVERY"].top - normal["SCM_FIELD_DELIVERY"].top == growth
    assert (
        wrapped["SCM_FIELD_SELLING_POINTS"].top
        - normal["SCM_FIELD_SELLING_POINTS"].top
        == growth
    )
    visible_bottom = (
        specification.top
        + ppt_official_template_v5._required_text_height(specification)
        - specification.text_frame.margin_bottom
    )
    assert visible_bottom <= wrapped["SCM_FIELD_PACKAGING"].top


@pytest.mark.parametrize(
    ("characters", "expected_lines"),
    ((10, 1), (50, 2), (160, 5)),
)
def test_v5_specification_shape_clears_packaging_at_real_line_height(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    characters: int,
    expected_lines: int,
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["product_specification"] = "宽" * characters  # type: ignore[index]
    result = PptRenderer().render(
        source,
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    reopened = Presentation(BytesIO(result))
    shapes = {shape.name: shape for shape in reopened.slides[1].shapes}
    specification = shapes["SCM_FIELD_SPECIFICATION"]
    assert ppt_official_template_v5._shape_line_count(specification) == expected_lines
    required_height = ppt_official_template_v5._required_text_height(specification)
    assert specification.height >= required_height
    assert (
        specification.top + required_height - specification.text_frame.margin_bottom
        <= shapes["SCM_FIELD_PACKAGING"].top
    )
    assert specification.text_frame.paragraphs[0].line_spacing == 1.0
    assert all(
        run.font.name == "Microsoft YaHei"
        and run.font.size is not None
        and run.font.size.pt == 12
        for run in specification.text_frame.paragraphs[0].runs
    )


def test_v5_renderer_accumulates_title_and_sku_growth_without_overlap(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    product = source["single_products"][0]["product"]  # type: ignore[index]
    product.update(product_name="宽" * 45, sku="W" * 55)
    baseline = Presentation(runtime).slides[1]
    page = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    ).slides[1]
    before = {shape.name: shape for shape in baseline.shapes}
    after = {shape.name: shape for shape in page.shapes}
    assert after["SCM_FIELD_TITLE"].height > before["SCM_FIELD_TITLE"].height
    assert after["SCM_FIELD_REFERENCE"].top > before["SCM_FIELD_REFERENCE"].top
    assert after["SCM_FIELD_REFERENCE"].height > before["SCM_FIELD_REFERENCE"].height
    assert after["SCM_FIELD_BRAND"].top > before["SCM_FIELD_BRAND"].top
    assert (
        ppt_official_template_v5._visible_text_bottom(
            after["SCM_FIELD_SELLING_POINTS"]
        )
        <= Inches(7.5)
    )


def test_v5_multiple_wrapped_attributes_move_every_linked_object_once(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    product = source["single_products"][0]["product"]  # type: ignore[index]
    product.update(
        brand="宽" * 45,
        product_specification="宽" * 75,
        packaging_list="宽" * 75,
        shipping_courier="宽" * 45,
    )
    before_page = Presentation(runtime).slides[1]
    after_page = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
                layout_preview=True,
            )
        )
    ).slides[1]
    before = {shape.name: shape for shape in before_page.shapes}
    after = {shape.name: shape for shape in after_page.shapes}
    growth = {
        name: after[name].height - before[name].height
        for name in (
            "SCM_FIELD_BRAND",
            "SCM_FIELD_SPECIFICATION",
            "SCM_FIELD_PACKAGING",
            "SCM_FIELD_DELIVERY",
        )
    }
    assert all(value > 0 for value in growth.values())
    divider_before, divider_after = before["Rectangle 6"], after["Rectangle 6"]
    customization_before = before["Rectangle 11"]
    customization_after = after["Rectangle 11"]
    assert divider_after.top - divider_before.top == growth["SCM_FIELD_BRAND"]
    assert after["SCM_FIELD_PACKAGING"].top - before["SCM_FIELD_PACKAGING"].top == (
        growth["SCM_FIELD_BRAND"]
        + growth["SCM_FIELD_SPECIFICATION"]
    )
    assert after["SCM_FIELD_DELIVERY"].top - before["SCM_FIELD_DELIVERY"].top == sum(
        growth[name]
        for name in ("SCM_FIELD_BRAND", "SCM_FIELD_SPECIFICATION", "SCM_FIELD_PACKAGING")
    )
    expected_lower_shift = sum(growth.values())
    assert customization_after.top - customization_before.top == expected_lower_shift
    assert (
        after["SCM_FIELD_SELLING_POINTS"].top - before["SCM_FIELD_SELLING_POINTS"].top
        == expected_lower_shift
    )
    for current, following in (
        ("SCM_FIELD_SPECIFICATION", "SCM_FIELD_PACKAGING"),
        ("SCM_FIELD_PACKAGING", "SCM_FIELD_DELIVERY"),
        ("SCM_FIELD_DELIVERY", "Rectangle 11"),
    ):
        shape = after[current]
        visible_bottom = (
            shape.top
            + ppt_official_template_v5._required_text_height(shape)
            - shape.text_frame.margin_bottom
            - int(
                ppt_official_template_v5._OFFICE_TEXT_HEIGHT_SAFETY_PT * 12_700
            )
        )
        assert visible_bottom <= after[following].top


def test_v5_specification_page_overflow_reports_field_geometry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    # This many wrapped lines put visible glyphs beyond the page, rather than only
    # moving the text-box outer boundary through its harmless bottom inset.
    source["single_products"][0]["product"]["product_specification"] = "宽" * 450  # type: ignore[index]
    with pytest.raises(AppError, match="specification") as error:
        render_jd_detail_red_v5(
            source,
            images={},
            template=get_ppt_template("JD_DETAIL_RED", "5"),
            manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
        )
    assert error.value.code == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"
    assert "SCM_FIELD_SPECIFICATION:行" in error.value.message
    assert "可见内容底部" in error.value.message
    assert "页高 540.0pt" in error.value.message


def test_v5_layout_preview_generates_title_and_ten_line_specification(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    product = source["single_products"][0]["product"]  # type: ignore[index]
    product.update(product_name="虚构长标题" + "宽" * 100, product_specification="宽" * 325)
    renderer = PptRenderer(layout_preview=True)
    result = renderer.render(
        source,
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    deck = Presentation(BytesIO(result))
    page = deck.slides[1]
    shapes = {shape.name: shape for shape in page.shapes}
    assert shapes["SCM_FIELD_TITLE"].text == product["product_name"]
    assert "宽" * 325 in shapes["SCM_FIELD_SPECIFICATION"].text
    diagnostic_text = "\n".join(
        shape.text
        for slide in list(deck.slides)[2:]
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "仅用于排版调试" in diagnostic_text
    assert renderer.layout_diagnostics_added is True
    assert "[商品规格]" in diagnostic_text
    assert diagnostic_text.count("宽") >= 425
    assert "PPT v5 layout preview bypassed page overflow" in caplog.text


@pytest.mark.parametrize("specification_length", (50, 90, 125, 160))
def test_v5_normal_two_to_five_line_wraps_stay_clean_in_explicit_preview_mode(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    specification_length: int,
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    product = source["single_products"][0]["product"]  # type: ignore[index]
    product["product_specification"] = "宽" * specification_length
    product["packaging_list"] = "包" * 40
    renderer = PptRenderer(layout_preview=True)
    result = renderer.render(
        source,
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    deck = Presentation(BytesIO(result))
    assert len(deck.slides) == 2
    assert renderer.layout_diagnostics_added is False
    assert "排版诊断版（仅用于排版调试）" not in "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


def test_v5_page_boundary_uses_visible_text_not_harmless_outer_box_inset(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["product_specification"] = "宽" * 325  # type: ignore[index]
    result = PptRenderer(layout_preview=True).render(
        source,
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    deck = Presentation(BytesIO(result))
    page = deck.slides[1]
    shapes = {shape.name: shape for shape in page.shapes}
    selling = shapes["SCM_FIELD_SELLING_POINTS"]
    # Reproduce the reported 540.6pt outer boundary exactly.  The bottom inset
    # remains below the final visible line, so this is not a content overflow.
    selling.top = int(deck.slide_height + Pt(0.6) - selling.height)
    linked = {
        name: shape for name, shape in shapes.items() if shape.left >= Inches(6)
    }
    visible_bottom, outer_bottom = ppt_official_template_v5._layout_required_bottom(
        shapes, linked
    )
    assert (outer_bottom - deck.slide_height) / 12_700 == pytest.approx(0.6, abs=0.01)
    assert visible_bottom <= deck.slide_height


def test_v5_visible_text_overlap_still_fails_safely(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    page = deck.slides[1]
    shapes = {shape.name: shape for shape in page.shapes}
    shapes["SCM_FIELD_PACKAGING"].top = shapes["SCM_FIELD_SPECIFICATION"].top + Pt(5)
    deck.save(runtime)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    with pytest.raises(AppError, match="可见文字") as error:
        render_jd_detail_red_v5(
            _source(),
            images={},
            template=get_ppt_template("JD_DETAIL_RED", "5"),
            manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
        )
    assert error.value.code == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"


def test_v5_renderer_accepts_normal_fixed_row_values(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    product = source["single_products"][0]["product"]  # type: ignore[index]
    product.update({"sku": "SCM-2026", "brand": "虚构品牌", "shipping_courier": "京东物流"})
    result = render_jd_detail_red_v5(
        source,
        images={},
        template=get_ppt_template("JD_DETAIL_RED", "5"),
        manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
    )
    assert len(Presentation(BytesIO(result)).slides) == 2


def test_v5_renderer_preserves_runtime_price_heights_and_vertical_baseline(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    baseline_page = Presentation(runtime).slides[1]
    baseline = {shape.name: shape for shape in baseline_page.shapes}
    result = PptRenderer().render(
        _source(),
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    reopened = Presentation(BytesIO(result))
    page = reopened.slides[1]
    rendered = {shape.name: shape for shape in page.shapes}
    for name in ("SCM_FIELD_REFERENCE", "SCM_FIELD_PRICES"):
        assert rendered[name].top == baseline[name].top
        assert rendered[name].height == baseline[name].height
    for name in (
        "SCM_FIELD_BRAND",
        "SCM_FIELD_SPECIFICATION",
        "SCM_FIELD_PACKAGING",
        "SCM_FIELD_DELIVERY",
        "SCM_FIELD_SELLING_POINTS",
    ):
        assert rendered[name].top == baseline[name].top
    assert rendered["SCM_SLOT_IMAGE_MAIN"].top == baseline["SCM_SLOT_IMAGE_MAIN"].top


def test_v5_renderer_expands_runtime_price_height_for_wrapped_sku(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["sku"] = "JD-PRO-" + "A1B2C3D4-" * 16  # type: ignore[index]
    baseline = {shape.name: shape for shape in Presentation(runtime).slides[1].shapes}
    result = render_jd_detail_red_v5(
        source,
        images={},
        template=get_ppt_template("JD_DETAIL_RED", "5"),
        manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
        layout_preview=True,
    )
    reopened = Presentation(BytesIO(result))
    rendered = {shape.name: shape for shape in reopened.slides[1].shapes}
    growth = rendered["SCM_FIELD_REFERENCE"].height - baseline["SCM_FIELD_REFERENCE"].height
    assert growth > 0
    assert rendered["SCM_FIELD_REFERENCE"].text_frame.paragraphs[1].text.endswith(
        "A1B2C3D4-" * 16
    )
    assert rendered["SCM_FIELD_BRAND"].top == (
        baseline["SCM_FIELD_BRAND"].top + growth
    )
    assert all(
        run.font.size is not None and run.font.size.pt == 14
        for run in rendered["SCM_FIELD_REFERENCE"].text_frame.paragraphs[1].runs
    )


def test_v5_full_renderer_dispatch_preserves_fixed_property_contract(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"].update(  # type: ignore[index]
        product_name="虚构中英文商品 JD-PRO-2026 " + "宽" * 24,
        product_specification="宽" * 75,
        packaging_list="包装一、包装二、说明书、电源适配器",
    )
    result = PptRenderer().render(
        source,
        template_code="JD_DETAIL_RED",
        template_version="5",
        product_images={},
    )
    deck = Presentation(BytesIO(result))
    assert len(deck.slides) == 2
    page = deck.slides[1]
    shapes = {shape.name: shape for shape in page.shapes}
    assert "虚构中英文商品 JD-PRO-2026" in shapes["SCM_FIELD_TITLE"].text
    assert "宽" * 75 in shapes["SCM_FIELD_SPECIFICATION"].text
    assert shapes["SCM_FIELD_PACKAGING"].top > shapes["SCM_FIELD_SPECIFICATION"].top
    selling = shapes["SCM_FIELD_SELLING_POINTS"]
    assert selling.top + selling.height <= Inches(7.5)


def test_v5_renderer_accepts_normal_mixed_model_title_without_false_wrap(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["product_name"] = (  # type: ignore[index]
        "JD-2026-PRO-500ML-A1B2C3D4E5F6G7H8J9K0L1M2N3P4"
    )
    result = render_jd_detail_red_v5(
        source,
        images={},
        template=get_ppt_template("JD_DETAIL_RED", "5"),
        manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
    )
    assert len(Presentation(BytesIO(result)).slides) == 2


@pytest.mark.parametrize(
    ("value", "expected"),
    (
        (0, "0%"),
        (1, "100%"),
        ("0.5", "50%"),
        ("0.85", "85%"),
        ("1.0", "100%"),
        (None, ""),
        ("", ""),
        ("not-a-rate", "not-a-rate"),
    ),
)
def test_v5_discount_formatter_preserves_integer_zeros(value: object, expected: str) -> None:
    assert ppt_official_template_v5._format_discount(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    (
        (889, "889.00"),
        ("699.50", "699.50"),
        ("￥889", "889.00"),
        ("¥ 699.5", "699.50"),
        (None, ""),
        ("", ""),
    ),
)
def test_v5_price_formatter_outputs_numeric_value_without_currency_symbol(
    value: object, expected: str
) -> None:
    assert ppt_official_template_v5._format_price(value) == expected


def test_v5_supplement_slots_accept_distinct_fictional_images_without_stretching(
    tmp_path: Path,
) -> None:
    page = Presentation(_runtime(tmp_path)).slides[1]
    fixtures = (
        ("SCM_SLOT_IMAGE_SUPPLEMENT_01", _png((200, 10, 10), (80, 160))),
        ("SCM_SLOT_IMAGE_SUPPLEMENT_02", _png((10, 200, 10), (160, 80))),
        ("SCM_SLOT_IMAGE_SUPPLEMENT_03", _png((10, 10, 200), (100, 100))),
    )
    hashes: list[str] = []
    for name, image in fixtures:
        slot = next(shape for shape in page.shapes if shape.name == name)
        ppt_official_template_v5._replace_picture(slot, image)
        hashes.append(sha256(slot.image.blob).hexdigest())
        assert slot.crop_left or slot.crop_top
    assert len(set(hashes)) == 3


def test_v5_renderer_centers_horizontal_and_vertical_main_image_crops(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    fixtures = (
        (_png((200, 10, 10), (240, 80)), True),
        (_png((10, 10, 200), (80, 240)), False),
    )
    for image, horizontal in fixtures:
        page = Presentation(
            BytesIO(
                render_jd_detail_red_v5(
                    _source(),
                    images={"fictional-1": image},
                    template=get_ppt_template("JD_DETAIL_RED", "5"),
                    manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
                )
            )
        ).slides[1]
        main = next(shape for shape in page.shapes if shape.name == "SCM_SLOT_IMAGE_MAIN")
        assert main.crop_left == pytest.approx(main.crop_right)
        assert main.crop_top == pytest.approx(main.crop_bottom)
        if horizontal:
            assert main.crop_left > 0 and main.crop_top == 0
        else:
            assert main.crop_top > 0 and main.crop_left == 0


def test_v5_clone_does_not_duplicate_slide_layout_placeholders() -> None:
    deck = Presentation()
    source = deck.slides.add_slide(deck.slide_layouts[1])
    source_count = len(source.shapes)
    clone = ppt_official_template_v5._clone_slide(deck, source)
    assert len(clone.shapes) == source_count


@pytest.mark.parametrize("corruption", ("token", "horizontal_geometry"))
def test_v5_renderer_rejects_runtime_token_or_horizontal_contract_mismatch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, corruption: str
) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    page = deck.slides[1]
    title = next(shape for shape in page.shapes if shape.name == "SCM_FIELD_TITLE")
    if corruption == "token":
        title.text_frame.paragraphs[0].runs[0].text = "{{wrong_product_name}}"
    else:
        title.left = Inches(6.2)
    deck.save(runtime)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    with pytest.raises(AppError) as error:
        render_jd_detail_red_v5(
            _source(),
            images={},
            template=get_ppt_template("JD_DETAIL_RED", "5"),
            manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
        )
    assert error.value.code == "PPT_TEMPLATE_ASSET_STRUCTURE_INVALID"


@pytest.mark.parametrize("packaging_list", ("", "主机、说明书"))
def test_v5_renderer_reads_short_packaging_height_and_top_from_runtime_asset(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    packaging_list: str,
) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    package = next(
        shape for shape in deck.slides[1].shapes if shape.name == "SCM_FIELD_PACKAGING"
    )
    package.height = Pt(21)
    package.top += Pt(2)
    authored_top = package.top
    authored_anchor = package.text_frame.vertical_anchor
    deck.save(runtime)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["packaging_list"] = packaging_list  # type: ignore[index]
    rendered = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    )
    output = next(
        shape
        for shape in rendered.slides[1].shapes
        if shape.name == "SCM_FIELD_PACKAGING"
    )
    assert output.height == Pt(21)
    assert output.top == authored_top
    assert output.text_frame.vertical_anchor == authored_anchor


def test_v5_renderer_expands_authored_packaging_height_and_moves_following_fields_once(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    baseline = {shape.name: shape for shape in deck.slides[1].shapes}
    baseline["SCM_FIELD_PACKAGING"].height = Pt(21)
    deck.save(runtime)
    authored_delivery_top = baseline["SCM_FIELD_DELIVERY"].top
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )

    heights: list[int] = []
    for value, expected_lines in (("宽" * 40, 2), ("宽" * 80, 3)):
        source = _source()
        source["single_products"][0]["product"]["packaging_list"] = value  # type: ignore[index]
        rendered = Presentation(
            BytesIO(
                render_jd_detail_red_v5(
                    source,
                    images={},
                    template=get_ppt_template("JD_DETAIL_RED", "5"),
                    manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
                )
            )
        )
        shapes = {shape.name: shape for shape in rendered.slides[1].shapes}
        package = shapes["SCM_FIELD_PACKAGING"]
        growth = package.height - Pt(21)
        assert ppt_official_template_v5._shape_line_count(package) == expected_lines
        assert growth > 0
        assert shapes["SCM_FIELD_DELIVERY"].top == authored_delivery_top + growth
        assert shapes["SCM_FIELD_SELLING_POINTS"].top >= (
            package.top + package.height
        )
        heights.append(package.height)
    assert heights[1] - heights[0] == pytest.approx(14.4 * 12_700, abs=1)


def test_v5_renderer_rejects_nonpositive_runtime_text_height(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    package = next(
        shape for shape in deck.slides[1].shapes if shape.name == "SCM_FIELD_PACKAGING"
    )
    package.height = 0
    deck.save(runtime)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    with pytest.raises(AppError) as error:
        render_jd_detail_red_v5(
            _source(),
            images={},
            template=get_ppt_template("JD_DETAIL_RED", "5"),
            manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
        )
    assert error.value.code == "PPT_TEMPLATE_ASSET_STRUCTURE_INVALID"


def test_v5_renderer_rejects_missing_runtime_slot_with_controlled_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    deck = Presentation(runtime)
    page = deck.slides[1]
    missing = next(shape for shape in page.shapes if shape.name == "SCM_SLOT_IMAGE_MAIN")
    missing._element.getparent().remove(missing._element)
    deck.save(runtime)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    with pytest.raises(AppError) as error:
        render_jd_detail_red_v5(
            _source(),
            images={},
            template=get_ppt_template("JD_DETAIL_RED", "5"),
            manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
        )
    assert error.value.code == "PPT_TEMPLATE_ASSET_STRUCTURE_INVALID"


def test_private_builder_preserves_existing_runtime_when_candidate_validation_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    original = tmp_path / "source.pptx"
    runtime = tmp_path / "runtime.pptx"
    original.write_bytes(b"private-source")
    runtime.write_bytes(b"approved-runtime")

    def fail_standardization(*_args: object, **_kwargs: object) -> None:
        raise AppError("PPT_TEMPLATE_ASSET_STRUCTURE_INVALID", "invalid", 409)

    monkeypatch.setattr(
        build_private_official_ppt_template_v5, "standardize_jd_detail_red_v5", fail_standardization
    )
    with pytest.raises(AppError, match="invalid"):
        build_runtime_asset(original, runtime)
    assert runtime.read_bytes() == b"approved-runtime"
    assert not list(tmp_path.glob("tmp*.pptx"))


def test_v5_render_preserves_fixed_cover_and_uses_frozen_values(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    result = render_jd_detail_red_v5(
        _source(),
        images={"fictional-1": _png((10, 20, 30))},
        template=get_ppt_template("JD_DETAIL_RED", "5"),
        manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
    )
    deck = Presentation(BytesIO(result))
    assert len(deck.slides) == 2
    assert deck.slides[0].shapes[0].text == "固定封面"
    page = deck.slides[1]
    text = "\n".join(shape.text for shape in page.shapes if shape.has_text_frame)
    assert "商品名称" not in text and "虚构单品名称" in text
    assert "市场参考价：199.50" in text and "折扣率：85%" in text
    assert "京东大客户价：159.25" in text and "None" not in text
    assert "¥" not in text and "￥" not in text
    picture = next(shape for shape in page.shapes if shape.name == "SCM_SLOT_IMAGE_MAIN")
    assert sha256(picture.image.blob).hexdigest() == sha256(_png((10, 20, 30))).hexdigest()
    assert picture.crop_left == picture.crop_right == pytest.approx(
        (1 - (2.2 / 1.5) / 2) / 2, abs=0.00001
    )


def test_v5_missing_values_leave_labels_and_empty_images(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"] = {"id": "fictional-1"}  # type: ignore[index]
    source["single_products"][0]["prices"] = {}
    source["single_products"][0]["manual"] = {}  # type: ignore[index]
    deck = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images={},
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    )
    text = "\n".join(shape.text for shape in deck.slides[1].shapes if shape.has_text_frame)
    assert "商品规格：" in text and "SKU：" in text and "None" not in text
    picture = next(shape for shape in deck.slides[1].shapes if shape.name == "SCM_SLOT_IMAGE_MAIN")
    assert Image.open(BytesIO(picture.image.blob)).size == (1, 1)


def test_v5_uses_only_approved_fields_and_four_images_keep_slot_order(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    item = source["single_products"][0]  # type: ignore[index]
    item["product"].update(  # type: ignore[union-attr]
        model="NEVER-RENDER-MODEL",
        internal_json='{"never":"render"}',
    )
    item["renderer_supplementary_image_ids"] = ["image-1", "image-2", "image-3"]  # type: ignore[index]
    images = {
        "fictional-1": _png((200, 10, 10)),
        "image-1": _png((10, 200, 10)),
        "image-2": _png((10, 10, 200)),
        "image-3": _png((200, 200, 10)),
    }
    page = Presentation(
        BytesIO(
            render_jd_detail_red_v5(
                source,
                images=images,
                template=get_ppt_template("JD_DETAIL_RED", "5"),
                manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
            )
        )
    ).slides[1]
    text = "\n".join(shape.text for shape in page.shapes if shape.has_text_frame)
    assert "NEVER-RENDER-MODEL" not in text and '"never":"render"' not in text
    for name, identifier in zip(
        (
            "SCM_SLOT_IMAGE_MAIN",
            "SCM_SLOT_IMAGE_SUPPLEMENT_01",
            "SCM_SLOT_IMAGE_SUPPLEMENT_02",
            "SCM_SLOT_IMAGE_SUPPLEMENT_03",
        ),
        images,
        strict=True,
    ):
        picture = next(shape for shape in page.shapes if shape.name == name)
        assert picture.image.blob == images[identifier]


def test_v5_rejects_unsafe_long_value_before_writing_partial_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    source = _source()
    source["single_products"][0]["product"]["product_name"] = "超长虚构商品名称" * 80  # type: ignore[index]
    with pytest.raises(AppError, match="页面安全区域") as error:
        render_jd_detail_red_v5(
            source,
            images={},
            template=get_ppt_template("JD_DETAIL_RED", "5"),
            manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
        )
    assert error.value.code == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"


def test_v5_uses_explicit_combination_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runtime = _runtime(tmp_path)
    monkeypatch.setattr(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    )
    with pytest.raises(AppError, match="仅支持单品页"):
        render_jd_detail_red_v5(
            {"packages": [{"items": []}]},
            images={},
            template=get_ppt_template("JD_DETAIL_RED", "5"),
            manifest=get_pptx_manifest("JD_DETAIL_RED", "5"),
        )
