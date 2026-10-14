from base64 import b64decode
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from app.common.contracts import AppError
from app.modules.recommendation.application import (
    ppt_renderer,
    ppt_renderer_v4,
    ppt_template_registry,
)
from app.modules.recommendation.application.ppt_renderer import PptRenderer
from app.modules.recommendation.application.ppt_template_manifest import (
    PPTX_SLOT_MANIFESTS,
    PPTX_VERSIONED_MANIFESTS,
    SLOT_BODY,
    SLOT_IMAGE_PREFIX,
    SLOT_PRICE,
    SLOT_TITLE,
)
from app.modules.recommendation.application.ppt_template_registry import (
    get_ppt_template,
    list_ppt_templates,
)

_PNG = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9J7VIAAAAASUVORK5CYII="
)


def _add_static_shape(
    slide: object, *, left: float, top: float, width: float, height: float
) -> None:
    shape = slide.shapes.add_shape(  # type: ignore[attr-defined]
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(230, 230, 230)
    shape.line.fill.background()


def _add_text_slot(
    slide: object, name: str, *, left: float, top: float, width: float, height: float
) -> None:
    slot = slide.shapes.add_textbox(  # type: ignore[attr-defined]
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    slot.name = name
    slot.text = "虚构模板槽位"


def _write_business_template_fixture(path: Path, template_code: str) -> None:
    """Create the minimum fictional clean deck required by the runtime slot contract."""
    manifest = PPTX_SLOT_MANIFESTS[template_code]
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)

    if manifest.cover_slide_index is not None:
        cover = deck.slides.add_slide(deck.slide_layouts[6])
        _add_static_shape(cover, left=0, top=0, width=13.333333, height=7.5)
        _add_text_slot(cover, "SCM_FIXTURE_COVER", left=1, top=3, width=11, height=0.8)

    product = deck.slides.add_slide(deck.slide_layouts[6])
    _add_static_shape(product, left=0, top=0, width=13.333333, height=0.2)
    _add_text_slot(product, SLOT_TITLE, left=6, top=0.5, width=6, height=0.7)
    _add_text_slot(product, SLOT_BODY, left=6, top=1.4, width=6, height=3.8)
    if SLOT_PRICE in manifest.source_text_slots:
        _add_text_slot(product, SLOT_PRICE, left=6, top=5.4, width=6, height=1.1)
    if template_code == "JD_FESTIVE_RED":
        # A fictional static background image must survive cloning while product
        # images are replaced in the named dynamic slots.
        product.shapes.add_picture(BytesIO(_PNG), 0, 0, deck.slide_width, deck.slide_height)
    image_count = len(manifest.source_image_slots)
    for index in range(image_count):
        _add_text_slot(
            product,
            f"{SLOT_IMAGE_PREFIX}{index + 1}",
            left=0.6 + (index % 2) * 2.5,
            top=1.0 + (index // 2) * 2.2,
            width=2.2,
            height=2.0,
        )
    deck.save(path)


@pytest.fixture
def business_template_assets(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Path]:
    """Point the renderer's directly imported resolver at fictional tmp assets."""
    paths: dict[str, Path] = {}
    for template in list_ppt_templates():
        if template.code == "SYSTEM_DEFAULT":
            continue
        path = tmp_path / f"{template.code.lower()}-fixture.pptx"
        _write_business_template_fixture(path, template.code)
        paths[template.code] = path

    # ppt_renderer imports this function directly, so patching the registry alone
    # would leave the actual render call bound to the production resolver.
    monkeypatch.setattr(
        ppt_renderer,
        "resolve_ppt_template_asset",
        lambda template: paths[template.code],
    )
    return paths


def _source() -> dict[str, object]:
    return {
        "project": {
            "project_code": "PPT-20260930-001",
            "project_name": "工会中秋福利方案",
            "buyer_name": "示例工会",
            "requirement": "中秋福利，预算 500 元，适用于职工慰问，支持一件代发。",
        },
        "packages": [
            {
                "name": "500 元节日套装",
                "total_price": "488.00",
                "reason": "人工确认的节日慰问组合。",
                "items": [
                    {
                        "quantity": 2,
                        "unit_price": "88.00",
                        "line_total": "176.00",
                        "product_snapshot": {
                            "id": "package-product-1",
                            "product_name": "玻璃保鲜盒",
                            "brand": "示例品牌",
                            "model": "B-100",
                        },
                    },
                    {
                        "quantity": 1,
                        "unit_price": "312.00",
                        "line_total": "312.00",
                        "product_snapshot": {
                            "id": "package-product-2",
                            "product_name": "家用小家电",
                        },
                    },
                    {
                        "quantity": 1,
                        "unit_price": "0.00",
                        "line_total": "0.00",
                        "product_snapshot": {
                            "id": "package-product-3",
                            "product_name": "第三件虚构商品",
                        },
                    },
                ],
            }
        ],
        "single_products": [
            {
                "product": {
                    "id": "single-product",
                    "product_name": "不锈钢保温杯",
                    "brand": "示例品牌",
                    "model": "C-200",
                    "product_specification": "450ml，食品级不锈钢",
                    "selling_points": "虚构测试卖点",
                    "shipping_courier": "测试快递",
                    "warranty_period": "一年",
                },
                "prices": {
                    "market_price": "129.00",
                    "jd_price": "109.00",
                    "agreement_price": "99.00",
                },
                "manual": {
                    "campaign_price": "89.00",
                    "delivery_status": "一件代发",
                    "evidence": "人工确认可供货",
                },
            }
        ],
    }


def _all_text(deck: Presentation) -> str:
    return "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text")
    )


def _slide_text(slide: object) -> str:
    return "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")).replace(  # type: ignore[attr-defined]
        "\u00a0", " "
    )


def _distinct_png(color: tuple[int, int, int]) -> bytes:
    image = Image.new("RGB", (4, 4), color)
    content = BytesIO()
    image.save(content, format="PNG")
    return content.getvalue()


def _sized_png(color: tuple[int, int, int], width: int, height: int) -> bytes:
    image = Image.new("RGB", (width, height), color)
    content = BytesIO()
    image.save(content, format="PNG")
    return content.getvalue()


def _write_jd_v4_fixture(path: Path) -> None:
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text_slot(cover, "SCM_SLOT_COVER_TITLE", left=1, top=1, width=8, height=0.5)
    cover.shapes[-1].text = "{{project_title}}"
    product = deck.slides.add_slide(deck.slide_layouts[6])
    static = product.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.2))
    static.name = "SCM_STATIC_HEADER"
    names = PPTX_VERSIONED_MANIFESTS[("JD_DETAIL_RED", "4")].source_text_slots
    for index, name in enumerate(names):
        _add_text_slot(product, name, left=6, top=0.35 + index * 0.35, width=6, height=0.3)
    slots = {shape.name: shape for shape in product.shapes}
    slots["SCM_FIELD_TITLE"].text = "{{title}}"
    slots["SCM_FIELD_SPECIFICATION"].text = "规格：{{product_specification}}"
    slots["SCM_FIELD_PACKAGING"].text = "包装：{{packaging_list}}"
    slots["SCM_FIELD_SELLING_POINTS"].text = (
        "{{detail_heading}}\n{{selling_points:0}}\n{{selling_points:1}}\n"
        "{{selling_points:2}}\n{{selling_points:3}}\n{{selling_points:4}}\n{{selling_points:5}}"
    )
    slots["SCM_FIELD_DELIVERY"].text = "配送：{{delivery}}"
    slots["SCM_FIELD_BRAND_MODEL"].text = "品牌 / 型号：{{brand_model}}"
    slots["SCM_FIELD_REFERENCE"].text = (
        "市场参考价：{{market_price}}\nSKU：{{sku}}\n履约：{{fulfillment}}"
    )
    price_frame = slots["SCM_FIELD_PRICES"].text_frame
    price_frame.clear()
    price_rows = (
        ("京东大客户价：", "{{jd_customer_price}}"),
        ("协议价：", "{{agreement_price}}"),
        ("{{offer_label}}", "{{campaign_price}}"),
    )
    for index, (label, token) in enumerate(price_rows):
        paragraph = price_frame.paragraphs[0] if index == 0 else price_frame.add_paragraph()
        paragraph.add_run().text = label
        paragraph.add_run().text = token
    slots["SCM_FIELD_IMAGE_NOTICE"].text = "商品主图："
    slots["SCM_FIELD_SERVICE"].text = "质保：{{warranty}}\n备注：{{evidence}}"
    for index in range(4):
        product.shapes.add_picture(
            BytesIO(_PNG), Inches(0.5 + index), Inches(4.5), Inches(0.8), Inches(0.8)
        ).name = f"SCM_SLOT_IMAGE_{index + 1:02d}"
    deck.save(path)


def _package_source(item_count: int) -> tuple[dict[str, object], dict[str, bytes]]:
    colors = ((210, 80, 80), (80, 160, 100), (100, 110, 210), (200, 160, 80))
    items = []
    images: dict[str, bytes] = {}
    for index in range(1, item_count + 1):
        product_id = f"combination-product-{index}"
        items.append(
            {
                "quantity": 1,
                "product_snapshot": {
                    "id": product_id,
                    "product_name": f"完整商品名称-{index}",
                    "sku": f"SKU-{index:03d}",
                    "model": f"MODEL-{index:03d}",
                },
            }
        )
        images[product_id] = _distinct_png(colors[index - 1])
    return (
        {
            "project": {"project_name": "虚构组合测试"},
            "single_products": [],
            "packages": [
                {
                    "name": f"{item_count} 件虚构组合",
                    "total_price": "399.00",
                    "reason": "虚构组合说明",
                    "items": items,
                }
            ],
        },
        images,
    )


def test_renderer_creates_editable_default_pptx(tmp_path: Path) -> None:
    output = tmp_path / "default.pptx"
    output.write_bytes(PptRenderer().render(_source(), product_images={"single-product": _PNG}))

    deck = Presentation(output)
    assert len(deck.slides) == 2
    text = _all_text(deck)
    assert "商品名称：500 元节日套装" in text
    assert "500 元节日套装" in text
    assert "不锈钢保温杯" in text
    assert any(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes)


def test_renderer_preserves_customer_template_and_appends_editable_pages(tmp_path: Path) -> None:
    template = Presentation()
    template.slides.add_slide(template.slide_layouts[6]).shapes.add_textbox(
        0, 0, 1000000, 1000000
    ).text = "甲方模板封面"
    template_path = tmp_path / "customer-template.pptx"
    template.save(template_path)

    deck = Presentation(BytesIO(PptRenderer().render(_source(), template_path=template_path)))
    assert len(deck.slides) == 3
    assert "甲方模板封面" in _all_text(deck)
    assert "商品名称：500 元节日套装" in _all_text(deck)


def test_each_builtin_template_creates_an_editable_valid_pptx_with_three_package_images(
    business_template_assets: dict[str, Path],
) -> None:
    renderer = PptRenderer()
    images = {
        "single-product": _PNG,
        "package-product-1": _PNG,
        "package-product-2": _PNG,
        "package-product-3": _PNG,
    }
    for template in list_ppt_templates():
        deck = Presentation(
            BytesIO(renderer.render(_source(), template_code=template.code, product_images=images))
        )
        expected_slide_count = 3 if template.code in {"JD_DETAIL_RED", "JD_FESTIVE_RED"} else 2
        assert len(deck.slides) == expected_slide_count
        text = _all_text(deck)
        assert "不锈钢保温杯" in text
        assert "89.00" in text
        assert "109.00" in text
        assert "虚构测试卖点" in text
        package_slide = deck.slides[-2]
        # Festive has one retained full-slide background image in addition to the
        # three dynamic package images; the other source templates do not.
        expected_pictures = 4 if template.code == "JD_FESTIVE_RED" else 3
        assert sum(shape.shape_type == 13 for shape in package_slide.shapes) == expected_pictures
        assert "500 元节日套装" in _slide_text(package_slide)
        assert "不锈钢保温杯" not in _slide_text(package_slide)
        for slide in deck.slides:
            for shape in slide.shapes:
                assert shape.left >= 0 and shape.top >= 0
                assert shape.left + shape.width <= deck.slide_width
                assert shape.top + shape.height <= deck.slide_height


def test_public_builtin_assets_render_without_private_asset_directory(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        ppt_template_registry,
        "get_settings",
        lambda: SimpleNamespace(ppt_template_asset_dir=None, app_env="production"),
    )
    for template in list_ppt_templates():
        if template.code == "SYSTEM_DEFAULT":
            continue
        deck = Presentation(BytesIO(PptRenderer().render(_source(), template_code=template.code)))
        assert "不锈钢保温杯" in _all_text(deck)


@pytest.mark.parametrize(
    ("code", "filename"),
    (
        ("JD_DETAIL_RED", "JD_DETAIL_RED_OFFICIAL_CLEAN_v5.pptx"),
        ("JD_FESTIVE_RED", "JD_FESTIVE_RED_OFFICIAL_CLEAN_v5.pptx"),
    ),
)
def test_public_v5_assets_resolve_without_explicit_directory(
    monkeypatch: pytest.MonkeyPatch, code: str, filename: str
) -> None:
    monkeypatch.setattr(
        ppt_template_registry,
        "get_settings",
        lambda: SimpleNamespace(ppt_template_asset_dir=None, app_env="production"),
    )

    asset = ppt_template_registry.resolve_ppt_template_asset(get_ppt_template(code, "5"))

    assert asset.name == filename
    assert asset.parent.name == "ppt-templates"
    assert asset.is_file()


@pytest.mark.parametrize("code", ("JD_DETAIL_RED", "JD_FESTIVE_RED"))
def test_public_v5_assets_render_without_layout_diagnostic_page(code: str) -> None:
    source = _source()
    source["packages"] = []  # type: ignore[index]

    renderer = PptRenderer(layout_preview=False)
    deck = Presentation(
        BytesIO(renderer.render(source, template_code=code, template_version="5"))
    )

    assert len(deck.slides) == 2
    assert "不锈钢保温杯" in _all_text(deck)
    assert "排版诊断版（仅用于排版调试）" not in _all_text(deck)
    assert renderer.layout_diagnostics_added is False


def test_long_details_use_continuation_pages_without_dropping_prices_or_name() -> None:
    source = _source()
    product = source["single_products"][0]["product"]  # type: ignore[index]
    product["product_name"] = "超长虚构商品名称" * 16  # type: ignore[index]
    product["product_specification"] = "长规格" * 220  # type: ignore[index]
    product["selling_points"] = "长卖点" * 220  # type: ignore[index]

    rendered = PptRenderer().render(source, product_images={"single-product": _PNG})
    deck = Presentation(BytesIO(rendered))
    text = _all_text(deck)
    assert "详情续页" in text
    assert "¥ 89.00" in text
    assert "¥ 109.00" in text
    assert "长规格" * 20 in text


def test_business_price_slot_avoids_body_price_and_missing_images_show_one_notice(
    business_template_assets: dict[str, Path],
) -> None:
    source = _source()
    source["packages"] = []  # type: ignore[index]
    for template in list_ppt_templates():
        if template.code == "SYSTEM_DEFAULT":
            continue
        deck = Presentation(BytesIO(PptRenderer().render(source, template_code=template.code)))
        text = _all_text(deck)
        assert text.count("协议价：¥ 99.00") == 1
        assert text.count("人工确认活动价：¥ 89.00") == 1
        assert text.count("商品图片待补充") == 1
        assert "快递：待确认" not in text
        assert "质保：待确认" not in text


def test_business_templates_retain_source_canvas_without_source_business_text(
    business_template_assets: dict[str, Path],
) -> None:
    renderer = PptRenderer()
    for template in list_ppt_templates():
        if template.code == "SYSTEM_DEFAULT":
            continue
        source_deck = Presentation(business_template_assets[template.code])
        source_strings = {
            shape.text.strip()
            for slide in source_deck.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        }
        rendered = renderer.render(_source(), template_code=template.code)
        output = Presentation(BytesIO(rendered))
        assert (output.slide_width, output.slide_height) == (
            source_deck.slide_width,
            source_deck.slide_height,
        )
        rendered_text = _all_text(output)
        assert not source_strings.intersection({rendered_text})
        assert template.code in PPTX_SLOT_MANIFESTS
        with ZipFile(BytesIO(rendered)) as archive:
            assert not any("notesSlides" in name for name in archive.namelist())


def test_real_business_asset_missing_still_fails_explicitly(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        ppt_template_registry,
        "get_settings",
        lambda: SimpleNamespace(ppt_template_asset_dir=tmp_path, app_env="production"),
    )

    with pytest.raises(AppError) as raised:
        ppt_template_registry.resolve_ppt_template_asset(get_ppt_template("JD_DETAIL_RED"))

    assert getattr(raised.value, "code", None) == "PPT_TEMPLATE_ASSET_UNAVAILABLE"


def test_real_business_asset_version_mismatch_is_rejected(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template = get_ppt_template("JD_DETAIL_RED")
    manifest = PPTX_SLOT_MANIFESTS[template.code]
    monkeypatch.setitem(
        PPTX_SLOT_MANIFESTS,
        template.code,
        replace(manifest, asset_version="fixture-version-mismatch"),
    )

    with pytest.raises(AppError) as raised:
        PptRenderer().render(
            _source(), template_code=template.code, template_version=template.version
        )

    assert getattr(raised.value, "code", None) == "PPT_TEMPLATE_ASSET_VERSION_MISMATCH"


def test_invalid_builtin_template_is_rejected() -> None:
    try:
        PptRenderer().render(_source(), template_code="../../not-a-template")
    except Exception as exc:
        assert getattr(exc, "code", None) == "PPT_TEMPLATE_INVALID"
    else:
        raise AssertionError("an unknown template code must be rejected")


def test_renderer_rejects_a_task_with_a_frozen_unavailable_template_version() -> None:
    try:
        PptRenderer().render(_source(), template_code="SYSTEM_DEFAULT", template_version="2")
    except Exception as exc:
        assert getattr(exc, "code", None) == "PPT_TEMPLATE_VERSION_UNAVAILABLE"
    else:
        raise AssertionError("an unavailable frozen version must be rejected")


def test_jd_v4_uses_tmp_fixture_and_preserves_named_static_shape(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """v4 is real rendering, but CI never needs a private customer deck."""
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text_slot(cover, "SCM_FIELD_PROJECT", left=1, top=1, width=8, height=0.5)
    cover.shapes[-1].text = "项目：{{project_name}}"
    product = deck.slides.add_slide(deck.slide_layouts[6])
    static = product.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.2))
    static.name = "SCM_STATIC_HEADER"
    for index, name in enumerate(
        PPTX_VERSIONED_MANIFESTS[("JD_DETAIL_RED", "4")].source_text_slots
    ):
        _add_text_slot(product, name, left=6, top=0.3 + index * 0.35, width=6, height=0.3)
        product.shapes[-1].text = "{{" + name.removeprefix("SCM_FIELD_").lower() + "}}"
    for index in range(4):
        product.shapes.add_picture(
            BytesIO(_PNG), Inches(0.5 + index), Inches(1), Inches(0.8), Inches(0.8)
        ).name = f"SCM_SLOT_IMAGE_{index + 1:02d}"
    slots = {shape.name: shape for shape in product.shapes}
    slots["SCM_FIELD_IMAGE_NOTICE"].text = "组套后或内外包装图片："
    slots["SCM_FIELD_SELLING_POINTS"].text = "产品卖点（介绍详细）\n{{selling_points:0}}"
    slots["SCM_FIELD_REFERENCE"].text = "履约：{{fulfillment}}"
    slots["SCM_FIELD_SERVICE"].text = "备注：{{evidence}}"
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4.pptx"
    deck.save(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)

    source = _source()
    source["packages"] = []  # type: ignore[index]
    source["single_products"][0]["manual"] = {}  # type: ignore[index]
    output = Presentation(
        BytesIO(
            PptRenderer().render(source, template_code="JD_DETAIL_RED", template_version="4")
        )
    )
    assert len(output.slides) == 2
    shapes = {shape.name: shape for shape in output.slides[-1].shapes}
    assert shapes["SCM_STATIC_HEADER"].left == 0
    assert "不锈钢保温杯" in _slide_text(output.slides[-1])
    assert "商品图片待补充" not in _slide_text(output.slides[-1])
    assert "商品主图：" in _slide_text(output.slides[-1])
    assert "产品卖点（介绍详细）" in _slide_text(output.slides[-1])
    assert "履约：" not in _slide_text(output.slides[-1])
    assert "备注：" not in _slide_text(output.slides[-1])


def test_jd_v4_keeps_sku_out_of_the_brand_model_field() -> None:
    item = _source()["single_products"][0]  # type: ignore[index]
    item["product"]["sku"] = "SKU-200"  # type: ignore[index]
    values = ppt_renderer_v4._fields(item)

    assert values["brand_model"] == "示例品牌 / C-200"
    assert values["sku"] == "SKU-200"


def test_jd_v4_runtime_contract_requires_the_final_text_fix_candidate() -> None:
    manifest = PPTX_VERSIONED_MANIFESTS[("JD_DETAIL_RED", "4")]

    assert manifest.clean_filename == "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"


@pytest.mark.parametrize("item_count", (2, 3, 4))
def test_jd_v4_renders_each_frozen_combination_item_once_without_name_continuations(
    item_count: int, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """JD v4 has a full package page, unlike the old single-products-only path."""
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    fixture = Presentation(asset)
    price_token = next(
        shape for shape in fixture.slides[-1].shapes if shape.name == "SCM_FIELD_PRICES"
    ).text_frame.paragraphs[0].runs[1]
    price_token.font.size = Pt(18)
    price_token.font.color.rgb = RGBColor(255, 0, 0)
    fixture.save(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source, images = _package_source(item_count)

    output = Presentation(
        BytesIO(
            PptRenderer().render(
                source,
                product_images=images,
                template_code="JD_DETAIL_RED",
                template_version="4",
            )
        )
    )

    assert len(output.slides) == 2
    page = output.slides[-1]
    text = _slide_text(page)
    assert "套装组成" in text
    assert "套装组成（续）" not in text
    for index in range(1, item_count + 1):
        assert f"完整商品名称-{index}" in text
        assert f"SKU-{index:03d}" in text
        assert f"MODEL-{index:03d}" in text

    image_slots = sorted(
        (shape for shape in page.shapes if shape.name.startswith("SCM_SLOT_IMAGE_")),
        key=lambda shape: shape.name,
    )
    assert len(image_slots) == 4
    rendered_hashes = [sha256(shape.image.blob).hexdigest() for shape in image_slots[:item_count]]
    assert len(set(rendered_hashes)) == item_count


def test_jd_v4_paginates_long_detail_without_growing_text_boxes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source = _source()
    source["packages"] = []
    item = source["single_products"][0]  # type: ignore[index]
    product = item["product"]  # type: ignore[index]
    # 28 full-width units: deliberately near, but inside, the audited
    # single-line title-bar capacity.
    title = "超长标题" * 7
    specification = "超长规格字段" * 14
    sku = "SKU-" + "X" * 48
    product.update(
        {
            "product_name": title,
            "product_specification": specification,
            "sku": sku,
            "selling_points": "；".join(
                f"超长卖点 {index} " + "说明" * 22 for index in range(1, 9)
            ),
        }
    )

    output = Presentation(
        BytesIO(
            PptRenderer().render(
                source,
                product_images={"single-product": _sized_png((90, 120, 160), 160, 80)},
                template_code="JD_DETAIL_RED",
                template_version="4",
            )
        )
    )

    assert len(output.slides) > 2
    rendered_text = _all_text(output)
    assert title in rendered_text
    text_without_line_breaks = rendered_text.replace("\n", "")
    assert specification in text_without_line_breaks
    assert "SKU-" in rendered_text
    assert rendered_text.count("X") >= len(sku) - len("SKU-")
    assert "产品详情（续页 2）" in rendered_text
    for slide in list(output.slides)[1:]:
        for shape in slide.shapes:
            if shape.name.startswith("SCM_FIELD_") and shape.has_text_frame:
                assert shape.text_frame.auto_size == MSO_AUTO_SIZE.NONE


def test_jd_v4_rejects_a_title_that_cannot_fit_safely(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source = _source()
    source["packages"] = []
    source["single_products"][0]["product"]["product_name"] = "超长标题" * 20  # type: ignore[index]

    with pytest.raises(AppError) as raised:
        PptRenderer().render(source, template_code="JD_DETAIL_RED", template_version="4")

    assert getattr(raised.value, "code", None) == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"


def test_jd_v4_rejects_a_two_line_title_instead_of_shrinking_its_template_font(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source = _source()
    source["packages"] = []
    source["single_products"][0]["product"]["product_name"] = "标题容量边界" * 6  # type: ignore[index]

    with pytest.raises(AppError, match="商品名称超出") as raised:
        PptRenderer().render(source, template_code="JD_DETAIL_RED", template_version="4")

    assert getattr(raised.value, "code", None) == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"


def test_jd_v4_moves_long_sku_and_brand_rows_to_a_non_overlapping_continuation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source = _source()
    source["packages"] = []
    item = source["single_products"][0]  # type: ignore[index]
    item["product"].update(  # type: ignore[index]
        {
            "sku": "SKU-" + "X" * 48,
            "brand": "超长虚构品牌名称" * 5,
            "model": "MODEL-" + "Y" * 40,
        }
    )
    item["manual"]["fulfillment_cycle"] = "预计 30 个工作日内分批交付并完成签收回传"  # type: ignore[index]

    output = Presentation(
        BytesIO(PptRenderer().render(source, template_code="JD_DETAIL_RED", template_version="4"))
    )
    first = output.slides[1]
    continuation = output.slides[2]
    reference = next(shape for shape in first.shapes if shape.name == "SCM_FIELD_REFERENCE")
    brand_model = next(shape for shape in first.shapes if shape.name == "SCM_FIELD_BRAND_MODEL")
    assert "SKU-" not in reference.text
    assert "超长虚构品牌" not in brand_model.text
    assert "SKU-" in _slide_text(continuation)
    assert "超长虚构品牌" in _slide_text(continuation)
    # The short reference area and lower brand row never grow into one another.
    assert reference.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert brand_model.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert (
        reference.top + reference.height <= brand_model.top
        or brand_model.top + brand_model.height <= reference.top
    )


def test_jd_v4_moves_reference_rows_before_an_ascii_sku_can_wrap(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Protect the real 2.53in reference box, including its static labels."""
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source = _source()
    source["packages"] = []
    item = source["single_products"][0]  # type: ignore[index]
    sku = "FICT-BOUNDARY-LONG-01"
    brand = "虚构超长品牌名称用于参考字段验证"
    item["product"].update({"sku": sku, "brand": brand, "model": "MODEL-REFERENCE-01"})  # type: ignore[index]
    item["manual"]["fulfillment_cycle"] = "预计三十个工作日内分批交付并签收回传"  # type: ignore[index]

    output = Presentation(
        BytesIO(PptRenderer().render(source, template_code="JD_DETAIL_RED", template_version="4"))
    )
    first = output.slides[1]
    continuation_text = _slide_text(output.slides[2])
    reference = next(shape for shape in first.shapes if shape.name == "SCM_FIELD_REFERENCE")
    brand_model = next(shape for shape in first.shapes if shape.name == "SCM_FIELD_BRAND_MODEL")
    assert sku not in reference.text
    assert brand not in brand_model.text
    assert sku in continuation_text
    assert brand in continuation_text
    assert reference.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert brand_model.text_frame.auto_size == MSO_AUTO_SIZE.NONE


def test_jd_v4_keeps_all_four_package_images_on_first_page_when_details_continue(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source, images = _package_source(4)
    for item in source["packages"][0]["items"]:  # type: ignore[index]
        product = item["product_snapshot"]
        product["product_name"] = "超长虚构组合商品名称" * 2
        product["sku"] = "SKU-" + "Q" * 25
        product["model"] = "MODEL-" + "M" * 25

    output = Presentation(
        BytesIO(
            PptRenderer().render(
                source, product_images=images, template_code="JD_DETAIL_RED", template_version="4"
            )
        )
    )
    assert len(output.slides) == 3
    first = output.slides[1]
    second = output.slides[2]
    first_slots = sorted(
        (shape for shape in first.shapes if shape.name.startswith("SCM_SLOT_IMAGE_")),
        key=lambda shape: shape.name,
    )
    second_slots = sorted(
        (shape for shape in second.shapes if shape.name.startswith("SCM_SLOT_IMAGE_")),
        key=lambda shape: shape.name,
    )
    assert len({sha256(shape.image.blob).hexdigest() for shape in first_slots}) == 4
    assert len({(shape.width, shape.height) for shape in first_slots}) == 1
    assert all(Image.open(BytesIO(shape.image.blob)).size == (1, 1) for shape in second_slots)
    assert "续页，第 4–4 项" in _slide_text(second)


def test_jd_v4_crops_landscape_and_portrait_images_without_stretching(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source = _source()
    source["packages"] = []

    def main_image(content: bytes) -> object:
        output = Presentation(
            BytesIO(
                PptRenderer().render(
                    source,
                    product_images={"single-product": content},
                    template_code="JD_DETAIL_RED",
                    template_version="4",
                )
            )
        )
        return next(
            shape
            for shape in output.slides[-1].shapes
            if shape.name == "SCM_SLOT_IMAGE_01"
        )

    landscape = main_image(_sized_png((200, 80, 80), 160, 80))
    portrait = main_image(_sized_png((80, 150, 100), 80, 160))

    assert landscape.crop_left == pytest.approx(0.25)
    assert landscape.crop_right == pytest.approx(0.25)
    assert landscape.crop_top == landscape.crop_bottom == 0
    assert portrait.crop_top == pytest.approx(0.25)
    assert portrait.crop_bottom == pytest.approx(0.25)
    assert portrait.crop_left == portrait.crop_right == 0


def test_jd_v4_preserves_distinct_price_meanings_and_package_total_label(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    fixture = Presentation(asset)
    price_token = next(
        shape for shape in fixture.slides[-1].shapes if shape.name == "SCM_FIELD_PRICES"
    ).text_frame.paragraphs[0].runs[1]
    price_token.font.size = Pt(18)
    price_token.font.color.rgb = RGBColor(255, 0, 0)
    fixture.save(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source = _source()
    source["packages"] = []
    source["single_products"][0]["manual"]["jd_customer_price"] = "95.00"  # type: ignore[index]
    output = Presentation(
        BytesIO(PptRenderer().render(source, template_code="JD_DETAIL_RED", template_version="4"))
    )
    single_text = _slide_text(output.slides[-1])
    assert "京东大客户价：¥ 95.00" in single_text
    assert "协议价：¥ 99.00" in single_text
    assert "人工活动价：¥ 89.00" in single_text
    rendered_price_token = next(
        shape for shape in output.slides[-1].shapes if shape.name == "SCM_FIELD_PRICES"
    ).text_frame.paragraphs[0].runs[1]
    assert rendered_price_token.font.size == Pt(18)
    assert rendered_price_token.font.color.rgb == RGBColor(255, 0, 0)

    package_source, images = _package_source(2)
    package_output = Presentation(
        BytesIO(
            PptRenderer().render(
                package_source,
                product_images=images,
                template_code="JD_DETAIL_RED",
                template_version="4",
            )
        )
    )
    package_text = _slide_text(package_output.slides[-1])
    assert "套装总价：¥ 399.00" in package_text
    package_price_shape = next(
        shape for shape in package_output.slides[-1].shapes if shape.name == "SCM_FIELD_PRICES"
    )
    package_total = next(
        paragraph.runs[-1]
        for paragraph in package_price_shape.text_frame.paragraphs
        if "套装总价" in paragraph.text
    )
    assert package_total.font.size == Pt(18)
    assert package_total.font.color.rgb == RGBColor(255, 0, 0)
    assert package_total.text == "¥\u00a0399.00"
    assert "\n" not in package_total.text


def test_jd_v4_package_grid_crops_mixed_aspect_images_without_stretching(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    asset = tmp_path / "JD_DETAIL_RED_CLEAN_v4-text-fix.pptx"
    _write_jd_v4_fixture(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)
    source, images = _package_source(4)
    for index, product_id in enumerate(images, start=1):
        images[product_id] = _sized_png(
            (40 * index, 60 * index, 80), 160 if index % 2 else 80, 80 if index % 2 else 160
        )

    output = Presentation(
        BytesIO(
            PptRenderer().render(
                source,
                product_images=images,
                template_code="JD_DETAIL_RED",
                template_version="4",
            )
        )
    )
    slots = sorted(
        (shape for shape in output.slides[-1].shapes if shape.name.startswith("SCM_SLOT_IMAGE_")),
        key=lambda shape: shape.name,
    )

    assert len({(shape.width, shape.height) for shape in slots}) == 1
    assert slots[0].crop_left == slots[0].crop_right == pytest.approx(0.25)
    assert slots[1].crop_top == slots[1].crop_bottom == pytest.approx(0.25)
    assert "组合商品图（全套 4 件）：" in _slide_text(output.slides[-1])


def test_festive_v4_preserves_static_labels_and_original_token_run_format(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The public fixture proves v4 replaces values, not original typography."""
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text_slot(cover, "SCM_SLOT_COVER_TITLE", left=1, top=1, width=8, height=0.5)
    cover.shapes[-1].text = "{{project_name}}"
    product = deck.slides.add_slide(deck.slide_layouts[6])
    names = PPTX_VERSIONED_MANIFESTS[("JD_FESTIVE_RED", "4")].source_text_slots
    for index, name in enumerate(names):
        _add_text_slot(product, name, left=6, top=0.5 + index * 0.4, width=6, height=0.35)
        shape = product.shapes[-1]
        shape.text_frame.paragraphs[0].add_run().font.size = Pt(20)
        shape.text_frame.paragraphs[0].runs[-1].font.bold = True
        shape.text_frame.paragraphs[0].runs[-1].font.color.rgb = RGBColor(255, 255, 255)
    shapes = {shape.name: shape for shape in product.shapes}
    shapes["SCM_FIELD_SELLING_POINTS"].text = (
        "产品卖点（介绍详细）\n{{selling_points:0}}\n{{selling_points:1}}\n{{selling_points:2}}"
    )
    shapes["SCM_FIELD_REFERENCE"].text = (
        "市场参考价：{{market_reference_value}} 京东大客户价：{{jd_customer_price_value}}\n"
        "货期：{{fulfillment}} 发货物流：{{delivery}}"
    )
    shapes["SCM_FIELD_TITLE"].text = "{{title}}"
    shapes["SCM_STATIC_PROJECT_NUMBER"].text = "方案编号"
    shapes["SCM_STATIC_IMAGE_CAPTION"].text = "组套后或内外包装图片："
    title_run = shapes["SCM_FIELD_TITLE"].text_frame.paragraphs[0].runs[0]
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(255, 255, 255)
    for index in range(5):
        product.shapes.add_picture(
            BytesIO(_PNG), Inches(0.5 + index), Inches(1.5), Inches(0.8), Inches(0.8)
        ).name = f"SCM_SLOT_IMAGE_{index + 1:02d}"
    asset = tmp_path / "JD_FESTIVE_RED_CLEAN_v4.pptx"
    deck.save(asset)
    monkeypatch.setattr(ppt_renderer_v4, "resolve_ppt_template_asset", lambda _template: asset)

    source = _source()
    source["single_products"][0]["manual"] = {"delivery_status": "虚构配送"}  # type: ignore[index]
    output = Presentation(
        BytesIO(PptRenderer().render(source, template_code="JD_FESTIVE_RED", template_version="4"))
    )
    rendered = {shape.name: shape for shape in output.slides[-1].shapes}
    assert "方案编号" in _slide_text(output.slides[-1])
    assert "组套后或内外包装图片：" in _slide_text(output.slides[-1])
    assert "京东大客户价：¥" not in _slide_text(output.slides[-1])
    rendered_title = rendered["SCM_FIELD_TITLE"].text_frame.paragraphs[0].runs[0]
    assert rendered_title.font.size == Pt(20)
    assert rendered_title.font.bold is True
    assert rendered_title.font.color.rgb == RGBColor(255, 255, 255)
