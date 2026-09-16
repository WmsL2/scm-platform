# ruff: noqa: E501, E701, E702
import uuid
from decimal import Decimal
from io import BytesIO

from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, event, select

from app.common.contracts import AppError
from app.core.database import SessionLocal
from app.main import app
from app.modules.auth.security import create_token, hash_password
from app.modules.catalog.application.category_service import CategoryService
from app.modules.catalog.infrastructure.models import Category
from app.modules.system.models import Permission, Role, RolePermission, User, UserRole

HEADERS = ("一级类目ID", "一级类目名称", "二级类目ID", "二级类目名称", "三级类目ID", "三级类目名称", "有效标记", "上下柜标记", "主营事业部")
CATEGORY_PERMISSIONS = (
    "category:list",
    "category:detail",
    "category:create",
    "category:update",
    "category:delete",
)

async def auth(permission_codes: tuple[str, ...] = CATEGORY_PERMISSIONS + ("product:list", "product:import")) -> tuple[uuid.UUID, dict[str, str]]:
    uid, rid = uuid.uuid4(), uuid.uuid4()
    async with SessionLocal() as s:
        ps = list((await s.scalars(select(Permission).where(Permission.permission_code.in_(permission_codes)))).all())
        assert {permission.permission_code for permission in ps} == set(permission_codes)
        s.add_all([User(id=uid, username=f"cat-{uid}", password_hash=hash_password("test")), Role(id=rid, role_code=f"cat-{rid}", role_name="cat"), UserRole(user_id=uid, role_id=rid), *[RolePermission(role_id=rid, permission_id=p.id) for p in ps]])
        await s.commit()
    return uid, {"Authorization": f"Bearer {create_token(uid, 1)}"}

def row(ident: object, active: object = 1) -> dict[str, object]: return dict(zip(HEADERS, ("1", "一级", "2", "二级", ident, f"三级{ident}", active, "", ""), strict=True))
def book(rows: list[dict[str, object]], headers: tuple[str, ...] = HEADERS) -> bytes:
    wb = Workbook(); ws = wb.active; assert ws; ws.append(headers)
    for item in rows: ws.append([item.get(h) for h in headers])
    out = BytesIO(); wb.save(out); return out.getvalue()
async def post(c: AsyncClient, h: dict[str, str], content: bytes, rate: str = "8"):
    return await c.post("/api/v1/categories/imports", headers=h, data={"deduction_rate_percent": rate}, files={"file": ("categories.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
async def clean(uid: uuid.UUID) -> None:
    async with SessionLocal() as s:
        await s.execute(delete(Category).where(Category.level3_external_id.like("cat-test-%")))
        rids = list((await s.scalars(select(UserRole.role_id).where(UserRole.user_id == uid))).all()); await s.execute(delete(RolePermission).where(RolePermission.role_id.in_(rids))); await s.execute(delete(UserRole).where(UserRole.user_id == uid)); await s.execute(delete(Role).where(Role.id.in_(rids))); await s.execute(delete(User).where(User.id == uid)); await s.commit()

async def test_category_import_chinese_template_and_rates() -> None:
    uid, h = await auth()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            template = await c.get("/api/v1/categories/imports/template", headers=h); ws = load_workbook(BytesIO(template.content)).active; assert ws and tuple(x.value for x in ws[1]) == HEADERS and ws.freeze_panes == "A2"
            for suffix, rate, expected in (("1", "5", "0.0500"), ("2", "8", "0.0800"), ("3", "7.5", "0.0750"), ("4", "6.25", "0.0625")):
                assert (await post(c, h, book([row("cat-test-" + suffix)]), rate)).status_code == 200
                async with SessionLocal() as s: assert (await s.scalar(select(Category.deduction_rate).where(Category.level3_external_id == "cat-test-" + suffix))) == Decimal(expected)
            assert (await post(c, h, book([row("cat-test-5", "是"), row("cat-test-6", "否")]))).status_code == 200
            for rate in ("6.255", "-1", "100.01", "101", "abc", "NaN", "Infinity"):
                assert (await post(c, h, book([row("cat-test-bad" + rate.replace(".", ""))]), rate)).status_code == 422
    finally: await clean(uid)

async def test_category_import_headers_identity_skip_and_ids() -> None:
    uid, h = await auth()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            assert (await post(c, h, book([row("cat-test-00123")], tuple(reversed(HEADERS))))).status_code == 200
            duplicate = await post(c, h, book([row("cat-test-dup"), row("cat-test-dup")]))
            assert duplicate.json()["data"]["errors"][0]["field"] == "level3_external_id"
            assert (await post(c, h, book([row(123.45)]))).json()["data"]["failed"] == 1
            assert (await post(c, h, book([row("cat-test-skip")]))).status_code == 200
            assert (await post(c, h, book([row("cat-test-skip")]))).json()["data"]["skipped"] == 1
            assert (await post(c, h, book([row("cat-test-skip")]), "5")).json()["data"]["failed"] == 1
            for bad in (HEADERS[:-1], HEADERS + ("测试",), HEADERS[:-1] + ("三级类目ID",)):
                assert (await post(c, h, book([row("cat-test-header")], bad))).status_code == 422
    finally: await clean(uid)

async def test_category_import_write_failure_rolls_back_all_rows() -> None:
    uid, _ = await auth()
    try:
        async with SessionLocal() as session:
            injected = False
            def inject(sync: object, *_: object) -> None:
                nonlocal injected
                if not injected:
                    injected = True
                    sync.add(Category(source_type="MALL_LEVEL3", level1_name="一级", level2_name="二级", level3_name="冲突", level3_external_id="cat-test-r2", deduction_rate=Decimal("0.0800")))  # type: ignore[attr-defined]
            event.listen(session.sync_session, "before_flush", inject)
            try:
                try: await CategoryService(session).import_xlsx("rollback.xlsx", book([row("cat-test-r1"), row("cat-test-r2")]), "8", uid)
                except AppError as exc: assert exc.code == "CATEGORY_IMPORT_CONFLICT"
            finally: event.remove(session.sync_session, "before_flush", inject)
        async with SessionLocal() as session:
            remaining = list((await session.scalars(select(Category.id).where(Category.level3_external_id.in_(("cat-test-r1", "cat-test-r2"))))).all())
            assert remaining == []
    finally: await clean(uid)


async def test_category_list_server_side_filters_pagination_and_selection() -> None:
    uid, h = await auth()
    try:
        categories = [
            Category(source_type="MALL_LEVEL3", level1_name="筛测个人护理", level2_name="筛测假发", level3_name="筛测配件甲", level3_external_id="cat-test-filter-a", deduction_rate=Decimal("0.0500"), is_active=True, business_unit="筛测京东零售-大商超事业部"),
            Category(source_type="MALL_LEVEL3", level1_name="筛测个人防护", level2_name="筛测假发", level3_name="筛测配件乙", level3_external_id="cat-test-filter-b", deduction_rate=Decimal("0.0800"), is_active=False, business_unit="筛测京东零售-家电事业部"),
            Category(source_type="MALL_LEVEL3", level1_name="筛测办公", level2_name="筛测收纳", level3_name="筛测桌面收纳", level3_external_id="cat-test-filter-c", deduction_rate=Decimal("0.0500"), is_active=True, business_unit="筛测企业服务"),
            *[Category(source_type="MALL_LEVEL3", level1_name="分页测试", level2_name="二级", level3_name=f"三级{i:02}", level3_external_id=f"cat-test-filter-page-{i}", deduction_rate=Decimal("0.0500"), is_active=True, business_unit="分页事业部") for i in range(21)],
        ]
        async with SessionLocal() as s:
            s.add_all(categories)
            await s.commit()
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            async def listed(query: str) -> dict[str, object]:
                response = await c.get(f"/api/v1/categories?{query}", headers=h)
                assert response.status_code == 200
                return response.json()["data"]

            assert (await listed("level1_name=%20%E7%AD%9B%E6%B5%8B%E4%B8%AA%E4%BA%BA%20"))["total"] == 2
            assert (await listed("level2_name=%E7%AD%9B%E6%B5%8B%E5%81%87%E5%8F%91"))["total"] == 2
            assert (await listed("level3_name=%E7%AD%9B%E6%B5%8B%E9%85%8D%E4%BB%B6"))["total"] == 2
            assert (await listed("level1_name=%E7%AD%9B%E6%B5%8B&deduction_rate=0.05"))["total"] == 2
            assert (await listed("level1_name=%E7%AD%9B%E6%B5%8B&is_active=true"))["total"] == 2
            inactive = await listed("level1_name=%E7%AD%9B%E6%B5%8B&is_active=false")
            assert inactive["total"] == 1 and inactive["items"][0]["level3_name"] == "筛测配件乙"  # type: ignore[index]
            assert (await listed("business_unit=%E7%AD%9B%E6%B5%8B%E4%BA%AC%E4%B8%9C%E9%9B%B6%E5%94%AE"))["total"] == 2
            combined = await listed("level1_name=%E7%AD%9B%E6%B5%8B%E4%B8%AA%E4%BA%BA&level2_name=%E7%AD%9B%E6%B5%8B%E5%81%87%E5%8F%91&level3_name=%E7%94%B2&deduction_rate=0.05&is_active=true&business_unit=%E5%A4%A7%E5%95%86%E8%B6%85")
            assert combined["total"] == 1 and combined["items"][0]["level3_name"] == "筛测配件甲"  # type: ignore[index]
            first = await listed("page=1&page_size=20&level1_name=%E5%88%86%E9%A1%B5%E6%B5%8B%E8%AF%95")
            second = await listed("page=2&page_size=20&level1_name=%E5%88%86%E9%A1%B5%E6%B5%8B%E8%AF%95")
            first_ids = {item["id"] for item in first["items"]}  # type: ignore[index]
            second_ids = {item["id"] for item in second["items"]}  # type: ignore[index]
            assert first["total"] == 21 and len(first_ids) == 20 and len(second_ids) == 1 and first_ids.isdisjoint(second_ids)
            selection = (await c.get("/api/v1/categories/selection", headers=h)).json()["data"]
            selected_ids = {item["id"] for item in selection}
            assert str(categories[0].id) in selected_ids and str(categories[1].id) not in selected_ids
    finally:
        await clean(uid)


async def test_category_crud_permissions_are_independent() -> None:
    actors: dict[str, tuple[uuid.UUID, dict[str, str]]] = {}
    for permission_code in CATEGORY_PERMISSIONS + ("product:list",):
        actors[permission_code] = await auth((permission_code,))
    viewer_headers = actors["category:list"][1]
    payload = {
        "source_type": "MALL_LEVEL3",
        "level1_external_id": "cat-test-rbac-l1",
        "level1_name": "权限一级",
        "level2_external_id": "cat-test-rbac-l2",
        "level2_name": "权限二级",
        "level3_external_id": "cat-test-rbac-crud",
        "level3_name": "权限三级",
        "deduction_rate": "0.0500",
        "is_active": True,
        "shelf_flag": None,
        "business_unit": "权限测试",
    }
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post("/api/v1/categories", headers=actors["category:create"][1], json=payload)
            assert created.status_code == 200
            category_id = created.json()["data"]["id"]

            assert (await client.get("/api/v1/categories", headers=viewer_headers)).status_code == 200
            assert (await client.get(f"/api/v1/categories/{category_id}", headers=viewer_headers)).status_code == 403
            assert (await client.post("/api/v1/categories", headers=viewer_headers, json=payload)).status_code == 403
            assert (await client.put(f"/api/v1/categories/{category_id}", headers=viewer_headers, json=payload)).status_code == 403
            assert (await client.delete(f"/api/v1/categories/{category_id}", headers=viewer_headers)).status_code == 403
            assert (await client.get("/api/v1/categories/selection", headers=viewer_headers)).status_code == 403

            product_headers = actors["product:list"][1]
            assert (await client.get("/api/v1/categories", headers=product_headers)).status_code == 403
            assert (await client.get("/api/v1/categories/selection", headers=product_headers)).status_code == 200

            detail_headers = actors["category:detail"][1]
            assert (await client.get(f"/api/v1/categories/{category_id}", headers=detail_headers)).status_code == 200
            updated_payload = {**payload, "level3_name": "权限三级已更新"}
            updated = await client.put(
                f"/api/v1/categories/{category_id}",
                headers=actors["category:update"][1],
                json=updated_payload,
            )
            assert updated.status_code == 200
            assert updated.json()["data"]["level3_name"] == "权限三级已更新"
            deleted = await client.delete(
                f"/api/v1/categories/{category_id}",
                headers=actors["category:delete"][1],
            )
            assert deleted.status_code == 200
    finally:
        for uid, _ in actors.values():
            await clean(uid)


async def test_category_permissions_are_seeded_and_existing_boss_is_granted() -> None:
    async with SessionLocal() as session:
        permission_codes = set(
            (
                await session.scalars(
                    select(Permission.permission_code).where(
                        Permission.permission_code.in_(CATEGORY_PERMISSIONS)
                    )
                )
            ).all()
        )
        assert permission_codes == set(CATEGORY_PERMISSIONS)

        boss_id = await session.scalar(
            select(Role.id).where(Role.role_code == "boss", Role.is_deleted.is_(False))
        )
        if boss_id is None:
            return

        codes = set(
            (
                await session.scalars(
                    select(Permission.permission_code)
                    .join(RolePermission, RolePermission.permission_id == Permission.id)
                    .where(
                        RolePermission.role_id == boss_id,
                        Permission.permission_code.in_(CATEGORY_PERMISSIONS),
                    )
                )
            ).all()
        )
    assert codes == set(CATEGORY_PERMISSIONS)
