from openpyxl import load_workbook

from app.modules.catalog.api.router import PRODUCT_IMPORT_TEMPLATE_PATH
from app.modules.catalog.application.failed_rows_export import export_failed_rows_workbook
from app.modules.catalog.application.import_service import PRODUCT_IMPORT_HEADERS
from app.modules.catalog.infrastructure.models import ProductImportRow


def test_failed_rows_export_reuses_product_master_template_styles(tmp_path) -> None:
    source_data = {header: f"值-{index}" for index, header in enumerate(PRODUCT_IMPORT_HEADERS)}
    source_data["sku"] = "SKU-FAILED"
    source_data["供应商"] = "待修正供应商"
    rows = [
        ProductImportRow(
            source_row_number=420,
            source_data=source_data,
            calculated_data=source_data,
            supplier_name_raw="待修正供应商",
            is_valid=False,
            error_message="供应商未匹配",
        ),
        ProductImportRow(
            source_row_number=421,
            source_data={**source_data, "sku": "SKU-FAILED-2"},
            calculated_data=source_data,
            supplier_name_raw="待修正供应商",
            is_valid=False,
            error_message="供应商未匹配",
        ),
    ]
    output = tmp_path / "failed.xlsx"

    export_failed_rows_workbook(PRODUCT_IMPORT_HEADERS, rows, output)

    exported = load_workbook(output, data_only=False)
    template = load_workbook(PRODUCT_IMPORT_TEMPLATE_PATH, data_only=False)
    exported_sheet = exported["Sheet1"]
    template_sheet = template["Sheet1"]
    assert tuple(cell.value for cell in exported_sheet[1]) == PRODUCT_IMPORT_HEADERS
    assert exported_sheet.max_row == 3
    assert exported_sheet.row_dimensions[1].height == template_sheet.row_dimensions[1].height
    assert exported_sheet.row_dimensions[2].height == template_sheet.row_dimensions[2].height
    assert exported_sheet.row_dimensions[3].height == template_sheet.row_dimensions[2].height
    for column in range(1, len(PRODUCT_IMPORT_HEADERS) + 1):
        assert exported_sheet.cell(1, column)._style == template_sheet.cell(1, column)._style
        assert exported_sheet.cell(2, column)._style == template_sheet.cell(2, column)._style
        assert exported_sheet.cell(3, column)._style == template_sheet.cell(2, column)._style
        letter = exported_sheet.cell(1, column).column_letter
        assert exported_sheet.column_dimensions[letter].width == (
            template_sheet.column_dimensions[letter].width
        )
    assert exported_sheet["F2"].value == "SKU-FAILED"
    assert exported_sheet["F3"].value == "SKU-FAILED-2"
    assert exported["错误说明"]["B2"].value == 420
    assert exported["错误说明"]["B3"].value == 421
    template.close()
    exported.close()
