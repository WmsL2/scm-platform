import uuid
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi import UploadFile
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from PIL import Image
from sqlalchemy import delete, select

from app.core.database import SessionLocal
from app.main import app
from app.modules.auth.security import create_token, hash_password
from app.modules.catalog.api.router import (
    PRODUCT_IMPORT_UPLOAD_CHUNK_BYTES,
    _stage_product_import_upload,
)
from app.modules.catalog.application.import_service import PRODUCT_IMPORT_HEADERS
from app.modules.catalog.domain.lifecycle import ProductStatus
from app.modules.catalog.infrastructure.models import (
    Category,
    Product,
    ProductImportRow,
    ProductImportSupplierMatch,
    ProductImportTask,
)
from app.modules.supplier.infrastructure.models import Supplier
from app.modules.system.models import Permission, Role, RolePermission, User, UserRole

IMPORT_PERMISSIONS = ("product:import", "product:import:resolve")


async def test_product_import_upload_is_staged_in_chunks() -> None:
    class ChunkRecordingFile:
        def __init__(self, content: bytes) -> None:
            self.content = content
            self.offset = 0
            self.read_sizes: list[int] = []

        def read(self, size: int = -1) -> bytes:
            self.read_sizes.append(size)
            if self.offset >= len(self.content):
                return b""
            chunk = self.content[self.offset : self.offset + size]
            self.offset += len(chunk)
            return chunk

        def close(self) -> None:
            return None

    content = b"x" * (PRODUCT_IMPORT_UPLOAD_CHUNK_BYTES * 2 + 1)
    source = ChunkRecordingFile(content)
    file = UploadFile(file=source, filename="large-product-import.xlsx")
    staged_path, file_size = await _stage_product_import_upload(file)
    try:
        assert file_size == len(content)
        assert staged_path.read_bytes() == content
        assert source.read_sizes == [PRODUCT_IMPORT_UPLOAD_CHUNK_BYTES] * 4
    finally:
        staged_path.unlink(missing_ok=True)
        await file.close()


class _RecordingStorage:
    def __init__(self) -> None:
        self.deleted: list[str] = []

    async def save(self, name: str, content: bytes) -> str:
        del content
        return name

    async def read(self, key: str) -> bytes:
        del key
        raise FileNotFoundError

    async def delete(self, key: str) -> None:
        self.deleted.append(key)


async def _create_import_user() -> tuple[uuid.UUID, dict[str, str]]:
    user_id, role_id = uuid.uuid4(), uuid.uuid4()
    async with SessionLocal() as session:
        permissions = list(
            (
                await session.scalars(
                    select(Permission).where(Permission.permission_code.in_(IMPORT_PERMISSIONS))
                )
            ).all()
        )
        assert len(permissions) == len(IMPORT_PERMISSIONS)
        session.add_all(
            [
                User(
                    id=user_id,
                    username=f"product-import-test-{user_id}",
                    password_hash=hash_password("secret"),
                ),
                Role(id=role_id, role_code=f"product-import-{role_id}", role_name="product import"),
                UserRole(user_id=user_id, role_id=role_id),
                *[
                    RolePermission(role_id=role_id, permission_id=permission.id)
                    for permission in permissions
                ],
            ]
        )
        await session.commit()
    return user_id, {"Authorization": f"Bearer {create_token(user_id, 1)}"}


async def _cleanup_import_user(user_id: uuid.UUID) -> None:
    async with SessionLocal() as session:
        role_ids = list(
            (
                await session.scalars(
                    select(UserRole.role_id).where(UserRole.user_id == user_id)
                )
            ).all()
        )
        await session.execute(delete(RolePermission).where(RolePermission.role_id.in_(role_ids)))
        await session.execute(delete(UserRole).where(UserRole.user_id == user_id))
        await session.execute(delete(Role).where(Role.id.in_(role_ids)))
        await session.execute(delete(User).where(User.id == user_id))
        await session.commit()


async def _create_category(
    path: tuple[str, str, str], *, is_active: bool = True
) -> uuid.UUID:
    category_id = uuid.uuid4()
    async with SessionLocal() as session:
        session.add(
            Category(
                id=category_id,
                source_type="MALL_LEVEL3",
                level1_external_id=f"L1-{category_id}",
                level1_name=path[0],
                level2_external_id=f"L2-{category_id}",
                level2_name=path[1],
                level3_external_id=f"L3-{category_id}",
                level3_name=path[2],
                deduction_rate=Decimal("0.0800"),
                is_active=is_active,
            )
        )
        await session.commit()
    return category_id


async def _create_references() -> tuple[uuid.UUID, uuid.UUID]:
    supplier_id = uuid.uuid4()
    async with SessionLocal() as session:
        session.add(
            Supplier(
                id=supplier_id,
                supplier_code=f"IMP{str(supplier_id)[:8]}",
                supplier_name="导入测试供应商",
                main_brands="测试品牌",
                advantage="测试优势",
                archive_status="ARCHIVED",
                cooperation_status="NORMAL",
            )
        )
        await session.commit()
    category_id = await _create_category(("测试一级", "测试二级", "测试三级"))
    return supplier_id, category_id


def _workbook_bytes(
    *supplier_names: str,
    image_value: str = "https://example.test/image.png",
    category_path: tuple[str, str, str] = ("测试一级", "测试二级", "测试三级"),
    product_name: str = "测试商品",
    cost_price: str = "100",
    sku_override: str | None = None,
    value_overrides: dict[str, object] | None = None,
    number_formats: dict[str, str] | None = None,
    styled_blank_columns_after_template: bool = False,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(PRODUCT_IMPORT_HEADERS)
    for index, supplier_name in enumerate(supplier_names, start=1):
        values: dict[str, object] = {
            "所属公司": "测试公司", "上架日期": "2026-09-10", "品牌": "测试品牌",
            "图片": image_value, "型号": "型号", "sku": sku_override or f"SKU-{index}",
            "商品名称": product_name, "一级类目": category_path[0], "二级类目": category_path[1],
            "三级类目": category_path[2], "货号": "货号", "链接": "https://example.test/item",
            "成本价": cost_price, "市场价": "999", "京东价": "201", "协议价": "199.9",
            "协议价采购价": "155.55", "利润": "55.55", "京东价毛利（30-50）": "12.34%",
            "扣点复核": "11.11%", "毛利率": "27.78%", "采销员": "采销员",
            "供应商": supplier_name, "69码": "6900000000000", "3c编码": "3C-TEST",
            "产品规格": "规格", "卖点": "卖点", "包装清单": "包装", "质保期": "一年",
            "限售区域": "限售区域", "京东自营前台价": "180", "自营旗舰店/官方旗舰店": "官方旗舰店",
            "参考链接": "https://example.test/ref", "销量": "10", "好评率": "95%",
            "折扣率": "88.88%", "价格虚高比例": "-11.11%", "税收编码": "TAX-1",
            "开票名称": "测试商品", "税收分类": "测试分类", "发货快递": "京东物流",
            "售后政策": "七天无理由", "备注": "备注",
        }
        values.update(value_overrides or {})
        worksheet.append([values[header] for header in PRODUCT_IMPORT_HEADERS])
        for header, number_format in (number_formats or {}).items():
            worksheet.cell(
                row=worksheet.max_row,
                column=PRODUCT_IMPORT_HEADERS.index(header) + 1,
            ).number_format = number_format
    if styled_blank_columns_after_template:
        worksheet.cell(row=1, column=16_384).fill = PatternFill(
            fill_type="solid", fgColor="FFFFFF"
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _with_cached_formula_results(
    content: bytes, results: dict[str, tuple[str, str]]
) -> bytes:
    source = BytesIO(content)
    output = BytesIO()
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(source) as source_archive, ZipFile(
        output, "w", compression=ZIP_DEFLATED
    ) as output_archive:
        for entry in source_archive.infolist():
            data = source_archive.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data)
                for cell in root.findall(f".//{{{namespace}}}c"):
                    result = results.get(cell.attrib.get("r", ""))
                    if result is None:
                        continue
                    cell_type, cached_value = result
                    cell.set("t", cell_type)
                    value_node = cell.find(f"{{{namespace}}}v")
                    if value_node is None:
                        value_node = ElementTree.SubElement(cell, f"{{{namespace}}}v")
                    value_node.text = cached_value
                data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            output_archive.writestr(entry, data)
    return output.getvalue()


def _with_wps_cell_image(content: bytes, *, image_id: str = "ID_PRODUCT") -> bytes:
    image_output = BytesIO()
    Image.new("RGB", (2, 2), color=(30, 120, 210)).save(image_output, format="PNG")
    source = BytesIO(content)
    output = BytesIO()
    with ZipFile(source) as source_archive, ZipFile(
        output, "w", compression=ZIP_DEFLATED
    ) as output_archive:
        for entry in source_archive.infolist():
            output_archive.writestr(entry, source_archive.read(entry.filename))
        output_archive.writestr(
            "xl/cellimages.xml",
            "<etc:cellImages "
            'xmlns:etc="http://www.wps.cn/officeDocument/2017/etCustomData" '
            'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            "<etc:cellImage><xdr:pic><xdr:nvPicPr>"
            f'<xdr:cNvPr name="{image_id}"/>'
            "</xdr:nvPicPr><xdr:blipFill><a:blip r:embed=\"rId1\"/>"
            "</xdr:blipFill></xdr:pic></etc:cellImage></etc:cellImages>",
        )
        output_archive.writestr(
            "xl/_rels/cellimages.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            'Target="media/image1.png"/></Relationships>',
        )
        output_archive.writestr("xl/media/image1.png", image_output.getvalue())
    return output.getvalue()


async def test_product_import_accepts_blank_styled_columns_after_approved_headers() -> None:
    user_id, headers = await _create_import_user()
    supplier_id, category_id = await _create_references()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "styled-blank-columns.xlsx",
                        _workbook_bytes(
                            "导入测试供应商", styled_blank_columns_after_template=True
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert preview.status_code == 200
            assert preview.json()["data"]["total_rows"] == 1
    finally:
        await _cleanup_import_data(supplier_id, user_id, (category_id,))
        await _cleanup_import_user(user_id)


async def _cleanup_import_data(
    supplier_id: uuid.UUID, user_id: uuid.UUID, category_ids: tuple[uuid.UUID, ...]
) -> None:
    async with SessionLocal() as session:
        task_ids = list(
            (
                await session.scalars(
                    select(ProductImportTask.id).where(ProductImportTask.created_by == user_id)
                )
            ).all()
        )
        await session.execute(delete(Product).where(Product.created_by == user_id))
        if task_ids:
            await session.execute(
                delete(ProductImportRow).where(
                    ProductImportRow.import_task_id.in_(task_ids)
                )
            )
            await session.execute(
                delete(ProductImportSupplierMatch).where(
                    ProductImportSupplierMatch.import_task_id.in_(task_ids)
                )
            )
            await session.execute(
                delete(ProductImportTask).where(ProductImportTask.id.in_(task_ids))
            )
        await session.execute(delete(Supplier).where(Supplier.id == supplier_id))
        await session.execute(delete(Category).where(Category.id.in_(category_ids)))
        await session.commit()


async def test_product_import_saves_required_category_text_without_category_binding(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    storage = _RecordingStorage()
    monkeypatch.setattr(
        "app.modules.catalog.application.import_service.get_object_storage", lambda: storage
    )
    user_id, headers = await _create_import_user()
    supplier_id, category_id = await _create_references()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            unresolved = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "products.xlsx",
                        _workbook_bytes("未知供应商"),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert unresolved.status_code == 200
            assert unresolved.json()["data"]["status"] == "NEEDS_RESOLUTION"
            assert unresolved.json()["data"]["invalid_rows"] == 1

            preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "products.xlsx",
                        _workbook_bytes("导入测试供应商"),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert preview.status_code == 200
            data = preview.json()["data"]
            assert data["status"] == "READY_TO_CONFIRM"
            assert data["valid_rows"] == 1
            assert data["rows"][0]["category_path"] == "测试一级 / 测试二级 / 测试三级"
            confirmed = await client.post(
                f"/api/v1/products/imports/{data['id']}/confirm", headers=headers
            )
            assert confirmed.status_code == 200
            assert confirmed.json()["data"]["imported_count"] == 1

            async with SessionLocal() as session:
                imported = await session.scalar(
                    select(Product).where(Product.created_by == user_id)
                )
                assert imported is not None
                imported.image_reference = "local-media/product-images/old.png"
                await session.commit()

            duplicate_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "duplicate-products.xlsx",
                        _workbook_bytes(
                            "导入测试供应商", product_name="更新后商品", cost_price="222"
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert duplicate_preview.status_code == 200
            duplicate_data = duplicate_preview.json()["data"]
            assert duplicate_data["status"] == "READY_TO_CONFIRM"
            assert duplicate_data["valid_rows"] == 0
            assert duplicate_data["update_rows"] == 1
            assert duplicate_data["invalid_rows"] == 0
            assert duplicate_data["rows"][0]["write_action"] == "UPDATE"
            assert {"商品名称", "成本价"}.issubset(duplicate_data["rows"][0]["changed_fields"])
            updated = await client.post(
                f"/api/v1/products/imports/{duplicate_data['id']}/confirm", headers=headers
            )
            assert updated.status_code == 200
            assert updated.json()["data"]["created_count"] == 0
            assert updated.json()["data"]["updated_count"] == 1
            assert "product-images/old.png" in storage.deleted
            updated_preview = await client.get(
                f"/api/v1/products/imports/{duplicate_data['id']}", headers=headers
            )
            assert updated_preview.status_code == 200
            assert updated_preview.json()["data"]["rows"][0]["is_imported"] is True
            assert updated_preview.json()["data"]["rows"][0]["write_action"] == "UPDATE"

            duplicate_rows_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "duplicate-rows-products.xlsx",
                        _workbook_bytes(
                            "导入测试供应商",
                            "导入测试供应商",
                            sku_override="SKU-1",
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert duplicate_rows_preview.status_code == 200
            duplicate_rows_data = duplicate_rows_preview.json()["data"]
            assert duplicate_rows_data["update_rows"] == 1
            assert duplicate_rows_data["invalid_rows"] == 1
            assert "与Excel第2行的来源供应商与SKU重复" in (
                duplicate_rows_data["rows"][1]["error_message"]
            )

        async with SessionLocal() as session:
            deleted_product = await session.scalar(
                select(Product).where(Product.created_by == user_id)
            )
            assert deleted_product is not None
            disabled_product_id = deleted_product.id
            assert deleted_product.product_name == "更新后商品"
            assert deleted_product.cost_price == Decimal("222.0000")
            deleted_product.product_name = "停用前商品名称"
            deleted_product.cost_price = Decimal("321.0000")
            deleted_product.status = ProductStatus.DISABLED
            deleted_product.disabled_by = user_id
            deleted_product.disabled_at = datetime.now()
            await session.commit()

        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            disabled_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "disabled-product.xlsx",
                        _workbook_bytes("导入测试供应商"),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert disabled_preview.status_code == 200
            disabled_data = disabled_preview.json()["data"]
            assert disabled_data["status"] == "VALIDATED"
            assert disabled_data["valid_rows"] == 0
            assert "商品已停用，请启用或永久删除后再导入" in (
                disabled_data["rows"][0]["error_message"]
            )
            blocked_confirm = await client.post(
                f"/api/v1/products/imports/{disabled_data['id']}/confirm", headers=headers
            )
            assert blocked_confirm.status_code == 409
            assert blocked_confirm.json()["code"] == "PRODUCT_IMPORT_NO_VALID_ROWS"

        async with SessionLocal() as session:
            product = await session.get(Product, disabled_product_id)
            assert product is not None
            assert product.status == ProductStatus.DISABLED
            assert product.product_name == "停用前商品名称"
            assert product.cost_price == Decimal("321.0000")
            assert product.source_supplier_id == supplier_id
            assert product.category_level3_name == "测试三级"
            assert product.market_price == Decimal("999.0000")
            assert product.agreement_price == Decimal("199.9000")
            assert str(product.jd_margin) == "0.1234"
            assert product.company_name == "测试公司"
            assert product.sales_volume == 10
            assert product.positive_rating == Decimal("0.9500")
    finally:
        await _cleanup_import_data(supplier_id, user_id, (category_id,))
        await _cleanup_import_user(user_id)


async def test_product_import_accepts_category_paths_without_category_master_binding() -> None:
    user_id, headers = await _create_import_user()
    supplier_id, default_category_id = await _create_references()
    ambiguous_path = ("重复一级", "重复二级", "重复三级")
    inactive_path = ("停用一级", "停用二级", "停用三级")
    ambiguous_category_ids = (
        await _create_category(ambiguous_path),
        await _create_category(ambiguous_path),
    )
    inactive_category_id = await _create_category(inactive_path, is_active=False)
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            missing = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "missing-category.xlsx",
                        _workbook_bytes(
                            "导入测试供应商",
                            category_path=("未知一级", "未知二级", "未知三级"),
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert missing.status_code == 200
            missing_row = missing.json()["data"]["rows"][0]
            assert missing_row["is_valid"] is True
            assert missing_row["error_message"] is None

            ambiguous = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "ambiguous-category.xlsx",
                        _workbook_bytes("导入测试供应商", category_path=ambiguous_path),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert ambiguous.status_code == 200
            ambiguous_row = ambiguous.json()["data"]["rows"][0]
            assert ambiguous_row["is_valid"] is True
            assert ambiguous_row["error_message"] is None

            inactive = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "inactive-category.xlsx",
                        _workbook_bytes("导入测试供应商", category_path=inactive_path),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert inactive.status_code == 200
            inactive_row = inactive.json()["data"]["rows"][0]
            assert inactive_row["is_valid"] is True
            assert inactive_row["error_message"] is None
    finally:
        await _cleanup_import_data(
            supplier_id,
            user_id,
            (default_category_id, *ambiguous_category_ids, inactive_category_id),
        )
        await _cleanup_import_user(user_id)


async def test_product_import_confirms_valid_rows_and_retains_failed_rows() -> None:
    user_id, headers = await _create_import_user()
    supplier_id, category_id = await _create_references()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "mixed-products.xlsx",
                        _workbook_bytes("导入测试供应商", "未知供应商"),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert preview.status_code == 200
            data = preview.json()["data"]
            assert data["status"] == "NEEDS_RESOLUTION"
            assert data["valid_rows"] == 1
            assert data["invalid_rows"] == 1
            assert data["imported_rows"] == 0

            partial = await client.post(
                f"/api/v1/products/imports/{data['id']}/confirm", headers=headers
            )
            assert partial.status_code == 200
            partial_data = partial.json()["data"]
            assert partial_data == {
                "id": data["id"],
                "status": "PARTIALLY_CONFIRMED",
                "imported_count": 1,
                "created_count": 1,
                "updated_count": 0,
                "imported_rows": 1,
                "valid_rows": 0,
                "update_rows": 0,
                "invalid_rows": 1,
            }

            refreshed = await client.get(f"/api/v1/products/imports/{data['id']}", headers=headers)
            assert refreshed.status_code == 200
            refreshed_data = refreshed.json()["data"]
            assert refreshed_data["rows"][0]["is_imported"] is True
            assert refreshed_data["rows"][1]["is_imported"] is False

            unknown_match = next(
                match
                for match in refreshed_data["supplier_matches"]
                if match["supplier_name_normalized"] == "未知供应商"
            )
            resolved = await client.post(
                f"/api/v1/products/imports/{data['id']}/supplier-matches/{unknown_match['id']}/resolve",
                headers=headers,
                json={"supplier_id": str(supplier_id)},
            )
            assert resolved.status_code == 200
            assert resolved.json()["data"]["status"] == "PARTIALLY_CONFIRMED"
            assert resolved.json()["data"]["valid_rows"] == 1

            confirmed = await client.post(
                f"/api/v1/products/imports/{data['id']}/confirm", headers=headers
            )
            assert confirmed.status_code == 200
            assert confirmed.json()["data"]["status"] == "CONFIRMED"
            assert confirmed.json()["data"]["imported_rows"] == 2

            template = await client.get("/api/v1/products/imports/template", headers=headers)
            assert template.status_code == 200
            assert template.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert template.content.startswith(b"PK")

        async with SessionLocal() as session:
            products = list(
                (await session.scalars(select(Product).where(Product.created_by == user_id))).all()
            )
            assert sorted(product.sku for product in products) == ["SKU-1", "SKU-2"]
    finally:
        await _cleanup_import_data(supplier_id, user_id, (category_id,))
        await _cleanup_import_user(user_id)


async def test_product_import_defers_formula_image_storage_until_confirm() -> None:
    user_id, headers = await _create_import_user()
    supplier_id, category_id = await _create_references()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "formula-image-products.xlsx",
                        _with_wps_cell_image(
                            _workbook_bytes(
                                "导入测试供应商",
                                image_value='=_xlfn.DISPIMG("ID_PRODUCT",1)',
                            )
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert preview.status_code == 200
            data = preview.json()["data"]
            assert data["rows"][0]["image_saved"] is False
            assert data["rows"][0]["image_pending_save"] is True

            async with SessionLocal() as session:
                task = await session.get(ProductImportTask, data["id"])
                row = await session.scalar(
                    select(ProductImportRow).where(ProductImportRow.import_task_id == data["id"])
                )
                assert task is not None and task.source_file_storage_key is not None
                assert row is not None and row.image_storage_key is None

            confirmed = await client.post(
                f"/api/v1/products/imports/{data['id']}/confirm", headers=headers
            )
            assert confirmed.status_code == 200
            assert confirmed.json()["data"]["status"] == "CONFIRMED"

            async with SessionLocal() as session:
                task = await session.get(ProductImportTask, data["id"])
                assert task is not None and task.source_file_storage_key is None
    finally:
        await _cleanup_import_data(supplier_id, user_id, (category_id,))
        await _cleanup_import_user(user_id)


async def test_product_import_rejects_formula_when_embedded_image_is_missing() -> None:
    user_id, headers = await _create_import_user()
    supplier_id, category_id = await _create_references()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "missing-formula-image.xlsx",
                        _workbook_bytes(
                            "导入测试供应商",
                            image_value='=_xlfn.DISPIMG("ID_MISSING",1)',
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )

            assert preview.status_code == 200
            data = preview.json()["data"]
            assert data["invalid_rows"] == 1
            assert data["rows"][0]["image_pending_save"] is True
            assert "图片无法读取" in data["rows"][0]["error_message"]
            confirmed = await client.post(
                f"/api/v1/products/imports/{data['id']}/confirm", headers=headers
            )
            assert confirmed.status_code == 409
            assert confirmed.json()["code"] == "PRODUCT_IMPORT_NO_VALID_ROWS"
    finally:
        await _cleanup_import_data(supplier_id, user_id, (category_id,))
        await _cleanup_import_user(user_id)


async def test_product_import_normalizes_rates_and_rejects_invalid_numeric_text() -> None:
    user_id, headers = await _create_import_user()
    supplier_id, category_id = await _create_references()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            valid_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "normalized-rates.xlsx",
                        _workbook_bytes(
                            "导入测试供应商",
                            value_overrides={
                                "利润": None,
                                "京东价毛利（30-50）": "74.32%",
                                "毛利率": 0.2778,
                                "折扣率": None,
                                "价格虚高比例": "46.25%",
                            },
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert valid_preview.status_code == 200
            valid_data = valid_preview.json()["data"]
            assert valid_data["valid_rows"] == 1

            async with SessionLocal() as session:
                staged_row = await session.scalar(
                    select(ProductImportRow).where(
                        ProductImportRow.import_task_id == valid_data["id"]
                    )
                )
                assert staged_row is not None
                assert staged_row.normalized_data is not None
                assert staged_row.normalized_data["price_inflation_rate"] == "0.4625"
                assert staged_row.normalized_data["jd_margin"] == "0.7432"
                assert staged_row.normalized_data["gross_margin"] == "0.2778"
                assert staged_row.normalized_data["discount_rate"] is None
                assert staged_row.normalized_data["profit"] is None

            confirmed = await client.post(
                f"/api/v1/products/imports/{valid_data['id']}/confirm", headers=headers
            )
            assert confirmed.status_code == 200

            display_workbook = _workbook_bytes(
                "导入测试供应商",
                sku_override="SKU-DISPLAY-PRECISION",
                value_overrides={
                    "利润": "=Q2-M2",
                    "京东价毛利（30-50）": "=(O2-Q2)/O2",
                    "毛利率": "=(Q2-M2)/Q2",
                    "折扣率": "=P2/O2",
                    "价格虚高比例": "=(P2-AE2)/AE2",
                },
                number_formats={
                    "利润": "0.00",
                    "京东价毛利（30-50）": "0.00%",
                    "毛利率": "0.00%",
                    "折扣率": "0%",
                    "价格虚高比例": "0%",
                },
            )
            display_workbook = _with_cached_formula_results(
                display_workbook,
                {
                    "R2": ("n", "4.40000000000001"),
                    "S2": ("n", "0.290393013100437"),
                    "U2": ("n", "0.160646714826281"),
                    "AJ2": ("n", "0.748908296943231"),
                    "AK2": ("n", "-0.138190954773869"),
                },
            )
            display_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "formula-display-precision.xlsx",
                        display_workbook,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert display_preview.status_code == 200
            display_data = display_preview.json()["data"]
            assert display_data["valid_rows"] == 1

            async with SessionLocal() as session:
                display_row = await session.scalar(
                    select(ProductImportRow).where(
                        ProductImportRow.import_task_id == display_data["id"]
                    )
                )
                assert display_row is not None
                assert display_row.normalized_data is not None
                assert display_row.normalized_data["profit"] == "4.40"
                assert display_row.normalized_data["jd_margin"] == "0.2904"
                assert display_row.normalized_data["gross_margin"] == "0.1606"
                assert display_row.normalized_data["discount_rate"] == "0.7500"
                assert display_row.normalized_data["price_inflation_rate"] == "-0.1400"

            div_zero_workbook = _workbook_bytes(
                "导入测试供应商",
                sku_override="SKU-DIV-ZERO",
                value_overrides={"毛利率": "=(Q2-M2)/Q2"},
                number_formats={"毛利率": "0.00%"},
            )
            div_zero_workbook = _with_cached_formula_results(
                div_zero_workbook, {"U2": ("e", "#DIV/0!")}
            )
            div_zero_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "formula-div-zero.xlsx",
                        div_zero_workbook,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert div_zero_preview.status_code == 200
            div_zero_row = div_zero_preview.json()["data"]["rows"][0]
            assert div_zero_row["is_valid"] is False
            assert "毛利率必须是数字" in div_zero_row["error_message"]

            price_precision_workbook = _workbook_bytes(
                "导入测试供应商",
                sku_override="SKU-PRICE-PRECISION",
                value_overrides={"成本价": "=1+1"},
                number_formats={"成本价": "0.00"},
            )
            price_precision_workbook = _with_cached_formula_results(
                price_precision_workbook, {"M2": ("n", "100.12345")}
            )
            price_precision_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "formula-price-precision.xlsx",
                        price_precision_workbook,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert price_precision_preview.status_code == 200
            price_precision_data = price_precision_preview.json()["data"]
            assert price_precision_data["valid_rows"] == 1
            async with SessionLocal() as session:
                price_precision_row = await session.scalar(
                    select(ProductImportRow).where(
                        ProductImportRow.import_task_id == price_precision_data["id"]
                    )
                )
                assert price_precision_row is not None
                assert price_precision_row.normalized_data is not None
                assert price_precision_row.normalized_data["cost_price"] == "100.12345"

            price_precision_confirm = await client.post(
                f"/api/v1/products/imports/{price_precision_data['id']}/confirm",
                headers=headers,
            )
            assert price_precision_confirm.status_code == 200

            async with SessionLocal() as session:
                exact_price_product = await session.scalar(
                    select(Product).where(Product.sku == "SKU-PRICE-PRECISION")
                )
                assert exact_price_product is not None
                assert exact_price_product.cost_price == Decimal("100.12345")

            maximum_price = "100.123456789012345678901234567891"
            maximum_precision_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "maximum-price-precision.xlsx",
                        _workbook_bytes(
                            "导入测试供应商",
                            sku_override="SKU-MAX-PRICE-PRECISION",
                            cost_price=maximum_price,
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert maximum_precision_preview.status_code == 200
            maximum_precision_data = maximum_precision_preview.json()["data"]
            assert maximum_precision_data["valid_rows"] == 1
            maximum_precision_confirm = await client.post(
                f"/api/v1/products/imports/{maximum_precision_data['id']}/confirm",
                headers=headers,
            )
            assert maximum_precision_confirm.status_code == 200

            async with SessionLocal() as session:
                maximum_precision_product = await session.scalar(
                    select(Product).where(Product.sku == "SKU-MAX-PRICE-PRECISION")
                )
                assert maximum_precision_product is not None
                assert maximum_precision_product.cost_price == Decimal(maximum_price)

            invalid_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "invalid-numbers.xlsx",
                        _workbook_bytes(
                            "导入测试供应商",
                            sku_override="SKU-INVALID",
                            value_overrides={"利润": "46.25%", "价格虚高比例": "5000+"},
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            invalid_row = invalid_preview.json()["data"]["rows"][0]
            assert invalid_row["is_valid"] is False
            assert "利润必须是数字" in invalid_row["error_message"]
            assert "价格虚高比例必须是数字" in invalid_row["error_message"]

            formula_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "formula-without-cache.xlsx",
                        _workbook_bytes(
                            "导入测试供应商",
                            sku_override="SKU-FORMULA",
                            value_overrides={"成本价": "=1+1"},
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            formula_row = formula_preview.json()["data"]["rows"][0]
            assert formula_row["is_valid"] is False
            assert "成本价公式没有可用计算结果" in formula_row["error_message"]

        async with SessionLocal() as session:
            product = await session.scalar(select(Product).where(Product.sku == "SKU-1"))
            assert product is not None
            assert product.price_inflation_rate == Decimal("0.4625")
            assert product.jd_margin == Decimal("0.7432")
            assert product.gross_margin == Decimal("0.2778")
            assert product.discount_rate is None
            assert product.profit is None
    finally:
        await _cleanup_import_data(supplier_id, user_id, (category_id,))
        await _cleanup_import_user(user_id)


async def test_product_import_preview_rows_are_server_paginated() -> None:
    user_id, headers = await _create_import_user()
    supplier_id, category_id = await _create_references()
    try:
        workbook = _workbook_bytes(*(["导入测试供应商"] * 105))
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "paginated-products.xlsx",
                        workbook,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert preview.status_code == 200
            data = preview.json()["data"]
            assert data["row_total"] == 105
            assert data["page"] == 1
            assert data["page_size"] == 50
            assert len(data["rows"]) == 50

            page_three = await client.get(
                f"/api/v1/products/imports/{data['id']}",
                headers=headers,
                params={"page": 3, "page_size": 50, "row_status": "PASSED"},
            )
            assert page_three.status_code == 200
            page_data = page_three.json()["data"]
            assert page_data["row_total"] == 105
            assert page_data["page"] == 3
            assert len(page_data["rows"]) == 5
            assert page_data["rows"][0]["source_row_number"] == 102
    finally:
        await _cleanup_import_data(supplier_id, user_id, (category_id,))
        await _cleanup_import_user(user_id)


async def test_product_import_confirm_rejects_stale_concurrent_preview() -> None:
    user_id, headers = await _create_import_user()
    supplier_id, category_id = await _create_references()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            preview_ids: list[str] = []
            for filename in ("concurrent-a.xlsx", "concurrent-b.xlsx"):
                preview = await client.post(
                    "/api/v1/products/imports/preview",
                    headers=headers,
                    files={
                        "file": (
                            filename,
                            _workbook_bytes("导入测试供应商"),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        )
                    },
                )
                assert preview.status_code == 200
                preview_ids.append(preview.json()["data"]["id"])

            first_confirm = await client.post(
                f"/api/v1/products/imports/{preview_ids[0]}/confirm", headers=headers
            )
            assert first_confirm.status_code == 200
            stale_create = await client.post(
                f"/api/v1/products/imports/{preview_ids[1]}/confirm", headers=headers
            )
            assert stale_create.status_code == 409
            assert stale_create.json()["code"] == "PRODUCT_IMPORT_STALE_PREVIEW"

            update_preview = await client.post(
                "/api/v1/products/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "stale-update.xlsx",
                        _workbook_bytes(
                            "导入测试供应商", product_name="来自过期预览的名称"
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            update_id = update_preview.json()["data"]["id"]

            async with SessionLocal() as session:
                product = await session.scalar(
                    select(Product).where(Product.created_by == user_id)
                )
                assert product is not None
                product.product_name = "其他并发操作已更新"
                product.updated_at = datetime.now() + timedelta(seconds=5)
                await session.commit()

            stale_update = await client.post(
                f"/api/v1/products/imports/{update_id}/confirm", headers=headers
            )
            assert stale_update.status_code == 409
            assert stale_update.json()["code"] == "PRODUCT_IMPORT_STALE_PREVIEW"

        async with SessionLocal() as session:
            products = list(
                (await session.scalars(select(Product).where(Product.created_by == user_id))).all()
            )
            assert len(products) == 1
            assert products[0].product_name == "其他并发操作已更新"
    finally:
        await _cleanup_import_data(supplier_id, user_id, (category_id,))
        await _cleanup_import_user(user_id)
