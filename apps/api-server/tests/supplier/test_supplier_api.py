import uuid
from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, func, select

from app.core.database import SessionLocal
from app.main import app
from app.modules.auth.security import create_token, hash_password
from app.modules.supplier.application.import_service import SupplierImportService
from app.modules.supplier.application.service import SupplierService
from app.modules.supplier.infrastructure.models import (
    Supplier,
    SupplierContact,
    SupplierCooperationRecord,
    SupplierImportBatch,
    SupplierImportRow,
    SupplierQualification,
)
from app.modules.supplier.schemas import SupplierCreateRequest, SupplierUpdateRequest
from app.modules.system.models import Permission, Role, RolePermission, User, UserRole

SUPPLIER_PERMISSIONS = (
    "supplier:list",
    "supplier:detail",
    "supplier:create",
    "supplier:update",
    "supplier:submit",
    "supplier:archive",
    "supplier:stop",
    "supplier:blacklist",
    "supplier:resume",
    "supplier:unblacklist",
    "supplier:delete",
)


async def create_user_with_permissions(
    permission_codes: tuple[str, ...]
) -> tuple[uuid.UUID, dict[str, str]]:
    user_id, role_id = uuid.uuid4(), uuid.uuid4()
    async with SessionLocal() as session:
        permissions = list(
            (
                await session.scalars(
                    select(Permission).where(Permission.permission_code.in_(permission_codes))
                )
            ).all()
        )
        assert len(permissions) == len(permission_codes)
        session.add_all(
            [
                User(
                    id=user_id,
                    username=f"supplier-test-{user_id}",
                    password_hash=hash_password("secret"),
                ),
                Role(id=role_id, role_code=f"supplier-test-{role_id}", role_name="supplier test"),
                UserRole(user_id=user_id, role_id=role_id),
                *[
                    RolePermission(role_id=role_id, permission_id=permission.id)
                    for permission in permissions
                ],
            ]
        )
        await session.commit()
    return user_id, {"Authorization": f"Bearer {create_token(user_id, 1)}"}


async def cleanup_user(user_id: uuid.UUID) -> None:
    async with SessionLocal() as session:
        role_ids = list(
            (
                await session.scalars(select(UserRole.role_id).where(UserRole.user_id == user_id))
            ).all()
        )
        await session.execute(delete(RolePermission).where(RolePermission.role_id.in_(role_ids)))
        await session.execute(delete(UserRole).where(UserRole.user_id == user_id))
        await session.execute(delete(Role).where(Role.id.in_(role_ids)))
        await session.execute(delete(User).where(User.id == user_id))
        await session.commit()


async def cleanup_suppliers(supplier_ids: list[str]) -> None:
    if not supplier_ids:
        return
    async with SessionLocal() as session:
        await session.execute(
            delete(SupplierContact).where(SupplierContact.supplier_id.in_(supplier_ids))
        )
        await session.execute(
            delete(SupplierQualification).where(SupplierQualification.supplier_id.in_(supplier_ids))
        )
        await session.execute(
            delete(SupplierCooperationRecord).where(
                SupplierCooperationRecord.supplier_id.in_(supplier_ids)
            )
        )
        await session.execute(delete(Supplier).where(Supplier.id.in_(supplier_ids)))
        await session.commit()


async def cleanup_import_batches(batch_ids: list[str]) -> None:
    if not batch_ids:
        return
    async with SessionLocal() as session:
        await session.execute(
            delete(SupplierImportRow).where(SupplierImportRow.batch_id.in_(batch_ids))
        )
        await session.execute(
            delete(SupplierImportBatch).where(SupplierImportBatch.id.in_(batch_ids))
        )
        await session.commit()


def workbook_bytes(rows: list[tuple[object, object, object, object, object]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["供应商名称", "主营品牌", "主要优势", "联系人", "联系电话"])
    for row in rows:
        worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def test_supplier_api_enforces_permissions_and_lifecycle() -> None:
    user_id, headers = await create_user_with_permissions(SUPPLIER_PERMISSIONS)
    supplier_ids: list[str] = []
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            unauthorized = await client.get("/api/v1/suppliers")
            assert unauthorized.status_code == 401

            create = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={
                    "supplier_name": "众诚测试供应商",
                    "main_brands": "品牌 A、品牌 B",
                    "advantage": "本地测试数据",
                    "contacts": [{"contact_name": "李四", "contact_phone": "13800000000"}],
                },
            )
            assert create.status_code == 201
            supplier = create.json()["data"]
            supplier_id = supplier["id"]
            supplier_ids.append(supplier_id)
            assert supplier["supplier_code"].startswith("SUP")
            assert supplier["archive_status"] == "DRAFT"
            assert supplier["cooperation_status"] == "NORMAL"
            assert supplier["contacts"][0]["contact_name"] == "李四"

            selected_status_create = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={
                    "supplier_name": "初始归档供应商",
                    "main_brands": "品牌",
                    "advantage": "优势",
                    "archive_status": "ARCHIVED",
                },
            )
            assert selected_status_create.status_code == 201
            selected_supplier = selected_status_create.json()["data"]
            supplier_ids.append(selected_supplier["id"])
            assert selected_supplier["archive_status"] == "ARCHIVED"
            assert selected_supplier["archived_by"] == str(user_id)
            assert selected_supplier["archived_at"] is not None

            supplied_code = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={
                    "supplier_code": "SUP00000001",
                    "supplier_name": "不能指定编码",
                    "main_brands": "品牌",
                    "advantage": "优势",
                },
            )
            assert supplied_code.status_code == 422

            update = await client.patch(
                f"/api/v1/suppliers/{supplier_id}",
                headers=headers,
                json={"supplier_name": "更新后的供应商", "contacts": []},
            )
            assert update.status_code == 200
            assert update.json()["data"]["supplier_name"] == "更新后的供应商"
            assert update.json()["data"]["contacts"] == []

            status_edit = await client.patch(
                f"/api/v1/suppliers/{supplier_id}",
                headers=headers,
                json={"archive_status": "PENDING"},
            )
            assert status_edit.status_code == 200
            assert status_edit.json()["data"]["archive_status"] == "PENDING"
            assert status_edit.json()["data"]["archived_by"] is None
            assert status_edit.json()["data"]["archived_at"] is None

            archive = await client.post(
                f"/api/v1/suppliers/{supplier_id}/commands/archive", headers=headers
            )
            assert archive.status_code == 200
            assert archive.json()["data"]["archive_status"] == "ARCHIVED"
            assert archive.json()["data"]["archived_by"] == str(user_id)
            assert archive.json()["data"]["archived_at"] is not None

            draft_edit = await client.patch(
                f"/api/v1/suppliers/{supplier_id}",
                headers=headers,
                json={"archive_status": "DRAFT"},
            )
            assert draft_edit.status_code == 200
            assert draft_edit.json()["data"]["archive_status"] == "DRAFT"
            assert draft_edit.json()["data"]["archived_by"] is None
            assert draft_edit.json()["data"]["archived_at"] is None

            stop = await client.post(
                f"/api/v1/suppliers/{supplier_id}/commands/stop",
                headers=headers,
                json={"reason": "合作终止"},
            )
            assert stop.status_code == 200
            assert stop.json()["data"]["cooperation_status"] == "STOPPED"

            blacklist_after_stop = await client.post(
                f"/api/v1/suppliers/{supplier_id}/commands/blacklist",
                headers=headers,
                json={"reason": "不能跳过状态机"},
            )
            assert blacklist_after_stop.status_code == 409

            listing = await client.get("/api/v1/suppliers", headers=headers)
            assert listing.status_code == 200
            assert any(item["id"] == supplier_id for item in listing.json()["data"]["items"])

        async with SessionLocal() as session:
            record = await session.scalar(
                select(SupplierCooperationRecord).where(
                    SupplierCooperationRecord.supplier_id == supplier_id
                )
            )
            assert record is not None
            assert record.from_status == "NORMAL"
            assert record.to_status == "STOPPED"
            assert record.reason == "合作终止"
            assert record.actor_id == user_id
    finally:
        await cleanup_suppliers(supplier_ids)
        await cleanup_user(user_id)


async def test_supplier_api_returns_403_for_authenticated_user_without_supplier_permissions(
) -> None:
    user_id, headers = await create_user_with_permissions(())
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.get("/api/v1/suppliers", headers=headers)
            assert response.status_code == 403
            assert response.json()["code"] == "AUTH_FORBIDDEN"
    finally:
        await cleanup_user(user_id)


async def test_supplier_cooperation_recovery_preserves_history_and_eligibility() -> None:
    user_id, headers = await create_user_with_permissions(SUPPLIER_PERMISSIONS)
    limited_user, limited_headers = await create_user_with_permissions(("supplier:list",))
    supplier_ids: list[str] = []
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={"supplier_name": "恢复合作测试", "main_brands": "品牌", "advantage": "优势"},
            )
            assert created.status_code == 201, created.text
            supplier_id = created.json()["data"]["id"]
            supplier_ids.append(supplier_id)
            resume_from_normal = await client.post(
                f"/api/v1/suppliers/{supplier_id}/commands/resume",
                headers=headers,
                json={"reason": "不允许正常状态恢复"},
            )
            assert resume_from_normal.status_code == 409, resume_from_normal.text
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/unblacklist",
                    headers=headers,
                    json={"reason": "不允许正常状态移出黑名单"},
                )
            ).status_code == 409
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/stop",
                    headers=headers,
                    json={"reason": "   "},
                )
            ).status_code == 422
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/stop",
                    headers=headers,
                    json={"reason": "停止合作原因"},
                )
            ).status_code == 200
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/unblacklist",
                    headers=headers,
                    json={"reason": "状态不匹配"},
                )
            ).status_code == 409
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/resume",
                    headers=limited_headers,
                    json={"reason": "没有恢复权限"},
                )
            ).status_code == 403
            resumed = await client.post(
                f"/api/v1/suppliers/{supplier_id}/commands/resume",
                headers=headers,
                json={"reason": "恢复合作原因"},
            )
            assert resumed.status_code == 200
            assert resumed.json()["data"]["cooperation_status"] == "NORMAL"
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/resume",
                    headers=headers,
                    json={"reason": "不能重复恢复"},
                )
            ).status_code == 409
            await client.post(f"/api/v1/suppliers/{supplier_id}/commands/submit", headers=headers)
            await client.post(f"/api/v1/suppliers/{supplier_id}/commands/archive", headers=headers)
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/blacklist",
                    headers=headers,
                    json={"reason": "加入黑名单原因"},
                )
            ).status_code == 200
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/resume",
                    headers=headers,
                    json={"reason": "状态不匹配"},
                )
            ).status_code == 409
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/unblacklist",
                    headers=limited_headers,
                    json={"reason": "没有移出权限"},
                )
            ).status_code == 403
            unblacklisted = await client.post(
                f"/api/v1/suppliers/{supplier_id}/commands/unblacklist",
                headers=headers,
                json={"reason": "移出黑名单原因"},
            )
            assert unblacklisted.status_code == 200
            assert unblacklisted.json()["data"]["cooperation_status"] == "NORMAL"
            assert (
                await client.post(
                    f"/api/v1/suppliers/{supplier_id}/commands/unblacklist",
                    headers=headers,
                    json={"reason": "不能重复移出"},
                )
            ).status_code == 409

        async with SessionLocal() as session:
            records = list(
                (
                    await session.scalars(
                        select(SupplierCooperationRecord)
                        .where(SupplierCooperationRecord.supplier_id == supplier_id)
                    )
                ).all()
            )
            assert {
                (record.from_status, record.to_status, record.reason) for record in records
            } == {
                ("NORMAL", "STOPPED", "停止合作原因"),
                ("STOPPED", "NORMAL", "恢复合作原因"),
                ("NORMAL", "BLACKLIST", "加入黑名单原因"),
                ("BLACKLIST", "NORMAL", "移出黑名单原因"),
            }
            assert all(record.actor_id == user_id for record in records)
            eligible = await SupplierService(session).repository.eligible_source_suppliers()
            assert supplier_id in {str(supplier.id) for supplier in eligible}
    finally:
        await cleanup_suppliers(supplier_ids)
        await cleanup_user(limited_user)
        await cleanup_user(user_id)


async def test_supplier_delete_is_logical_and_requires_permission() -> None:
    authorized_user, authorized_headers = await create_user_with_permissions(
        ("supplier:create", "supplier:list", "supplier:detail", "supplier:delete")
    )
    unprivileged_user, unprivileged_headers = await create_user_with_permissions(
        ("supplier:list",)
    )
    supplier_ids: list[str] = []
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/api/v1/suppliers",
                headers=authorized_headers,
                json={
                    "supplier_name": "删除测试供应商",
                    "main_brands": "品牌",
                    "advantage": "优势",
                    "contacts": [{"contact_name": "联系人", "contact_phone": None}],
                },
            )
            assert created.status_code == 201
            supplier_id = created.json()["data"]["id"]
            supplier_ids.append(supplier_id)

            forbidden = await client.delete(
                f"/api/v1/suppliers/{supplier_id}", headers=unprivileged_headers
            )
            assert forbidden.status_code == 403

            deleted = await client.delete(
                f"/api/v1/suppliers/{supplier_id}", headers=authorized_headers
            )
            assert deleted.status_code == 200
            assert deleted.json()["data"] == {"id": supplier_id, "status": "deleted"}
            detail_after_delete = await client.get(
                f"/api/v1/suppliers/{supplier_id}", headers=authorized_headers
            )
            assert detail_after_delete.status_code == 404
            listing = await client.get("/api/v1/suppliers", headers=authorized_headers)
            assert supplier_id not in {item["id"] for item in listing.json()["data"]["items"]}

        async with SessionLocal() as session:
            supplier = await session.get(Supplier, supplier_id)
            assert supplier is not None
            assert supplier.is_deleted is True
            assert supplier.deleted_by == authorized_user
            assert supplier.deleted_at is not None
            contacts = list(
                (
                    await session.scalars(
                        select(SupplierContact).where(SupplierContact.supplier_id == supplier_id)
                    )
                ).all()
            )
            assert contacts and all(contact.is_deleted for contact in contacts)
    finally:
        await cleanup_suppliers(supplier_ids)
        await cleanup_user(authorized_user)
        await cleanup_user(unprivileged_user)


async def test_supplier_name_is_unique_and_deleted_supplier_is_restored() -> None:
    user_id, headers = await create_user_with_permissions(
        ("supplier:create", "supplier:detail", "supplier:delete")
    )
    supplier_ids: list[str] = []
    supplier_name = f"恢复供应商-{uuid.uuid4()}"
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={
                    "supplier_name": supplier_name,
                    "main_brands": "旧品牌",
                    "advantage": "旧优势",
                    "contacts": [{"contact_name": "旧联系人", "contact_phone": "13000000000"}],
                },
            )
            assert created.status_code == 201
            original = created.json()["data"]
            supplier_ids.append(original["id"])

            deleted = await client.delete(f"/api/v1/suppliers/{original['id']}", headers=headers)
            assert deleted.status_code == 200

            restored = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={
                    "supplier_name": supplier_name,
                    "main_brands": "新品牌",
                    "advantage": "新优势",
                    "contacts": [{"contact_name": "新联系人", "contact_phone": "13100000000"}],
                },
            )
            assert restored.status_code == 201
            restored_data = restored.json()["data"]
            assert restored_data["id"] == original["id"]
            assert restored_data["supplier_code"] == original["supplier_code"]
            assert restored_data["main_brands"] == "新品牌"
            assert restored_data["archive_status"] == "DRAFT"
            assert restored_data["cooperation_status"] == "NORMAL"
            assert restored_data["contacts"] == [
                {
                    "id": restored_data["contacts"][0]["id"],
                    "contact_name": "新联系人",
                    "contact_phone": "13100000000",
                }
            ]

            duplicate = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={
                    "supplier_name": supplier_name,
                    "main_brands": "重复品牌",
                    "advantage": "重复优势",
                },
            )
            assert duplicate.status_code == 409
            assert duplicate.json()["code"] == "SUPPLIER_NAME_EXISTS"
            assert duplicate.json()["message"] == "该供应商已存在"

        async with SessionLocal() as session:
            supplier = await session.get(Supplier, original["id"])
            assert supplier is not None
            assert supplier.is_deleted is False
            assert supplier.deleted_by is None
            assert supplier.deleted_at is None
    finally:
        await cleanup_suppliers(supplier_ids)
        await cleanup_user(user_id)


async def test_supplier_excel_preview_and_confirm() -> None:
    user_id, headers = await create_user_with_permissions(
        ("supplier:create", "supplier:list", "supplier:delete")
    )
    another_user, another_headers = await create_user_with_permissions(("supplier:create",))
    batch_ids: list[str] = []
    supplier_ids: list[str] = []
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            template = await client.get("/api/v1/suppliers/imports/template", headers=headers)
            assert template.status_code == 200
            template_workbook = load_workbook(BytesIO(template.content), read_only=True)
            template_sheet = template_workbook.active
            assert template_sheet is not None
            assert tuple(cell.value for cell in next(template_sheet.iter_rows(max_row=1))) == (
                "供应商名称",
                "主营品牌",
                "主要优势",
                "联系人",
                "联系电话",
            )

            invalid_preview = await client.post(
                "/api/v1/suppliers/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "invalid.xlsx",
                        workbook_bytes([("错误供应商", "品牌", None, None, None)]),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert invalid_preview.status_code == 200
            invalid_batch = invalid_preview.json()["data"]
            batch_ids.append(invalid_batch["id"])
            assert invalid_batch["invalid_rows"] == 1
            assert invalid_batch["rows"][0]["error_message"] == "主要优势不能为空"
            invalid_confirm = await client.post(
                f"/api/v1/suppliers/imports/{invalid_batch['id']}/confirm", headers=headers
            )
            assert invalid_confirm.status_code == 409

            existing = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={
                    "supplier_name": "已存在供应商",
                    "main_brands": "品牌",
                    "advantage": "优势",
                },
            )
            assert existing.status_code == 201
            supplier_ids.append(existing.json()["data"]["id"])
            duplicate_preview = await client.post(
                "/api/v1/suppliers/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "duplicate.xlsx",
                        workbook_bytes(
                            [
                                ("已存在供应商", "品牌", "优势", None, None),
                                ("Excel 重复供应商", "品牌", "优势", None, None),
                                ("Excel 重复供应商", "品牌", "优势", None, None),
                            ]
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert duplicate_preview.status_code == 200
            duplicate_batch = duplicate_preview.json()["data"]
            batch_ids.append(duplicate_batch["id"])
            assert duplicate_batch["invalid_rows"] == 2
            assert duplicate_batch["rows"][0]["error_message"] == "该供应商已存在"
            assert duplicate_batch["rows"][2]["error_message"] == "与 Excel 第 3 行供应商名称重复"

            valid_preview = await client.post(
                "/api/v1/suppliers/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "valid.xlsx",
                        workbook_bytes([("导入供应商", "品牌 A", "现货", "王五", "13800000000")]),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert valid_preview.status_code == 200
            valid_batch = valid_preview.json()["data"]
            batch_ids.append(valid_batch["id"])
            assert valid_batch["valid_rows"] == 1
            assert valid_batch["invalid_rows"] == 0

            wrong_uploader = await client.post(
                f"/api/v1/suppliers/imports/{valid_batch['id']}/confirm", headers=another_headers
            )
            assert wrong_uploader.status_code == 403

            confirmed = await client.post(
                f"/api/v1/suppliers/imports/{valid_batch['id']}/confirm",
                headers=headers,
                json={"archive_status": "PENDING"},
            )
            assert confirmed.status_code == 200
            assert confirmed.json()["data"]["imported_count"] == 1
            repeated = await client.post(
                f"/api/v1/suppliers/imports/{valid_batch['id']}/confirm", headers=headers
            )
            assert repeated.status_code == 409

            deleted = await client.post(
                "/api/v1/suppliers",
                headers=headers,
                json={
                    "supplier_name": "导入恢复供应商",
                    "main_brands": "旧品牌",
                    "advantage": "旧优势",
                },
            )
            assert deleted.status_code == 201
            deleted_supplier = deleted.json()["data"]
            supplier_ids.append(deleted_supplier["id"])
            assert (
                await client.delete(f"/api/v1/suppliers/{deleted_supplier['id']}", headers=headers)
            ).status_code == 200
            restore_preview = await client.post(
                "/api/v1/suppliers/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "restore.xlsx",
                        workbook_bytes([("导入恢复供应商", "新品牌", "新优势", "赵六", None)]),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert restore_preview.status_code == 200
            restore_batch = restore_preview.json()["data"]
            batch_ids.append(restore_batch["id"])
            assert restore_batch["invalid_rows"] == 0
            restored = await client.post(
                f"/api/v1/suppliers/imports/{restore_batch['id']}/confirm", headers=headers
            )
            assert restored.status_code == 200

            integrity_preview = await client.post(
                "/api/v1/suppliers/imports/preview",
                headers=headers,
                files={
                    "file": (
                        "integrity.xlsx",
                        workbook_bytes([("完整性校验供应商", "品牌", "优势", None, None)]),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert integrity_preview.status_code == 200
            integrity_batch = integrity_preview.json()["data"]
            batch_ids.append(integrity_batch["id"])
            async with SessionLocal() as session:
                batch = await session.get(SupplierImportBatch, integrity_batch["id"])
                assert batch is not None
                batch.valid_rows = 0
                await session.commit()
            integrity_confirm = await client.post(
                f"/api/v1/suppliers/imports/{integrity_batch['id']}/confirm", headers=headers
            )
            assert integrity_confirm.status_code == 409
            assert integrity_confirm.json()["code"] == "SUPPLIER_IMPORT_BATCH_INTEGRITY_ERROR"

        async with SessionLocal() as session:
            restored_supplier = await session.get(Supplier, deleted_supplier["id"])
            assert restored_supplier is not None
            assert restored_supplier.is_deleted is False
            assert restored_supplier.main_brands == "新品牌"
            assert restored_supplier.archive_status == "ARCHIVED"
            integrity_batch_after_failure = await session.get(
                SupplierImportBatch, integrity_batch["id"]
            )
            assert integrity_batch_after_failure is not None
            assert integrity_batch_after_failure.status == "VALIDATED"
            assert await session.scalar(
                select(func.count())
                .select_from(Supplier)
                .where(Supplier.supplier_name == "完整性校验供应商")
            ) == 0

        async with SessionLocal() as session:
            imported = list(
                (
                    await session.scalars(
                        select(Supplier).where(
                            Supplier.created_by == user_id,
                            Supplier.supplier_name == "导入供应商",
                        )
                    )
                ).all()
            )
            assert len(imported) == 1
            supplier_ids.append(str(imported[0].id))
            assert imported[0].archive_status == "PENDING"
            assert imported[0].cooperation_status == "NORMAL"
            assert imported[0].supplier_code.startswith("SUP")
    finally:
        await cleanup_suppliers(supplier_ids)
        await cleanup_import_batches(batch_ids)
        await cleanup_user(user_id)
        await cleanup_user(another_user)


@pytest.mark.parametrize(
    ("method", "path", "payload"),
    [
        ("get", "/api/v1/suppliers/not-a-uuid", None),
        ("patch", "/api/v1/suppliers/not-a-uuid", {"supplier_name": "valid"}),
        ("delete", "/api/v1/suppliers/not-a-uuid", None),
        ("post", "/api/v1/suppliers/not-a-uuid/commands/submit", None),
        ("post", "/api/v1/suppliers/not-a-uuid/commands/archive", None),
        ("post", "/api/v1/suppliers/not-a-uuid/commands/stop", {"reason": "valid"}),
        (
            "post",
            "/api/v1/suppliers/not-a-uuid/commands/blacklist",
            {"reason": "valid"},
        ),
        ("post", "/api/v1/suppliers/imports/not-a-uuid/confirm", None),
    ],
)
async def test_supplier_uuid_path_validation(
    method: str, path: str, payload: dict[str, str] | None
) -> None:
    user_id, headers = await create_user_with_permissions(SUPPLIER_PERMISSIONS)
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.request(method, path, headers=headers, json=payload)
            assert response.status_code == 422
            assert response.json()["code"] == "VALIDATION_ERROR"
    finally:
        await cleanup_user(user_id)


async def test_supplier_services_participate_in_caller_transactions() -> None:
    actor_id = uuid.uuid4()
    supplier_name = f"rollback-supplier-{uuid.uuid4()}"
    import_filename = f"rollback-import-{uuid.uuid4()}.xlsx"
    async with SessionLocal() as session:
        try:
            async with session.begin():
                await SupplierService(session).create(
                    SupplierCreateRequest(
                        supplier_name=supplier_name,
                        main_brands="brand",
                        advantage="advantage",
                        contacts=[],
                    ),
                    actor_id,
                )
                raise RuntimeError("caller rollback")
        except RuntimeError as exc:
            assert str(exc) == "caller rollback"
    async with SessionLocal() as session:
        assert await session.scalar(
            select(func.count())
            .select_from(Supplier)
            .where(Supplier.supplier_name == supplier_name)
        ) == 0

    async with SessionLocal() as session:
        try:
            async with session.begin():
                await SupplierImportService(session).preview(
                    import_filename,
                    workbook_bytes([("rollback import", "brand", "advantage", None, None)]),
                    actor_id,
                )
                raise RuntimeError("caller rollback")
        except RuntimeError as exc:
            assert str(exc) == "caller rollback"
    async with SessionLocal() as session:
        assert await session.scalar(
            select(func.count())
            .select_from(SupplierImportBatch)
            .where(SupplierImportBatch.original_filename == import_filename)
        ) == 0


async def test_supplier_standalone_sequential_writes_do_not_leave_transaction_active() -> None:
    actor_id = uuid.uuid4()
    supplier_id: str | None = None
    supplier_name = f"standalone-supplier-{uuid.uuid4()}"
    updated_name = f"updated-{supplier_name}"
    try:
        async with SessionLocal() as session:
            service = SupplierService(session)
            created = await service.create(
                SupplierCreateRequest(
                    supplier_name=supplier_name,
                    main_brands="brand",
                    advantage="advantage",
                    contacts=[],
                ),
                actor_id,
            )
            supplier_id = str(created.id)
            assert session.in_transaction() is False

            updated = await service.update(
                created.id,
                SupplierUpdateRequest(supplier_name=updated_name),
                actor_id,
            )
            assert updated.supplier_name == updated_name
            assert session.in_transaction() is False

        async with SessionLocal() as session:
            supplier = await session.get(Supplier, supplier_id)
            assert supplier is not None
            assert supplier.supplier_name == updated_name
    finally:
        await cleanup_suppliers([supplier_id] if supplier_id else [])
