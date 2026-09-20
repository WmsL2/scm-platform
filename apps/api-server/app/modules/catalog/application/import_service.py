"""Fixed-template Product Master import workflow."""

from __future__ import annotations

import asyncio
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from tempfile import NamedTemporaryFile, TemporaryDirectory
from typing import Literal, cast

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.contracts import AppError, PageParams
from app.core.config import get_settings
from app.core.transaction import transaction_scope
from app.infrastructure.adapters import ObjectStorage, get_object_storage
from app.modules.catalog.application.excel_images import (
    DispimgImageArchive,
    DispimgImageError,
    ExcelImage,
    dispimg_image_id,
    inspect_dispimg_references,
)
from app.modules.catalog.domain.lifecycle import ProductStatus
from app.modules.catalog.infrastructure.models import (
    Product,
    ProductImportRow,
    ProductImportSupplierMatch,
    ProductImportTask,
)
from app.modules.catalog.infrastructure.repository import ProductRepository
from app.modules.catalog.schemas import (
    ProductImportConfirmResponse,
    ProductImportDiscardResponse,
    ProductImportPreviewResponse,
    ProductImportResolveSupplierRequest,
    ProductImportRowResponse,
    ProductImportSupplierCandidateResponse,
    ProductImportSupplierMatchResponse,
    ProductUpdateRequest,
)
from app.modules.supplier.domain.matching import (
    SupplierMatchCandidate,
    SupplierMatchStatus,
    classify_supplier_name_match,
    is_eligible_source_supplier,
    normalize_supplier_name,
)
from app.modules.supplier.infrastructure.models import Supplier
from app.modules.supplier.infrastructure.repository import SupplierRepository

PRODUCT_IMPORT_HEADERS = (
    "所属公司",
    "上架日期",
    "品牌",
    "图片",
    "型号",
    "sku",
    "商品名称",
    "一级类目",
    "二级类目",
    "三级类目",
    "货号",
    "链接",
    "成本价",
    "市场价",
    "京东价",
    "协议价",
    "协议价采购价",
    "利润",
    "京东价毛利（30-50）",
    "扣点复核",
    "毛利率",
    "采销员",
    "供应商",
    "69码",
    "3c编码",
    "产品规格",
    "卖点",
    "包装清单",
    "质保期",
    "限售区域",
    "京东自营前台价",
    "自营旗舰店/官方旗舰店",
    "参考链接",
    "销量",
    "好评率",
    "折扣率",
    "价格虚高比例",
    "税收编码",
    "开票名称",
    "税收分类",
    "发货快递",
    "售后政策",
    "备注",
)

_IMAGE_VALIDATION_ERROR_KEY = "__image_validation_error"


_DIRECT_DECIMAL_HEADERS = (
    "成本价",
    "市场价",
    "京东价",
    "协议价",
    "协议价采购价",
    "利润",
    "京东价毛利（30-50）",
    "扣点复核",
    "毛利率",
    "京东自营前台价",
    "折扣率",
    "价格虚高比例",
    "好评率",
)
_PERCENT_HEADERS = {
    "京东价毛利（30-50）",
    "扣点复核",
    "毛利率",
    "好评率",
    "折扣率",
    "价格虚高比例",
}
_DISPLAY_FORMAT_DECIMAL_HEADERS = _PERCENT_HEADERS | {"利润"}
_DATABASE_DECIMAL_QUANTUM = Decimal("0.0001")
PRODUCT_IMPORT_DB_BATCH_SIZE = 500
PRODUCT_IMPORT_DEFAULT_PAGE_SIZE = 50
_WORKBOOK_SEMAPHORES: dict[int, asyncio.Semaphore] = {}


def _workbook_semaphore() -> asyncio.Semaphore:
    limit = get_settings().product_import_max_concurrent_workbooks
    return _WORKBOOK_SEMAPHORES.setdefault(limit, asyncio.Semaphore(limit))

_IMPORT_UPDATE_FIELD_LABELS = {
    "company_name": "所属公司",
    "listed_at": "上架日期",
    "brand": "品牌",
    "image_reference": "图片",
    "model": "型号",
    "product_name": "商品名称",
    "category_level1_name": "一级类目",
    "category_level2_name": "二级类目",
    "category_level3_name": "三级类目",
    "item_number": "货号",
    "jd_same_product_url": "链接",
    "cost_price": "成本价",
    "market_price": "市场价",
    "jd_price": "京东价",
    "agreement_price": "协议价",
    "agreement_purchase_price": "协议价采购价",
    "profit": "利润",
    "jd_margin": "京东价毛利",
    "deduction_review": "扣点复核",
    "gross_margin": "毛利率",
    "purchasing_agent": "采销员",
    "barcode_text": "69码",
    "certification_3c_code": "3c编码",
    "product_specification": "产品规格",
    "selling_points": "卖点",
    "packaging_list": "包装清单",
    "warranty_period": "质保期",
    "restricted_regions": "限售区域",
    "jd_self_operated_price": "京东自营前台价",
    "reference_url": "参考链接",
    "storefront_type": "自营旗舰店/官方旗舰店",
    "discount_rate": "折扣率",
    "price_inflation_rate": "价格虚高比例",
    "sales_volume": "销量",
    "positive_rating": "好评率",
    "tax_code": "税收编码",
    "invoice_name": "开票名称",
    "tax_category": "税收分类",
    "shipping_courier": "发货快递",
    "after_sales_policy": "售后政策",
    "remark": "备注",
}


class ProductImportService:
    def __init__(self, session: AsyncSession, storage: ObjectStorage | None = None) -> None:
        self.session = session
        self.storage = storage or get_object_storage()
        self.repository = ProductRepository(session)
        self.supplier_repository = SupplierRepository(session)

    async def preview(
        self,
        filename: str,
        file_path: Path,
        file_size: int,
        actor_id: uuid.UUID,
    ) -> ProductImportPreviewResponse:
        async with _workbook_semaphore():
            await self._cleanup_expired_temp_media()
            self._validate_upload(filename, file_size)
            rows = await asyncio.to_thread(self._parse_rows, file_path)
            has_embedded_image_formula = any(
                dispimg_image_id(row.source_data.get("图片")) is not None for row in rows
            )
            candidates = await self.supplier_repository.eligible_source_suppliers()
            task = ProductImportTask(
                id=uuid.uuid4(),
                original_filename=filename[:255],
                status="VALIDATED",
                total_rows=len(rows),
                valid_rows=0,
                update_rows=0,
                invalid_rows=0,
                imported_rows=0,
                created_by=actor_id,
            )
            self._create_supplier_matches(task, candidates, rows)
            stored_source_key: str | None = None
            try:
                if has_embedded_image_formula:
                    stored_source_key = await self.storage.save_file(
                        f"product-import-sources/{task.id}.xlsx", file_path
                    )
                    task.source_file_storage_key = stored_source_key
                async with transaction_scope(self.session):
                    self.session.add(task)
                    await self.session.flush()
                    matches_by_name = {
                        match.supplier_name_normalized: match
                        for match in task.supplier_matches
                    }
                    for offset in range(0, len(rows), PRODUCT_IMPORT_DB_BATCH_SIZE):
                        batch = rows[offset : offset + PRODUCT_IMPORT_DB_BATCH_SIZE]
                        for row in batch:
                            row.import_task_id = task.id
                            if row.supplier_name_raw:
                                match = matches_by_name.get(
                                    normalize_supplier_name(row.supplier_name_raw)
                                )
                                row.supplier_match_id = match.id if match is not None else None
                        self.session.add_all(batch)
                        await self.session.flush()
                    self.session.expire(task, ["rows", "supplier_matches"])
                    reloaded_task = await self.repository.import_task_by_id(task.id)
                    if reloaded_task is None:
                        raise RuntimeError("The new Product Import task could not be reloaded")
                    task = reloaded_task
                    await self._refresh_validation(task)
                    response = await self._preview_response(task)
            except Exception:
                if stored_source_key is not None:
                    await self.storage.delete(stored_source_key)
                raise
            return response

    async def get_preview(
        self,
        task_id: uuid.UUID,
        actor_id: uuid.UUID,
        page_params: PageParams,
        *,
        row_status: str,
    ) -> ProductImportPreviewResponse:
        task = await self.repository.import_task_summary_by_id(task_id)
        if task is None:
            raise AppError("PRODUCT_IMPORT_TASK_NOT_FOUND", "Import task not found", 404)
        self._assert_viewable_task(task, actor_id)
        rows, row_total = await self.repository.import_rows_page(
            task_id, page_params, row_status=row_status
        )
        return await self._preview_response(
            task,
            rows=rows,
            row_total=row_total,
            page_params=page_params,
        )

    async def supplier_candidates(self) -> list[ProductImportSupplierCandidateResponse]:
        suppliers = await self.supplier_repository.eligible_source_suppliers()
        return [
            ProductImportSupplierCandidateResponse(
                id=supplier.id,
                supplier_code=supplier.supplier_code,
                supplier_name=supplier.supplier_name,
                main_brands=supplier.main_brands,
            )
            for supplier in suppliers
        ]

    async def resolve_supplier(
        self,
        task_id: uuid.UUID,
        match_id: uuid.UUID,
        payload: ProductImportResolveSupplierRequest,
        actor_id: uuid.UUID,
    ) -> ProductImportPreviewResponse:
        async with transaction_scope(self.session):
            task = await self.repository.import_task_by_id_for_update(task_id)
            self._assert_editable_task(task, actor_id)
            assert task is not None
            match = next((item for item in task.supplier_matches if item.id == match_id), None)
            if match is None:
                raise AppError("PRODUCT_IMPORT_MATCH_NOT_FOUND", "Supplier match not found", 404)
            supplier = await self.supplier_repository.active_by_id(payload.supplier_id)
            if supplier is None or not self._supplier_is_eligible(supplier):
                raise AppError(
                    "PRODUCT_IMPORT_SUPPLIER_INELIGIBLE",
                    "Only archived, normal and non-deleted suppliers may be selected",
                    409,
                )
            match.match_status = SupplierMatchStatus.MATCHED
            match.match_method = "MANUAL"
            match.matched_supplier_id = supplier.id
            match.resolved_by = actor_id
            match.resolved_at = datetime.now()
            await self._refresh_validation(task)
            response = await self._preview_response(task)
        return response

    async def confirm(
        self,
        task_id: uuid.UUID,
        actor_id: uuid.UUID,
    ) -> ProductImportConfirmResponse:
        staged_image_keys: list[str] = []
        retired_image_keys: set[str] = set()
        completed_source: tuple[uuid.UUID, str] | None = None
        try:
            await self._cleanup_expired_temp_media()
            async with transaction_scope(self.session):
                task = await self.repository.import_task_by_id_for_update(task_id)
                self._assert_editable_task(task, actor_id)
                assert task is not None
                await self._refresh_validation(task, preserve_target_snapshots=True)
                rows_to_write = [row for row in task.rows if row.is_valid and not row.is_imported]
                if not rows_to_write:
                    raise AppError(
                        "PRODUCT_IMPORT_NO_VALID_ROWS",
                        "There are no validated rows available to import",
                        409,
                    )
                matches = {match.id: match for match in task.supplier_matches}
                staged_image_keys = await self._stage_confirmed_row_images(task, rows_to_write)
                try:
                    async with self.session.begin_nested():
                        row_supplier_sku_keys: dict[int, tuple[uuid.UUID, str]] = {}
                        for row in rows_to_write:
                            match = (
                                matches.get(row.supplier_match_id)
                                if row.supplier_match_id
                                else None
                            )
                            if match is None or match.matched_supplier_id is None:
                                raise RuntimeError(
                                    "A ready Product Import row is missing a resolved relation"
                                )
                            sku = self._optional(row.source_data["sku"])
                            if sku is None:
                                raise RuntimeError("A ready Product Import row is missing an SKU")
                            row_supplier_sku_keys[row.source_row_number] = (
                                match.matched_supplier_id,
                                sku,
                            )
                        supplier_ids = {
                            supplier_id for supplier_id, _ in row_supplier_sku_keys.values()
                        }
                        get_suppliers = (
                            self.supplier_repository.eligible_source_suppliers_by_ids_for_update
                        )
                        suppliers = await get_suppliers(supplier_ids)
                        if set(suppliers) != supplier_ids:
                            raise AppError(
                                "PRODUCT_IMPORT_REFERENCE_CHANGED",
                                "A source supplier changed after preview; preview the batch again",
                                409,
                            )
                        existing_products = await self.repository.products_by_supplier_sku(
                            set(row_supplier_sku_keys.values()), for_update=True
                        )
                        resolved_rows: list[tuple[ProductImportRow, uuid.UUID]] = []
                        for row in rows_to_write:
                            match = (
                                matches.get(row.supplier_match_id)
                                if row.supplier_match_id
                                else None
                            )
                            if match is None or match.matched_supplier_id is None:
                                raise RuntimeError(
                                    "A ready Product Import row is missing a resolved relation"
                                )
                            supplier_sku_key = row_supplier_sku_keys[row.source_row_number]
                            existing_product = existing_products.get(supplier_sku_key)
                            if (
                                existing_product is not None
                                and existing_product.status == ProductStatus.DISABLED
                            ):
                                raise AppError(
                                    "PRODUCT_IMPORT_DISABLED_PRODUCT",
                                    "A disabled product with the same source supplier and SKU "
                                    "must be enabled or permanently deleted before importing",
                                    409,
                                )
                            if existing_product is None and row.target_product_id is not None:
                                raise AppError(
                                    "PRODUCT_IMPORT_STALE_PREVIEW",
                                    "This product was deleted after preview; "
                                    "upload and preview the workbook again",
                                    409,
                                )
                            if existing_product is None and row.write_action == "UPDATE":
                                raise AppError(
                                    "PRODUCT_IMPORT_STALE_PREVIEW",
                                    "A product changed after preview; "
                                    "upload and preview the workbook again",
                                    409,
                                )
                            if existing_product is not None and row.write_action != "UPDATE":
                                raise AppError(
                                    "PRODUCT_IMPORT_STALE_PREVIEW",
                                    "Another import created this product after preview; "
                                    "upload and preview the workbook again",
                                    409,
                                )
                            if existing_product is not None and (
                                row.target_product_id != existing_product.id
                                or row.target_product_updated_at != existing_product.updated_at
                            ):
                                raise AppError(
                                    "PRODUCT_IMPORT_STALE_PREVIEW",
                                    "This product was updated after preview; "
                                    "upload and preview the workbook again",
                                    409,
                                )
                            resolved_rows.append((row, match.matched_supplier_id))
                        created_count = 0
                        updated_count = 0
                        for row, source_supplier_id in resolved_rows:
                            supplier_sku_key = row_supplier_sku_keys[row.source_row_number]
                            existing_product = existing_products.get(supplier_sku_key)
                            if existing_product is None:
                                self.session.add(
                                    self._product_from_row(row, source_supplier_id, actor_id)
                                )
                                created_count += 1
                            else:
                                previous_image_reference = existing_product.image_reference
                                self._update_product_from_row(
                                    existing_product, row, source_supplier_id, actor_id
                                )
                                if previous_image_reference != existing_product.image_reference:
                                    previous_key = self._local_media_key(previous_image_reference)
                                    if previous_key is not None:
                                        retired_image_keys.add(previous_key)
                                updated_count += 1
                        imported_count = created_count + updated_count
                        await self.session.flush()
                except IntegrityError as exc:
                    if "uq_scm_product_source_supplier_sku" in str(exc.orig).lower():
                        raise AppError(
                            "PRODUCT_IMPORT_STALE_PREVIEW",
                            "Another import created the same source supplier and SKU; "
                            "upload and preview the workbook again",
                            409,
                        ) from exc
                    raise
                task.confirmed_by = actor_id
                task.confirmed_at = datetime.now()
                for row in rows_to_write:
                    row.is_imported = True
                    row.imported_by = actor_id
                    row.imported_at = task.confirmed_at
                await self._refresh_validation(task)
                if task.status == "CONFIRMED" and task.source_file_storage_key:
                    completed_source = (task.id, task.source_file_storage_key)
                response = ProductImportConfirmResponse(
                    id=task.id,
                    status=task.status,
                    imported_count=imported_count,
                    created_count=created_count,
                    updated_count=updated_count,
                    imported_rows=task.imported_rows,
                    valid_rows=task.valid_rows,
                    update_rows=task.update_rows,
                    invalid_rows=task.invalid_rows,
                )
        except Exception:
            for key in staged_image_keys:
                await self.storage.delete(key)
            raise
        if completed_source is not None:
            await self._delete_completed_source(*completed_source)
        for key in retired_image_keys:
            try:
                await self.storage.delete(key)
            except Exception:
                pass
        return response

    async def discard(
        self, task_id: uuid.UUID, actor_id: uuid.UUID
    ) -> ProductImportDiscardResponse:
        source_key: str | None = None
        image_keys: list[str] = []
        async with transaction_scope(self.session):
            task = await self.repository.import_task_by_id_for_update(task_id)
            self._assert_viewable_task(task, actor_id)
            assert task is not None
            if task.status == "CONFIRMED":
                raise AppError(
                    "PRODUCT_IMPORT_TASK_ALREADY_CONFIRMED",
                    "A confirmed import cannot be discarded",
                    409,
                )
            source_key = task.source_file_storage_key
            image_keys = [
                row.image_storage_key
                for row in task.rows
                if row.image_storage_key is not None and not row.is_imported
            ]
            task.status = "EXPIRED"
            response = ProductImportDiscardResponse(id=task.id, status="EXPIRED")
        await self._delete_task_temp_media(task_id, source_key, image_keys)
        return response

    def _validate_upload(self, filename: str, file_size: int) -> None:
        if not filename.lower().endswith(".xlsx"):
            raise AppError(
                "PRODUCT_IMPORT_FILE_TYPE_INVALID", "Only .xlsx files are supported", 422
            )
        if file_size <= 0:
            raise AppError("PRODUCT_IMPORT_FILE_EMPTY", "The import file is empty", 422)
        max_file_mb = get_settings().product_import_max_file_mb
        if file_size > max_file_mb * 1024 * 1024:
            raise AppError(
                "PRODUCT_IMPORT_FILE_TOO_LARGE",
                f"The import file exceeds {max_file_mb} MB",
                422,
            )

    def _parse_rows(self, file_path: Path) -> list[ProductImportRow]:
        try:
            formula_workbook = load_workbook(file_path, read_only=True, data_only=False)
            cached_workbook = load_workbook(file_path, read_only=True, data_only=True)
        except Exception as exc:
            raise AppError(
                "PRODUCT_IMPORT_FILE_INVALID", "The file is not a valid .xlsx workbook", 422
            ) from exc
        try:
            formula_sheet = formula_workbook.active
            cached_sheet = cached_workbook.active
            assert formula_sheet is not None and cached_sheet is not None
            header_cells = next(
                formula_sheet.iter_rows(max_row=1, max_col=len(PRODUCT_IMPORT_HEADERS))
            )
            headers = tuple(self._normalized(cell.value) for cell in header_cells)
            extra_header_cells = next(
                formula_sheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    min_col=len(PRODUCT_IMPORT_HEADERS) + 1,
                    max_col=formula_sheet.max_column,
                ),
                (),
            )
            has_extra_header = any(self._normalized(cell.value) for cell in extra_header_cells)
            if headers != PRODUCT_IMPORT_HEADERS or has_extra_header:
                raise AppError(
                    "PRODUCT_IMPORT_TEMPLATE_INVALID",
                    "Template headers must exactly match the approved "
                    "43-column Product Master template",
                    422,
                )
            max_rows = get_settings().product_import_max_rows
            rows: list[ProductImportRow] = []
            for source_row_number, (formula_cells, cached_cells) in enumerate(
                zip(
                    formula_sheet.iter_rows(min_row=2, max_col=len(PRODUCT_IMPORT_HEADERS)),
                    cached_sheet.iter_rows(min_row=2, max_col=len(PRODUCT_IMPORT_HEADERS)),
                    strict=True,
                ),
                start=2,
            ):
                if not any(cell.value is not None for cell in formula_cells):
                    continue
                if len(rows) >= max_rows:
                    raise AppError(
                        "PRODUCT_IMPORT_TOO_MANY_ROWS",
                        f"An import may contain at most {max_rows} data rows",
                        422,
                    )
                source_data = {
                    header: self._import_cell_text(header, cell.value, cell.number_format)
                    for header, cell in zip(PRODUCT_IMPORT_HEADERS, formula_cells)
                }
                calculated_data = {
                    header: self._import_cell_text(
                        header, cached_cell.value, formula_cell.number_format
                    )
                    for header, formula_cell, cached_cell in zip(
                        PRODUCT_IMPORT_HEADERS, formula_cells, cached_cells, strict=True
                    )
                }
                rows.append(
                    ProductImportRow(
                        source_row_number=source_row_number,
                        source_data=source_data,
                        calculated_data=calculated_data,
                        supplier_name_raw=self._optional(source_data["供应商"]),
                        is_valid=False,
                    )
                )
            if not rows:
                raise AppError("PRODUCT_IMPORT_NO_DATA_ROWS", "The workbook has no data rows", 422)
            image_ids = {
                image_id
                for row in rows
                if (image_id := dispimg_image_id(row.source_data.get("图片"))) is not None
            }
            if image_ids:
                max_image_bytes = get_settings().product_import_max_image_mb * 1024 * 1024
                try:
                    image_errors = inspect_dispimg_references(
                        file_path,
                        image_ids,
                        max_image_bytes=max_image_bytes,
                    )
                except DispimgImageError as exc:
                    image_errors = {image_id: str(exc) for image_id in image_ids}
                for row in rows:
                    image_id = dispimg_image_id(row.source_data.get("图片"))
                    if image_id is not None and image_id in image_errors:
                        row.calculated_data[_IMAGE_VALIDATION_ERROR_KEY] = (
                            f"图片无法读取：{image_errors[image_id]}"
                        )
            return rows
        finally:
            formula_workbook.close()
            cached_workbook.close()

    def _create_supplier_matches(
        self,
        task: ProductImportTask,
        candidates: list[Supplier],
        rows: list[ProductImportRow],
    ) -> None:
        candidate_models = [
            SupplierMatchCandidate(
                id=supplier.id,
                supplier_name=supplier.supplier_name,
                archive_status=supplier.archive_status,
                cooperation_status=supplier.cooperation_status,
                is_deleted=supplier.is_deleted,
            )
            for supplier in candidates
        ]
        names = sorted(
            {
                normalize_supplier_name(row.supplier_name_raw)
                for row in rows
                if row.supplier_name_raw and normalize_supplier_name(row.supplier_name_raw)
            }
        )
        for normalized_name in names:
            result = classify_supplier_name_match(normalized_name, candidate_models)
            task.supplier_matches.append(
                ProductImportSupplierMatch(
                    supplier_name_normalized=normalized_name,
                    match_status=result.status,
                    match_method=result.match_method,
                    matched_supplier_id=result.matched_supplier_id,
                )
            )

    async def _refresh_validation(
        self, task: ProductImportTask, *, preserve_target_snapshots: bool = False
    ) -> None:
        matches = {match.id: match for match in task.supplier_matches}
        row_supplier_sku_keys: dict[int, tuple[uuid.UUID, str]] = {}
        for row in task.rows:
            if row.is_imported:
                continue
            match = (
                matches.get(row.supplier_match_id)
                if row.supplier_match_id is not None
                else None
            )
            sku = self._optional(row.source_data["sku"])
            if match is not None and match.matched_supplier_id is not None and sku is not None:
                row_supplier_sku_keys[row.source_row_number] = (match.matched_supplier_id, sku)
        existing_products = await self.repository.products_by_supplier_sku(
            set(row_supplier_sku_keys.values())
        )
        valid_rows = 0
        update_rows = 0
        imported_rows = 0
        unresolved_match = False
        first_excel_row_for_key: dict[tuple[uuid.UUID, str], int] = {}
        for row in task.rows:
            if row.is_imported:
                imported_rows += 1
                continue
            supplier_sku_key = row_supplier_sku_keys.get(row.source_row_number)
            first_excel_row = (
                first_excel_row_for_key.get(supplier_sku_key)
                if supplier_sku_key is not None
                else None
            )
            errors, warnings = await self._row_messages(
                row,
                matches,
                existing_products=existing_products,
                first_excel_row=first_excel_row,
                preserve_target_snapshot=preserve_target_snapshots,
            )
            row.error_message = "；".join(errors) or None
            row.warning_message = "；".join(warnings) or None
            row.is_valid = not errors
            if row.is_valid:
                if row.write_action == "UPDATE":
                    update_rows += 1
                else:
                    valid_rows += 1
            match = (
                matches.get(row.supplier_match_id) if row.supplier_match_id is not None else None
            )
            if match is not None and match.match_status != SupplierMatchStatus.MATCHED:
                unresolved_match = True
            if supplier_sku_key is not None and first_excel_row is None:
                first_excel_row_for_key[supplier_sku_key] = row.source_row_number
        task.valid_rows = valid_rows
        task.update_rows = update_rows
        task.imported_rows = imported_rows
        task.invalid_rows = task.total_rows - imported_rows - valid_rows - update_rows
        if imported_rows == task.total_rows:
            task.status = "CONFIRMED"
        elif imported_rows:
            task.status = "PARTIALLY_CONFIRMED"
        elif valid_rows + update_rows == task.total_rows:
            task.status = "READY_TO_CONFIRM"
        elif unresolved_match:
            task.status = "NEEDS_RESOLUTION"
        else:
            task.status = "VALIDATED"

    async def _row_messages(
        self,
        row: ProductImportRow,
        matches: dict[uuid.UUID, ProductImportSupplierMatch],
        *,
        existing_products: dict[tuple[uuid.UUID, str], Product],
        first_excel_row: int | None,
        preserve_target_snapshot: bool,
    ) -> tuple[list[str], list[str]]:
        values = row.source_data
        errors: list[str] = []
        warnings: list[str] = []
        row.write_action = "CREATE"
        row.changed_fields = None
        row.normalized_data = None
        if not preserve_target_snapshot:
            row.target_product_id = None
            row.target_product_updated_at = None
        existing_product: Product | None = None
        image_validation_error = row.calculated_data.get(_IMAGE_VALIDATION_ERROR_KEY)
        if image_validation_error:
            errors.append(image_validation_error)
        cost_price = self._decimal_or_none(self._import_value(row, "成本价"))
        if cost_price is not None and cost_price <= 0:
            errors.append("成本价必须大于0")
        for header in ("一级类目", "二级类目", "三级类目"):
            if self._optional(self._import_value(row, header)) is None:
                errors.append(f"{header}不能为空")
        sku = self._optional(values["sku"])
        if sku is None:
            errors.append("SKU不能为空")
        for header in _DIRECT_DECIMAL_HEADERS:
            if header == "成本价":
                continue
        for header in ("好评率", "折扣率"):
            rate_value = self._decimal_or_none(
                self._import_value(row, header), percentage=True
            )
            if rate_value is not None and not Decimal("0") <= rate_value <= Decimal("1"):
                errors.append(f"{header}必须在0%到100%之间，裸数字请填写0到1")
        sales_value = self._optional(self._import_value(row, "销量"))
        parsed_sales_value = self._integer_or_none(sales_value)
        if parsed_sales_value is not None and parsed_sales_value < 0:
            errors.append("销量必须是大于等于0的整数")
        match = matches.get(row.supplier_match_id) if row.supplier_match_id is not None else None
        if not row.supplier_name_raw:
            errors.append("供应商不能为空")
        elif match is None or match.match_status != SupplierMatchStatus.MATCHED:
            errors.append("来源供应商尚未解析")
        elif sku is not None and match.matched_supplier_id is not None:
            supplier_sku_key = (match.matched_supplier_id, sku)
            existing_product = existing_products.get(supplier_sku_key)
            if existing_product is not None and existing_product.status == ProductStatus.DISABLED:
                errors.append("该来源供应商与SKU组合的商品已停用，请启用或永久删除后再导入")
            elif first_excel_row is not None:
                errors.append(f"与Excel第{first_excel_row}行的来源供应商与SKU重复")
            elif existing_product is not None:
                row.write_action = "UPDATE"
        if values["上架日期"] and self._optional_date(values["上架日期"]) is None:
            warnings.append("上架日期无法确定年份，正式商品将暂不写入上架日期")
        if not errors and match is not None and match.matched_supplier_id is not None:
            try:
                self._store_normalized_product_values(row, match.matched_supplier_id)
            except ValidationError as exc:
                field = str(exc.errors()[0]["loc"][0]) if exc.errors() else "商品字段"
                label = _IMPORT_UPDATE_FIELD_LABELS.get(field, field)
                errors.append(f"{label}格式、长度或精度不符合要求")
            if not errors and existing_product is not None:
                if not preserve_target_snapshot:
                    row.target_product_id = existing_product.id
                    row.target_product_updated_at = existing_product.updated_at
                row.changed_fields = self._changed_import_fields(
                    existing_product, row, match.matched_supplier_id
                )
        return errors, warnings

    def _product_from_row(
        self,
        row: ProductImportRow,
        source_supplier_id: uuid.UUID,
        actor_id: uuid.UUID,
    ) -> Product:
        return Product(
            **self._product_values_from_row(row, source_supplier_id),
            created_by=actor_id,
            updated_by=actor_id,
        )

    def _update_product_from_row(
        self,
        product: Product,
        row: ProductImportRow,
        source_supplier_id: uuid.UUID,
        actor_id: uuid.UUID,
    ) -> None:
        product_values = self._product_values_from_row(row, source_supplier_id)
        for field, value in product_values.items():
            if field not in {"source_supplier_id", "sku"}:
                setattr(product, field, value)
        product.updated_by = actor_id
        product.updated_at = datetime.now()

    def _changed_import_fields(
        self, product: Product, row: ProductImportRow, source_supplier_id: uuid.UUID
    ) -> list[str]:
        planned_values = self._product_values_from_row(row, source_supplier_id)
        return [
            label
            for field, label in _IMPORT_UPDATE_FIELD_LABELS.items()
            if getattr(product, field) != planned_values[field]
        ]

    def _product_values_from_row(
        self,
        row: ProductImportRow,
        source_supplier_id: uuid.UUID,
    ) -> dict[str, object]:
        if row.normalized_data is None:
            raise RuntimeError("A ready Product Import row is missing normalized values")
        data = row.normalized_data
        payload = ProductUpdateRequest.model_validate(
            {name: data.get(name) for name in ProductUpdateRequest.model_fields}
        )
        sku = data.get("sku")
        if not isinstance(sku, str) or not sku:
            raise RuntimeError("A ready Product Import row is missing an SKU")
        product_values = payload.model_dump()
        product_values.update(
            {
                "image_reference": self._optional(self._image_reference(row)),
                "sku": sku,
                "source_supplier_id": source_supplier_id,
                "deduction_rate": None,
            }
        )
        return product_values

    def _store_normalized_product_values(
        self, row: ProductImportRow, source_supplier_id: uuid.UUID
    ) -> None:
        values = row.source_data
        candidate = {
            "company_name": self._optional(values["所属公司"]),
            "listed_at": self._optional_date(values["上架日期"]),
            "brand": self._optional(values["品牌"]),
            "model": self._optional(values["型号"]),
            "product_name": self._optional(values["商品名称"]),
            "category_level1_name": self._required_text(row, "一级类目"),
            "category_level2_name": self._required_text(row, "二级类目"),
            "category_level3_name": self._required_text(row, "三级类目"),
            "item_number": self._optional(values["货号"]),
            "jd_same_product_url": self._optional(values["链接"]),
            "cost_price": self._decimal_or_none(self._import_value(row, "成本价")),
            "market_price": self._decimal_or_none(self._import_value(row, "市场价")),
            "jd_price": self._decimal_or_none(self._import_value(row, "京东价")),
            "agreement_price": self._decimal_or_none(self._import_value(row, "协议价")),
            "agreement_purchase_price": self._decimal_or_none(
                self._import_value(row, "协议价采购价")
            ),
            "profit": self._decimal_or_none(self._import_value(row, "利润")),
            "jd_margin": self._decimal_or_none(
                self._import_value(row, "京东价毛利（30-50）"), percentage=True
            ),
            "deduction_review": self._decimal_or_none(
                self._import_value(row, "扣点复核"), percentage=True
            ),
            "gross_margin": self._decimal_or_none(
                self._import_value(row, "毛利率"), percentage=True
            ),
            "purchasing_agent": self._optional(values["采销员"]),
            "barcode_text": self._optional(values["69码"]),
            "certification_3c_code": self._optional(values["3c编码"]),
            "product_specification": self._optional(values["产品规格"]),
            "selling_points": self._optional(values["卖点"]),
            "packaging_list": self._optional(values["包装清单"]),
            "warranty_period": self._optional(values["质保期"]),
            "remark": self._optional(values["备注"]),
            "discount_rate": self._decimal_or_none(
                self._import_value(row, "折扣率"), percentage=True
            ),
            "restricted_regions": self._optional(values["限售区域"]),
            "jd_self_operated_price": self._decimal_or_none(
                self._import_value(row, "京东自营前台价")
            ),
            "reference_url": self._optional(values["参考链接"]),
            "storefront_type": self._optional(values["自营旗舰店/官方旗舰店"]),
            "sales_volume": self._integer_or_none(self._import_value(row, "销量")),
            "positive_rating": self._decimal_or_none(
                self._import_value(row, "好评率"), percentage=True
            ),
            "price_inflation_rate": self._decimal_or_none(
                self._import_value(row, "价格虚高比例"), percentage=True
            ),
            "tax_code": self._optional(values["税收编码"]),
            "invoice_name": self._optional(values["开票名称"]),
            "tax_category": self._optional(values["税收分类"]),
            "shipping_courier": self._optional(values["发货快递"]),
            "after_sales_policy": self._optional(values["售后政策"]),
        }
        payload = ProductUpdateRequest.model_validate(candidate)
        normalized = payload.model_dump(mode="json")
        normalized.update(
            {
                "sku": self._optional(values["sku"]),
                "source_supplier_id": str(source_supplier_id),
                "category_level1_name": self._optional(values["一级类目"]),
                "category_level2_name": self._optional(values["二级类目"]),
                "category_level3_name": self._optional(values["三级类目"]),
            }
        )
        row.normalized_data = normalized

    async def _stage_confirmed_row_images(
        self, task: ProductImportTask, rows: list[ProductImportRow]
    ) -> list[str]:
        rows_by_image_id: dict[str, list[ProductImportRow]] = defaultdict(list)
        for row in rows:
            image_id = dispimg_image_id(row.source_data.get("图片"))
            if image_id is not None:
                rows_by_image_id[image_id].append(row)
        if not rows_by_image_id:
            return []
        if not task.source_file_storage_key:
            raise AppError(
                "PRODUCT_IMPORT_SOURCE_FILE_MISSING",
                "The temporary import file is unavailable; upload the workbook again",
                409,
            )
        async with _workbook_semaphore():
            temporary = NamedTemporaryFile(
                prefix="scm-product-import-source-", suffix=".xlsx", delete=False
            )
            source_path = Path(temporary.name)
            temporary.close()
            try:
                try:
                    await self.storage.copy_to(task.source_file_storage_key, source_path)
                except FileNotFoundError as exc:
                    raise AppError(
                        "PRODUCT_IMPORT_SOURCE_FILE_MISSING",
                        "The temporary import file is unavailable; upload the workbook again",
                        409,
                    ) from exc
                max_image_bytes = get_settings().product_import_max_image_mb * 1024 * 1024
                archive = await asyncio.to_thread(DispimgImageArchive, source_path)
                saved_keys: list[str] = []
                try:
                    with TemporaryDirectory(prefix="scm-product-import-images-") as directory:
                        temporary_directory = Path(directory)
                        for index, (image_id, image_rows) in enumerate(
                            rows_by_image_id.items(), start=1
                        ):
                            try:
                                image = await asyncio.to_thread(
                                    archive.extract_to,
                                    image_id,
                                    temporary_directory / f"image-{index}",
                                    max_image_bytes=max_image_bytes,
                                )
                            except DispimgImageError as exc:
                                row_numbers = ", ".join(
                                    str(row.source_row_number) for row in image_rows[:5]
                                )
                                raise AppError(
                                    "PRODUCT_IMPORT_IMAGE_INVALID",
                                    f"Excel row {row_numbers} has an unreadable "
                                    f"embedded image: {exc}",
                                    422,
                                ) from exc
                            try:
                                for row in image_rows:
                                    row.image_storage_key = await self.storage.save_file(
                                        f"product-images/{task.id}/{row.id}{image.extension}",
                                        image.path,
                                    )
                                    saved_keys.append(row.image_storage_key)
                            finally:
                                image.path.unlink(missing_ok=True)
                    return saved_keys
                except Exception:
                    for key in saved_keys:
                        try:
                            await self.storage.delete(key)
                        except Exception:
                            pass
                    raise
                finally:
                    await asyncio.to_thread(archive.close)
            finally:
                source_path.unlink(missing_ok=True)

    async def _cleanup_expired_temp_media(self) -> None:
        cutoff = datetime.now() - timedelta(
            days=get_settings().product_import_unconfirmed_retention_days
        )
        async with transaction_scope(self.session):
            tasks = await self.repository.temporary_media_cleanup_tasks_for_update(
                stale_before=cutoff
            )
            cleanup_targets: list[tuple[uuid.UUID, str | None, list[str]]] = []
            for task in tasks:
                if task.status in {
                    "VALIDATED",
                    "NEEDS_RESOLUTION",
                    "READY_TO_CONFIRM",
                    "PARTIALLY_CONFIRMED",
                }:
                    task.status = "EXPIRED"
                if task.status == "EXPIRED":
                    cleanup_targets.append(
                        (
                            task.id,
                            task.source_file_storage_key,
                            [
                                row.image_storage_key
                                for row in task.rows
                                if row.image_storage_key is not None and not row.is_imported
                            ],
                        )
                    )
                elif task.status == "CONFIRMED" and task.source_file_storage_key:
                    cleanup_targets.append((task.id, task.source_file_storage_key, []))
        for task_id, source_key, image_keys in cleanup_targets:
            await self._delete_task_temp_media(task_id, source_key, image_keys)

    async def _delete_completed_source(self, task_id: uuid.UUID, source_key: str) -> None:
        await self._delete_task_temp_media(task_id, source_key, [])

    async def _delete_task_temp_media(
        self, task_id: uuid.UUID, source_key: str | None, image_keys: list[str]
    ) -> None:
        deleted_source = False
        deleted_images: set[str] = set()
        if source_key is not None:
            try:
                await self.storage.delete(source_key)
                deleted_source = True
            except Exception:
                pass
        for image_key in image_keys:
            try:
                await self.storage.delete(image_key)
                deleted_images.add(image_key)
            except Exception:
                pass
        if not deleted_source and not deleted_images:
            return
        async with transaction_scope(self.session):
            task = await self.repository.import_task_by_id_for_update(task_id)
            if task is None or task.status not in {"EXPIRED", "CONFIRMED"}:
                return
            if deleted_source and task.source_file_storage_key == source_key:
                task.source_file_storage_key = None
            if task.status == "EXPIRED":
                for row in task.rows:
                    if not row.is_imported and row.image_storage_key in deleted_images:
                        row.image_storage_key = None

    async def _stage_images(
        self,
        task: ProductImportTask,
        embedded_images: dict[str, ExcelImage],
        rows: list[ProductImportRow] | None = None,
    ) -> list[str]:
        saved_keys: list[str] = []
        for row in rows or task.rows:
            image_id = dispimg_image_id(row.source_data.get("图片"))
            image = embedded_images.get(image_id or "")
            if image is None:
                continue
            row.image_storage_key = await self.storage.save(
                f"product-images/{task.id}/{row.id}{image.extension}", image.content
            )
            saved_keys.append(row.image_storage_key)
        return saved_keys

    @staticmethod
    def _image_reference(row: ProductImportRow) -> str | None:
        if row.image_storage_key:
            return f"local-media/{row.image_storage_key}"
        image_value = row.source_data.get("图片") or ""
        return None if image_value.startswith("=") else ProductImportService._optional(image_value)

    @staticmethod
    def _local_media_key(image_reference: str | None) -> str | None:
        prefix = "local-media/"
        if image_reference and image_reference.startswith(prefix):
            return image_reference[len(prefix) :]
        return None

    @staticmethod
    def _assert_viewable_task(task: ProductImportTask | None, actor_id: uuid.UUID) -> None:
        if task is None:
            raise AppError("PRODUCT_IMPORT_TASK_NOT_FOUND", "Import task not found", 404)
        if task.created_by != actor_id:
            raise AppError(
                "PRODUCT_IMPORT_TASK_FORBIDDEN",
                "Only the uploader can view or operate this import batch",
                403,
            )

    @classmethod
    def _assert_editable_task(
        cls, task: ProductImportTask | None, actor_id: uuid.UUID
    ) -> None:
        cls._assert_viewable_task(task, actor_id)
        assert task is not None
        if task.status == "EXPIRED":
            raise AppError(
                "PRODUCT_IMPORT_TASK_EXPIRED",
                "This unconfirmed import expired; upload the workbook again",
                409,
            )
        if task.status == "CONFIRMED":
            raise AppError(
                "PRODUCT_IMPORT_TASK_ALREADY_CONFIRMED", "Import task is already confirmed", 409
            )

    @staticmethod
    def _supplier_is_eligible(supplier: Supplier) -> bool:
        return is_eligible_source_supplier(
            SupplierMatchCandidate(
                id=supplier.id,
                supplier_name=supplier.supplier_name,
                archive_status=supplier.archive_status,
                cooperation_status=supplier.cooperation_status,
                is_deleted=supplier.is_deleted,
            )
        )

    async def _preview_response(
        self,
        task: ProductImportTask,
        *,
        rows: list[ProductImportRow] | None = None,
        row_total: int | None = None,
        page_params: PageParams | None = None,
    ) -> ProductImportPreviewResponse:
        if page_params is None:
            page_params = PageParams(page=1, page_size=PRODUCT_IMPORT_DEFAULT_PAGE_SIZE)
        if rows is None:
            all_rows = list(task.rows)
            row_total = len(all_rows)
            start = (page_params.page - 1) * page_params.page_size
            rows = all_rows[start : start + page_params.page_size]
        if row_total is None:
            row_total = len(rows)
        supplier_ids = {
            match.matched_supplier_id
            for match in task.supplier_matches
            if match.matched_supplier_id is not None
        }
        suppliers = await self.supplier_repository.active_by_ids(supplier_ids)
        supplier_names: dict[uuid.UUID, tuple[str, str]] = {}
        for match in task.supplier_matches:
            if match.matched_supplier_id is None:
                continue
            supplier = suppliers.get(match.matched_supplier_id)
            if supplier is not None:
                supplier_names[match.id] = (supplier.supplier_code, supplier.supplier_name)
        return ProductImportPreviewResponse(
            id=task.id,
            original_filename=task.original_filename,
            status=task.status,
            total_rows=task.total_rows,
            valid_rows=task.valid_rows,
            update_rows=task.update_rows,
            invalid_rows=task.invalid_rows,
            imported_rows=task.imported_rows,
            row_total=row_total,
            page=page_params.page,
            page_size=page_params.page_size,
            rows=[
                ProductImportRowResponse(
                    source_row_number=row.source_row_number,
                    product_name=self._optional(row.source_data["商品名称"]),
                    supplier_name_raw=row.supplier_name_raw,
                    image_saved=row.image_storage_key is not None,
                    image_pending_save=(
                        row.image_storage_key is None
                        and dispimg_image_id(row.source_data.get("图片")) is not None
                    ),
                    category_path=" / ".join(
                        self._normalized(row.source_data[name])
                        for name in ("一级类目", "二级类目", "三级类目")
                    ),
                    supplier_match_id=row.supplier_match_id,
                    is_valid=row.is_valid,
                    write_action=cast(Literal["CREATE", "UPDATE"], row.write_action),
                    changed_fields=row.changed_fields,
                    is_imported=row.is_imported,
                    error_message=row.error_message,
                    warning_message=row.warning_message,
                )
                for row in rows
            ],
            supplier_matches=[
                ProductImportSupplierMatchResponse(
                    id=match.id,
                    supplier_name_normalized=match.supplier_name_normalized,
                    match_status=match.match_status,
                    match_method=match.match_method,
                    matched_supplier_id=match.matched_supplier_id,
                    matched_supplier_code=supplier_names.get(match.id, (None, None))[0],
                    matched_supplier_name=supplier_names.get(match.id, (None, None))[1],
                )
                for match in task.supplier_matches
            ],
        )

    @staticmethod
    def _cell_text(value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

    @classmethod
    def _import_cell_text(
        cls, header: str, value: object, number_format: str
    ) -> str:
        if (
            header not in _DISPLAY_FORMAT_DECIMAL_HEADERS
            or isinstance(value, bool)
            or not isinstance(value, (int, float, Decimal))
        ):
            return cls._cell_text(value)
        try:
            decimal_value = Decimal(str(value))
        except (InvalidOperation, ValueError):
            return cls._cell_text(value)
        if not decimal_value.is_finite():
            return cls._cell_text(value)

        display_format = cls._display_number_format(decimal_value, number_format)
        decimal_places = cls._display_decimal_places(display_format)
        if decimal_places is None:
            return cls._cell_text(value)

        quantum = Decimal(1).scaleb(-decimal_places)
        if "%" in display_format:
            displayed = (decimal_value * Decimal("100")).quantize(
                quantum, rounding=ROUND_HALF_UP
            )
            return f"{format(displayed, f'.{decimal_places}f')}%"
        displayed = decimal_value.quantize(quantum, rounding=ROUND_HALF_UP)
        return format(displayed, f".{decimal_places}f")

    @staticmethod
    def _display_number_format(value: Decimal, number_format: str) -> str:
        sections = number_format.split(";")
        if value < 0 and len(sections) > 1:
            return sections[1]
        return sections[0]

    @staticmethod
    def _display_decimal_places(number_format: str) -> int | None:
        if not number_format or number_format.lower() == "general":
            return None
        normalized = re.sub(r'"[^"]*"', "", number_format)
        normalized = re.sub(r"\[[^\]]+\]", "", normalized)
        normalized = re.sub(r"\\.", "", normalized)
        match = re.search(r"[0#?]+\.([0#?]+)", normalized)
        if match is not None:
            return len(match.group(1))
        if re.search(r"[0#?]+", normalized):
            return 0
        return None

    @staticmethod
    def _normalized(value: object) -> str:
        return str(value or "").strip()

    @staticmethod
    def _optional(value: str | None) -> str | None:
        if value is None or not value.strip():
            return None
        return value

    def _required_text(self, row: ProductImportRow, header: str) -> str:
        value = self._optional(self._import_value(row, header))
        if value is None:
            raise RuntimeError(f"A ready Product Import row is missing {header}")
        return value

    @staticmethod
    def _decimal_or_none(value: str | None, *, percentage: bool = False) -> Decimal | None:
        if value is None or not value.strip() or value.startswith("="):
            return None
        try:
            normalized = value.strip()
            if percentage and normalized.endswith("%"):
                parsed = Decimal(normalized[:-1].strip()) / Decimal("100")
            else:
                parsed = Decimal(normalized)
            if not parsed.is_finite():
                return None
            if percentage:
                return parsed.quantize(_DATABASE_DECIMAL_QUANTUM, rounding=ROUND_HALF_UP)
            return parsed
        except (InvalidOperation, ValueError):
            return None

    @staticmethod
    def _integer_or_none(value: str | None) -> int | None:
        if value is None or not value.strip() or value.startswith("="):
            return None
        try:
            parsed = Decimal(value.strip())
        except (InvalidOperation, ValueError):
            return None
        if parsed != parsed.to_integral_value():
            return None
        return int(parsed)

    @staticmethod
    def _import_value(row: ProductImportRow, header: str) -> str | None:
        source = row.source_data.get(header)
        if source and source.startswith("="):
            return row.calculated_data.get(header)
        return source

    @staticmethod
    def _optional_date(value: str | None) -> date | None:
        if value is None or not value.strip():
            return None
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
