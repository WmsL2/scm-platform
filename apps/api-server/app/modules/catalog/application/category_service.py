# ruff: noqa: E501, E701, E702
from __future__ import annotations

import uuid
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Literal, cast

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.sql.elements import ColumnElement

from app.common.contracts import AppError, PageParams, PageResult
from app.core.transaction import transaction_scope
from app.modules.catalog.infrastructure.models import Category, Product
from app.modules.catalog.schemas import (
    CategoryImportError,
    CategoryImportResponse,
    CategoryWriteRequest,
)

IMPORT_HEADERS = ("一级类目ID", "一级类目名称", "二级类目ID", "二级类目名称", "三级类目ID", "三级类目名称", "有效标记", "上下柜标记", "主营事业部")
IMPORT_HEADER_MAP = dict(zip(IMPORT_HEADERS, ("level1_external_id", "level1_name", "level2_external_id", "level2_name", "level3_external_id", "level3_name", "is_active", "shelf_flag", "business_unit"), strict=True))


class CategoryService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def list_page(
        self,
        page_params: PageParams,
        active_only: bool = False,
        level1_name: str | None = None,
        level2_name: str | None = None,
        level3_name: str | None = None,
        deduction_rate: Decimal | None = None,
        is_active: bool | None = None,
        business_unit: str | None = None,
    ) -> PageResult[Category]:
        filters: list[ColumnElement[bool]] = []
        for column, value in (
            (Category.level1_name, level1_name),
            (Category.level2_name, level2_name),
            (Category.level3_name, level3_name),
            (Category.business_unit, business_unit),
        ):
            if value and (keyword := value.strip()):
                filters.append(column.ilike(f"%{keyword}%"))
        if deduction_rate is not None:
            filters.append(Category.deduction_rate == deduction_rate)
        effective_active = is_active if is_active is not None else (True if active_only else None)
        if effective_active is not None:
            filters.append(Category.is_active.is_(effective_active))

        statement = select(Category).where(*filters).order_by(
            Category.source_type, Category.level1_name, Category.level2_name, Category.level3_name, Category.id
        )
        count_statement = select(func.count()).select_from(Category).where(*filters)
        total = int(await self.session.scalar(count_statement) or 0)
        items = list((await self.session.scalars(statement.offset((page_params.page - 1) * page_params.page_size).limit(page_params.page_size))).all())
        return PageResult(items=items, total=total, page=page_params.page, page_size=page_params.page_size)

    async def list_active_selection(self) -> list[Category]:
        statement = select(Category).where(Category.is_active.is_(True)).order_by(Category.source_type, Category.level1_name, Category.level2_name, Category.level3_name, Category.id)
        return list((await self.session.scalars(statement)).all())

    async def list_filter_options(
        self,
        level: Literal["LEVEL1", "LEVEL2", "LEVEL3"],
        keyword: str | None,
        offset: int,
        limit: int,
    ) -> tuple[list[Category], bool]:
        statement = select(Category).where(
            Category.source_type == "MALL_LEVEL3", Category.is_active.is_(True)
        )
        if keyword and (query := keyword.strip()):
            if level == "LEVEL1":
                statement = statement.where(Category.level1_name.ilike(f"%{query}%"))
            elif level == "LEVEL2":
                statement = statement.where(or_(
                    Category.level1_name.ilike(f"%{query}%"),
                    Category.level2_name.ilike(f"%{query}%"),
                ))
            else:
                statement = statement.where(or_(
                    Category.level1_name.ilike(f"%{query}%"),
                    Category.level2_name.ilike(f"%{query}%"),
                    Category.level3_name.ilike(f"%{query}%"),
                ))
        statement = statement.order_by(
            Category.level1_name, Category.level2_name, Category.level3_name, Category.id
        )
        if level == "LEVEL3":
            rows = list(
                (
                    await self.session.scalars(statement.offset(offset).limit(limit + 1))
                ).all()
            )
            return rows[:limit], len(rows) > limit
        categories = list((await self.session.scalars(statement)).all())
        seen: set[str | tuple[str, str]] = set()
        options: list[Category] = []
        for category in categories:
            key: str | tuple[str, str] = (
                category.level1_name
                if level == "LEVEL1"
                else (category.level1_name, category.level2_name)
            )
            if key not in seen:
                seen.add(key)
                options.append(category)
            if len(options) == limit:
                break
        page = options[offset : offset + limit]
        return page, len(options) > offset + limit

    async def get(self, category_id: uuid.UUID) -> Category:
        category = await self.session.get(Category, category_id)
        if category is None:
            raise AppError("CATEGORY_NOT_FOUND", "Category not found", 404)
        return category

    async def _external_id_conflict(
        self, payload: CategoryWriteRequest, exclude: uuid.UUID | None = None
    ) -> Category | None:
        if payload.level3_external_id is None:
            return None
        statement = select(Category).where(
            Category.source_type == payload.source_type,
            Category.level3_external_id == payload.level3_external_id,
        )
        if exclude:
            statement = statement.where(Category.id != exclude)
        return cast(Category | None, await self.session.scalar(statement))

    async def create(self, payload: CategoryWriteRequest, actor_id: uuid.UUID) -> Category:
        if await self._external_id_conflict(payload):
            raise AppError(
                "CATEGORY_EXTERNAL_ID_EXISTS",
                "This source type and level-3 external ID are already bound to a category",
                409,
            )
        category = Category(**payload.model_dump(), created_by=actor_id, updated_by=actor_id)
        async with transaction_scope(self.session):
            self.session.add(category)
            await self.session.flush()
        return category

    async def update(
        self, category_id: uuid.UUID, payload: CategoryWriteRequest, actor_id: uuid.UUID
    ) -> Category:
        category = await self.get(category_id)
        if await self._external_id_conflict(payload, category_id):
            raise AppError(
                "CATEGORY_EXTERNAL_ID_EXISTS",
                "This source type and level-3 external ID are already bound to a category",
                409,
            )
        async with transaction_scope(self.session):
            for field, value in payload.model_dump().items():
                setattr(category, field, value)
            category.updated_by = actor_id
            await self.session.flush()
        return category

    async def delete(self, category_id: uuid.UUID) -> None:
        category = await self.get(category_id)
        used = await self.session.scalar(
            select(func.count()).select_from(Product).where(Product.category_id == category_id)
        )
        if used:
            raise AppError(
                "CATEGORY_IN_USE", "Category is referenced by products and cannot be deleted", 409
            )
        async with transaction_scope(self.session):
            await self.session.delete(category)

    def template(self) -> bytes:
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "类目导入"
        sheet.append(IMPORT_HEADERS)
        sheet.freeze_panes = "A2"
        output = BytesIO()
        book.save(output)
        return output.getvalue()

    async def import_xlsx(
        self, filename: str, content: bytes, deduction_rate_percent: str, actor_id: uuid.UUID
    ) -> CategoryImportResponse:
        if not filename.lower().endswith(".xlsx"):
            raise AppError(
                "CATEGORY_IMPORT_FILE_TYPE_INVALID", "Only .xlsx files are supported", 422
            )
        if not content:
            raise AppError("CATEGORY_IMPORT_FILE_EMPTY", "The import file is empty", 422)
        try:
            percent = Decimal(deduction_rate_percent)
            ratio = percent / Decimal("100")
            if not percent.is_finite():
                raise ValueError
            exponent = ratio.as_tuple().exponent
            if (
                percent < Decimal("0")
                or percent > Decimal("100")
                or not isinstance(exponent, int)
                or exponent < -4
            ):
                raise ValueError
            ratio = ratio.quantize(Decimal("0.0001"))
        except (InvalidOperation, ValueError):
            raise AppError("CATEGORY_IMPORT_DEDUCTION_RATE_INVALID", "deduction_rate_percent must be 0..100 and precise to 0.01%", 422)
        try:
            sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
            if sheet is None:
                raise ValueError("no worksheet")
            headers = tuple(
                str(cell.value or "").strip() for cell in next(sheet.iter_rows(max_row=1))
            )
        except Exception as exc:
            raise AppError(
                "CATEGORY_IMPORT_FILE_INVALID", "Unable to parse .xlsx file", 422
            ) from exc
        if len(headers) != len(IMPORT_HEADERS) or set(headers) != set(IMPORT_HEADERS):
            raise AppError(
                "CATEGORY_IMPORT_HEADER_INVALID",
                f"Required headers: {', '.join(IMPORT_HEADERS)}",
                422,
            )
        errors: list[CategoryImportError] = []
        rows: list[tuple[int, CategoryWriteRequest]] = []
        seen_external_ids: set[tuple[str, str]] = set()
        for number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None and str(value).strip() for value in cells):
                continue
            source = {header: value for header, value in zip(headers, cells, strict=True)}
            data: dict[str, object | None] = {target: (None if source[header] is None else str(source[header]).strip()) for header, target in IMPORT_HEADER_MAP.items()}
            data["source_type"] = "MALL_LEVEL3"
            data["deduction_rate"] = ratio
            for key in ("level1_external_id", "level2_external_id", "level3_external_id"):
                value = source[next(header for header, target in IMPORT_HEADER_MAP.items() if target == key)]
                if isinstance(value, float) and not value.is_integer():
                    errors.append(CategoryImportError(row_number=number, field=key, value=str(value), reason="External ID must be an integer or text")); continue
                if isinstance(value, float) and value.is_integer(): data[key] = str(int(value))
                elif isinstance(value, int): data[key] = str(value)
            if data["is_active"] in (True, "TRUE", "true", "1", "是"):
                data["is_active"] = True
            elif data["is_active"] in (False, "FALSE", "false", "0", "否"):
                data["is_active"] = False
            else:
                errors.append(CategoryImportError(row_number=number, field="is_active", value=str(data["is_active"]), reason="Invalid active flag")); continue
            try:
                payload = CategoryWriteRequest.model_validate(data)
            except Exception as exc:
                errors.append(CategoryImportError(row_number=number, reason=str(exc)))
                continue
            external_key = (
                (payload.source_type, payload.level3_external_id)
                if payload.level3_external_id is not None
                else None
            )
            if external_key is not None and external_key in seen_external_ids:
                errors.append(
                    CategoryImportError(
                        row_number=number,
                        field="level3_external_id",
                        value=payload.level3_external_id,
                        reason="Duplicate source type and level-3 external ID in Excel",
                    )
                )
                continue
            if external_key is not None:
                seen_external_ids.add(external_key)
            rows.append((number, payload))
        if errors:
            return CategoryImportResponse(
                total=len(rows) + len(errors),
                success=0,
                skipped=0,
                failed=len(errors),
                errors=errors,
            )
        existing = {
            (item.source_type, item.level3_external_id): item
            async for item in await self.session.stream_scalars(select(Category))
            if item.level3_external_id is not None
        }
        inserts: list[Category] = []
        skipped = 0
        for number, payload in rows:
            external_key = (
                (payload.source_type, payload.level3_external_id)
                if payload.level3_external_id is not None
                else None
            )
            prior = existing.get(external_key) if external_key is not None else None
            if prior:
                comparable = payload.model_dump()
                same = all(getattr(prior, name) == value for name, value in comparable.items())
                if same:
                    skipped += 1
                    continue
                errors.append(
                    CategoryImportError(
                        row_number=number,
                        field="level3_external_id",
                        value=payload.level3_external_id,
                        reason="This source type and level-3 external ID are already bound to different data",
                    )
                )
            else:
                inserts.append(
                    Category(**payload.model_dump(), created_by=actor_id, updated_by=actor_id)
                )
        if errors:
            return CategoryImportResponse(
                total=len(rows), success=0, skipped=0, failed=len(errors), errors=errors
            )
        try:
            async with transaction_scope(self.session):
                self.session.add_all(inserts)
                await self.session.flush()
        except IntegrityError as exc:
            raise AppError(
                "CATEGORY_IMPORT_CONFLICT", "Import conflicts with existing category data", 409
            ) from exc
        return CategoryImportResponse(
            total=len(rows), success=len(inserts), skipped=skipped, failed=0, errors=[]
        )
