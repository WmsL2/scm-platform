"""Build a reusable Product Master workbook containing only failed import rows."""

from __future__ import annotations

import copy
import posixpath
import shutil
import zipfile
from pathlib import Path, PurePosixPath
from typing import cast
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.catalog.application.excel_images import dispimg_image_id
from app.modules.catalog.infrastructure.models import ProductImportRow

_CONTENT_TYPE_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_MAIN_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CELL_IMAGE_REL_TYPE = "http://www.wps.cn/officeDocument/2020/cellImage"
_CELL_IMAGE_CONTENT_TYPE = "application/vnd.wps-officedocument.cellimage+xml"
_PRODUCT_MASTER_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1] / "resources" / "product-master-template.xlsx"
)


def export_failed_rows_workbook(
    headers: tuple[str, ...],
    rows: list[ProductImportRow],
    destination: Path,
    source_workbook: Path | None = None,
) -> None:
    """Write failed rows in the approved template order and retain referenced cell images."""
    workbook = load_workbook(_PRODUCT_MASTER_TEMPLATE_PATH, data_only=False)
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    if tuple(cell.value for cell in sheet[1][: len(headers)]) != headers:
        workbook.close()
        raise ValueError("The Product Master template headers do not match the approved headers")
    if sheet.max_row > 2:
        sheet.delete_rows(3, sheet.max_row - 2)

    template_row_height = sheet.row_dimensions[2].height
    error_sheet = workbook.create_sheet("错误说明")
    error_sheet.append(["导出行", "原Excel行", "错误原因"])
    _format_header(error_sheet, 3)
    error_sheet.column_dimensions["A"].width = 12
    error_sheet.column_dimensions["B"].width = 14
    error_sheet.column_dimensions["C"].width = 80

    referenced_image_ids: set[str] = set()
    for export_row_number, row in enumerate(rows, start=2):
        for column_number, header in enumerate(headers, start=1):
            value = row.source_data.get(header)
            if isinstance(value, str) and value.startswith("="):
                value = _translate_formula(
                    value,
                    column_number=column_number,
                    source_row_number=row.source_row_number,
                    export_row_number=export_row_number,
                )
            cell = sheet.cell(export_row_number, column_number, value)
            if export_row_number > 2:
                _copy_cell_style(cast(Cell, sheet.cell(2, column_number)), cell)
            if header == "图片":
                image_id = dispimg_image_id(value)
                if image_id is not None:
                    referenced_image_ids.add(image_id)
        sheet.row_dimensions[export_row_number].height = template_row_height
        error_sheet.append(
            [export_row_number, row.source_row_number, row.error_message or "校验不通过"]
        )

    workbook.save(destination)
    workbook.close()
    if source_workbook is not None and referenced_image_ids:
        _inject_referenced_cell_images(source_workbook, destination, referenced_image_ids)


def _format_header(sheet: Worksheet, column_count: int) -> None:
    for cell in sheet[1][:column_count]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="409EFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _copy_cell_style(source: Cell, target: Cell) -> None:
    target._style = copy.copy(source._style)  # type: ignore[attr-defined]


def _translate_formula(
    formula: str,
    *,
    column_number: int,
    source_row_number: int,
    export_row_number: int,
) -> str:
    column = get_column_letter(column_number)
    try:
        return cast(
            str,
            Translator(
                formula,
                origin=f"{column}{source_row_number}",
            ).translate_formula(f"{column}{export_row_number}"),
        )
    except (TypeError, ValueError):
        return formula


def _inject_referenced_cell_images(
    source_workbook: Path,
    destination: Path,
    image_ids: set[str],
) -> None:
    temporary = destination.with_name(f"{destination.stem}-with-images{destination.suffix}")
    try:
        with (
            zipfile.ZipFile(source_workbook) as source,
            zipfile.ZipFile(destination) as generated,
        ):
            selected = _selected_cell_image_parts(source, image_ids)
            if selected is None:
                return
            cell_images_xml, cell_image_relationships_xml, media_paths = selected
            workbook_relationships = _workbook_relationships_with_cell_images(generated)
            content_types = _content_types_with_cell_images(
                source,
                generated,
                media_paths,
            )
            replacements = {
                "[Content_Types].xml": content_types,
                "xl/_rels/workbook.xml.rels": workbook_relationships,
                "xl/cellimages.xml": cell_images_xml,
                "xl/_rels/cellimages.xml.rels": cell_image_relationships_xml,
            }
            with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as output:
                for entry in generated.infolist():
                    if entry.filename in replacements or entry.filename in media_paths:
                        continue
                    output.writestr(entry, generated.read(entry.filename))
                for name, content in replacements.items():
                    output.writestr(name, content)
                for media_path in media_paths:
                    with source.open(media_path) as source_file, output.open(
                        source.getinfo(media_path), "w"
                    ) as target_file:
                        shutil.copyfileobj(source_file, target_file, length=1024 * 1024)
        temporary.replace(destination)
    except (KeyError, zipfile.BadZipFile, ElementTree.ParseError):
        temporary.unlink(missing_ok=True)


def _selected_cell_image_parts(
    source: zipfile.ZipFile,
    image_ids: set[str],
) -> tuple[bytes, bytes, set[str]] | None:
    cell_images = ElementTree.fromstring(source.read("xl/cellimages.xml"))
    relationships = ElementTree.fromstring(source.read("xl/_rels/cellimages.xml.rels"))
    selected_pictures = []
    relationship_ids: set[str] = set()
    for cell_image in list(cell_images):
        properties = cell_image.find(f".//{{{_DRAWING_NS}}}cNvPr")
        blip = cell_image.find(f".//{{{_MAIN_NS}}}blip")
        if properties is None or blip is None or properties.attrib.get("name") not in image_ids:
            continue
        relationship_id = blip.attrib.get(f"{{{_REL_NS}}}embed")
        if relationship_id:
            selected_pictures.append(copy.deepcopy(cell_image))
            relationship_ids.add(relationship_id)
    if not selected_pictures:
        return None

    selected_relationships = []
    media_paths: set[str] = set()
    for relationship in relationships.findall(f"{{{_PACKAGE_REL_NS}}}Relationship"):
        if relationship.attrib.get("Id") not in relationship_ids:
            continue
        media_path = _media_path(relationship.attrib.get("Target"))
        if media_path is None:
            continue
        source.getinfo(media_path)
        selected_relationships.append(copy.deepcopy(relationship))
        media_paths.add(media_path)
    if len(selected_relationships) != len(relationship_ids):
        return None

    cell_images.clear()
    cell_images.extend(selected_pictures)
    relationships.clear()
    relationships.extend(selected_relationships)
    return (
        ElementTree.tostring(cell_images, encoding="utf-8", xml_declaration=True),
        ElementTree.tostring(relationships, encoding="utf-8", xml_declaration=True),
        media_paths,
    )


def _workbook_relationships_with_cell_images(generated: zipfile.ZipFile) -> bytes:
    path = "xl/_rels/workbook.xml.rels"
    root = ElementTree.fromstring(generated.read(path))
    relationships = root.findall(f"{{{_PACKAGE_REL_NS}}}Relationship")
    if not any(item.attrib.get("Type") == _CELL_IMAGE_REL_TYPE for item in relationships):
        ids = {item.attrib.get("Id") for item in relationships}
        index = 1
        while f"rId{index}" in ids:
            index += 1
        ElementTree.SubElement(
            root,
            f"{{{_PACKAGE_REL_NS}}}Relationship",
            {"Id": f"rId{index}", "Type": _CELL_IMAGE_REL_TYPE, "Target": "cellimages.xml"},
        )
    return cast(bytes, ElementTree.tostring(root, encoding="utf-8", xml_declaration=True))


def _content_types_with_cell_images(
    source: zipfile.ZipFile,
    generated: zipfile.ZipFile,
    media_paths: set[str],
) -> bytes:
    root = ElementTree.fromstring(generated.read("[Content_Types].xml"))
    source_root = ElementTree.fromstring(source.read("[Content_Types].xml"))
    overrides = {
        item.attrib.get("PartName")
        for item in root.findall(f"{{{_CONTENT_TYPE_NS}}}Override")
    }
    if "/xl/cellimages.xml" not in overrides:
        ElementTree.SubElement(
            root,
            f"{{{_CONTENT_TYPE_NS}}}Override",
            {"PartName": "/xl/cellimages.xml", "ContentType": _CELL_IMAGE_CONTENT_TYPE},
        )

    extensions = {
        PurePosixPath(path).suffix.removeprefix(".").lower() for path in media_paths
    }
    existing_defaults = {
        item.attrib.get("Extension", "").lower()
        for item in root.findall(f"{{{_CONTENT_TYPE_NS}}}Default")
    }
    for item in source_root.findall(f"{{{_CONTENT_TYPE_NS}}}Default"):
        extension = item.attrib.get("Extension", "").lower()
        if extension in extensions and extension not in existing_defaults:
            root.append(copy.deepcopy(item))
            existing_defaults.add(extension)

    selected_parts = {f"/{path}" for path in media_paths}
    for item in source_root.findall(f"{{{_CONTENT_TYPE_NS}}}Override"):
        part_name = item.attrib.get("PartName")
        if part_name in selected_parts and part_name not in overrides:
            root.append(copy.deepcopy(item))
            overrides.add(part_name)
    return cast(bytes, ElementTree.tostring(root, encoding="utf-8", xml_declaration=True))


def _media_path(target: str | None) -> str | None:
    if not target:
        return None
    candidate = target.replace("\\", "/")
    candidate = candidate[1:] if candidate.startswith("/") else candidate
    if not candidate.startswith("xl/"):
        candidate = f"xl/{candidate}"
    normalized = posixpath.normpath(candidate)
    if not normalized.startswith("xl/media/") or ".." in PurePosixPath(normalized).parts:
        return None
    return normalized
