# ruff: noqa: E501,E701,E702
# mypy: ignore-errors
import asyncio
import colorsys
import logging
from decimal import ROUND_HALF_UP, Decimal, localcontext
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import xlsxwriter
from openpyxl import load_workbook
from openpyxl.styles.colors import Color
from PIL import Image as PillowImage
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.contracts import AppError
from app.infrastructure.adapters import get_object_storage
from app.modules.catalog.application.media import local_media_storage_key
from app.modules.catalog.application.product_export_columns import (
    PRICE_EXPORT_COLUMN_KEYS,
    PRODUCT_EXPORT_COLUMNS,
)
from app.modules.catalog.domain.lifecycle import ProductStatus
from app.modules.catalog.infrastructure.models import Product
from app.modules.supplier.domain.rules import CooperationStatus
from app.modules.supplier.infrastructure.models import Supplier

logger = logging.getLogger(__name__)
_MAX_IMAGE_DIMENSION = 80
_IMAGE_ROW_HEIGHT = 60
_IMAGE_COLUMN_WIDTH = 16
_PRODUCT_MASTER_TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "resources" / "product-master-template.xlsx"
_THEME_COLOR_NAMES = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink")
_DRAWING_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_PRICE_EXPORT_QUANTUM = Decimal("0.01")
_MAX_EXACT_CENTS_AS_EXCEL_NUMBER = Decimal("10000000000000")
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_WPS_CELL_IMAGE_NS = "http://www.wps.cn/officeDocument/2017/etCustomData"
_SPREADSHEET_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_WPS_CELL_IMAGE_RELATIONSHIP = "http://www.wps.cn/officeDocument/2020/cellImage"
_WPS_CELL_IMAGE_CONTENT_TYPE = "application/vnd.wps-officedocument.cellimage+xml"


def _export_price(value: object) -> Decimal:
    amount = Decimal(str(value))
    with localcontext() as context:
        context.prec = max(65, len(amount.as_tuple().digits) + 2)
        return amount.quantize(_PRICE_EXPORT_QUANTUM, rounding=ROUND_HALF_UP)


def _template_color(color: Color, theme_colors: dict[str, str]) -> str | None:
    if color.type == "rgb" and isinstance(color.rgb, str):
        return f"#{color.rgb[-6:]}"
    if color.type != "theme" or color.theme is None:
        return None
    base = theme_colors.get(_THEME_COLOR_NAMES[color.theme])
    if base is None:
        return None
    tint = color.tint or 0
    if not tint:
        return f"#{base}"
    red, green, blue = (int(base[index:index + 2], 16) / 255 for index in (0, 2, 4))
    hue, lightness, saturation = colorsys.rgb_to_hls(red, green, blue)
    lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
    return "#%02X%02X%02X" % tuple(round(channel * 255) for channel in colorsys.hls_to_rgb(hue, lightness, saturation))


def _template_header_styles(columns: list) -> tuple[list[dict], float | None]:
    template = load_workbook(_PRODUCT_MASTER_TEMPLATE_PATH, read_only=False)
    try:
        sheet = template.active
        theme = ElementTree.fromstring(template.loaded_theme)
        scheme = theme.find(f".//{_DRAWING_NS}clrScheme")
        theme_colors = {
            element.tag.removeprefix(_DRAWING_NS): element[0].get("val", element[0].get("lastClr", ""))
            for element in scheme if len(element)
        } if scheme is not None else {}
        headers = {cell.value: cell for cell in sheet[1]}
        styles = []
        for column in columns:
            cell = headers[column.header]
            style = {
                "font_name": cell.font.name or "宋体",
                "font_size": cell.font.sz or 11,
                "bold": bool(cell.font.bold),
                "align": cell.alignment.horizontal or "center",
                "valign": "vcenter",
                "text_wrap": bool(cell.alignment.wrap_text),
                "border": 1,
            }
            fill = _template_color(cell.fill.fgColor, theme_colors)
            if cell.fill.patternType == "solid" and fill:
                style.update({"pattern": 1, "fg_color": fill})
            styles.append(style)
        return styles, sheet.row_dimensions[1].height
    finally:
        template.close()


class ProductExportService:
    def __init__(self, session: AsyncSession) -> None: self.session = session
    async def export(self, product_ids: list, column_keys: list[str], can_export_disabled: bool) -> bytes:
        rows = []
        for offset in range(0, len(product_ids), 500):
            rows.extend((await self.session.execute(select(Product, Supplier).join(Supplier).where(Product.id.in_(product_ids[offset:offset + 500])))).all())
        by_id = {product.id: (product, supplier) for product, supplier in rows}
        selected = [by_id.get(product_id) for product_id in product_ids]
        if len(by_id) != len(product_ids) or any(item is None for item in selected):
            raise AppError("PRODUCT_EXPORT_SELECTION_STALE", "所选商品已发生变化，请刷新列表后重新选择", 409)
        for product, supplier in selected:
            if (product.status == ProductStatus.DISABLED and not can_export_disabled) or (product.status == ProductStatus.ACTIVE and (supplier.is_deleted or supplier.cooperation_status != CooperationStatus.NORMAL)):
                raise AppError("PRODUCT_EXPORT_SELECTION_STALE", "所选商品已发生变化，请刷新列表后重新选择", 409)
        columns = [column for column in PRODUCT_EXPORT_COLUMNS if column.key in column_keys]
        image_bytes_by_reference = await self._read_export_images(columns, selected)
        return await asyncio.to_thread(self._build, columns, selected, image_bytes_by_reference)

    @staticmethod
    async def _read_export_images(columns, rows) -> dict[str, bytes | None]:
        if not any(column.key == "image_reference" for column in columns):
            return {}
        storage = get_object_storage()
        image_bytes_by_reference: dict[str, bytes | None] = {}
        references = dict.fromkeys(
            product.image_reference for product, _supplier in rows if product.image_reference
        )
        for reference in references:
            key = local_media_storage_key(reference)
            if key is None:
                image_bytes_by_reference[reference] = None
                continue
            try:
                image_bytes_by_reference[reference] = await storage.read(key)
            except Exception:
                logger.warning("Unable to read product export image: %s", reference, exc_info=True)
                image_bytes_by_reference[reference] = None
        return image_bytes_by_reference

    @staticmethod
    def _build(columns, rows, image_bytes_by_reference: dict[str, bytes | None]) -> bytes:
        output = BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        sheet = workbook.add_worksheet("商品主数据")
        sheet.freeze_panes(1, 0)
        header_styles, header_height = _template_header_styles(columns)
        if header_height is not None:
            sheet.set_row(0, header_height)
        for index, (column, style) in enumerate(zip(columns, header_styles)):
            sheet.write(0, index, column.header, workbook.add_format(style))
        image_column_index = next(
            (index for index, column in enumerate(columns) if column.key == "image_reference"),
            None,
        )
        cell_format = workbook.add_format({"border": 1})
        price_format = workbook.add_format({"border": 1, "num_format": "0.00"})
        image_format = workbook.add_format({"border": 1, "align": "center", "valign": "vcenter"})
        date_format = workbook.add_format({"border": 1, "num_format": "yyyy-mm-dd"})
        wps_cell_images: list[tuple[str, bytes, str]] = []
        for row_index, (product, supplier) in enumerate(rows, start=1):
            for column_index, column in enumerate(columns):
                if column.key == "image_reference":
                    continue
                value = column.value(product, supplier)
                if column.key == "listed_at" and value is not None:
                    sheet.write_datetime(row_index, column_index, value, date_format)
                elif column.key in PRICE_EXPORT_COLUMN_KEYS and value is not None:
                    price = _export_price(value)
                    if abs(price) < _MAX_EXACT_CENTS_AS_EXCEL_NUMBER:
                        sheet.write_number(row_index, column_index, float(price), price_format)
                    else:
                        sheet.write_string(row_index, column_index, f"{price:.2f}", cell_format)
                elif value is not None:
                    sheet.write(row_index, column_index, value, cell_format)
                else:
                    sheet.write_blank(row_index, column_index, None, cell_format)
            if image_column_index is not None:
                image_source = ProductExportService._prepare_embedded_image(
                    image_bytes_by_reference.get(product.image_reference)
                )
                if image_source is not None:
                    image_id = f"ID_{uuid4().hex.upper()}"
                    sheet.write_formula(
                        row_index,
                        image_column_index,
                        f'=_xlfn.DISPIMG("{image_id}",1)',
                        image_format,
                        f'=DISPIMG("{image_id}",1)',
                    )
                    wps_cell_images.append(
                        (
                            image_id,
                            image_source.getvalue(),
                            ProductExportService._wps_image_extension(image_source),
                        )
                    )
                    sheet.set_row(row_index, _IMAGE_ROW_HEIGHT)
                else:
                    sheet.write_blank(row_index, image_column_index, None, image_format)
        sheet.autofilter(0, 0, len(rows), max(len(columns) - 1, 0))
        sheet.set_column(0, max(len(columns) - 1, 0), 18)
        if image_column_index is not None:
            sheet.set_column_pixels(image_column_index, image_column_index, 90)
        workbook.close()
        return ProductExportService._add_wps_cell_images(output.getvalue(), wps_cell_images)

    @staticmethod
    def _wps_image_extension(content: BytesIO) -> str:
        with PillowImage.open(BytesIO(content.getvalue())) as image:
            return ".jpeg" if image.format == "JPEG" else ".png"

    @staticmethod
    def _add_wps_cell_images(
        content: bytes, images: list[tuple[str, bytes, str]]
    ) -> bytes:
        if not images:
            return content
        with ZipFile(BytesIO(content)) as source:
            content_types = ElementTree.fromstring(source.read("[Content_Types].xml"))
            workbook_relationships = ElementTree.fromstring(
                source.read("xl/_rels/workbook.xml.rels")
            )
            ElementTree.SubElement(
                content_types,
                f"{{{_CONTENT_TYPES_NS}}}Override",
                {
                    "PartName": "/xl/cellimages.xml",
                    "ContentType": _WPS_CELL_IMAGE_CONTENT_TYPE,
                },
            )
            relationship_ids = {
                relationship.attrib.get("Id", "")
                for relationship in workbook_relationships.findall(
                    f"{{{_PACKAGE_REL_NS}}}Relationship"
                )
            }
            relationship_number = 1
            while f"rId{relationship_number}" in relationship_ids:
                relationship_number += 1
            ElementTree.SubElement(
                workbook_relationships,
                f"{{{_PACKAGE_REL_NS}}}Relationship",
                {
                    "Id": f"rId{relationship_number}",
                    "Type": _WPS_CELL_IMAGE_RELATIONSHIP,
                    "Target": "cellimages.xml",
                },
            )
            cell_images = ProductExportService._wps_cell_images_xml(images)
            relationships = ProductExportService._wps_cell_image_relationships_xml(images)
            replacements = {
                "[Content_Types].xml": ElementTree.tostring(
                    content_types, encoding="utf-8", xml_declaration=True
                ),
                "xl/_rels/workbook.xml.rels": ElementTree.tostring(
                    workbook_relationships, encoding="utf-8", xml_declaration=True
                ),
            }
            output = BytesIO()
            with ZipFile(output, "w", ZIP_DEFLATED) as target:
                for info in source.infolist():
                    target_info = ZipInfo(info.filename)
                    target_info.date_time = info.date_time
                    target_info.external_attr = info.external_attr
                    target_info.extra = info.extra
                    target_info.comment = info.comment
                    target_info.compress_type = info.compress_type
                    target.writestr(
                        target_info, replacements.get(info.filename, source.read(info.filename))
                    )
                target.writestr("xl/cellimages.xml", cell_images)
                target.writestr("xl/_rels/cellimages.xml.rels", relationships)
                for index, (_image_id, image_content, extension) in enumerate(images, start=1):
                    target.writestr(f"xl/media/image{index}{extension}", image_content)
        return output.getvalue()

    @staticmethod
    def _wps_cell_images_xml(images: list[tuple[str, bytes, str]]) -> bytes:
        ElementTree.register_namespace("etc", _WPS_CELL_IMAGE_NS)
        ElementTree.register_namespace("xdr", _SPREADSHEET_DRAWING_NS)
        ElementTree.register_namespace("a", "http://schemas.openxmlformats.org/drawingml/2006/main")
        ElementTree.register_namespace("r", _OFFICE_REL_NS)
        drawing_main_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
        root = ElementTree.Element(f"{{{_WPS_CELL_IMAGE_NS}}}cellImages")
        for index, (image_id, _image_content, _extension) in enumerate(images, start=1):
            cell_image = ElementTree.SubElement(root, f"{{{_WPS_CELL_IMAGE_NS}}}cellImage")
            picture = ElementTree.SubElement(cell_image, f"{{{_SPREADSHEET_DRAWING_NS}}}pic")
            non_visual = ElementTree.SubElement(picture, f"{{{_SPREADSHEET_DRAWING_NS}}}nvPicPr")
            ElementTree.SubElement(
                non_visual,
                f"{{{_SPREADSHEET_DRAWING_NS}}}cNvPr",
                {"id": str(index + 1), "name": image_id},
            )
            non_visual_picture = ElementTree.SubElement(
                non_visual, f"{{{_SPREADSHEET_DRAWING_NS}}}cNvPicPr"
            )
            ElementTree.SubElement(
                non_visual_picture,
                f"{{{drawing_main_ns}}}picLocks",
                {"noChangeAspect": "1"},
            )
            blip_fill = ElementTree.SubElement(picture, f"{{{_SPREADSHEET_DRAWING_NS}}}blipFill")
            ElementTree.SubElement(
                blip_fill,
                f"{{{drawing_main_ns}}}blip",
                {f"{{{_OFFICE_REL_NS}}}embed": f"rId{index}"},
            )
            stretch = ElementTree.SubElement(blip_fill, f"{{{drawing_main_ns}}}stretch")
            ElementTree.SubElement(stretch, f"{{{drawing_main_ns}}}fillRect")
            shape_properties = ElementTree.SubElement(
                picture, f"{{{_SPREADSHEET_DRAWING_NS}}}spPr"
            )
            transform = ElementTree.SubElement(shape_properties, f"{{{drawing_main_ns}}}xfrm")
            ElementTree.SubElement(transform, f"{{{drawing_main_ns}}}off", {"x": "0", "y": "0"})
            ElementTree.SubElement(
                transform, f"{{{drawing_main_ns}}}ext", {"cx": "762000", "cy": "762000"}
            )
            geometry = ElementTree.SubElement(
                shape_properties, f"{{{drawing_main_ns}}}prstGeom", {"prst": "rect"}
            )
            ElementTree.SubElement(geometry, f"{{{drawing_main_ns}}}avLst")
        return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)

    @staticmethod
    def _wps_cell_image_relationships_xml(images: list[tuple[str, bytes, str]]) -> bytes:
        root = ElementTree.Element(f"{{{_PACKAGE_REL_NS}}}Relationships")
        for index, (_image_id, _image_content, extension) in enumerate(images, start=1):
            ElementTree.SubElement(
                root,
                f"{{{_PACKAGE_REL_NS}}}Relationship",
                {
                    "Id": f"rId{index}",
                    "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                    "Target": f"media/image{index}{extension}",
                },
            )
        return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)

    @staticmethod
    def _prepare_embedded_image(content: bytes | None) -> BytesIO | None:
        if not content:
            return None
        try:
            with PillowImage.open(BytesIO(content)) as image:
                image.load()
                if image.format in {"PNG", "JPEG"} and (
                    image.width >= _MAX_IMAGE_DIMENSION or image.height >= _MAX_IMAGE_DIMENSION
                ):
                    return BytesIO(content)
                prepared = image.convert("RGBA")
                if image.width >= _MAX_IMAGE_DIMENSION or image.height >= _MAX_IMAGE_DIMENSION:
                    prepared.thumbnail(
                        (_MAX_IMAGE_DIMENSION, _MAX_IMAGE_DIMENSION),
                        PillowImage.Resampling.LANCZOS,
                    )
                canvas = PillowImage.new("RGBA", (_MAX_IMAGE_DIMENSION, _MAX_IMAGE_DIMENSION))
                canvas.alpha_composite(
                    prepared,
                    ((canvas.width - prepared.width) // 2, (canvas.height - prepared.height) // 2),
                )
                output = BytesIO()
                canvas.save(output, format="PNG")
                output.seek(0)
                return output
        except Exception:
            logger.warning("Unable to decode product export image", exc_info=True)
            return None
