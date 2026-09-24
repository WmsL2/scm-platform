from __future__ import annotations

import hashlib
from copy import copy
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, cast
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.contracts import AppError
from app.core.transaction import transaction_scope
from app.infrastructure.adapters import ObjectStorage, get_object_storage
from app.modules.bid.domain.lifecycle import BidFileType
from app.modules.bid.infrastructure.models import BidProjectEvent, BidProjectFile
from app.modules.bid.schemas import BidProjectFileResponse
from app.modules.recommendation.domain.lifecycle import ensure_transition
from app.modules.recommendation.infrastructure.models import (
    RecommendationCandidate,
    RecommendationConfirmation,
    RecommendationExport,
)
from app.modules.recommendation.infrastructure.repository import RecommendationRepository
from app.modules.recommendation.template.schemas import RecommendationRunStatus

_DECIMAL_FIELDS = {
    "cost_price",
    "market_price",
    "jd_price",
    "agreement_price",
    "agreement_purchase_price",
    "profit",
    "jd_margin",
    "deduction_review",
    "gross_margin",
    "jd_self_operated_price",
    "positive_rating",
    "discount_rate",
    "price_inflation_rate",
}
_FACTORY_DIRECT_LABELS = {"PENDING": "待确认", "YES": "是", "NO": "否"}


class RecommendationExportService:
    """Build a traceable export from manually confirmed recommendation snapshots."""

    def __init__(self, session: AsyncSession, storage: ObjectStorage | None = None) -> None:
        self.session = session
        self.storage = storage or get_object_storage()
        self.repository = RecommendationRepository(session)

    async def export(
        self, project_id: UUID, run_id: UUID, actor_id: UUID
    ) -> BidProjectFileResponse:
        saved_key: str | None = None
        try:
            async with transaction_scope(self.session):
                project = await self.repository.free_project_for_update(project_id)
                if project is None:
                    raise AppError(
                        "FREE_RECOMMENDATION_PROJECT_NOT_FOUND", "自由推品项目不存在", 404
                    )
                run = await self.repository.run_for_update(run_id)
                if run is None or run.project_id != project.id:
                    raise AppError("RECOMMENDATION_RUN_NOT_FOUND", "推品任务不存在", 404)
                if run.status not in {
                    RecommendationRunStatus.CONFIRMED.value,
                    RecommendationRunStatus.EXPORTED.value,
                }:
                    raise AppError(
                        "RECOMMENDATION_EXPORT_NOT_ALLOWED",
                        "当前推品任务尚无可导出的已确认候选",
                        409,
                    )
                template_row = await self.repository.latest_template_mapping_for_update(project.id)
                if template_row is None:
                    raise AppError("RECOMMENDATION_TEMPLATE_NOT_FOUND", "推荐模板或映射不存在", 409)
                mapping, template_file = template_row
                if mapping.confirmed_by is None or mapping.confirmed_at is None:
                    raise AppError(
                        "RECOMMENDATION_TEMPLATE_MAPPING_NOT_CONFIRMED",
                        "请先确认最新推品模板映射",
                        409,
                    )
                rows = await self.repository.confirmed_candidates_for_export(run.id)
                if not rows:
                    raise AppError(
                        "RECOMMENDATION_EXPORT_CONFIRMATION_REQUIRED",
                        "至少确认一条推荐候选后才能导出",
                        409,
                    )
                content = self._build_workbook(
                    await self.storage.read(template_file.storage_key),
                    mapping.sheet_name,
                    mapping.header_row,
                    mapping.data_start_row,
                    mapping.mapping_json,
                    rows,
                )
                version = await self.repository.next_export_version(project.id)
                filename = f"{project.project_code}-自由推品结果-R{str(run.id)[:8]}-V{version}.xlsx"
                saved_key = await self.storage.save(
                    f"bid-projects/{project.id}/recommendation-exports/run-{run.id}-v{version}.xlsx",
                    content,
                )
                export_file = BidProjectFile(
                    project_id=project.id,
                    file_type=BidFileType.RECOMMENDATION_EXPORT.value,
                    version_no=version,
                    original_filename=filename,
                    storage_key=saved_key,
                    file_size=len(content),
                    sha256=hashlib.sha256(content).hexdigest(),
                    created_by=actor_id,
                )
                self.session.add(export_file)
                await self.session.flush()
                self.session.add(
                    RecommendationExport(
                        project_id=project.id,
                        run_id=run.id,
                        template_file_id=template_file.id,
                        template_mapping_id=mapping.id,
                        export_file_id=export_file.id,
                        mapping_snapshot={
                            "template_sha256": mapping.template_sha256,
                            "sheet_name": mapping.sheet_name,
                            "header_row": mapping.header_row,
                            "data_start_row": mapping.data_start_row,
                            "mapping_json": mapping.mapping_json,
                        },
                        version_no=version,
                        exported_by=actor_id,
                        exported_at=datetime.now(UTC).replace(tzinfo=None),
                    )
                )
                project.updated_by = actor_id
                self.session.add(
                    BidProjectEvent(
                        project_id=project.id,
                        actor_id=actor_id,
                        event_type="RECOMMENDATION_EXPORTED",
                        from_status=project.status,
                        to_status=project.status,
                        note=f"生成自由推品确认结果第{version}版（Run {run.id}）",
                    )
                )
                if run.status == RecommendationRunStatus.CONFIRMED.value:
                    ensure_transition(run.status, RecommendationRunStatus.EXPORTED)
                    run.status = RecommendationRunStatus.EXPORTED.value
                await self.session.flush()
                await self.session.refresh(export_file)
        except Exception:
            if saved_key is not None:
                await self.storage.delete(saved_key)
            raise
        return BidProjectFileResponse.model_validate(export_file, from_attributes=True)

    async def download_export(self, run_id: UUID, file_id: UUID) -> tuple[BidProjectFile, bytes]:
        file = await self.repository.export_file_for_run(run_id, file_id)
        if file is None:
            raise AppError("RECOMMENDATION_EXPORT_NOT_FOUND", "推荐导出文件不存在", 404)
        return file, await self.storage.read(file.storage_key)

    @classmethod
    def _build_workbook(
        cls,
        source: bytes,
        sheet_name: str,
        header_row: int,
        data_start_row: int,
        mapping_json: dict[str, str],
        rows: list[tuple[RecommendationCandidate, RecommendationConfirmation]],
    ) -> bytes:
        workbook = load_workbook(BytesIO(source), data_only=False)
        try:
            if sheet_name not in workbook.sheetnames:
                raise AppError(
                    "RECOMMENDATION_EXPORT_TEMPLATE_CHANGED", "推荐模板工作表已变更", 409
                )
            sheet = workbook[sheet_name]
            header_columns = {
                cls._text(cell.value): cell.column
                for cell in sheet[header_row]
                if cls._text(cell.value)
            }
            columns: dict[str, int] = {}
            for field, header in mapping_json.items():
                column = header_columns.get(header)
                if column is None:
                    raise AppError(
                        "RECOMMENDATION_EXPORT_TEMPLATE_CHANGED",
                        f"推荐模板缺少已确认的字段列：{header}",
                        409,
                    )
                if column in columns.values():
                    raise AppError(
                        "RECOMMENDATION_EXPORT_MAPPING_DUPLICATE",
                        "同一模板列不能映射多个导出字段",
                        409,
                    )
                columns[field] = column
            if not columns:
                raise AppError(
                    "RECOMMENDATION_EXPORT_MAPPING_EMPTY", "推荐模板未配置任何导出字段", 409
                )
            for offset, (candidate, confirmation) in enumerate(rows):
                target_row = data_start_row + offset
                if target_row > data_start_row:
                    cls._copy_template_row(sheet, data_start_row, target_row)
                for field, column in columns.items():
                    sheet.cell(target_row, column).value = cast(
                        Any, cls._value(field, candidate, confirmation)
                    )
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()

    @staticmethod
    def _copy_template_row(sheet: Any, source_row: int, target_row: int) -> None:
        for column in range(1, sheet.max_column + 1):
            source = sheet.cell(source_row, column)
            target = sheet.cell(target_row, column)
            if isinstance(source.value, str) and source.value.startswith("="):
                try:
                    target.value = Translator(
                        source.value, origin=source.coordinate
                    ).translate_formula(target.coordinate)
                except ValueError:
                    target.value = source.value
            else:
                target.value = source.value
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.alignment:
                target.alignment = copy(source.alignment)
            if source.protection:
                target.protection = copy(source.protection)
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height

    @classmethod
    def _value(
        cls,
        field: str,
        candidate: RecommendationCandidate,
        confirmation: RecommendationConfirmation,
    ) -> object | None:
        if field == "factory_direct":
            return _FACTORY_DIRECT_LABELS.get(confirmation.factory_direct, "待确认")
        if field == "supplier_name":
            return candidate.supplier_snapshot.get(field)
        if field not in _DECIMAL_FIELDS:
            return candidate.product_snapshot.get(
                field,
                candidate.price_snapshot.get(field, candidate.supplier_snapshot.get(field)),
            )
        value = candidate.price_snapshot.get(field)
        if value is None:
            return value
        try:
            return Decimal(str(value))
        except InvalidOperation as exc:
            raise AppError(
                "RECOMMENDATION_EXPORT_SNAPSHOT_INVALID", "推荐候选价格快照格式无效", 409
            ) from exc

    @staticmethod
    def _text(value: object) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None
