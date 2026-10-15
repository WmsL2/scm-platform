"""Independent renderer for the reviewed UNION_QUOTE_WHITE v5 customer deck."""

from __future__ import annotations

from collections.abc import Mapping
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from json import dumps
from math import ceil
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit
from zipfile import BadZipFile

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.exc import PackageNotFoundError
from pptx.util import Pt

from app.common.contracts import AppError
from app.modules.recommendation.application.ppt_template_manifest import (
    PptxSlotManifest,
    get_pptx_manifest,
)
from app.modules.recommendation.application.ppt_template_registry import (
    PptTemplateDefinition,
    resolve_ppt_template_asset,
)

_SOURCE_SHA256 = "6851d36ccc981df665601c6c172599c07a7ebf1e6ee3ea088e5d141084f36899"
_SOURCE_TXT_SHA256 = "3215768c7b472d9e34db7dbc0003fd8d8908710b7ae51d1b5ca07b3299490a19"
_SOURCE_STRUCTURE_FINGERPRINT = "5b85ad2dd8463bf985dace593a0cd2c5f01c85f28f5841b2481202f2f2eec632"
_ASSET_MARKER = "UNION_QUOTE_WHITE:5:template4:" + _SOURCE_SHA256
_FONT = "Microsoft YaHei"
_BLACK = RGBColor(0, 0, 0)  # type: ignore[no-untyped-call]
_RED = RGBColor(255, 0, 0)  # type: ignore[no-untyped-call]
_TRANSPARENT_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc````"
    b"\x00\x00\x00\x05\x00\x01\xa5\xf6E@\x00\x00\x00\x00IEND\xaeB`\x82"
)
_GEOMETRY = {
    "SCM_STATIC_DIVIDER": (5_721_719, 765_621, 0, 5_634_283),
    "SCM_FIELD_UNION_BODY": (294_640, 1_668_145, 5_331_460, 4_731_385),
    "SCM_FIELD_TITLE": (150_495, 235_585, 9_812_020, 755_650),
    "SCM_SLOT_IMAGE_01": (6_572_250, 1_419_225, 4_680_000, 4_680_000),
}
_TITLE_MAX_UNITS = 45
_BODY_LINE_UNITS = 34
_BODY_MAX_LINES = 16


def _error(message: str) -> AppError:
    return AppError("PPT_TEMPLATE_ASSET_STRUCTURE_INVALID", message, 409)


def _capacity_error(field: str, identifier: str) -> AppError:
    return AppError(
        "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED",
        f"UNION_QUOTE_WHITE v5 {field}超出固定槽位安全容量（商品标识 {identifier or '未知'}）",
        422,
    )


def _mapping(value: object) -> Mapping[str, object]:
    return value if isinstance(value, Mapping) else {}


def _mappings(value: object) -> list[Mapping[str, object]]:
    return [_mapping(item) for item in value] if isinstance(value, list) else []


def _text(values: Mapping[str, object], key: str) -> str:
    value = values.get(key)
    return str(value).strip() if value is not None else ""


def _unique(values: list[str]) -> str:
    return "；".join(dict.fromkeys(value for value in values if value))


def _decimal(value: object) -> Decimal | None:
    text = str(value).strip() if value is not None else ""
    if not text:
        return None
    try:
        return Decimal(text.replace(",", ""))
    except InvalidOperation:
        return None


def _money(value: object) -> str:
    amount = _decimal(value)
    return f"¥ {amount:.2f}" if amount is not None else ""


def _display_units(value: str) -> float:
    return sum(1.0 if ord(character) > 0xFF else 0.55 for character in value)


def _visual_lines(value: str) -> int:
    lines = value.splitlines() or [""]
    return sum(max(1, ceil(_display_units(line) / _BODY_LINE_UNITS)) for line in lines)


def _compact_multiline(value: str) -> str:
    """Keep every specification fragment while removing spreadsheet-only line breaks."""
    return "；".join(part.strip() for part in value.splitlines() if part.strip())


def _clean_reference(value: str) -> str:
    """Drop tracking query/fragment data that has no value in a printable proposal."""
    try:
        parsed = urlsplit(value)
    except ValueError:
        return value
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return value
    return urlunsplit((parsed.scheme, parsed.netloc, parsed.path, "", ""))


def _load(path: Path) -> Any:
    try:
        return Presentation(str(path))
    except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError, SyntaxError) as exc:
        raise _error("PPT 模板无法读取") from exc


def _structure_fingerprint(deck: Any) -> str:
    payload: list[list[tuple[int, int, int, int, int, str, bool]]] = []
    for slide in deck.slides:
        payload.append(
            [
                (
                    int(shape.shape_type),
                    int(shape.left),
                    int(shape.top),
                    int(shape.width),
                    int(shape.height),
                    str(shape.name),
                    bool(getattr(shape, "has_text_frame", False)),
                )
                for shape in slide.shapes
            ]
        )
    return sha256(
        dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _remove_notes(slide: Any) -> None:
    for relationship in list(slide.part.rels.values()):
        if "notesSlide" in relationship.reltype:
            slide.part.drop_rel(relationship.rId)


def _remove_slide(deck: Any, slide: Any) -> None:
    slide_id = deck.slides._sldIdLst[list(deck.slides).index(slide)]
    deck.part.drop_rel(slide_id.rId)
    deck.slides._sldIdLst.remove(slide_id)


def _replace_with_picture_slot(slide: Any, source: Any) -> None:
    picture = slide.shapes.add_picture(
        BytesIO(_TRANSPARENT_PNG), source.left, source.top, source.width, source.height
    )
    picture.name = "SCM_SLOT_IMAGE_01"
    parent = source._element.getparent()
    parent.remove(picture._element)
    parent.replace(source._element, picture._element)


def _manifest() -> PptxSlotManifest:
    manifest = get_pptx_manifest("UNION_QUOTE_WHITE", "5")
    if manifest is None:
        raise _error("模板4 v5 manifest 不可用")
    return manifest


def standardize_union_quote_white_v5(
    source: Path,
    specification: Path,
    destination: Path,
    *,
    expected_source_sha256: str | None = _SOURCE_SHA256,
    expected_specification_sha256: str | None = _SOURCE_TXT_SHA256,
    expected_source_fingerprint: str = _SOURCE_STRUCTURE_FINGERPRINT,
) -> None:
    """Create the private clean runtime asset from the approved original."""
    if (
        expected_source_sha256 is not None
        and sha256(source.read_bytes()).hexdigest() != expected_source_sha256
    ):
        raise _error("模板4原件 SHA-256 与已确认版本不匹配")
    if (
        expected_specification_sha256 is not None
        and sha256(specification.read_bytes()).hexdigest() != expected_specification_sha256
    ):
        raise _error("模板4回填说明 TXT 与已确认版本不匹配")
    deck = _load(source)
    if len(deck.slides) != 2 or _structure_fingerprint(deck) != expected_source_fingerprint:
        raise _error("模板4原件页面、形状或几何指纹不匹配")
    prototype = deck.slides[0]
    shapes = list(prototype.shapes)
    if len(shapes) != 4 or not shapes[1].has_text_frame or not shapes[2].has_text_frame:
        raise _error("模板4原件槽位结构不匹配")
    shapes[0].name = "SCM_STATIC_DIVIDER"
    shapes[1].name = "SCM_FIELD_UNION_BODY"
    shapes[2].name = "SCM_FIELD_TITLE"
    shapes[1].text = ""
    shapes[2].text = ""
    _replace_with_picture_slot(prototype, shapes[3])
    _remove_slide(deck, deck.slides[1])
    _remove_notes(prototype)
    deck.core_properties.identifier = _ASSET_MARKER
    destination.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(destination))
    validate_union_quote_white_v5_runtime(destination, _manifest())


def validate_union_quote_white_v5_runtime(path: Path, manifest: PptxSlotManifest) -> None:
    deck = _load(path)
    if manifest != _manifest() or len(deck.slides) != 1:
        raise _error("模板4运行资产版本或页数不匹配")
    if deck.core_properties.identifier != _ASSET_MARKER:
        raise _error("模板4运行资产来源标记不匹配")
    names = {shape.name: shape for shape in deck.slides[0].shapes}
    if set(_GEOMETRY) - set(names):
        raise _error("模板4运行资产缺少固定槽位")
    for name, geometry in _GEOMETRY.items():
        shape = names[name]
        if (shape.left, shape.top, shape.width, shape.height) != geometry:
            raise _error(f"模板4运行资产槽位几何不匹配：{name}")
    if names["SCM_SLOT_IMAGE_01"].shape_type != MSO_SHAPE_TYPE.PICTURE:
        raise _error("模板4运行资产图片槽位类型不匹配")


def verify_no_sample_media(source: Path, runtime: Path) -> None:
    original = _load(source)
    sample_hashes = {
        sha256(shape.image.blob).hexdigest()
        for shape in original.slides[1].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    }
    clean = _load(runtime)
    runtime_hashes = {
        sha256(shape.image.blob).hexdigest()
        for slide in clean.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    }
    if sample_hashes & runtime_hashes:
        raise _error("模板4运行资产仍包含原件示例商品图片")


def _clone_slide(deck: Any, source: Any) -> Any:
    destination = deck.slides.add_slide(source.slide_layout)
    # The source author deleted the layout's title/body/page placeholders from the
    # actual slide. ``add_slide`` recreates them, so remove those fresh instances
    # before copying the four reviewed source-slide objects.
    for shape in list(destination.shapes):
        shape._element.getparent().remove(shape._element)
    relationship_map: dict[str, str] = {}
    for relationship in source.part.rels.values():
        if "notesSlide" in relationship.reltype:
            continue
        target = relationship.target_ref if relationship.is_external else relationship.target_part
        relationship_map[relationship.rId] = destination.part.relate_to(
            target, relationship.reltype, is_external=relationship.is_external
        )
    for shape in source.shapes:
        element = deepcopy(shape.element)
        for node in element.iter():
            for key, value in list(node.attrib.items()):
                if value in relationship_map:
                    node.set(key, relationship_map[value])
        destination.shapes._spTree.insert_element_before(element, "p:extLst")
    return destination


def _style_paragraph(paragraph: Any, text: str, *, red: bool = False) -> None:
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.name = _FONT
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = _RED if red else _BLACK
    paragraph.line_spacing = 1.8
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)


def _set_title(shape: Any, value: str, identifier: str) -> None:
    if _display_units(value) > _TITLE_MAX_UNITS or "\n" in value:
        raise _capacity_error("商品名称", identifier)
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = value
    run.font.name = _FONT
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.underline = True
    run.font.color.rgb = _BLACK
    paragraph.line_spacing = 1.8
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)


def _set_body(shape: Any, lines: list[tuple[str, bool]], identifier: str) -> None:
    if sum(_visual_lines(text) for text, _red in lines) > _BODY_MAX_LINES:
        raise _capacity_error("商品配置/报价信息", identifier)
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for index, (text, red) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _style_paragraph(paragraph, text, red=red)


def _image_bytes(value: bytes | None) -> bytes:
    if not value:
        return _TRANSPARENT_PNG
    try:
        with Image.open(BytesIO(value)) as image:
            image.verify()
        return value
    except (UnidentifiedImageError, OSError, ValueError):
        return _TRANSPARENT_PNG


def _add_contained_picture(
    slide: Any, blob: bytes, left: int, top: int, width: int, height: int, name: str
) -> None:
    data = _image_bytes(blob)
    with Image.open(BytesIO(data)) as image:
        image_width, image_height = image.size
    scale = min(width / image_width, height / image_height)
    target_width, target_height = int(image_width * scale), int(image_height * scale)
    picture = slide.shapes.add_picture(
        BytesIO(data),
        left + (width - target_width) // 2,
        top + (height - target_height) // 2,
        target_width,
        target_height,
    )
    picture.name = name


def _set_images(slide: Any, blobs: list[bytes]) -> None:
    slots = [shape for shape in slide.shapes if shape.name == "SCM_SLOT_IMAGE_01"]
    if len(slots) != 1:
        raise _error("模板4运行页图片槽位不匹配")
    slot = slots[0]
    left, top, width, height = slot.left, slot.top, slot.width, slot.height
    slot._element.getparent().remove(slot._element)
    count = max(1, len(blobs))
    columns = 1 if count == 1 else 2
    rows = ceil(count / columns)
    gap = 76_200
    cell_width = (width - gap * (columns - 1)) // columns
    cell_height = (height - gap * (rows - 1)) // rows
    for index in range(count):
        row, column = divmod(index, columns)
        blob = blobs[index] if index < len(blobs) else _TRANSPARENT_PNG
        _add_contained_picture(
            slide,
            blob,
            left + column * (cell_width + gap),
            top + row * (cell_height + gap),
            cell_width,
            cell_height,
            f"SCM_RENDERED_IMAGE_{index + 1:02d}",
        )


def _single_entry(
    item: Mapping[str, object], images: Mapping[str, bytes]
) -> tuple[str, list[tuple[str, bool]], list[bytes], str]:
    product, prices, manual = (_mapping(item.get(key)) for key in ("product", "prices", "manual"))
    identifier = _text(product, "sku") or _text(product, "id")
    reference = _clean_reference(
        _text(product, "reference_url") or _text(product, "jd_same_product_url")
    )
    validity = _text(product, "price_validity") or _text(manual, "price_validity")
    delivery = _text(product, "shipping_courier") or _text(manual, "delivery_status")
    remark = _unique([_text(product, "remark"), _text(manual, "evidence")])
    lines = [
        (f"sku：{_text(product, 'sku')}", False),
        ("商品配置：", False),
        (_compact_multiline(_text(product, "product_specification")), False),
        ("", False),
        (f"京东价：{_money(prices.get('jd_price'))}", False),
        (f"协议价：{_money(prices.get('agreement_price'))}", True),
        (f"参考链接：{reference}", False),
        (f"价格有效期：{validity}", False),
        (f"配送物流：{delivery}", False),
        (f"备注：{remark}", False),
        (f"采销：{_text(product, 'purchasing_agent')}", False),
    ]
    blob = images.get(_text(product, "id"), _TRANSPARENT_PNG)
    return _text(product, "product_name") or "商品推荐", lines, [blob], identifier


def _sum_item_price(items: list[Mapping[str, object]], key: str) -> str:
    total = Decimal("0")
    for item in items:
        price = _decimal(_mapping(item.get("price_snapshot")).get(key))
        quantity = _decimal(item.get("quantity")) or Decimal("1")
        if price is None:
            return ""
        total += price * quantity
    return _money(total)


def _package_entry(
    package: Mapping[str, object], images: Mapping[str, bytes]
) -> tuple[str, list[tuple[str, bool]], list[bytes], str]:
    items = _mappings(package.get("items"))
    if not 2 <= len(items) <= 4:
        raise AppError("PPT_COMBINATION_INVALID", "组合商品数量必须为 2 至 4 件", 409)
    products = [_mapping(item.get("product_snapshot")) for item in items]
    identifiers = [_text(product, "id") for product in products]
    if any(not value for value in identifiers) or len(set(identifiers)) != len(identifiers):
        raise AppError("PPT_COMBINATION_INVALID", "组合商品必须为不同的已冻结商品", 409)
    configuration = "\n".join(
        f"{index}. {_text(product, 'product_name') or '\u672a\u547d\u540d\u5546\u54c1'} "
        f"× {_text(item, 'quantity') or '1'}"
        for index, (item, product) in enumerate(zip(items, products, strict=True), start=1)
    )
    sku = "、".join(_text(product, "sku") for product in products if _text(product, "sku"))
    links = _unique(
        [
            _text(product, "reference_url") or _text(product, "jd_same_product_url")
            for product in products
        ]
    )
    validity = _unique([_text(product, "price_validity") for product in products])
    delivery = _unique([_text(product, "shipping_courier") for product in products])
    remark = _unique(
        [_text(package, "reason")] + [_text(product, "remark") for product in products]
    )
    purchasing = _unique([_text(product, "purchasing_agent") for product in products])
    lines = [
        (f"sku：{sku}", False),
        ("商品配置：", False),
        (configuration, False),
        ("", False),
        (f"京东价：{_sum_item_price(items, 'jd_price')}", False),
        (f"协议价：{_sum_item_price(items, 'agreement_price')}", True),
        (f"参考链接：{links}", False),
        (f"价格有效期：{validity}", False),
        (f"配送物流：{delivery}", False),
        (f"备注：{remark}", False),
        (f"采销：{purchasing}", False),
    ]
    blobs = [images.get(identifier, _TRANSPARENT_PNG) for identifier in identifiers]
    return _text(package, "name") or "商品组合方案", lines, blobs, _text(package, "id") or sku


def _populate(slide: Any, entry: tuple[str, list[tuple[str, bool]], list[bytes], str]) -> None:
    title, lines, blobs, identifier = entry
    names = {shape.name: shape for shape in slide.shapes}
    try:
        _set_title(names["SCM_FIELD_TITLE"], title, identifier)
        _set_body(names["SCM_FIELD_UNION_BODY"], lines, identifier)
    except KeyError as exc:
        raise _error("模板4运行页缺少文字槽位") from exc
    _set_images(slide, blobs)


def render_union_quote_white_v5(
    source: Mapping[str, object],
    *,
    images: Mapping[str, bytes],
    template: PptTemplateDefinition,
    manifest: PptxSlotManifest,
) -> bytes:
    path = resolve_ppt_template_asset(template)
    validate_union_quote_white_v5_runtime(path, manifest)
    deck = _load(path)
    prototype = deck.slides[manifest.product_slide_index]
    entries = [_package_entry(package, images) for package in _mappings(source.get("packages"))]
    entries.extend(_single_entry(item, images) for item in _mappings(source.get("single_products")))
    if not entries:
        entries.append(("商品推荐方案", [("sku：", False), ("商品配置：", False)], [], "empty"))
    for entry in entries:
        slide = _clone_slide(deck, prototype)
        _populate(slide, entry)
    _remove_slide(deck, prototype)
    for slide in deck.slides:
        _remove_notes(slide)
    output = BytesIO()
    deck.save(output)
    return output.getvalue()
