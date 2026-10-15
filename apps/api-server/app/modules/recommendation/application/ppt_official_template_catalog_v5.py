"""In-place renderer and package sanitizer for the audited CATALOG_MINIMAL v5 deck.

The runtime asset is derived from the customer's empty first slide. It keeps
the two approved right-side decorative pictures, removes the sample slide and
historical package metadata, and replaces the left instruction rectangle with
one transparent ``contain`` image slot. Every output page is cloned from that
clean prototype.
"""

from __future__ import annotations

import logging
import posixpath
import re
from collections.abc import Mapping
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from json import dumps
from pathlib import Path, PurePosixPath
from re import split
from tempfile import NamedTemporaryFile
from typing import Any
from urllib.parse import urlparse
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from lxml import etree  # type: ignore[import-untyped]
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.exc import PackageNotFoundError
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.common.contracts import AppError
from app.modules.recommendation.application.ppt_layout_preview import (
    LayoutPreviewIssue,
    add_layout_preview_pages,
)
from app.modules.recommendation.application.ppt_template_manifest import (
    PptxSlotManifest,
    get_pptx_manifest,
)
from app.modules.recommendation.application.ppt_template_registry import (
    PptTemplateDefinition,
    resolve_ppt_template_asset,
)

logger = logging.getLogger(__name__)

_SOURCE_SHA256 = "2d9da8beee187e6a0f457a07c0447945f80605fedb15c5a413b38e2639062321"
_SOURCE_TXT_SHA256 = "561c7b08a569e343947c1ea61eedfc21ef7e72c5b06f20197677e5820c1f7d24"
_SOURCE_STRUCTURE_FINGERPRINT = "747997d01c664eb24c38b2aec0fcbed494869fda4fa6b36fd97a9aca67cd89dc"
_ASSET_MARKER = "CATALOG_MINIMAL:5:one-page:" + _SOURCE_SHA256
_FONT = "Microsoft YaHei"
_EMU_PER_POINT = 12_700
_MODEL_COLOUR = RGBColor(157, 144, 139)  # type: ignore[no-untyped-call]
_BLACK = RGBColor(0, 0, 0)  # type: ignore[no-untyped-call]
_DARK_GREY = RGBColor(89, 89, 89)  # type: ignore[no-untyped-call]
_RED = RGBColor(255, 0, 0)  # type: ignore[no-untyped-call]
_TRANSPARENT_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc````"
    b"\x00\x00\x00\x05\x00\x01\xa5\xf6E@\x00\x00\x00\x00IEND\xaeB`\x82"
)
_TEXT_GEOMETRY = (6.365, 0.692, 5.366, 6.600)
_IMAGE_GEOMETRY = (0.556, 0.678, 5.446, 6.153)
_DECORATION_GEOMETRIES = (
    (11.508, 0.789, 1.191, 1.226),
    (11.508, 2.971, 1.191, 1.303),
)
# WPS does not break a long ASCII model code at hyphens.  This width is lower
# than the physical text box width so the renderer inserts a complete second
# line before the code can enter the right-side decoration area.
_SAFE_TEXT_WIDTH_PT = 320.0
_TEXT_HEIGHT_SAFETY_PT = 6.0
_MIN_SPEC_LINES = 5
_MIN_SELLING_LINES = 6
_JD_LINK_LABEL = "京东链接："
_OFFICE_RELATIONSHIP_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
_APP_PROPERTIES_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
)
_EMAIL_PATTERN = re.compile(rb"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.IGNORECASE)
_FORBIDDEN_PART_PREFIXES = (
    "ppt/comments/",
    "ppt/commentauthors",
    "ppt/persons/",
    "ppt/threadedcomments/",
    "ppt/notesslides/",
    "ppt/notesmasters/",
    "ppt/tags/",
    "customxml/",
    "ppt/embeddings/",
    "ppt/activex/",
    "ppt/oleobjects/",
    "ppt/controls/",
)
_FORBIDDEN_RELATIONSHIP_NAMES = {
    "commentauthors",
    "comments",
    "custom-properties",
    "customxml",
    "notesmaster",
    "notesslide",
    "oleobject",
    "package",
    "person",
    "tags",
    "threadedcomment",
}
_CORE_PROPERTY_ALLOWLIST = {"identifier", "keywords"}
_APP_PROPERTY_ALLOWLIST = {
    "Application",
    "PresentationFormat",
    "Slides",
    "Notes",
    "HiddenSlides",
    "MMClips",
    "ScaleCrop",
    "LinksUpToDate",
    "SharedDoc",
    "HyperlinksChanged",
    "AppVersion",
}


def _error(message: str) -> AppError:
    return AppError("PPT_TEMPLATE_ASSET_STRUCTURE_INVALID", message, 409)


def _capacity_error(product_identifier: str, required: float, available: float) -> AppError:
    return AppError(
        "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED",
        "CATALOG_MINIMAL v5 动态文字超出客户固定页面安全容量"
        f"（商品标识 {product_identifier or '未知'}；预计 {required:.1f}pt，"
        f"可用 {available:.1f}pt，超出 {max(0.0, required - available):.1f}pt）",
        422,
    )


def _mapping(value: object) -> Mapping[str, object]:
    return value if isinstance(value, Mapping) else {}


def _mappings(value: object) -> list[Mapping[str, object]]:
    return [_mapping(item) for item in value] if isinstance(value, list) else []


def _text(values: Mapping[str, object], key: str) -> str:
    value = values.get(key)
    return str(value).strip() if value is not None else ""


def _raw_text(values: Mapping[str, object], key: str) -> str:
    value = values.get(key)
    return "" if value is None else str(value)


def _load(path: Path) -> Any:
    try:
        return Presentation(str(path))
    except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError, SyntaxError) as error:
        raise _error("运行模板无法读取") from error


def _structure_fingerprint(presentation: Any) -> str:
    payload = {
        "slides": len(presentation.slides),
        "shapes": [
            [(int(s.shape_type), s.left, s.top, s.width, s.height) for s in slide.shapes]
            for slide in presentation.slides
        ],
    }
    return sha256(dumps(payload, separators=(",", ":")).encode("utf-8")).hexdigest()


def _contract_digest(presentation: Any) -> str:
    def node_contract(node: Any) -> object:
        return (
            node.tag,
            sorted(node.attrib.items()),
            (node.text or "") if node.tag.endswith("}t") else None,
            [node_contract(child) for child in node],
        )

    payload = {
        "slides": [node_contract(slide._element) for slide in presentation.slides],
        "static_parts": sorted(
            (
                str(part.partname),
                dumps(node_contract(part._element))
                if hasattr(part, "_element")
                else sha256(part.blob).hexdigest(),
            )
            for part in presentation.part.package.iter_parts()
            if str(part.partname).startswith(
                ("/ppt/media/", "/ppt/theme/", "/ppt/slideMasters/", "/ppt/slideLayouts/")
            )
        ),
    }
    return sha256(dumps(payload, ensure_ascii=False).encode("utf-8")).hexdigest()


def _remove_notes(slide: Any) -> None:
    for rel in list(slide.part.rels.values()):
        if "notesSlide" in rel.reltype:
            slide.part.drop_rel(rel.rId)


def _remove_slide(presentation: Any, slide: Any) -> None:
    slide_id = presentation.slides._sldIdLst[list(presentation.slides).index(slide)]
    presentation.part.drop_rel(slide_id.rId)
    presentation.slides._sldIdLst.remove(slide_id)


def _is_forbidden_part(name: str) -> bool:
    normalized = name.lstrip("/").lower()
    return normalized == "docprops/custom.xml" or normalized == "ppt/vbaproject.bin" or any(
        normalized.startswith(prefix) for prefix in _FORBIDDEN_PART_PREFIXES
    )


def _relationship_source(rels_name: str) -> str:
    normalized_name = posixpath.normpath(rels_name.lstrip("/"))
    if normalized_name == "_rels/.rels":
        return ""
    rels_path = PurePosixPath(normalized_name)
    source_name = rels_path.name.removesuffix(".rels")
    source = rels_path.parent.parent / source_name
    return source.as_posix()


def _relationship_target(source: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    source_directory = posixpath.dirname(source)
    return posixpath.normpath(posixpath.join(source_directory, target))


def _remove_relationship_references(root: Any, removed_ids: set[str]) -> None:
    if not removed_ids:
        return
    relationship_prefix = f"{{{_OFFICE_RELATIONSHIP_NS}}}"
    for node in list(root.iter()):
        if node is root:
            continue
        if any(
            key.startswith(relationship_prefix) and value in removed_ids
            for key, value in node.attrib.items()
        ):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    for node in reversed(list(root.iter())):
        if node is root or len(node) or (node.text or "").strip() or node.attrib:
            continue
        if etree.QName(node).localname in {
            "custDataLst",
            "extLst",
            "notesMasterIdLst",
        }:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)


def _sanitize_catalog_package(path: Path) -> None:
    """Remove identity/history parts and every corresponding OOXML reference."""
    with ZipFile(path) as archive:
        infos = archive.infolist()
        payloads = {info.filename: archive.read(info.filename) for info in infos}
    removed_parts = {name for name in payloads if _is_forbidden_part(name)}
    removed_relationships: dict[str, set[str]] = {}
    for name, payload in list(payloads.items()):
        if not name.endswith(".rels") or name in removed_parts:
            continue
        root = etree.fromstring(payload)
        source = _relationship_source(name)
        removed_ids: set[str] = set()
        for relationship in list(root):
            target_mode = relationship.get("TargetMode")
            relation_name = relationship.get("Type", "").rsplit("/", 1)[-1].lower()
            target = relationship.get("Target", "")
            target_part = _relationship_target(source, target) if target else ""
            if (
                target_mode == "External"
                or relation_name in _FORBIDDEN_RELATIONSHIP_NAMES
                or _is_forbidden_part(target_part)
            ):
                removed_ids.add(relationship.get("Id", ""))
                root.remove(relationship)
        if removed_ids:
            removed_relationships[source] = removed_ids
            payloads[name] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
    for source, removed_ids in removed_relationships.items():
        if not source or source not in payloads or not source.endswith(".xml"):
            continue
        root = etree.fromstring(payloads[source])
        _remove_relationship_references(root, removed_ids)
        payloads[source] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
    content_types = "[Content_Types].xml"
    if content_types in payloads:
        root = etree.fromstring(payloads[content_types])
        for override in list(root):
            if etree.QName(override).localname == "Override" and _is_forbidden_part(
                override.get("PartName", "")
            ):
                root.remove(override)
        payloads[content_types] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
    if "docProps/core.xml" in payloads:
        root = etree.fromstring(payloads["docProps/core.xml"])
        for child in list(root):
            if etree.QName(child).localname not in _CORE_PROPERTY_ALLOWLIST:
                root.remove(child)
        payloads["docProps/core.xml"] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
    if "docProps/app.xml" in payloads:
        root = etree.fromstring(payloads["docProps/app.xml"])
        for child in list(root):
            if etree.QName(child).localname not in _APP_PROPERTY_ALLOWLIST:
                root.remove(child)
        values = {etree.QName(child).localname: child for child in root}
        statistics = {
            "Slides": sum(
                bool(re.fullmatch(r"ppt/slides/slide\d+\.xml", name))
                for name in payloads
                if name not in removed_parts
            ),
            "Notes": 0,
            "HiddenSlides": 0,
            "MMClips": sum(
                name.startswith("ppt/media/") and not name.endswith("/")
                for name in payloads
                if name not in removed_parts
            ),
        }
        for name, value in statistics.items():
            node = values.get(name)
            if node is None:
                node = etree.SubElement(root, f"{{{_APP_PROPERTIES_NS}}}{name}")
            node.text = str(value)
        payloads["docProps/app.xml"] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(dir=path.parent, suffix=".pptx", delete=False) as handle:
        candidate = Path(handle.name)
    try:
        with ZipFile(candidate, "w", compression=ZIP_DEFLATED) as archive:
            for info in infos:
                if info.filename not in removed_parts:
                    archive.writestr(info, payloads[info.filename])
        candidate.replace(path)
    finally:
        candidate.unlink(missing_ok=True)


def _validate_catalog_package(path: Path) -> None:
    """Reject identity metadata, dangling relationships and unsafe package parts."""
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        payloads = {name: archive.read(name) for name in names if not name.endswith("/")}
    forbidden = sorted(name for name in names if _is_forbidden_part(name))
    if forbidden:
        raise _error("运行模板包含批注、备注、自定义元数据或嵌入对象")
    for name, payload in payloads.items():
        if name.endswith((".xml", ".rels")) and _EMAIL_PATTERN.search(payload):
            raise _error("运行模板包含邮箱形式的身份元数据")
    relationship_ids: dict[str, set[str]] = {}
    for name, payload in payloads.items():
        if not name.endswith(".rels"):
            continue
        root = etree.fromstring(payload)
        source = _relationship_source(name)
        ids: set[str] = set()
        for relationship in root:
            relation_name = relationship.get("Type", "").rsplit("/", 1)[-1].lower()
            if relationship.get("TargetMode") == "External":
                raise _error("运行模板包含未批准的外部关系")
            if relation_name in _FORBIDDEN_RELATIONSHIP_NAMES:
                raise _error("运行模板包含禁止保留的元数据关系")
            target = relationship.get("Target", "")
            if not target or _relationship_target(source, target) not in names:
                raise _error("运行模板包含悬空的 OOXML 关系")
            ids.add(relationship.get("Id", ""))
        relationship_ids[source] = ids
    relationship_prefix = f"{{{_OFFICE_RELATIONSHIP_NS}}}"
    for name, payload in payloads.items():
        if not name.endswith(".xml"):
            continue
        root = etree.fromstring(payload)
        expected_ids = relationship_ids.get(name, set())
        for node in root.iter():
            for key, value in node.attrib.items():
                if key.startswith(relationship_prefix) and value not in expected_ids:
                    raise _error("运行模板 XML 引用了不存在的关系")
    core = etree.fromstring(payloads["docProps/core.xml"])
    if {etree.QName(child).localname for child in core} - _CORE_PROPERTY_ALLOWLIST:
        raise _error("运行模板核心属性包含身份或编辑历史")
    app = etree.fromstring(payloads["docProps/app.xml"])
    app_values = {etree.QName(child).localname: child.text for child in app}
    slide_count = sum(
        bool(re.fullmatch(r"ppt/slides/slide\d+\.xml", name)) for name in names
    )
    expected = {
        "Slides": str(slide_count),
        "Notes": "0",
        "HiddenSlides": "0",
    }
    if any(app_values.get(key) != value for key, value in expected.items()):
        raise _error("运行模板应用统计与实际页面/备注不一致")
    presentation = etree.fromstring(payloads["ppt/presentation.xml"])
    namespace = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    if presentation.xpath(
        './/p:sldId[@show="0" or @show="false"]', namespaces=namespace
    ):
        raise _error("运行模板包含隐藏页面")
    if any(
        etree.fromstring(payload).get("show") in {"0", "false"}
        for name, payload in payloads.items()
        if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
    ):
        raise _error("运行模板包含隐藏页面")


def _set_run_style(
    run: Any, *, size: float, colour: RGBColor, bold: bool | None
) -> None:
    run.font.name = _FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    properties = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        node = properties.find(qn(tag))
        if node is None:
            from pptx.oxml.xmlchemy import OxmlElement

            node = OxmlElement(tag)
            properties.append(node)
        node.set("typeface", _FONT)


def _reset_paragraph(
    paragraph: Any,
    runs: tuple[tuple[str, RGBColor, bool | None], ...],
    *,
    size: float,
) -> None:
    paragraph.clear()
    paragraph.line_spacing = 1.5
    for value, colour, bold in runs:
        run = paragraph.add_run()
        run.text = value
        _set_run_style(run, size=size, colour=colour, bold=bold)


def _single_run(paragraph: Any, value: str, *, size: float, colour: RGBColor, bold: bool) -> None:
    _reset_paragraph(paragraph, ((value, colour, bold),), size=size)


def _picture_slot(slide: Any, source: Any) -> None:
    picture = slide.shapes.add_picture(
        BytesIO(_TRANSPARENT_PNG), source.left, source.top, source.width, source.height
    )
    picture.name = "SCM_SLOT_IMAGE_MAIN"
    parent = source._element.getparent()
    parent.remove(picture._element)
    parent.replace(source._element, picture._element)


def _set_hyperlink_theme_black(deck: Any) -> None:
    """Keep WPS from overriding the approved black link Run with theme blue."""
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for part in deck.part.package.iter_parts():
        if not str(part.partname).startswith("/ppt/theme/"):
            continue
        root = etree.fromstring(part.blob)
        for name in ("hlink", "folHlink"):
            colour = root.find(f".//a:clrScheme/a:{name}", namespaces=namespace)
            if colour is None:
                raise _error("原件主题缺少超链接颜色合同")
            for child in list(colour):
                colour.remove(child)
            value = etree.SubElement(colour, qn("a:srgbClr"))
            value.set("val", "000000")
        part._blob = etree.tostring(  # noqa: SLF001 - python-pptx has no theme write API
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )


def _manifest() -> PptxSlotManifest:
    manifest = get_pptx_manifest("CATALOG_MINIMAL", "5")
    if manifest is None:
        raise _error("v5 manifest 不可用")
    return manifest


def standardize_catalog_minimal_v5(
    source: Path,
    destination: Path,
    *,
    expected_source_fingerprint: str = _SOURCE_STRUCTURE_FINGERPRINT,
    expected_source_sha256: str | None = _SOURCE_SHA256,
) -> None:
    """Build the public-safe clean prototype without modifying customer files."""
    if (
        expected_source_sha256 is not None
        and sha256(source.read_bytes()).hexdigest() != expected_source_sha256
    ):
        raise _error("原件 SHA-256 与已确认正式版本不匹配")
    if expected_source_sha256 is not None:
        description = source.with_suffix(".txt")
        if (
            not description.is_file()
            or sha256(description.read_bytes()).hexdigest() != _SOURCE_TXT_SHA256
        ):
            raise _error("字体/回填说明 TXT 与已确认版本不匹配")
    deck = _load(source)
    if len(deck.slides) != 2 or _structure_fingerprint(deck) != expected_source_fingerprint:
        raise _error("原件页面、形状类型/数量或几何指纹不匹配")
    prototype = deck.slides[0]
    shapes = list(prototype.shapes)
    if (
        len(shapes) != 4
        or not shapes[1].has_text_frame
        or len(shapes[1].text_frame.paragraphs) != 21
    ):
        raise _error("原件空白商品页结构或段落数量不匹配")
    _remove_slide(deck, deck.slides[1])
    shapes = list(prototype.shapes)
    _picture_slot(prototype, shapes[0])
    content = shapes[1]
    content.name = "SCM_FIELD_CATALOG_CONTENT"
    content.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    content.text_frame.word_wrap = True
    paragraphs = content.text_frame.paragraphs
    _single_run(paragraphs[0], "{{model}}", size=18, colour=_MODEL_COLOUR, bold=True)
    _single_run(paragraphs[1], "{{product_name}}", size=18, colour=_BLACK, bold=False)
    _single_run(paragraphs[3], "规格参数：", size=9, colour=_BLACK, bold=False)
    _single_run(paragraphs[4], "{{specification}}", size=9, colour=_BLACK, bold=False)
    _reset_paragraph(
        paragraphs[9],
        (
            ("零  售  价：", _DARK_GREY, True),
            ("{{retail_price}}", _MODEL_COLOUR, True),
            ("元", _MODEL_COLOUR, True),
        ),
        size=9,
    )
    _reset_paragraph(
        paragraphs[10],
        (("团购价：", _RED, True), ("{{group_price}}", _RED, True), ("元", _RED, True)),
        size=9,
    )
    _single_run(paragraphs[12], "产品特点：", size=9, colour=_BLACK, bold=True)
    _single_run(paragraphs[13], "{{selling_points}}", size=9, colour=_BLACK, bold=False)
    _reset_paragraph(
        paragraphs[19],
        (("京东链接：", _BLACK, True), ("{{jd_link}}", _BLACK, True)),
        size=9,
    )
    for index in (2, 5, 6, 7, 8, 11, 14, 15, 16, 17, 18, 20):
        paragraphs[index].clear()
        paragraphs[index].line_spacing = 1.5
    shapes[2].name = "SCM_STATIC_DECORATION_01"
    shapes[3].name = "SCM_STATIC_DECORATION_02"
    _set_hyperlink_theme_black(deck)
    for slide in deck.slides:
        _remove_notes(slide)
    deck.core_properties.identifier = _ASSET_MARKER
    deck.core_properties.keywords = ""
    destination.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(destination))
    _sanitize_catalog_package(destination)
    sanitized = _load(destination)
    sanitized.core_properties.identifier = _ASSET_MARKER
    sanitized.core_properties.keywords = _contract_digest(sanitized)
    sanitized.save(str(destination))
    _sanitize_catalog_package(destination)
    validate_catalog_minimal_v5_runtime(destination, _manifest())


def _geometry_matches(shape: Any, expected: tuple[float, float, float, float]) -> bool:
    return all(
        abs(actual - Inches(value)) <= Inches(0.01)
        for actual, value in zip(
            (shape.left, shape.top, shape.width, shape.height), expected, strict=True
        )
    )


def _validate_runtime_structure(deck: Any, manifest: PptxSlotManifest) -> None:
    if (
        manifest != _manifest()
        or len(deck.slides) != 1
        or deck.core_properties.identifier != _ASSET_MARKER
        or deck.core_properties.keywords != _contract_digest(deck)
    ):
        raise _error("运行模板版本、页数或结构/Run 样式合同不匹配")
    shapes = {shape.name: shape for shape in deck.slides[0].shapes}
    required = {
        "SCM_SLOT_IMAGE_MAIN",
        "SCM_FIELD_CATALOG_CONTENT",
        "SCM_STATIC_DECORATION_01",
        "SCM_STATIC_DECORATION_02",
    }
    if len(deck.slides[0].shapes) != 4 or set(shapes) != required:
        raise _error("运行模板缺少、重复或增加了正式槽位")
    if shapes["SCM_SLOT_IMAGE_MAIN"].shape_type != MSO_SHAPE_TYPE.PICTURE:
        raise _error("商品主图槽类型不匹配")
    if not _geometry_matches(
        shapes["SCM_SLOT_IMAGE_MAIN"], _IMAGE_GEOMETRY
    ) or not _geometry_matches(shapes["SCM_FIELD_CATALOG_CONTENT"], _TEXT_GEOMETRY):
        raise _error("商品图片或文字槽几何不匹配")
    for index, expected in enumerate(_DECORATION_GEOMETRIES, start=1):
        if not _geometry_matches(shapes[f"SCM_STATIC_DECORATION_{index:02d}"], expected):
            raise _error("右侧静态装饰几何不匹配")
    content = shapes["SCM_FIELD_CATALOG_CONTENT"]
    if len(content.text_frame.paragraphs) != 21:
        raise _error("动态文字段落合同不匹配")
    tokens = {
        "{{model}}",
        "{{product_name}}",
        "{{specification}}",
        "{{retail_price}}",
        "{{group_price}}",
        "{{selling_points}}",
        "{{jd_link}}",
    }
    actual = {
        run.text
        for paragraph in content.text_frame.paragraphs
        for run in paragraph.runs
        if run.text.startswith("{{")
    }
    if actual != tokens:
        raise _error("动态字段 Token 合同不匹配")


def validate_catalog_minimal_v5_runtime(path: Path, manifest: PptxSlotManifest) -> None:
    _validate_runtime_structure(_load(path), manifest)
    _validate_catalog_package(path)


def verify_no_sample_media(source: Path, runtime: Path) -> None:
    """Allow only the two approved decorations and the transparent image slot."""
    source_deck = _load(source)
    allowed_hashes = {
        sha256(source_deck.slides[0].shapes[index].image.blob).digest() for index in (2, 3)
    }
    allowed_hashes.add(sha256(_TRANSPARENT_PNG).digest())
    sample_hash = sha256(source_deck.slides[1].shapes[4].image.blob).digest()
    with ZipFile(runtime) as archive:
        actual = {
            sha256(archive.read(name)).digest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        }
        names = set(archive.namelist())
    if not actual <= allowed_hashes or sample_hash in actual or any(
        name.startswith("ppt/slides/slide2.") for name in names
    ):
        raise _error("运行模板携带样例商品媒体、样例页或非授权媒体")


def _clone_slide(deck: Any, source: Any) -> Any:
    slide = deck.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    rels: dict[str, str] = {}
    for rel in source.part.rels.values():
        if "notesSlide" not in rel.reltype:
            rels[rel.rId] = slide.part.relate_to(
                rel.target_ref if rel.is_external else rel.target_part,
                rel.reltype,
                is_external=rel.is_external,
            )
    copied = deepcopy(source._element)
    relationship_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for node in copied.iter():
        for key, value in list(node.attrib.items()):
            if key.startswith(relationship_namespace) and value in rels:
                node.set(key, rels[value])
    slide._element.attrib.clear()
    slide._element.attrib.update(copied.attrib)
    for child in list(slide._element):
        slide._element.remove(child)
    for child in copied:
        slide._element.append(child)
    slide.__dict__.pop("shapes", None)
    return slide


def _estimated_width(value: str, size: float) -> float:
    return size * sum(
        0.30
        if character.isspace() or character in ".,:;!ilI|'"
        else 1.0
        if ord(character) >= 0x2E80
        else 0.90
        if character in "MWmw@#%&"
        else 0.68
        if character.isupper()
        else 0.56
        if character.isdigit()
        else 0.35
        if character in "-_()/[]{}"
        else 0.56
        for character in value
    )


def _line_count(value: str, size: float, *, first_line_prefix: str = "") -> int:
    """Return the exact number of paragraphs that the renderer will emit."""
    return len(_wrapped_chunks(value, size, first_line_prefix=first_line_prefix))


def _wrapped_chunks(
    value: str, size: float, *, first_line_prefix: str = ""
) -> list[str]:
    """Break even unspaced ASCII identifiers before the decoration boundary."""
    first_line_width = _SAFE_TEXT_WIDTH_PT - _estimated_width(first_line_prefix, size)
    if first_line_width <= 0:
        raise _error("固定文本已超出动态文字安全宽度")
    chunks: list[str] = []
    is_first_output_line = True
    for raw_line in value.splitlines() or [""]:
        if not raw_line:
            chunks.append("")
            is_first_output_line = False
            continue
        start = 0
        width = 0.0
        for index, character in enumerate(raw_line):
            advance = _estimated_width(character, size)
            allowed_width = (
                first_line_width if is_first_output_line else _SAFE_TEXT_WIDTH_PT
            )
            if index > start and width + advance > allowed_width:
                chunks.append(raw_line[start:index])
                is_first_output_line = False
                start = index
                width = 0.0
            width += advance
        chunks.append(raw_line[start:])
        is_first_output_line = False
    return chunks


def _logical_lines(value: str, *, semicolon_separator: bool = False) -> list[str]:
    lines = [line.strip() for line in value.splitlines() if line.strip()]
    if semicolon_separator and len(lines) == 1:
        lines = [item.strip() for item in split(r"[;；]", lines[0]) if item.strip()]
    return lines


def _paragraph_copy(paragraph: Any, value: str) -> Any:
    copied = deepcopy(paragraph._p)
    runs = copied.findall(qn("a:r"))
    if not runs:
        raise _error("运行模板动态段落缺少预留 Run")
    for run in runs[1:]:
        copied.remove(run)
    for child in list(copied):
        if child.tag in {qn("a:br"), qn("a:fld")}:
            copied.remove(child)
    text = runs[0].find(qn("a:t"))
    if text is None:
        raise _error("运行模板动态 Run 缺少文本节点")
    text.text = value
    return copied


def _fixed_paragraph_copy(paragraph: Any) -> Any:
    return deepcopy(paragraph._p)


def _replace_body_paragraphs(frame: Any, paragraphs: list[Any]) -> None:
    body = frame._txBody
    for node in list(body):
        if node.tag == qn("a:p"):
            body.remove(node)
    for paragraph in paragraphs:
        body.append(paragraph)


def _format_price(value: object) -> str:
    if value is None or not str(value).strip():
        return ""
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return ""
    if not number.is_finite() or number < 0:
        return ""
    # The approved price fields are amounts, so retain two decimal places.
    # In particular, Decimal("0") must remain a value rather than looking missing.
    return format(number, ".2f")


def _valid_link(value: str) -> str:
    if not value:
        return ""
    invalid_character = any(
        character.isspace() or ord(character) < 0x20 or ord(character) == 0x7F
        for character in value
    )
    if invalid_character:
        raise AppError("PPT_TEMPLATE_LINK_INVALID", "京东链接必须是完整的 http/https 地址", 422)
    try:
        parsed = urlparse(value)
        hostname = parsed.hostname
        _ = parsed.port
    except ValueError as error:
        raise AppError(
            "PPT_TEMPLATE_LINK_INVALID", "京东链接必须是完整的 http/https 地址", 422
        ) from error
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc or not hostname:
        raise AppError("PPT_TEMPLATE_LINK_INVALID", "京东链接必须是完整的 http/https 地址", 422)
    return value


def _visible_text_bottom(content: Any, *, high_lines: int, low_lines: int) -> int:
    """Estimate the bottom of the emitted text using its actual line contract."""
    return int(
        content.top
        + content.text_frame.margin_top
        + (high_lines * 18 * 1.5 + low_lines * 9 * 1.5) * _EMU_PER_POINT
    )


def _fill_text(
    content: Any,
    values: Mapping[str, str],
    product_identifier: str,
    *,
    slide_height: int,
    allow_overflow: bool = False,
) -> tuple[float, float] | None:
    templates = content.text_frame.paragraphs
    if len(templates) != 21:
        raise _error("运行模板动态段落合同不匹配")
    model_chunks = _wrapped_chunks(values["model"], 18)
    name_chunks = _wrapped_chunks(values["product_name"], 18)
    model_lines = len(model_chunks)
    name_lines = len(name_chunks)
    specifications = [
        chunk
        for value in (_logical_lines(values["specification"]) or [""])
        for chunk in _wrapped_chunks(value, 9)
    ]
    specification_lines = len(specifications)
    selling_points = [
        chunk
        for value in (
            _logical_lines(values["selling_points"], semicolon_separator=True) or [""]
        )
        for chunk in _wrapped_chunks(value, 9)
    ]
    selling_lines = len(selling_points)
    link_chunks = _wrapped_chunks(
        values["jd_link"], 9, first_line_prefix=_JD_LINK_LABEL
    )
    link_lines = _line_count(
        values["jd_link"], 9, first_line_prefix=_JD_LINK_LABEL
    )
    if link_lines != len(link_chunks):
        raise _error("京东链接容量与分行合同不一致")
    high_lines = model_lines + name_lines + 1
    low_lines = (
        1
        + max(_MIN_SPEC_LINES, specification_lines)
        + 2
        + 1
        + 1
        + max(_MIN_SELLING_LINES, selling_lines)
        + link_lines
    )
    required_height = high_lines * 18 * 1.5 + low_lines * 9 * 1.5 + _TEXT_HEIGHT_SAFETY_PT
    frame = content.text_frame
    available_height = (content.height - frame.margin_top - frame.margin_bottom) / _EMU_PER_POINT
    overflow = required_height > available_height
    if overflow and not allow_overflow:
        raise _capacity_error(product_identifier, required_height, available_height)

    body: list[Any] = [
        *(_paragraph_copy(templates[0], value) for value in model_chunks),
        *(_paragraph_copy(templates[1], value) for value in name_chunks),
        _fixed_paragraph_copy(templates[2]),
        _fixed_paragraph_copy(templates[3]),
    ]
    body.extend(_paragraph_copy(templates[4], value) for value in specifications)
    body.extend(
        _paragraph_copy(templates[4], "")
        for _ in range(max(0, _MIN_SPEC_LINES - specification_lines))
    )
    retail = _fixed_paragraph_copy(templates[9])
    retail_runs = retail.findall(qn("a:r"))
    retail_runs[1].find(qn("a:t")).text = values["retail_price"]
    retail_runs[2].find(qn("a:t")).text = "元" if values["retail_price"] else ""
    group = _fixed_paragraph_copy(templates[10])
    group_runs = group.findall(qn("a:r"))
    group_runs[1].find(qn("a:t")).text = values["group_price"]
    group_runs[2].find(qn("a:t")).text = "元" if values["group_price"] else ""
    body.extend(
        (
            retail,
            group,
            _fixed_paragraph_copy(templates[11]),
            _fixed_paragraph_copy(templates[12]),
        )
    )
    body.extend(_paragraph_copy(templates[13], value) for value in selling_points)
    body.extend(
        _paragraph_copy(templates[13], "")
        for _ in range(max(0, _MIN_SELLING_LINES - selling_lines))
    )
    for line_number, chunk in enumerate(link_chunks):
        link = _fixed_paragraph_copy(templates[19])
        link_runs = link.findall(qn("a:r"))
        # Only the first physical link line contains the fixed customer label.
        # Every value Run retains the complete original target relationship below.
        link_runs[0].find(qn("a:t")).text = _JD_LINK_LABEL if line_number == 0 else ""
        link_runs[1].find(qn("a:t")).text = chunk
        body.append(link)
    _replace_body_paragraphs(frame, body)

    # This verifies text, rather than merely the enclosing textbox, stays on
    # the slide.  It uses the same generated paragraph counts as the capacity
    # calculation above, including every long URL chunk.
    visible_bottom = _visible_text_bottom(
        content, high_lines=high_lines, low_lines=low_lines
    )
    if visible_bottom > slide_height and not allow_overflow:
        raise _capacity_error(product_identifier, required_height, available_height)
    if values["jd_link"]:
        for paragraph in content.text_frame.paragraphs[-link_lines:]:
            rendered_link = paragraph.runs[1]
            rendered_link.hyperlink.address = values["jd_link"]
            # WPS otherwise applies the theme's blue/underline hyperlink style even
            # though the customer explicitly requires a black, bold link value.
            _set_run_style(rendered_link, size=9, colour=_BLACK, bold=True)
            rendered_link.font.underline = False
    return (required_height, available_height) if overflow else None


def _fill_image(slide: Any, slot: Any, image_bytes: bytes | None) -> None:
    if image_bytes is None:
        return
    try:
        with Image.open(BytesIO(image_bytes)) as image:
            width, height = image.size
    except (OSError, UnidentifiedImageError) as error:
        raise AppError("PPT_TEMPLATE_IMAGE_INVALID", "商品图片无法读取", 422) from error
    if width <= 0 or height <= 0:
        raise AppError("PPT_TEMPLATE_IMAGE_INVALID", "商品图片尺寸无效", 422)
    scale = min(slot.width / width, slot.height / height)
    rendered_width, rendered_height = int(width * scale), int(height * scale)
    picture = slide.shapes.add_picture(
        BytesIO(image_bytes),
        slot.left + (slot.width - rendered_width) // 2,
        slot.top + (slot.height - rendered_height) // 2,
        width=rendered_width,
        height=rendered_height,
    )
    picture.name = slot.name
    parent = slot._element.getparent()
    parent.remove(picture._element)
    parent.replace(slot._element, picture._element)


def _item_values(item: Mapping[str, object]) -> dict[str, str]:
    product, prices = (_mapping(item.get(key)) for key in ("product", "prices"))
    return {
        "model": _text(product, "model"),
        "product_name": _text(product, "product_name"),
        "specification": _text(product, "product_specification"),
        "retail_price": _format_price(prices.get("jd_price")),
        "group_price": _format_price(prices.get("agreement_price")),
        "selling_points": _text(product, "selling_points"),
        # Unlike ordinary text fields, URLs must be validated before any trim or
        # normalization so invalid whitespace/control characters cannot turn into
        # a different, accepted address.
        "jd_link": _valid_link(_raw_text(product, "jd_same_product_url")),
    }


def render_catalog_minimal_v5(
    source: Mapping[str, object],
    *,
    images: Mapping[str, bytes],
    template: PptTemplateDefinition,
    manifest: PptxSlotManifest,
    layout_preview: bool = False,
) -> bytes:
    if (template.code, template.version) != ("CATALOG_MINIMAL", "5"):
        raise _error("渲染入口必须明确为 CATALOG_MINIMAL + 5")
    deck = _load(resolve_ppt_template_asset(template))
    _validate_runtime_structure(deck, manifest)
    if _mappings(source.get("packages")):
        raise AppError(
            "PPT_TEMPLATE_COMBINATION_UNSUPPORTED",
            "CATALOG_MINIMAL v5 原件只有一个商品、价格、链接和主图合同，不支持组合页",
            422,
        )
    prototype = deck.slides[0]
    preview_issues: list[LayoutPreviewIssue] = []
    for product_number, item in enumerate(_mappings(source.get("single_products")), start=1):
        product = _mapping(item.get("product"))
        product_identifier = _text(product, "id")
        values = _item_values(item)
        slide = _clone_slide(deck, prototype)
        slots = {shape.name: shape for shape in slide.shapes}
        overflow = _fill_text(
            slots["SCM_FIELD_CATALOG_CONTENT"],
            values,
            product_identifier,
            slide_height=deck.slide_height,
            allow_overflow=layout_preview,
        )
        if overflow is not None:
            required_height, available_height = overflow
            error = _capacity_error(
                product_identifier, required_height, available_height
            )
            logger.warning(
                "PPT v5 layout preview bypassed capacity: template=%s product=%s",
                template.code,
                product_identifier or "unknown",
            )
            preview_issues.append(
                LayoutPreviewIssue(
                    product_number=product_number,
                    product_name=values["product_name"],
                    summary=error.message,
                    values={key: value for key, value in values.items() if value},
                )
            )
        _fill_image(slide, slots["SCM_SLOT_IMAGE_MAIN"], images.get(product_identifier))
    if layout_preview and preview_issues:
        add_layout_preview_pages(deck, template.code, preview_issues)
    _remove_slide(deck, prototype)
    deck.core_properties.keywords = ""
    output = BytesIO()
    deck.save(output)
    return output.getvalue()
