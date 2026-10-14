"""High-fidelity, versioned business-template rendering.

This module intentionally handles only an audited v4 contract.  It replaces text
inside named runs and image relationship targets in place; it never redraws a
business template's static design.
"""

from __future__ import annotations

from collections.abc import Mapping
from copy import deepcopy
from io import BytesIO
from math import ceil
from typing import Any

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

from app.common.contracts import AppError
from app.modules.recommendation.application.ppt_template_manifest import PptxSlotManifest
from app.modules.recommendation.application.ppt_template_registry import resolve_ppt_template_asset

_TRANSPARENT_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc````"
    b"\x00\x00\x00\x05\x00\x01\xa5\xf6E@\x00\x00\x00\x00IEND\xaeB`\x82"
)

# The audited title bar is only 0.34 in high.  It is intentionally a
# single-line field: shrinking its 14pt template typography made a title look
# as though it fitted in OOXML, while LibreOffice wrapped it outside the red
# bar.  Reject before rendering instead of producing a partially-visible name.
_TITLE_MAX_UNITS = 30
_DETAIL_LINE_UNITS = 36
_DETAIL_PAGE_LINES = 6
_ATOMIC_FIELD_UNITS = {
    # These values are deliberately below the measured one-line capacity of
    # each fixed-height JD_DETAIL_RED field.  Anything longer moves, whole, to
    # a continuation page; it must not wrap into a neighbouring static row.
    "brand_model": 26,
    # SCM_FIELD_REFERENCE is only 2.53 in wide at 14pt and contains the
    # market/SKU/fulfilment rows. These limits reserve width for its static
    # labels and inner margins; otherwise even a short-looking ASCII SKU can
    # wrap to a fourth line and intrude into the brand/model row below.
    "sku": 10,
    "fulfillment": 10,
    "product_specification": 34,
    "packaging_list": 34,
    "delivery": 34,
    "warranty": 28,
    "evidence": 28,
}
_PRICE_UNITS = 18


def _mapping(value: object) -> Mapping[str, object]:
    return value if isinstance(value, Mapping) else {}


def _mappings(value: object) -> list[Mapping[str, object]]:
    return [_mapping(item) for item in value] if isinstance(value, list) else []


def _text(values: Mapping[str, object], name: str) -> str:
    value = values.get(name)
    return str(value).strip() if value is not None and str(value).strip() else ""


def _money(value: object) -> str:
    text = str(value or "").strip()
    return f"¥ {text}" if text else ""


def _plain_amount(value: object) -> str:
    """Return a value for a template that owns its currency unit in a static run."""
    return str(value or "").strip()


def _fields(item: Mapping[str, object]) -> dict[str, str]:
    product, manual, prices = (_mapping(item.get(key)) for key in ("product", "manual", "prices"))
    title = _text(product, "product_name")
    brand_model = " / ".join(
        value for value in (_text(product, "brand"), _text(product, "model")) if value
    )
    price_lines = [
        ("市场参考价", _money(prices.get("market_price"))),
        ("京东价格", _money(prices.get("jd_price"))),
        ("协议价格", _money(prices.get("agreement_price"))),
        ("活动价格", _money(manual.get("campaign_price"))),
    ]
    specification = _text(product, "product_specification")
    packaging = _text(product, "packaging_list")
    selling_points = _text(product, "selling_points")
    delivery = _text(manual, "delivery_status") or _text(product, "shipping_courier")
    warranty = _text(product, "warranty_period")
    evidence = _text(manual, "evidence")
    market_price = _money(prices.get("market_price"))
    jd_price = _money(prices.get("jd_price"))
    agreement_price = _money(prices.get("agreement_price"))
    campaign_price = _money(manual.get("campaign_price"))
    jd_customer_price = _money(manual.get("jd_customer_price"))
    return {
        "title": title,
        "detail_heading": "产品卖点（介绍详细）",
        "brand_model": brand_model,
        "product_specification": specification,
        "packaging_list": packaging,
        "selling_points": selling_points,
        "delivery": delivery,
        "warranty": warranty,
        "evidence": evidence,
        "market_price": market_price,
        "jd_price": jd_price,
        "jd_customer_price": jd_customer_price,
        # These slots retain the original festive price typography and units.
        # There is currently no confirmed `manual.jd_customer_price` input in
        # the persisted task contract, so this intentionally stays empty rather
        # than falling back to a differently-scoped ordinary JD price.
        "market_reference_value": _plain_amount(prices.get("market_price")),
        "jd_customer_price_value": _plain_amount(manual.get("jd_customer_price")),
        "agreement_price_value": _plain_amount(prices.get("agreement_price")),
        "discount_rate": _plain_amount(manual.get("discount_rate")),
        "agreement_price": agreement_price,
        "campaign_price": campaign_price,
        "jd_customer_price_label": "京东大客户价：",
        "agreement_price_label": "协议价：",
        "offer_label": "人工活动价：" if campaign_price else "",
        "image_notice": "商品主图：",
        "sku": _text(product, "sku"),
        "fulfillment": _text(manual, "fulfillment_cycle"),
        "specification": specification,
        "packaging": packaging,
        "service": warranty,
        "reference": evidence,
        "prices": "\n".join(f"{label}：{value}" for label, value in price_lines if value),
        "category": _text(product, "third_category") or _text(product, "category"),
        "catalog_content": "\n".join(
            value
            for value in (title, specification, selling_points, _text(manual, "evidence"))
            if value
        ),
        "union_body": "\n".join(
            value
            for value in (
                f"SKU：{_text(product, 'sku')}" if _text(product, "sku") else "",
                f"商品配置：{specification}" if specification else "",
                f"市场参考价：{market_price}" if market_price else "",
                f"协议价：{agreement_price}" if agreement_price else "",
                f"发货物流：{delivery}" if delivery else "",
            )
            if value
        ),
    }


def _package_item_label(item: Mapping[str, object]) -> str:
    product = _mapping(item.get("product_snapshot"))
    name = _text(product, "product_name") or "未命名商品"
    sku = _text(product, "sku")
    model = _text(product, "model")
    # Tokenized selling-point rows split Chinese semicolons into rows. Keep an
    # item's SKU/model identity in the same row so it can never be separated
    # from its product name during package rendering.
    identity = "，".join(
        value
        for value in (f"SKU：{sku}" if sku else "", f"型号：{model}" if model else "")
        if value
    )
    quantity = _text(item, "quantity") or "1"
    return f"{name}（{identity}）× {quantity}" if identity else f"{name} × {quantity}"


def _package_fields(package: Mapping[str, object]) -> dict[str, str]:
    """Map one frozen 2--4 item combination to the audited JD page slots.

    The composition remains one complete row per product. Later pagination moves
    whole rows to a continuation page and never character-slices a row, because
    that could separate a product name from its SKU/model.
    """
    items = _mappings(package.get("items"))
    if not 2 <= len(items) <= 4:
        raise AppError("PPT_COMBINATION_INVALID", "组合商品数量必须为 2 至 4 件", 409)
    product_ids = [_text(_mapping(item.get("product_snapshot")), "id") for item in items]
    if any(not product_id for product_id in product_ids) or (
        len(set(product_ids)) != len(product_ids)
    ):
        raise AppError("PPT_COMBINATION_INVALID", "组合商品必须为不同的已冻结商品", 409)
    values = _fields({})
    values.update(
        {
            "title": _text(package, "name") or "商品组合方案",
            "detail_heading": "套装组成",
            "product_specification": f"共 {len(items)} 件不同商品",
            "selling_points": "\n".join(
                f"{index}. {_package_item_label(item)}" for index, item in enumerate(items, start=1)
            ),
            "agreement_price": _money(package.get("total_price")),
            "agreement_price_label": "套装总价：",
            "evidence": _text(package, "reason"),
            "image_notice": "组合商品图（按套装组成顺序）：",
        }
    )
    return values


def _display_units(value: str) -> float:
    """Approximate rendered width while treating CJK glyphs as full-width."""
    return sum(1 if ord(character) > 0xFF else 0.55 for character in value)


def _wrap_detail(value: str) -> list[str]:
    """Wrap complete source text into safe, fixed-height detail rows."""
    lines: list[str] = []
    for raw_line in value.replace("；", "\n").splitlines() or [value]:
        line = raw_line.strip()
        while line:
            used = 0.0
            end = 0
            for index, character in enumerate(line, start=1):
                width = 1 if ord(character) > 0xFF else 0.55
                if end and used + width > _DETAIL_LINE_UNITS:
                    break
                used += width
                end = index
            lines.append(line[:end])
            line = line[end:]
    return lines


def _require_title_capacity(title: str) -> None:
    if _display_units(title) > _TITLE_MAX_UNITS:
        raise AppError(
            "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED",
            "商品名称超出 JD_DETAIL_RED 模板可安全排版范围",
            409,
        )


def _require_price_capacity(values: Mapping[str, str]) -> None:
    for key in ("market_price", "jd_customer_price", "agreement_price", "campaign_price"):
        if _display_units(values.get(key, "")) > _PRICE_UNITS:
            raise AppError(
                "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED",
                "价格值超出 JD_DETAIL_RED 模板可安全排版范围",
                409,
            )


def _continuation_fields(
    title: str, number: int, lines: list[str], *, package: bool
) -> dict[str, str]:
    values = _fields({})
    values.update(
        {
            "title": title,
            "detail_heading": f"{'套装组成' if package else '产品详情'}（续页 {number}）",
            "selling_points": "\n".join(lines),
            # Continuations are information pages, not copies of the hero
            # image.  The rendering path clears their image slots and expands
            # the existing detail field into that safely-empty area.
            "image_notice": "",
        }
    )
    return values


def _single_product_pages(item: Mapping[str, object]) -> list[tuple[dict[str, str], bool]]:
    """Keep every string intact by placing overflow into fixed-height continuation pages."""
    values = _fields(item)
    _require_title_capacity(values["title"])
    _require_price_capacity(values)
    overflow: list[str] = []
    for key, label in (
        ("brand_model", "品牌 / 型号"),
        ("sku", "SKU"),
        ("fulfillment", "履约"),
        ("product_specification", "规格"),
        ("packaging_list", "包装"),
        ("delivery", "配送"),
        ("warranty", "质保"),
        ("evidence", "备注"),
    ):
        value = values[key]
        if _display_units(value) > _ATOMIC_FIELD_UNITS[key]:
            values[key] = ""
            overflow.extend(_wrap_detail(f"{label}：{value}"))

    selling_lines = [
        line for point in values["selling_points"].replace("；", "\n").splitlines()
        for line in _wrap_detail(point)
    ]
    values["selling_points"] = "\n".join(selling_lines[:_DETAIL_PAGE_LINES])
    overflow.extend(selling_lines[_DETAIL_PAGE_LINES:])
    pages = [(values, False)]
    while overflow:
        page_number = len(pages) + 1
        pages.append(
            (
                _continuation_fields(
                    values["title"],
                    page_number,
                    overflow[:_DETAIL_PAGE_LINES],
                    package=False,
                ),
                True,
            )
        )
        overflow = overflow[_DETAIL_PAGE_LINES:]
    return pages


def _package_pages(
    package: Mapping[str, object], *, jd_detail: bool = False
) -> list[tuple[dict[str, str], list[int]]]:
    values = _package_fields(package)
    _require_title_capacity(values["title"])
    _require_price_capacity(values)
    items = _mappings(package.get("items"))
    rows = [f"{index}. {_package_item_label(item)}" for index, item in enumerate(items, start=1)]
    if any(_display_units(row) > 2 * _DETAIL_LINE_UNITS for row in rows):
        raise AppError(
            "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED",
            "组合商品名称、SKU 或型号超出 JD_DETAIL_RED 模板可安全排版范围",
            409,
        )
    pages: list[tuple[dict[str, str], list[int]]] = []
    remaining = list(enumerate(rows))
    while remaining:
        current: list[str] = []
        current_indexes: list[int] = []
        used_lines = 0
        while remaining:
            item_index, row = remaining[0]
            row_lines = max(1, ceil(_display_units(row) / _DETAIL_LINE_UNITS))
            if current and used_lines + row_lines > _DETAIL_PAGE_LINES:
                break
            remaining.pop(0)
            current.append(row)
            current_indexes.append(item_index)
            used_lines += row_lines
        if not pages:
            if jd_detail and len(rows) == 4:
                values["detail_heading"] = f"套装组成（第 1–{len(current)} 项；配图为全套 4 件）"
                values["image_notice"] = "组合商品图（全套 4 件）："
            values["selling_points"] = "\n".join(current)
            pages.append((values, current_indexes))
        else:
            continuation = _continuation_fields(
                values["title"], len(pages) + 1, current, package=True
            )
            if jd_detail:
                continuation["detail_heading"] = (
                    f"套装组成（续页，第 {current_indexes[0] + 1}–{current_indexes[-1] + 1} 项）"
                )
            pages.append((continuation, current_indexes))
    return pages


def _replace_tokenized_runs(shape: Any, values: Mapping[str, str]) -> None:
    """Replace only token-bearing runs, retaining pPr/rPr and static labels."""
    if not getattr(shape, "has_text_frame", False):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            for token in list(_tokens(run.text)):
                value = _token_value(token, values)
                run.text = run.text.replace("{{" + token + "}}", value)


def _tokens(value: str) -> list[str]:
    import re

    return re.findall(r"\{\{([^{}]+)\}\}", value)


def _token_value(token: str, values: Mapping[str, str]) -> str:
    name, separator, index = token.partition(":")
    value = values.get(name, "")
    if not separator:
        return value
    lines = [line.strip() for line in value.replace("；", "\n").splitlines() if line.strip()]
    return lines[int(index)] if index.isdigit() and int(index) < len(lines) else ""


def _remove_empty_dynamic_paragraphs(shape: Any, values: Mapping[str, str]) -> None:
    """Remove a whole optional information row when all of its values are absent."""
    if not getattr(shape, "has_text_frame", False):
        return
    for paragraph in list(shape.text_frame.paragraphs):
        tokens = _tokens(paragraph.text)
        if tokens and all(not _token_value(token, values) for token in tokens):
            paragraph._element.getparent().remove(paragraph._element)


def _clone_slide(presentation: Any, source: Any) -> Any:
    destination = presentation.slides.add_slide(source.slide_layout)
    rels: dict[str, str] = {}
    for relationship in source.part.rels.values():
        if "notesSlide" in relationship.reltype:
            continue
        target = relationship.target_ref if relationship.is_external else relationship.target_part
        rels[relationship.rId] = destination.part.relate_to(
            target, relationship.reltype, is_external=relationship.is_external
        )
    # A blank layout can inject placeholders.  Replacing the content tree makes
    # cloned pages have exactly the reference shape order and geometry.
    source_c_sld = deepcopy(source._element.cSld)
    old_c_sld = destination._element.cSld
    old_c_sld.getparent().replace(old_c_sld, source_c_sld)
    for node in destination._element.cSld.iter():
        for key, value in list(node.attrib.items()):
            if value in rels:
                node.set(key, rels[value])
    destination.__dict__.pop("shapes", None)
    return destination


def _is_image(content: bytes) -> bool:
    try:
        with Image.open(BytesIO(content)) as image:
            image.verify()
        return True
    except (UnidentifiedImageError, OSError, ValueError):
        return False


def _replace_picture(shape: Any, content: bytes | None) -> None:
    """Swap the existing image relationship, preserving shape and z-order.

    PowerPoint pictures always fill their bounding rectangle.  Updating the
    crop after swapping the relationship is therefore essential: otherwise a
    portrait in a landscape frame (or the reverse) is stretched.
    """
    if content is None:
        # A transparent pixel removes old product media without adding any text.
        content = _TRANSPARENT_PNG
    _set_picture_crop(shape, content)
    _image_part, rel_id = shape.part.get_or_add_image_part(BytesIO(content))
    blip = next(node for node in shape._element.iter() if node.tag.endswith("}blip"))
    blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rel_id)


def _set_picture_crop(shape: Any, content: bytes) -> None:
    try:
        with Image.open(BytesIO(content)) as image:
            image_width, image_height = image.size
    except (UnidentifiedImageError, OSError, ValueError):
        return
    if image_width <= 1 or image_height <= 1:
        shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
        return
    image_ratio = image_width / image_height
    frame_ratio = shape.width / shape.height
    if image_ratio > frame_ratio:
        crop = (1 - frame_ratio / image_ratio) / 2
        shape.crop_left = shape.crop_right = crop
        shape.crop_top = shape.crop_bottom = 0
    else:
        crop = (1 - image_ratio / frame_ratio) / 2
        shape.crop_top = shape.crop_bottom = crop
        shape.crop_left = shape.crop_right = 0


def _replace_static_run(shape: Any, paragraph_index: int, value: str) -> None:
    """Change a controlled label without discarding its original run styling."""
    if paragraph_index >= len(shape.text_frame.paragraphs):
        # Public fixtures may intentionally use one generic price field. The
        # private audited asset has the three labelled price paragraphs.
        return
    paragraph = shape.text_frame.paragraphs[paragraph_index]
    if not paragraph.runs:
        raise AppError("PPT_TEMPLATE_SLOT_MISMATCH", "PPT 模板文字槽位不匹配", 409)
    first_run = paragraph.runs[0]
    if "{{" in first_run.text:
        first_run.text = value + first_run.text[first_run.text.index("{{") :]
    else:
        first_run.text = value


def _set_image_notice_and_price_labels(
    shapes: Mapping[str, Any], values: Mapping[str, str]
) -> None:
    _replace_static_run(shapes["SCM_FIELD_IMAGE_NOTICE"], 0, values["image_notice"])
    prices = shapes["SCM_FIELD_PRICES"]
    _replace_static_run(prices, 0, values["jd_customer_price_label"])
    _replace_static_run(prices, 1, values["agreement_price_label"])


def _copy_run_style(source: Any, target: Any) -> None:
    """Copy only an audited run style, retaining the target's text."""
    source_rpr = source._r.rPr
    if source_rpr is None:
        return
    target_rpr = target._r.rPr
    copied = deepcopy(source_rpr)
    if target_rpr is None:
        target._r.insert(0, copied)
    else:
        target_rpr.getparent().replace(target_rpr, copied)


def _apply_price_emphasis(shapes: Mapping[str, Any], values: Mapping[str, str]) -> None:
    """Reuse the private template's 18pt red price run for active prices.

    The source template stores that emphasis on the JD-customer-price token.
    A package total and an agreement-only price occupy a different token, so
    merely preserving each target run left their important values black/14pt.
    """
    prices = shapes["SCM_FIELD_PRICES"]
    paragraphs = prices.text_frame.paragraphs
    if not paragraphs or len(paragraphs[0].runs) < 2:
        return
    source = paragraphs[0].runs[-1]
    for index, key in ((1, "agreement_price"), (2, "campaign_price")):
        if values.get(key) and index < len(paragraphs) and paragraphs[index].runs:
            _copy_run_style(source, paragraphs[index].runs[-1])


def _jd_nonbreaking_price_values(values: Mapping[str, str]) -> dict[str, str]:
    """Bind the JD price glyph and amount without affecting other templates."""
    rendered = dict(values)
    for key in (
        "market_price",
        "jd_price",
        "jd_customer_price",
        "agreement_price",
        "campaign_price",
    ):
        # PowerPoint/LibreOffice treat the normal space as a wrapping point.
        # The non-breaking space has identical visual output in the template.
        rendered[key] = rendered[key].replace("¥ ", "¥\u00a0")
    return rendered


def _lock_text_boxes(shapes: Mapping[str, Any], values: Mapping[str, str]) -> None:
    """Freeze the original fixed geometry after capacity validation."""
    for name in shapes:
        if name.startswith("SCM_FIELD_") and getattr(shapes[name], "has_text_frame", False):
            shapes[name].text_frame.auto_size = MSO_AUTO_SIZE.NONE


def _apply_continuation_layout(shapes: Mapping[str, Any], image_slots: list[Any]) -> None:
    """Use the empty original hero region for a text continuation page.

    Static background/brand objects and all shape order remain untouched.  Only
    the existing dynamic detail field moves into the cleared product-media area,
    avoiding a repeated large image beside a handful of remaining lines.
    """
    if not image_slots:
        return
    detail = shapes["SCM_FIELD_SELLING_POINTS"]
    hero = image_slots[0]
    detail.left, detail.top, detail.width, detail.height = (
        hero.left,
        hero.top,
        hero.width,
        hero.height,
    )


def _apply_package_image_layout(image_slots: list[Any], item_count: int) -> None:
    """Use the existing picture objects as an equal-weight package grid.

    The single-product page keeps the original hero-plus-thumbnail placement.
    A package uses the original hero region only, so its 2--4 products have
    equal visual weight while the static background and shape z-order remain.
    """
    anchor = image_slots[0]
    left, top, width, height = anchor.left, anchor.top, anchor.width, anchor.height
    if item_count < 2:
        return
    rectangles: tuple[tuple[int, int, int, int], ...]
    if item_count == 2:
        rectangles = (
            (left, top, width, height // 2),
            (left, top + height // 2, width, height - height // 2),
        )
    else:
        half_width, half_height = width // 2, height // 2
        rectangles = (
            (left, top, half_width, half_height),
            (left + half_width, top, width - half_width, half_height),
            (left + width // 4, top + half_height, half_width, height - half_height),
        )
        if item_count == 4:
            rectangles = (
                (left, top, half_width, half_height),
                (left + half_width, top, width - half_width, half_height),
                (left, top + half_height, half_width, height - half_height),
                (
                    left + half_width,
                    top + half_height,
                    width - half_width,
                    height - half_height,
                ),
            )
    for shape, rectangle in zip(image_slots, rectangles, strict=False):
        shape.left, shape.top, shape.width, shape.height = rectangle


def _package_images(
    package: Mapping[str, object], images: Mapping[str, bytes]
) -> list[bytes | None]:
    """Return one image per distinct confirmed package item, in frozen order."""
    items = _mappings(package.get("items"))
    if not 2 <= len(items) <= 4:
        raise AppError("PPT_COMBINATION_INVALID", "组合商品数量必须为 2 至 4 件", 409)
    product_ids = [_text(_mapping(item.get("product_snapshot")), "id") for item in items]
    if any(not product_id for product_id in product_ids) or (
        len(set(product_ids)) != len(product_ids)
    ):
        raise AppError("PPT_COMBINATION_INVALID", "组合商品必须为不同的已冻结商品", 409)
    return [
        content if (content := images.get(product_id)) is not None and _is_image(content) else None
        for product_id in product_ids
    ]


def _populate_page(
    presentation: Any,
    prototype: Any,
    manifest: PptxSlotManifest,
    values: Mapping[str, str],
    images: list[bytes | None],
    *,
    package: bool = False,
    continuation: bool = False,
) -> None:
    slide = _clone_slide(presentation, prototype)
    shapes = {shape.name: shape for shape in slide.shapes}
    required = set(manifest.source_text_slots)
    if not required.issubset(shapes):
        raise AppError("PPT_TEMPLATE_SLOT_MISMATCH", "PPT 模板槽位不匹配", 409)
    is_jd_detail = {
        "SCM_FIELD_IMAGE_NOTICE",
        "SCM_FIELD_PRICES",
        "SCM_FIELD_TITLE",
    }.issubset(shapes)
    rendered_values = _jd_nonbreaking_price_values(values) if is_jd_detail else values
    if is_jd_detail:
        _set_image_notice_and_price_labels(shapes, rendered_values)
        _apply_price_emphasis(shapes, rendered_values)
    for shape in shapes.values():
        _remove_empty_dynamic_paragraphs(shape, rendered_values)
        _replace_tokenized_runs(shape, rendered_values)
    if is_jd_detail:
        _lock_text_boxes(shapes, rendered_values)
    image_slots = sorted(
        (shape for name, shape in shapes.items() if name.startswith("SCM_SLOT_IMAGE_")),
        key=lambda shape: shape.name,
    )
    if not image_slots:
        raise AppError("PPT_TEMPLATE_SLOT_MISMATCH", "PPT 模板图片槽位不匹配", 409)
    if package and is_jd_detail:
        _apply_package_image_layout(image_slots, len(images))
    if continuation and is_jd_detail:
        _apply_continuation_layout(shapes, image_slots)
    for index, shape in enumerate(image_slots):
        _replace_picture(shape, images[index] if index < len(images) else None)


def _remove_slide(presentation: Any, slide: Any) -> None:
    index = list(presentation.slides).index(slide)
    slide_id = presentation.slides._sldIdLst[index]
    presentation.part.drop_rel(slide_id.rId)
    del presentation.slides._sldIdLst[index]


def render_v4(
    source: Mapping[str, object],
    *,
    images: Mapping[str, bytes],
    template: Any,
    manifest: PptxSlotManifest,
) -> bytes:
    presentation = Presentation(str(resolve_ppt_template_asset(template)))
    prototype = presentation.slides[manifest.product_slide_index]
    cover = (
        presentation.slides[manifest.cover_slide_index]
        if manifest.cover_slide_index is not None
        else None
    )
    for slide in list(presentation.slides):
        if slide is not prototype and slide is not cover:
            _remove_slide(presentation, slide)
    if cover is not None:
        project = _mapping(source.get("project"))
        for shape in cover.shapes:
            project_name = _text(project, "project_name")
            _replace_tokenized_runs(
                shape,
                {"project_name": project_name, "project_title": project_name},
            )
    packages = source.get("packages")
    for package in packages if isinstance(packages, list) else []:
        package_source = _mapping(package)
        package_images = _package_images(package_source, images)
        is_jd_detail = template.code == "JD_DETAIL_RED"
        for page_number, (values, item_indexes) in enumerate(
            _package_pages(package_source, jd_detail=is_jd_detail)
        ):
            page_images = (
                package_images if page_number == 0 else []
            ) if is_jd_detail else [package_images[index] for index in item_indexes]
            _populate_page(
                presentation,
                prototype,
                manifest,
                values,
                page_images,
                package=True,
                continuation=is_jd_detail and page_number > 0,
            )
    items = source.get("single_products")
    for item in items if isinstance(items, list) else []:
        product_item = _mapping(item)
        product = _mapping(product_item.get("product"))
        product_id = _text(product, "id")
        image = images.get(product_id)
        image = image if image is not None and _is_image(image) else None
        if template.code == "JD_DETAIL_RED":
            pages = _single_product_pages(product_item)
        else:
            values = _fields(product_item)
            if template.code == "JD_FESTIVE_RED":
                selling_point_lines = [
                    line
                    for line in values["selling_points"].replace("；", "\n").splitlines()
                    if line.strip()
                ]
                if len(selling_point_lines) > 3:
                    raise AppError(
                        "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED",
                        "节庆模板卖点区域无法容纳全部商品卖点",
                        409,
                    )
            pages = [(values, False)]
        for values, continuation in pages:
            _populate_page(
                presentation, prototype, manifest, values, [image] if not continuation else [],
                continuation=continuation,
            )
    _remove_slide(presentation, prototype)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
