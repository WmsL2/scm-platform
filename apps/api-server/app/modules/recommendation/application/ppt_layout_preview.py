"""Local-only diagnostic pages for overflowing Type-5 v5 customer layouts."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class LayoutPreviewIssue:
    product_number: int
    product_name: str
    summary: str
    values: Mapping[str, str]


_FIELD_LABELS = {
    "product_name": "商品名称",
    "market_price": "市场参考价",
    "jd_customer_price": "京东大客户价",
    "discount_rate": "折扣率",
    "sku": "SKU",
    "fulfillment_cycle": "货期/备货时间",
    "brand": "商品品牌",
    "specification": "商品规格",
    "packaging_list": "包装清单",
    "shipping_courier": "发货物流/快递",
    "warranty_period": "质保",
    "selling_points": "商品卖点",
}
_BODY_LINES_PER_SLIDE = 25
_WRAP_CHARACTERS = 58


def _wrapped_lines(value: str) -> list[str]:
    source_lines = value.splitlines() or [""]
    result: list[str] = []
    for source in source_lines:
        if not source:
            result.append("")
            continue
        result.extend(
            source[index : index + _WRAP_CHARACTERS]
            for index in range(0, len(source), _WRAP_CHARACTERS)
        )
    return result


def _issue_lines(template_code: str, issue: LayoutPreviewIssue) -> list[str]:
    lines = [
        f"模板：{template_code}",
        f"商品序号：{issue.product_number}",
        f"商品名称：{issue.product_name or '（空）'}",
        f"容量警告：{issue.summary}",
        "",
        "以下为未截断的原始动态字段，仅供排版核对：",
    ]
    for field, value in issue.values.items():
        if not value:
            continue
        lines.append(f"[{_FIELD_LABELS.get(field, field)}]")
        lines.extend(_wrapped_lines(value))
        lines.append("")
    return lines


def add_layout_preview_pages(
    presentation: Any,
    template_code: str,
    issues: Sequence[LayoutPreviewIssue],
) -> None:
    """Append readable diagnostic pages without changing customer product pages."""
    for issue in issues:
        lines = _issue_lines(template_code, issue)
        page_count = max(1, (len(lines) + _BODY_LINES_PER_SLIDE - 1) // _BODY_LINES_PER_SLIDE)
        for page_index in range(page_count):
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[6]
                if len(presentation.slide_layouts) > 6
                else presentation.slide_layouts[-1]
            )
            title_box = slide.shapes.add_textbox(
                Inches(0.7), Inches(0.35), presentation.slide_width - Inches(1.4), Inches(0.55)
            )
            title_frame = title_box.text_frame
            title_frame.clear()
            title_run = title_frame.paragraphs[0].add_run()
            title_run.text = "排版诊断版（仅用于排版调试）"
            title_run.font.name = "Microsoft YaHei"
            title_run.font.size = Pt(24)
            title_run.font.bold = True
            title_run.font.color.rgb = RGBColor(192, 0, 0)  # type: ignore[no-untyped-call]

            body_box = slide.shapes.add_textbox(
                Inches(0.7), Inches(1.0), presentation.slide_width - Inches(1.4), Inches(6.05)
            )
            body_frame = body_box.text_frame
            body_frame.clear()
            body_frame.word_wrap = True
            body_frame.auto_size = MSO_AUTO_SIZE.NONE
            page_lines = lines[
                page_index * _BODY_LINES_PER_SLIDE : (page_index + 1)
                * _BODY_LINES_PER_SLIDE
            ]
            paragraph = body_frame.paragraphs[0]
            paragraph.text = "\n".join(page_lines)
            paragraph.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(12)
            if page_count > 1:
                footer = slide.shapes.add_textbox(
                    presentation.slide_width - Inches(2.0),
                    presentation.slide_height - Inches(0.4),
                    Inches(1.3),
                    Inches(0.25),
                )
                footer.text_frame.paragraphs[0].text = f"{page_index + 1}/{page_count}"
