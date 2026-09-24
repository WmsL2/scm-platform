from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

from app.common.contracts import AppError
from app.modules.catalog.application.product_export_columns import PRODUCT_EXPORT_COLUMNS

_AUTO_MAPPING = {
    **{column.header: column.key for column in PRODUCT_EXPORT_COLUMNS},
    "一级类目": "category_level1_name",
    "二级类目": "category_level2_name",
    "三级类目": "category_level3_name",
    "品牌": "brand",
    "SKU": "sku",
    "名称": "product_name",
    "京东价": "jd_price",
    "大客户协议价": "agreement_price",
    "折扣率": "discount_rate",
    "采销": "purchasing_agent",
    "是否厂直": "factory_direct",
}


@dataclass(frozen=True)
class TemplateAnalysis:
    sheet_name: str
    header_row: int
    data_start_row: int
    mapping_json: dict[str, str]


@dataclass(frozen=True)
class TemplateColumn:
    column_index: int
    header: str
    duplicate: bool


@dataclass(frozen=True)
class TemplateStructure:
    sheet_names: list[str]
    sheet_name: str
    header_row: int
    max_row: int
    columns: list[TemplateColumn]


def analyze_template(file_bytes: bytes) -> TemplateAnalysis:
    """Read only deterministic workbook structure; never infers business semantics from style."""
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=False, data_only=False)
    except Exception as exc:
        raise AppError("RECOMMENDATION_TEMPLATE_INVALID", "推荐模板无法读取", 422) from exc
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row, 50)), 1):
                headers = [
                    str(cell.value).strip() if cell.value is not None else "" for cell in row
                ]
                mapping = {
                    target: header
                    for header in headers
                    for label, target in _AUTO_MAPPING.items()
                    if header == label
                }
                if mapping:
                    return TemplateAnalysis(sheet.title, row_number, row_number + 1, mapping)
        sheet = workbook.worksheets[0]
        return TemplateAnalysis(sheet.title, 1, 2, {})
    finally:
        workbook.close()


def inspect_template_structure(
    file_bytes: bytes, *, sheet_name: str | None, header_row: int
) -> TemplateStructure:
    """Return the exact stored workbook headers for deterministic UI selection."""
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
    except Exception as exc:
        raise AppError("RECOMMENDATION_TEMPLATE_INVALID", "推荐模板无法读取", 422) from exc
    try:
        selected_sheet = sheet_name or workbook.sheetnames[0]
        if selected_sheet not in workbook.sheetnames:
            raise AppError("RECOMMENDATION_TEMPLATE_MAPPING_INVALID", "模板工作表不存在", 422)
        sheet = workbook[selected_sheet]
        if header_row < 1 or header_row > sheet.max_row:
            raise AppError("RECOMMENDATION_TEMPLATE_MAPPING_INVALID", "模板表头行不存在", 422)
        headers = [
            str(cell.value).strip() if cell.value is not None else "" for cell in sheet[header_row]
        ]
        counts = {header: headers.count(header) for header in set(headers) if header}
        columns = [
            TemplateColumn(
                column_index=index,
                header=header,
                duplicate=counts[header] > 1,
            )
            for index, header in enumerate(headers, start=1)
            if header
        ]
        return TemplateStructure(
            sheet_names=list(workbook.sheetnames),
            sheet_name=selected_sheet,
            header_row=header_row,
            max_row=sheet.max_row,
            columns=columns,
        )
    finally:
        workbook.close()


def validate_mapping_contract(
    file_bytes: bytes,
    *,
    sheet_name: str,
    header_row: int,
    data_start_row: int,
    mapping_json: dict[str, str],
) -> None:
    """Validate an explicit mapping against the exact stored workbook."""
    if data_start_row <= header_row:
        raise AppError(
            "RECOMMENDATION_TEMPLATE_MAPPING_INVALID",
            "数据起始行必须位于表头行之后",
            422,
        )
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
    except Exception as exc:
        raise AppError("RECOMMENDATION_TEMPLATE_MAPPING_INVALID", "推荐模板无法读取", 422) from exc
    try:
        if sheet_name not in workbook.sheetnames:
            raise AppError("RECOMMENDATION_TEMPLATE_MAPPING_INVALID", "模板工作表不存在", 422)
        sheet = workbook[sheet_name]
        if header_row > sheet.max_row:
            raise AppError("RECOMMENDATION_TEMPLATE_MAPPING_INVALID", "模板表头行不存在", 422)
        headers = [
            str(cell.value).strip() if cell.value is not None else "" for cell in sheet[header_row]
        ]
        for source_header in mapping_json.values():
            if headers.count(source_header) != 1:
                raise AppError(
                    "RECOMMENDATION_TEMPLATE_MAPPING_INVALID",
                    "映射源表头不存在或不唯一",
                    422,
                )
    finally:
        workbook.close()
