from __future__ import annotations

import json
import re
from collections.abc import Sequence
from io import BytesIO
from typing import Protocol, TypeVar

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from app.integrations.deepseek.client import (
    DeepSeekConfigurationError,
    DeepSeekStructuredOutputError,
)

StructuredModel = TypeVar("StructuredModel", bound=BaseModel)
TYPE2_PROMPT_VERSION = "type2-ai-v2"
_MAX_TEMPLATE_SHEETS = 8
_MAX_TEMPLATE_ROWS = 12
_MAX_TEMPLATE_COLUMNS = 40


class Type2StructuredProvider(Protocol):
    @property
    def provider(self) -> str: ...

    @property
    def model(self) -> str: ...

    @property
    def prompt_version(self) -> str: ...

    async def structured_completion(
        self,
        *,
        system_prompt: str,
        user_prompt: str,
        response_model: type[StructuredModel],
        max_tokens: int | None = None,
    ) -> StructuredModel: ...


class Type2AgentContractError(RuntimeError):
    pass


class Type2InputColumnMapping(BaseModel):
    model_config = ConfigDict(extra="forbid")

    product_name: str | None = None
    brand: str | None = None
    model: str | None = None
    specification: str | None = None
    category_text: str | None = None
    quantity: str | None = None
    unit: str | None = None
    max_price: str | None = None
    buyer_item_code: str | None = None
    color: str | None = None
    compatible_model: str | None = None
    capacity: str | None = None
    packaging: str | None = None
    row_note: str | None = None

    @model_validator(mode="after")
    def require_identity(self) -> "Type2InputColumnMapping":
        if not any(
            (
                self.buyer_item_code,
                self.brand and self.model,
                self.product_name,
                self.specification,
            )
        ):
            raise ValueError(
                "类型2模板必须包含商品编码、品牌加型号，或可提取品牌型号的商品名称/参数列"
            )
        return self


class Type2OutputColumnMapping(BaseModel):
    model_config = ConfigDict(extra="forbid")

    selected_unit_price: str = Field(min_length=1)
    discount_compliance: str | None = None
    price_valid_until: str | None = None
    note: str | None = None
    is_same_product: str | None = None
    product_name: str | None = None
    brand: str | None = None
    model: str | None = None
    sku: str | None = None


class Type2TemplateMappingSuggestion(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sheet_name: str = Field(min_length=1, max_length=128)
    header_row: int = Field(ge=1)
    data_start_row: int = Field(ge=2)
    input_mapping: Type2InputColumnMapping
    output_mapping: Type2OutputColumnMapping

    @model_validator(mode="after")
    def validate_rows(self) -> "Type2TemplateMappingSuggestion":
        if self.data_start_row <= self.header_row:
            raise ValueError("data_start_row must be after header_row")
        return self


class Type2CandidateContext(BaseModel):
    model_config = ConfigDict(extra="forbid")

    candidate_key: str = Field(pattern=r"^P\d{2}$")
    product_name: str | None = None
    brand: str | None = None
    model: str | None = None
    specification: str | None = None
    packaging: str | None = None
    category_path: str | None = None
    selling_points: str | None = None
    product_remark: str | None = None


class Type2RequirementContext(BaseModel):
    model_config = ConfigDict(extra="forbid")

    requirement_key: str = Field(pattern=r"^R[0-9a-f]{12}$")
    source_row: dict[str, object]
    normalized_fields: dict[str, str | None]
    important_fields: dict[str, str] = Field(default_factory=dict)
    allow_substitution: bool = False
    candidates: list[Type2CandidateContext] = Field(max_length=20)


class Type2SubstituteDecision(BaseModel):
    model_config = ConfigDict(extra="forbid")

    candidate_key: str = Field(pattern=r"^P\d{2}$")
    differences: list[str] = Field(default_factory=list, max_length=8)
    note: str = Field(min_length=4, max_length=500)


class Type2RequirementDecision(BaseModel):
    model_config = ConfigDict(extra="forbid")

    requirement_key: str = Field(pattern=r"^R[0-9a-f]{12}$")
    exact_candidate_keys: list[str] = Field(default_factory=list, max_length=20)
    substitute_candidates: list[Type2SubstituteDecision] = Field(
        default_factory=list, max_length=20
    )
    no_match_reason: str | None = Field(default=None, max_length=500)


class Type2BatchMatchResult(BaseModel):
    model_config = ConfigDict(extra="forbid")

    items: list[Type2RequirementDecision]


class Type2Agent:
    """Type-2-only AI boundary for workbook mapping and same-product adjudication."""

    MAX_PROVIDER_ATTEMPTS = 2

    def __init__(self, provider: Type2StructuredProvider) -> None:
        self.provider = provider

    async def suggest_template_mapping(
        self, file_bytes: bytes, *, project_remark: str | None
    ) -> Type2TemplateMappingSuggestion:
        structure = _workbook_structure(file_bytes)
        result = await self._complete(
            Type2TemplateMappingSuggestion,
            system_prompt=(
                "你是类型2指定商品报价的Excel字段识别器。这是类型2专属任务，不得套用自由推品、"
                "指定类目推品或PPT方案的业务逻辑。你只能从服务器提供的工作簿结构中选择真实存在的"
                "工作表、表头行和表头原文，不得创造列名。识别客户需求输入列与需要回填的空白结果列。"
                "固定业务语义以品牌、型号、商品名称、规格型号或参数为核心；若存在客户商品编码也要"
                "识别。品牌型号可能位于独立列，也可能写在商品名称或参数全文中。遇到包含品牌型号的"
                "组合参数列，应映射到specification供服务器后续提取，不得凭空拆出不存在的列。继续"
                "识别颜色、适用机型、容量或打印量、包装规格、行备注、数量、单位、最高限价等条件列。"
                "行备注和项目备注中的非空明确要求必须参考；空白的结果备注、是否同品、折扣结果等"
                "待回填列不是客户约束。输出列至少必须找到商品报价或等价的单价回填列。项目备注只能"
                "帮助理解业务语义，不能让你选择工作簿中不存在的列。同一表头物理重复时，第一个使用"
                "原表头或“表头#1”，第二个使用“表头#2”，依次编号。严格返回JSON。"
            ),
            user_prompt=json.dumps(
                {"project_remark": project_remark, "workbook": structure},
                ensure_ascii=False,
            ),
        )
        _validate_mapping_against_workbook(file_bytes, result)
        return result

    async def match_batch(
        self,
        requirements: Sequence[Type2RequirementContext],
        *,
        project_remark: str | None,
    ) -> Type2BatchMatchResult:
        if not requirements:
            return Type2BatchMatchResult(items=[])
        result = await self._complete(
            Type2BatchMatchResult,
            system_prompt=(
                "你是类型2指定商品报价的条件复核Agent，只处理服务器给出的客户需求行和候选商品。"
                "这是类型2专属提示词，不得采用类型1规格替代、类型3类目配额、类型4自由推品或类型5"
                "场景组套逻辑。你不能创造商品、SKU或candidate_key。判断时必须阅读完整需求行，包括"
                "商品通用名称、品牌、型号、规格参数、适用设备、颜色、打印量或容量、包装规格及行备注；"
                "important_fields是服务器从客户表识别的重要条件，必须逐项核对；项目备注可作为全局"
                "约束或偏好参考，但不能覆盖明确的行级事实。品牌型号或SKU的首轮召回已由服务器代码"
                "完成，你不得改变主匹配入口。"
                "客户商品编码精确一致，或受控归一后的品牌加型号一致，是同品主条件。商品名称、规格"
                "参数和非空备注用于复核已知事实是否冲突。候选商品大表有对应可比较信息且与客户明确"
                "要求冲突时，必须排除；候选没有该字段、字段为空或文本没有提及，只表示未知，不得仅因"
                "缺少证据而排除，也不得把未知描述成已经证明满足。即使采购编码一致，已有品牌、型号或"
                "商品属性出现明确冲突时仍不得判为EXACT。预计数量只用于采购数量，绝不参与是否同品"
                "判断；空白结果列、待回填备注、折扣结果和是否同品列也不是客户要求。"
                "只有allow_substitution为true且没有满足全部条件的EXACT候选时，才能返回SUBSTITUTE；"
                "否则substitute_candidates必须为空。SUBSTITUTE必须是同类型"
                "且用途/适用设备兼容的商品，颜色和单位不得冲突，容量或打印量不得低于明确要求。"
                "每个替代候选必须返回以“同类型商品替代：”开头的note，准确写明品牌、型号、包装、"
                "容量等实际差异；没有差异说明的替代结果无效。价格、限价、需求中明确的动态折扣率、供应商状态和最终最低"
                "价由服务器代码校验，你不得因价格更低把替代品标成EXACT。不能确认时不返回候选，并在"
                "no_match_reason中说明缺少的证据。每个输入requirement_key必须且只能返回一次。严格返回JSON。"
            ),
            user_prompt=json.dumps(
                {
                    "prompt_version": TYPE2_PROMPT_VERSION,
                    "project_remark": project_remark,
                    "requirements": [item.model_dump() for item in requirements],
                },
                ensure_ascii=False,
            ),
        )
        self._validate_match_result(requirements, result)
        return result

    async def _complete(
        self,
        response_model: type[StructuredModel],
        *,
        system_prompt: str,
        user_prompt: str,
    ) -> StructuredModel:
        last_error: Exception | None = None
        retry_prompt = user_prompt
        for attempt in range(self.MAX_PROVIDER_ATTEMPTS):
            try:
                return await self.provider.structured_completion(
                    system_prompt=system_prompt,
                    user_prompt=retry_prompt,
                    response_model=response_model,
                )
            except DeepSeekConfigurationError:
                raise
            except DeepSeekStructuredOutputError as exc:
                last_error = exc
                if attempt + 1 == self.MAX_PROVIDER_ATTEMPTS:
                    raise
                retry_prompt = (
                    f"{user_prompt}\n\n上一次输出未通过JSON Schema校验："
                    f"{exc.safe_validation_summary}。请重新返回完整合法JSON，不要解释。"
                )
        assert last_error is not None
        raise last_error

    @staticmethod
    def _validate_match_result(
        requirements: Sequence[Type2RequirementContext], result: Type2BatchMatchResult
    ) -> None:
        expected = {item.requirement_key: item for item in requirements}
        returned = [item.requirement_key for item in result.items]
        if len(returned) != len(set(returned)) or set(returned) != set(expected):
            raise Type2AgentContractError("类型2 AI 未完整返回全部需求行")
        for item in result.items:
            allowed = {
                candidate.candidate_key for candidate in expected[item.requirement_key].candidates
            }
            exact = item.exact_candidate_keys
            substitutes = [candidate.candidate_key for candidate in item.substitute_candidates]
            if len(exact) != len(set(exact)) or len(substitutes) != len(set(substitutes)):
                raise Type2AgentContractError("类型2 AI 返回了重复候选")
            if not set(exact + substitutes).issubset(allowed):
                raise Type2AgentContractError("类型2 AI 返回了候选池外商品")
            if set(exact) & set(substitutes):
                raise Type2AgentContractError("类型2 AI 将同一商品同时标记为同品和替代品")
            if substitutes and not expected[item.requirement_key].allow_substitution:
                raise Type2AgentContractError("类型2需求未允许替代商品")
            for substitute in item.substitute_candidates:
                if not substitute.note.startswith("同类型商品替代："):
                    raise Type2AgentContractError("类型2替代商品缺少强制替代备注")


def _workbook_structure(file_bytes: bytes) -> dict[str, object]:
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
    except Exception as exc:
        raise Type2AgentContractError("类型2 Excel 无法读取") from exc
    try:
        sheets: list[dict[str, object]] = []
        for sheet in workbook.worksheets[:_MAX_TEMPLATE_SHEETS]:
            rows: list[dict[str, object]] = []
            max_column = min(sheet.max_column, _MAX_TEMPLATE_COLUMNS)
            for row_number in range(1, min(sheet.max_row, _MAX_TEMPLATE_ROWS) + 1):
                values = [
                    _json_cell(sheet.cell(row_number, column).value)
                    for column in range(1, max_column + 1)
                ]
                if any(value not in (None, "") for value in values):
                    rows.append({"row": row_number, "values": values})
            sheets.append(
                {
                    "sheet_name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "sample_rows": rows,
                }
            )
        return {"sheets": sheets}
    finally:
        workbook.close()


def _validate_mapping_against_workbook(
    file_bytes: bytes, mapping: Type2TemplateMappingSuggestion
) -> None:
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
    try:
        if mapping.sheet_name not in workbook.sheetnames:
            raise Type2AgentContractError("类型2 AI 返回了不存在的工作表")
        sheet = workbook[mapping.sheet_name]
        if mapping.header_row > sheet.max_row:
            raise Type2AgentContractError("类型2 AI 返回的表头行不存在")
        raw_headers = [
            str(cell.value).strip()
            for cell in sheet[mapping.header_row]
            if cell.value is not None and str(cell.value).strip()
        ]
        totals = {header: raw_headers.count(header) for header in raw_headers}
        seen: dict[str, int] = {}
        headers: set[str] = set()
        for header in raw_headers:
            seen[header] = seen.get(header, 0) + 1
            headers.add(header)
            if totals[header] > 1:
                headers.add(f"{header}#{seen[header]}")
        selected_headers = {
            value
            for value in (
                *mapping.input_mapping.model_dump().values(),
                *mapping.output_mapping.model_dump().values(),
            )
            if value
        }
        missing = selected_headers - headers
        if missing:
            raise Type2AgentContractError(
                f"类型2 AI 返回了不存在的列：{'、'.join(sorted(missing))}"
            )
    finally:
        workbook.close()


def _json_cell(value: object) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def normalized_search_tokens(value: str | None) -> set[str]:
    if not value:
        return set()
    normalized = value.casefold()
    return {
        token
        for token in re.findall(r"[a-z0-9]{2,}|[\u3400-\u9fff]{2,}", normalized)
        if token
    }
