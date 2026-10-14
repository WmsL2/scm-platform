from __future__ import annotations

import json
from collections.abc import Sequence
from io import BytesIO
from typing import Protocol, TypeVar

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from app.integrations.deepseek.client import (
    DeepSeekConfigurationError,
    DeepSeekStructuredOutputError,
)

StructuredModel = TypeVar("StructuredModel", bound=BaseModel)
TYPE1_PROMPT_VERSION = "type1-ai-v1"
_MAX_TEMPLATE_SHEETS = 8
_MAX_TEMPLATE_ROWS = 12
_MAX_TEMPLATE_COLUMNS = 40


class Type1StructuredProvider(Protocol):
    @property
    def provider(self) -> str: ...

    @property
    def model(self) -> str: ...

    @property
    def prompt_version(self) -> str: ...

    async def structured_completion(
        self,
        *,
        system_prompt: str,
        user_prompt: str,
        response_model: type[StructuredModel],
        max_tokens: int | None = None,
    ) -> StructuredModel: ...


class Type1AgentContractError(RuntimeError):
    pass


class Type1InputColumnMapping(BaseModel):
    model_config = ConfigDict(extra="forbid")

    category_text: str | None = None
    product_name: str | None = None
    specification: str | None = None
    technical_requirements: str | None = None
    max_price: str | None = None
    row_note: str | None = None
    buyer_item_code: str | None = None
    brand: str | None = None
    model: str | None = None
    quantity: str | None = None
    unit: str | None = None

    @model_validator(mode="after")
    def require_matching_reference(self) -> "Type1InputColumnMapping":
        if not any(
            (
                self.category_text,
                self.product_name,
                self.specification,
                self.technical_requirements,
            )
        ):
            raise ValueError("类型1模板必须包含类目、商品名称、规格参数或主要技术要求")
        return self


class Type1OutputColumnMapping(BaseModel):
    model_config = ConfigDict(extra="forbid")

    agreement_price: str | None = None
    sku: str | None = None
    product_name: str | None = None
    brand: str | None = None
    model: str | None = None
    unit: str | None = None
    supplier_name: str | None = None
    note: str | None = None


class Type1TemplateMappingSuggestion(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sheet_name: str = Field(min_length=1, max_length=128)
    header_row: int = Field(ge=1)
    data_start_row: int = Field(ge=2)
    input_mapping: Type1InputColumnMapping
    output_mapping: Type1OutputColumnMapping

    @model_validator(mode="after")
    def validate_rows(self) -> "Type1TemplateMappingSuggestion":
        if self.data_start_row <= self.header_row:
            raise ValueError("data_start_row must be after header_row")
        return self


class Type1CandidateContext(BaseModel):
    model_config = ConfigDict(extra="forbid")

    candidate_key: str = Field(pattern=r"^P\d+$")
    product_name: str | None = None
    brand: str | None = None
    model: str | None = None
    specification: str | None = None
    packaging: str | None = None
    category_path: str | None = None
    selling_points: str | None = None
    product_remark: str | None = None


class Type1RequirementContext(BaseModel):
    model_config = ConfigDict(extra="forbid")

    requirement_key: str = Field(pattern=r"^R[0-9a-f]{12}$")
    source_row: dict[str, object]
    normalized_fields: dict[str, str | None]
    important_fields: dict[str, str] = Field(default_factory=dict)
    candidates: list[Type1CandidateContext] = Field(max_length=20)


class Type1RequirementDecision(BaseModel):
    model_config = ConfigDict(extra="forbid")

    requirement_key: str = Field(pattern=r"^R[0-9a-f]{12}$")
    candidate_keys: list[str] = Field(default_factory=list, max_length=20)
    no_match_reason: str | None = Field(default=None, max_length=500)


class Type1BatchMatchResult(BaseModel):
    model_config = ConfigDict(extra="forbid")

    items: list[Type1RequirementDecision]


class Type1Agent:
    """Type-1-only AI boundary for workbook mapping and specification review."""

    MAX_PROVIDER_ATTEMPTS = 2

    def __init__(self, provider: Type1StructuredProvider) -> None:
        self.provider = provider

    async def suggest_template_mapping(
        self, file_bytes: bytes, *, project_remark: str | None
    ) -> Type1TemplateMappingSuggestion:
        result = await self._complete(
            Type1TemplateMappingSuggestion,
            system_prompt=(
                "你是类型1已知规格参数报价的Excel字段识别器。这是类型1专属任务，不得套用类型2"
                "品牌型号同品、类型3类目配额、类型4自由推品或类型5场景组套逻辑。固定业务语义"
                "以三级类目、商品名称、规格型号或规格参数、主要技术要求、限价为核心，并识别客户"
                "已经填写的行备注或特殊备注。列名可以不同，规格和技术要求也可能分散在多个列。"
                "数量、单位、客户物料编码、品牌和型号可作为辅助输入，但不得替代上述核心语义。"
                "只选择服务器提供的真实工作表、表头行和表头原文，不得创造列名。优先准确找到"
                "商品报价、供应商报价或等价的单价回填列，并映射到商品主数据 "
                "agreement_price（协议价）；"
                "无法唯一确认的输出列返回null，交由人工确认，严禁猜测。"
                "如果商品报价列之后是连续的空白结果区域，应把商品报价列作为输出区域起点，并尽量"
                "识别其中商品编码、品牌、型号、单位、供应商和备注等可明确判断的回填列；不得把"
                "商品报价列之前的物料信息、限价、采购量等需求输入列当成回填列。"
                "空白的结果备注列是输出列，不是客户要求。重复表头"
                "按物理次序使用原表头或“表头#1”“表头#2”。严格返回JSON。"
            ),
            user_prompt=json.dumps(
                {
                    "project_remark": project_remark,
                    "workbook": _workbook_structure(file_bytes),
                },
                ensure_ascii=False,
            ),
        )
        _validate_mapping_against_workbook(file_bytes, result)
        return result

    async def match_batch(
        self,
        requirements: Sequence[Type1RequirementContext],
        *,
        project_remark: str | None,
    ) -> Type1BatchMatchResult:
        if not requirements:
            return Type1BatchMatchResult(items=[])
        system_prompt = (
            "你是类型1已知规格参数报价的候选复核Agent，只能判断服务器给出的客户需求行和"
            "候选商品，不得创造商品、SKU或candidate_key。类型1围绕三级类目、商品名称、"
            "规格型号或规格参数、主要技术要求以及客户已填写的行备注和项目备注判断。品牌默认"
            "可以替换；只有客户在需求或备注中明确限定品牌时，品牌才是硬条件。限价、正式商品"
            "状态、供应商资格和最终报价由服务器代码校验，不参与商品身份判断。"
            "候选商品只要已知信息没有与客户明确要求冲突，就可以保留。商品主数据没有某项"
            "属性、字段为空或文本未提及，只表示未知，不得因此排除。只有商品名称、类目、型号、"
            "规格、材质、尺寸、功能、包装或备注中的已知事实与客户要求明确冲突时才排除。"
            "数量只用于报价数量，不判断是否为同类商品。空白结果列和系统待回填备注也不是"
            "客户约束。candidate_keys按适配程度排序；不能确认存在合规候选时返回空列表并说明"
            "明确冲突或条件不足。每个requirement_key必须且只能返回一次。严格返回JSON。"
        )
        base_user_prompt = json.dumps(
            {
                "prompt_version": TYPE1_PROMPT_VERSION,
                "project_remark": project_remark,
                "requirements": [item.model_dump() for item in requirements],
            },
            ensure_ascii=False,
        )
        user_prompt = base_user_prompt
        for attempt in range(self.MAX_PROVIDER_ATTEMPTS):
            result = await self._complete(
                Type1BatchMatchResult,
                system_prompt=system_prompt,
                user_prompt=user_prompt,
            )
            try:
                self._validate_match_result(requirements, result)
            except Type1AgentContractError as exc:
                if attempt + 1 == self.MAX_PROVIDER_ATTEMPTS:
                    raise
                user_prompt = (
                    f"{base_user_prompt}\n\n上一次输出违反批次业务合同：{exc}。"
                    "必须逐一返回全部requirement_key，且只能引用各自行候选池中的candidate_key。"
                    "请重新返回完整合法JSON，不要解释。"
                )
                continue
            return result
        raise AssertionError("unreachable type1 contract retry")

    async def _complete(
        self,
        response_model: type[StructuredModel],
        *,
        system_prompt: str,
        user_prompt: str,
    ) -> StructuredModel:
        last_error: Exception | None = None
        retry_prompt = user_prompt
        for attempt in range(self.MAX_PROVIDER_ATTEMPTS):
            try:
                return await self.provider.structured_completion(
                    system_prompt=system_prompt,
                    user_prompt=retry_prompt,
                    response_model=response_model,
                )
            except DeepSeekConfigurationError:
                raise
            except DeepSeekStructuredOutputError as exc:
                last_error = exc
                if attempt + 1 == self.MAX_PROVIDER_ATTEMPTS:
                    raise
                retry_prompt = (
                    f"{user_prompt}\n\n上一次输出未通过JSON Schema校验："
                    f"{exc.safe_validation_summary}。请重新返回完整合法JSON，不要解释。"
                )
        assert last_error is not None
        raise last_error

    @staticmethod
    def _validate_match_result(
        requirements: Sequence[Type1RequirementContext], result: Type1BatchMatchResult
    ) -> None:
        expected = {item.requirement_key: item for item in requirements}
        returned = [item.requirement_key for item in result.items]
        if len(returned) != len(set(returned)) or set(returned) != set(expected):
            raise Type1AgentContractError("类型1 AI 未完整返回全部需求行")
        for item in result.items:
            allowed = {
                candidate.candidate_key for candidate in expected[item.requirement_key].candidates
            }
            if len(item.candidate_keys) != len(set(item.candidate_keys)):
                raise Type1AgentContractError("类型1 AI 返回了重复候选")
            if not set(item.candidate_keys).issubset(allowed):
                raise Type1AgentContractError("类型1 AI 返回了候选池外商品")


def _workbook_structure(file_bytes: bytes) -> dict[str, object]:
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
    except Exception as exc:
        raise Type1AgentContractError("类型1 Excel 无法读取") from exc
    try:
        sheets: list[dict[str, object]] = []
        for sheet in workbook.worksheets[:_MAX_TEMPLATE_SHEETS]:
            rows: list[dict[str, object]] = []
            max_column = min(sheet.max_column, _MAX_TEMPLATE_COLUMNS)
            for row_number in range(1, min(sheet.max_row, _MAX_TEMPLATE_ROWS) + 1):
                values = [
                    _json_cell(sheet.cell(row_number, column).value)
                    for column in range(1, max_column + 1)
                ]
                if any(value not in (None, "") for value in values):
                    rows.append({"row": row_number, "values": values})
            sheets.append(
                {
                    "sheet_name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "sample_rows": rows,
                }
            )
        return {"sheets": sheets}
    finally:
        workbook.close()


def _validate_mapping_against_workbook(
    file_bytes: bytes, mapping: Type1TemplateMappingSuggestion
) -> None:
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
    try:
        if mapping.sheet_name not in workbook.sheetnames:
            raise Type1AgentContractError("类型1 AI 返回了不存在的工作表")
        sheet = workbook[mapping.sheet_name]
        if mapping.header_row > sheet.max_row:
            raise Type1AgentContractError("类型1 AI 返回的表头行不存在")
        raw_headers = [
            str(cell.value).strip()
            for cell in sheet[mapping.header_row]
            if cell.value is not None and str(cell.value).strip()
        ]
        totals = {header: raw_headers.count(header) for header in raw_headers}
        seen: dict[str, int] = {}
        headers: set[str] = set()
        for header in raw_headers:
            seen[header] = seen.get(header, 0) + 1
            headers.add(header)
            if totals[header] > 1:
                headers.add(f"{header}#{seen[header]}")
        selected_headers = {
            value
            for value in (
                *mapping.input_mapping.model_dump().values(),
                *mapping.output_mapping.model_dump().values(),
            )
            if value
        }
        missing = selected_headers - headers
        if missing:
            raise Type1AgentContractError(
                f"类型1 AI 返回了不存在的列：{'、'.join(sorted(missing))}"
            )
    finally:
        workbook.close()


def _json_cell(value: object) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)
