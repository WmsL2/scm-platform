from __future__ import annotations

import hashlib
import json
import re
import uuid
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.contracts import AppError, PageParams, PageResult
from app.core.transaction import transaction_scope
from app.infrastructure.adapters import ObjectStorage, get_object_storage
from app.integrations.deepseek.client import (
    DeepSeekConfigurationError,
    DeepSeekProviderError,
    DeepSeekStructuredOutputError,
)
from app.modules.bid.domain.lifecycle import (
    BidFileType,
    BidImportStatus,
    BidItemStatus,
    BidProjectStatus,
    BidProjectType,
    BidRecommendationType,
    ensure_transition,
    project_type_for_recommendation_type,
)
from app.modules.bid.infrastructure.models import (
    BidProject,
    BidProjectEvent,
    BidProjectFile,
    BidProjectItem,
    BidTemplate,
)
from app.modules.bid.infrastructure.repository import BidProjectRepository
from app.modules.bid.schemas import (
    BidCurrentSelectionResponse,
    BidOutputMappingFieldResponse,
    BidOutputMappingResponse,
    BidOutputMappingUpdateRequest,
    BidProjectCreateResponse,
    BidProjectDetailResponse,
    BidProjectEventResponse,
    BidProjectFileResponse,
    BidProjectItemResponse,
    BidProjectListItem,
    BidProjectStatusResponse,
)
from app.modules.catalog.application.product_export_columns import PRODUCT_EXPORT_COLUMNS
from app.modules.matching.application.type1_agent import (
    Type1Agent,
    Type1AgentContractError,
    Type1TemplateMappingSuggestion,
)
from app.modules.matching.application.type2_agent import (
    Type2Agent,
    Type2AgentContractError,
    Type2TemplateMappingSuggestion,
)
from app.modules.recommendation.template.models import RecommendationTemplateMapping
from app.modules.recommendation.template.schemas import (
    RecommendationTemplateFileResponse,
    RecommendationTemplateMappingResponse,
    RecommendationTemplateMappingUpdateRequest,
    RecommendationTemplateStructureResponse,
)
from app.modules.system.service import BusinessSequenceService

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_PROJECT_ITEMS = 100_000
_PRODUCT_MASTER_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2] / "catalog" / "resources" / "product-master-template.xlsx"
)
_TYPE3_CATEGORY_PATTERN = re.compile(r"^\s*\d+(?:\.\d+){2}\s+.+")
_PRODUCT_OUTPUT_MAPPING_FIELDS = tuple(
    (column.key, column.header, False) for column in PRODUCT_EXPORT_COLUMNS
)
_OUTPUT_MAPPING_FIELDS: dict[BidRecommendationType, tuple[tuple[str, str, bool], ...]] = {
    BidRecommendationType.TYPE_1_SPECIFICATION: (
        *_PRODUCT_OUTPUT_MAPPING_FIELDS,
        ("selected_unit_price", "商品报价（单价）", True),
        ("unit", "计量单位", False),
        ("note", "匹配备注", False),
    ),
    BidRecommendationType.TYPE_2_IDENTIFIED_PRODUCT: (
        *_PRODUCT_OUTPUT_MAPPING_FIELDS,
        ("selected_unit_price", "商品报价（单价）", True),
        ("unit", "计量单位", False),
        ("is_same_product", "是否同品", False),
        ("discount_compliance", "折扣是否满足", False),
        ("price_valid_until", "价格有效期", False),
        ("note", "匹配／替代备注", False),
    ),
}


class BidProjectService:
    def __init__(
        self,
        session: AsyncSession,
        storage: ObjectStorage | None = None,
        type1_agent: Type1Agent | None = None,
        type2_agent: Type2Agent | None = None,
    ) -> None:
        self.session = session
        self.storage = storage or get_object_storage()
        self.repository = BidProjectRepository(session)
        self.type1_agent = type1_agent
        self.type2_agent = type2_agent

    async def create(
        self,
        *,
        project_name: str,
        buyer_name: str,
        start_at: datetime | None,
        deadline_at: datetime | None,
        remark: str | None,
        filename: str | None,
        file_bytes: bytes | None,
        recommendation_template_filename: str | None = None,
        recommendation_template_bytes: bytes | None = None,
        project_type: BidProjectType = BidProjectType.FILTER_RECOMMENDATION,
        recommendation_type: BidRecommendationType | None = None,
        actor_id: uuid.UUID,
    ) -> BidProjectCreateResponse:
        if (
            recommendation_type is not None
            and project_type_for_recommendation_type(recommendation_type) is not project_type
        ):
            raise AppError(
                "BID_RECOMMENDATION_TYPE_MISMATCH",
                "业务类型与项目处理流程不一致",
                422,
            )
        if start_at is not None and deadline_at is not None and start_at > deadline_at:
            raise AppError("BID_PROJECT_INVALID_TIME_RANGE", "项目开始时间不能晚于截止时间", 422)
        using_system_recommendation_template = False
        if recommendation_type is BidRecommendationType.TYPE_3_CATEGORY:
            # Type 3 must retain the customer's workbook.  Its category nodes are
            # selected explicitly in the dedicated Type-3 workspace; never turn
            # a guessed cell scan into the customer's remark/requirement.
            if recommendation_template_filename is None or recommendation_template_bytes is None:
                raise AppError(
                    "TYPE3_WORKBOOK_REQUIRED",
                    "类型 3 必须上传客户 Excel，以便选择 Sheet、类目及对应导出区域",
                    422,
                )
        if project_type == BidProjectType.FILTER_RECOMMENDATION:
            if (
                recommendation_template_filename is not None
                or recommendation_template_bytes is not None
            ):
                raise AppError(
                    "RECOMMENDATION_TEMPLATE_NOT_ALLOWED",
                    "筛选推品项目不接受自由推品模板",
                    422,
                )
            if filename is None or file_bytes is None:
                raise AppError("BID_EXCEL_REQUIRED", "筛选推品必须上传客户需求 Excel", 422)
            self._validate_upload(filename, file_bytes)
        elif project_type == BidProjectType.FREE_RECOMMENDATION:
            if not remark or len(remark.strip()) < 20:
                raise AppError(
                    "FREE_RECOMMENDATION_REMARK_REQUIRED",
                    "自由推品需求说明至少需要 20 个有效字符",
                    422,
                )
            if recommendation_template_filename is None or recommendation_template_bytes is None:
                raise AppError("RECOMMENDATION_TEMPLATE_REQUIRED", "自由推品必须上传推荐模板", 422)
            self._validate_recommendation_template(
                recommendation_template_filename, recommendation_template_bytes
            )
        else:
            if not remark or len(remark.strip()) < 20:
                raise AppError(
                    "PPT_SOLUTION_REMARK_REQUIRED",
                    "PPT 方案需求说明至少需要 20 个有效字符",
                    422,
                )
            if (
                recommendation_template_filename is not None
                or recommendation_template_bytes is not None
            ):
                raise AppError(
                    "PPT_TEMPLATE_NOT_ALLOWED",
                    "PPT 方案统一使用系统默认模板，不支持上传项目模板",
                    422,
                )
        project_id = uuid.uuid4()
        template: BidTemplate | None = None
        project_output_mapping: dict[str, str] | None = None
        parsed_rows: list[dict[str, Any]] = []
        import_status = BidImportStatus.MAPPING_REQUIRED.value
        import_error: str | None = "未识别到已启用模板，需要完成字段映射后再解析"
        if project_type == BidProjectType.FILTER_RECOMMENDATION:
            assert file_bytes is not None
            try:
                template = await self._recognize_template(file_bytes, recommendation_type)
                if (
                    template is None
                    and (
                        (
                            recommendation_type
                            is BidRecommendationType.TYPE_1_SPECIFICATION
                            and self.type1_agent is not None
                        )
                        or (
                            recommendation_type
                            is BidRecommendationType.TYPE_2_IDENTIFIED_PRODUCT
                            and self.type2_agent is not None
                        )
                    )
                ):
                    suggestion: (
                        Type1TemplateMappingSuggestion | Type2TemplateMappingSuggestion
                    )
                    if (
                        recommendation_type is BidRecommendationType.TYPE_1_SPECIFICATION
                        and self.type1_agent is not None
                    ):
                        suggestion = await self.type1_agent.suggest_template_mapping(
                            file_bytes, project_remark=remark
                        )
                        template_prefix = "TYPE1_AI"
                        template_name = "类型1·AI识别客户规格报价表"
                    elif (
                        recommendation_type
                        is BidRecommendationType.TYPE_2_IDENTIFIED_PRODUCT
                        and self.type2_agent is not None
                    ):
                        suggestion = await self.type2_agent.suggest_template_mapping(
                            file_bytes, project_remark=remark
                        )
                        template_prefix = "TYPE2_AI"
                        template_name = "类型2·AI识别客户品牌型号报价表"
                    else:  # pragma: no cover - guarded by the outer condition
                        raise AssertionError("unreachable recommendation type")
                    mapping_workbook = load_workbook(
                        BytesIO(file_bytes), read_only=True, data_only=False
                    )
                    try:
                        headers = [
                            self._text(cell.value)
                            for cell in mapping_workbook[suggestion.sheet_name][
                                suggestion.header_row
                            ]
                        ]
                    finally:
                        mapping_workbook.close()
                    fingerprint = self._fingerprint(
                        suggestion.sheet_name, suggestion.header_row, headers
                    )
                    template = BidTemplate(
                        template_code=f"{template_prefix}_{project_id.hex.upper()}",
                        template_name=template_name,
                        version=1,
                        sheet_name=suggestion.sheet_name,
                        header_row=suggestion.header_row,
                        data_start_row=suggestion.data_start_row,
                        import_mapping={
                            key: value
                            for key, value in suggestion.input_mapping.model_dump().items()
                            if value
                        },
                        export_mapping={
                            key: value
                            for key, value in suggestion.output_mapping.model_dump().items()
                            if value
                        },
                        fingerprint=fingerprint,
                        is_active=True,
                        created_by=actor_id,
                    )
                    project_output_mapping = dict(template.export_mapping)
                if template is not None:
                    parsed_rows = self._parse_rows(file_bytes, template)
                    import_status = BidImportStatus.PARSED.value
                    import_error = None
            except AppError as exc:
                import_status = BidImportStatus.FAILED.value
                import_error = exc.message
            except DeepSeekConfigurationError as exc:
                business_label = (
                    "类型1"
                    if recommendation_type is BidRecommendationType.TYPE_1_SPECIFICATION
                    else "类型2"
                )
                import_error = f"{business_label} AI 模板识别暂不可用：{exc}"
            except (
                DeepSeekStructuredOutputError,
                Type1AgentContractError,
                Type2AgentContractError,
            ):
                business_label = (
                    "类型1"
                    if recommendation_type is BidRecommendationType.TYPE_1_SPECIFICATION
                    else "类型2"
                )
                import_error = f"{business_label} AI 未能安全识别模板字段，请检查表头或稍后重试"
            except DeepSeekProviderError as exc:
                business_label = (
                    "类型1"
                    if recommendation_type is BidRecommendationType.TYPE_1_SPECIFICATION
                    else "类型2"
                )
                import_error = f"{business_label} AI 模板识别失败：{exc}"
        else:
            import_status = BidImportStatus.NOT_REQUIRED.value
            import_error = None

        analysis = None
        if project_type == BidProjectType.FREE_RECOMMENDATION:
            from app.modules.recommendation.template.analyzer import analyze_template

            assert recommendation_template_bytes is not None
            analysis = analyze_template(
                recommendation_template_bytes,
                exclude_sheet_title_terms=(
                    ("参考",)
                    if recommendation_type is BidRecommendationType.TYPE_3_CATEGORY
                    else ()
                ),
            )

        saved_keys: list[str] = []
        try:
            storage_key: str | None = None
            template_storage_key: str | None = None
            if project_type == BidProjectType.FILTER_RECOMMENDATION:
                assert file_bytes is not None
                storage_key = await self.storage.save(
                    f"bid-projects/{project_id}/original/source.xlsx", file_bytes
                )
                saved_keys.append(storage_key)
            elif project_type == BidProjectType.FREE_RECOMMENDATION:
                assert recommendation_template_bytes is not None
                template_storage_key = await self.storage.save(
                    f"bid-projects/{project_id}/recommendation-templates/v1.xlsx",
                    recommendation_template_bytes,
                )
                saved_keys.append(template_storage_key)
            async with transaction_scope(self.session):
                project_code = await BusinessSequenceService(self.session).issue_code("BID_PROJECT")
                if template is not None and template.id is None:
                    self.session.add(template)
                    await self.session.flush()
                project = BidProject(
                    id=project_id,
                    project_code=project_code,
                    project_name=project_name.strip(),
                    buyer_name=buyer_name.strip(),
                    start_at=start_at,
                    deadline_at=deadline_at,
                    remark=remark.strip() if remark else None,
                    template_id=template.id if template else None,
                    template_version=template.version if template else None,
                    status=BidProjectStatus.IMPORTED.value,
                    project_type=project_type.value,
                    recommendation_type=(
                        recommendation_type.value if recommendation_type else None
                    ),
                    import_status=import_status,
                    import_error=import_error,
                    output_mapping=(
                        dict(project_output_mapping or {})
                        if recommendation_type in _OUTPUT_MAPPING_FIELDS
                        else None
                    ),
                    total_item_count=len(parsed_rows),
                    processed_item_count=0,
                    created_by=actor_id,
                )
                self.session.add(project)
                if file_bytes is not None and filename is not None and storage_key is not None:
                    self.session.add(
                        BidProjectFile(
                            project_id=project_id,
                            file_type=BidFileType.ORIGINAL.value,
                            version_no=1,
                            original_filename=Path(filename).name[:255],
                            storage_key=storage_key,
                            file_size=len(file_bytes),
                            sha256=hashlib.sha256(file_bytes).hexdigest(),
                            created_by=actor_id,
                        )
                    )
                if (
                    template_storage_key is not None
                    and project_type == BidProjectType.FREE_RECOMMENDATION
                ):
                    assert recommendation_template_bytes is not None
                    assert recommendation_template_filename is not None
                    assert analysis is not None
                    template_file = BidProjectFile(
                        project_id=project_id,
                        file_type=BidFileType.RECOMMENDATION_TEMPLATE.value,
                        version_no=1,
                        original_filename=Path(recommendation_template_filename).name[:255],
                        storage_key=template_storage_key,
                        file_size=len(recommendation_template_bytes),
                        sha256=hashlib.sha256(recommendation_template_bytes).hexdigest(),
                        created_by=actor_id,
                    )
                    self.session.add(template_file)
                    await self.session.flush()
                    self.session.add(
                        RecommendationTemplateMapping(
                            project_id=project_id,
                            template_file_id=template_file.id,
                            template_sha256=template_file.sha256,
                            sheet_name=analysis.sheet_name,
                            header_row=analysis.header_row,
                            data_start_row=analysis.data_start_row,
                            mapping_json=analysis.mapping_json,
                            confirmed_by=(
                                actor_id if using_system_recommendation_template else None
                            ),
                            confirmed_at=(
                                datetime.now() if using_system_recommendation_template else None
                            ),
                        )
                    )
                for row in parsed_rows:
                    self.session.add(BidProjectItem(project_id=project_id, **row))
                self._add_event(
                    project_id,
                    actor_id,
                    "PROJECT_CREATED",
                    None,
                    BidProjectStatus.IMPORTED.value,
                    import_error,
                )
                await self.session.flush()
        except Exception:
            for saved_key in saved_keys:
                await self.storage.delete(saved_key)
            raise
        return BidProjectCreateResponse(
            id=project_id,
            project_code=project_code,
            status=BidProjectStatus.IMPORTED,
            import_status=BidImportStatus(import_status),
            project_type=project_type,
            recommendation_type=recommendation_type,
            import_error=import_error,
            template_id=template.id if template else None,
            template_version=template.version if template else None,
            total_item_count=len(parsed_rows),
        )

    async def list_projects(
        self, page_params: PageParams, *, keyword: str | None, status: BidProjectStatus | None
    ) -> PageResult[BidProjectListItem]:
        rows, total = await self.repository.project_page(
            page=page_params.page,
            page_size=page_params.page_size,
            keyword=keyword.strip() if keyword else None,
            status=status.value if status else None,
        )
        return PageResult(
            items=[self._list_item(row) for row in rows],
            total=total,
            page=page_params.page,
            page_size=page_params.page_size,
        )

    async def add_recommendation_template(
        self, project_id: uuid.UUID, filename: str, file_bytes: bytes, actor_id: uuid.UUID
    ) -> RecommendationTemplateFileResponse:
        self._validate_recommendation_template(filename, file_bytes)
        project = await self._project_or_404(project_id)
        self._ensure_free_project(project)
        from app.modules.recommendation.template.analyzer import analyze_template

        analysis = analyze_template(file_bytes)
        storage_key: str | None = None
        try:
            async with transaction_scope(self.session):
                project = await self._project_or_404(project_id, lock=True)
                self._ensure_free_project(project)
                version = await self._next_template_version(project_id)
                storage_key = await self.storage.save(
                    f"bid-projects/{project_id}/recommendation-templates/v{version}.xlsx",
                    file_bytes,
                )
                file = BidProjectFile(
                    project_id=project_id,
                    file_type=BidFileType.RECOMMENDATION_TEMPLATE.value,
                    version_no=version,
                    original_filename=Path(filename).name[:255],
                    storage_key=storage_key,
                    file_size=len(file_bytes),
                    sha256=hashlib.sha256(file_bytes).hexdigest(),
                    created_by=actor_id,
                )
                self.session.add(file)
                await self.session.flush()
                self.session.add(
                    RecommendationTemplateMapping(
                        project_id=project_id,
                        template_file_id=file.id,
                        template_sha256=file.sha256,
                        sheet_name=analysis.sheet_name,
                        header_row=analysis.header_row,
                        data_start_row=analysis.data_start_row,
                        mapping_json=analysis.mapping_json,
                    )
                )
                await self.session.flush()
                await self.session.refresh(file)
                response = RecommendationTemplateFileResponse(
                    **self._file_response(file).model_dump(), mapping_confirmed=False
                )
        except Exception:
            if storage_key is not None:
                await self.storage.delete(storage_key)
            raise
        return response

    async def recommendation_templates(
        self, project_id: uuid.UUID
    ) -> list[RecommendationTemplateFileResponse]:
        project = await self._project_or_404(project_id)
        self._ensure_free_project(project)
        rows = await self.session.execute(
            select(BidProjectFile, RecommendationTemplateMapping)
            .outerjoin(
                RecommendationTemplateMapping,
                RecommendationTemplateMapping.template_file_id == BidProjectFile.id,
            )
            .where(
                BidProjectFile.project_id == project_id,
                BidProjectFile.file_type == BidFileType.RECOMMENDATION_TEMPLATE.value,
            )
            .order_by(BidProjectFile.version_no.asc())
        )
        return [
            RecommendationTemplateFileResponse(
                **self._file_response(file).model_dump(),
                mapping_confirmed=mapping is not None and mapping.confirmed_at is not None,
            )
            for file, mapping in rows.all()
        ]

    async def recommendation_template_mapping(
        self, project_id: uuid.UUID, file_id: uuid.UUID
    ) -> RecommendationTemplateMappingResponse:
        file, mapping = await self._recommendation_mapping_or_error(project_id, file_id)
        return self._mapping_response(file, mapping)

    async def recommendation_template_structure(
        self,
        project_id: uuid.UUID,
        file_id: uuid.UUID,
        *,
        sheet_name: str | None,
        header_row: int,
    ) -> RecommendationTemplateStructureResponse:
        file, _mapping = await self._recommendation_mapping_or_error(project_id, file_id)
        content = await self.storage.read(file.storage_key)
        if hashlib.sha256(content).hexdigest() != file.sha256:
            raise AppError("RECOMMENDATION_TEMPLATE_MAPPING_INVALID", "模板摘要校验失败", 422)
        from app.modules.recommendation.template.analyzer import inspect_template_structure

        structure = inspect_template_structure(
            content,
            sheet_name=sheet_name.strip() if sheet_name else None,
            header_row=header_row,
        )
        return RecommendationTemplateStructureResponse.model_validate(
            structure, from_attributes=True
        )

    async def update_recommendation_template_mapping(
        self,
        project_id: uuid.UUID,
        file_id: uuid.UUID,
        payload: RecommendationTemplateMappingUpdateRequest,
        actor_id: uuid.UUID,
    ) -> RecommendationTemplateMappingResponse:
        async with transaction_scope(self.session):
            file, mapping = await self._recommendation_mapping_or_error(
                project_id, file_id, lock=True
            )
            mapping.sheet_name = payload.sheet_name.strip()
            mapping.header_row = payload.header_row
            mapping.data_start_row = payload.data_start_row
            mapping.mapping_json = payload.mapping_json
            content = await self.storage.read(file.storage_key)
            if (
                hashlib.sha256(content).hexdigest() != file.sha256
                or file.sha256 != mapping.template_sha256
            ):
                raise AppError("RECOMMENDATION_TEMPLATE_MAPPING_INVALID", "模板摘要校验失败", 422)
            from app.modules.recommendation.template.analyzer import validate_mapping_contract

            validate_mapping_contract(
                content,
                sheet_name=mapping.sheet_name,
                header_row=mapping.header_row,
                data_start_row=mapping.data_start_row,
                mapping_json=mapping.mapping_json,
            )
            mapping.confirmed_by = actor_id
            mapping.confirmed_at = datetime.now()
        return self._mapping_response(file, mapping)

    async def get(self, project_id: uuid.UUID) -> BidProjectDetailResponse:
        project = await self._project_or_404(project_id)
        files = await self.repository.files(project_id)
        events = await self.repository.events(project_id)
        return BidProjectDetailResponse(
            **self._list_item(project).model_dump(),
            remark=project.remark,
            template_id=project.template_id,
            template_version=project.template_version,
            import_error=project.import_error,
            submitted_file_id=project.submitted_file_id,
            files=[self._file_response(item) for item in files],
            events=[BidProjectEventResponse.model_validate(item) for item in events],
        )

    async def items(
        self,
        project_id: uuid.UUID,
        page_params: PageParams,
        *,
        status: BidItemStatus | None,
        keyword: str | None = None,
    ) -> PageResult[BidProjectItemResponse]:
        project = await self._project_or_404(project_id)
        rows, total = await self.repository.item_page(
            project_id,
            page=page_params.page,
            page_size=page_params.page_size,
            status=status.value if status else None,
            keyword=keyword.strip() or None if keyword else None,
            matched_first=(
                project.recommendation_type
                in {
                    BidRecommendationType.TYPE_1_SPECIFICATION.value,
                    BidRecommendationType.TYPE_2_IDENTIFIED_PRODUCT.value,
                }
            ),
        )
        return PageResult(
            items=[
                BidProjectItemResponse(
                    **BidProjectItemResponse.model_validate(item, from_attributes=True).model_dump(
                        exclude={"current_selection"}
                    ),
                    current_selection=(
                        BidCurrentSelectionResponse(
                            selection_id=selection.id,
                            product_id=selection.product_id,
                            supplier_id=selection.supplier_id,
                            selected_unit_price=selection.selected_unit_price,
                            requirement_snapshot=selection.requirement_snapshot,
                            product_snapshot=selection.product_snapshot,
                            supplier_snapshot=selection.supplier_snapshot,
                            price_snapshot=selection.price_snapshot,
                            note=selection.note,
                            created_at=selection.created_at,
                        )
                        if selection is not None
                        else None
                    ),
                )
                for item, selection in rows
            ],
            total=total,
            page=page_params.page,
            page_size=page_params.page_size,
        )

    async def files(self, project_id: uuid.UUID) -> list[BidProjectFileResponse]:
        await self._project_or_404(project_id)
        return [self._file_response(item) for item in await self.repository.files(project_id)]

    async def output_mapping(self, project_id: uuid.UUID) -> BidOutputMappingResponse:
        project = await self._project_or_404(project_id)
        recommendation_type, template = await self._output_mapping_context(project)
        original = await self.repository.original_file(project_id)
        if original is None:
            raise AppError("BID_PROJECT_ORIGINAL_NOT_FOUND", "项目原始文件不存在", 404)
        headers = self._workbook_headers(
            await self.storage.read(original.storage_key), template.sheet_name, template.header_row
        )
        configured_mapping = (
            dict(project.output_mapping)
            if project.output_mapping is not None
            else dict(template.export_mapping)
        )
        if (
            project.output_mapping is not None
            and project.output_mapping_confirmed_at is None
        ):
            configured_mapping = self._with_obvious_price_mapping(
                headers, configured_mapping
            )
        mapping = {
            field: resolved_header
            for field, configured_header in configured_mapping.items()
            if (
                resolved_header := self._resolve_export_header(
                    field, configured_header, headers
                )
            )
            is not None
        }
        legacy_confirmed = project.output_mapping is None
        return self._output_mapping_response(
            project,
            template,
            recommendation_type,
            headers,
            mapping,
            confirmed=legacy_confirmed or project.output_mapping_confirmed_at is not None,
        )

    async def update_output_mapping(
        self,
        project_id: uuid.UUID,
        payload: BidOutputMappingUpdateRequest,
        actor_id: uuid.UUID,
    ) -> BidOutputMappingResponse:
        async with transaction_scope(self.session):
            project = await self._project_or_404(project_id, lock=True)
            recommendation_type, template = await self._output_mapping_context(project)
            if project.status in {
                BidProjectStatus.SUBMITTED.value,
                BidProjectStatus.WON.value,
                BidProjectStatus.LOST.value,
            }:
                raise AppError(
                    "BID_OUTPUT_MAPPING_LOCKED", "项目已进入投标结果阶段，不能修改回填字段", 409
                )
            original = await self.repository.original_file(project_id)
            if original is None:
                raise AppError("BID_PROJECT_ORIGINAL_NOT_FOUND", "项目原始文件不存在", 404)
            headers = self._workbook_headers(
                await self.storage.read(original.storage_key),
                template.sheet_name,
                template.header_row,
            )
            available = set(self._header_columns(headers))
            field_definitions = _OUTPUT_MAPPING_FIELDS[recommendation_type]
            allowed = {key for key, _, _ in field_definitions}
            unknown = set(payload.mapping) - allowed
            if unknown:
                raise AppError(
                    "BID_OUTPUT_MAPPING_FIELD_INVALID",
                    f"包含不支持的回填字段：{'、'.join(sorted(unknown))}",
                    422,
                )
            mapping = {
                key: value.strip()
                for key, value in payload.mapping.items()
                if value is not None and value.strip()
            }
            missing_headers = sorted(set(mapping.values()) - available)
            if missing_headers:
                raise AppError(
                    "BID_OUTPUT_MAPPING_HEADER_INVALID",
                    f"回填列不存在：{'、'.join(missing_headers)}",
                    422,
                )
            if len(set(mapping.values())) != len(mapping):
                raise AppError(
                    "BID_OUTPUT_MAPPING_DUPLICATE", "同一个客户表列不能映射多个回填字段", 422
                )
            required = {key for key, _, required in field_definitions if required}
            missing_required = required - set(mapping)
            if missing_required:
                raise AppError(
                    "BID_OUTPUT_MAPPING_REQUIRED",
                    "必须确认商品报价（单价）回填列",
                    422,
                )
            project.output_mapping = mapping
            project.output_mapping_confirmed_by = actor_id
            project.output_mapping_confirmed_at = datetime.now()
            project.updated_by = actor_id
            self._add_event(
                project_id,
                actor_id,
                "OUTPUT_MAPPING_CONFIRMED",
                project.status,
                project.status,
                "已确认客户原表回填字段映射",
            )
            await self.session.flush()
            response = self._output_mapping_response(
                project,
                template,
                recommendation_type,
                headers,
                mapping,
                confirmed=True,
            )
        return response

    async def download(
        self, project_id: uuid.UUID, file_id: uuid.UUID
    ) -> tuple[BidProjectFile, bytes]:
        file = await self.repository.file_by_id(project_id, file_id)
        if file is None:
            raise AppError("BID_PROJECT_FILE_NOT_FOUND", "投标文件不存在", 404)
        return file, await self.storage.read(file.storage_key)

    async def export(self, project_id: uuid.UUID, actor_id: uuid.UUID) -> BidProjectFileResponse:
        async with transaction_scope(self.session):
            project = await self._project_or_404(project_id, lock=True)
            if project.status not in {
                BidProjectStatus.READY.value,
                BidProjectStatus.EXPORTED.value,
            }:
                raise AppError("BID_PROJECT_NOT_READY", "项目尚未完成选品，不能生成报价文件", 409)
            if project.template_id is None:
                raise AppError("BID_TEMPLATE_REQUIRED", "项目尚未绑定可导出模板", 409)
            template = await self.session.get(BidTemplate, project.template_id)
            if template is None:
                raise AppError("BID_TEMPLATE_NOT_FOUND", "项目模板不存在", 409)
            if (
                project.output_mapping is not None
                and project.output_mapping_confirmed_at is None
            ):
                raise AppError(
                    "BID_OUTPUT_MAPPING_NOT_CONFIRMED",
                    "请先确认客户原表的回填字段映射",
                    409,
                )
            resolved_statuses = [BidItemStatus.SELECTED.value, BidItemStatus.NO_QUOTE.value]
            if project.recommendation_type in {
                BidRecommendationType.TYPE_1_SPECIFICATION.value,
                BidRecommendationType.TYPE_2_IDENTIFIED_PRODUCT.value,
            }:
                # For Types 1 and 2 an unmatched row is an explicit automatic result;
                # it remains blank in the customer's workbook and does not block export.
                resolved_statuses.append(BidItemStatus.NO_MATCH.value)
            unresolved_count = await self.session.scalar(
                select(func.count())
                .select_from(BidProjectItem)
                .where(
                    BidProjectItem.project_id == project_id,
                    BidProjectItem.status.not_in(resolved_statuses),
                )
            )
            if unresolved_count:
                raise AppError("BID_PROJECT_ITEMS_UNRESOLVED", "仍有需求行未完成选品", 409)
            original = await self.repository.original_file(project_id)
            if original is None:
                raise AppError("BID_PROJECT_ORIGINAL_NOT_FOUND", "项目原始文件不存在", 409)
            is_type1 = (
                project.recommendation_type
                == BidRecommendationType.TYPE_1_SPECIFICATION.value
            )
            selections = (
                await self.repository.type1_export_rows(project_id)
                if is_type1
                else await self.repository.export_rows(project_id)
            )
            content = self._build_export(
                await self.storage.read(original.storage_key),
                template,
                selections,
                expand_multiple_rows=is_type1,
                export_mapping=project.output_mapping or template.export_mapping,
            )
            version = await self.repository.next_export_version(project_id)
            filename = f"{project.project_code}-报价文件-V{version}.xlsx"
            storage_key = await self.storage.save(
                f"bid-projects/{project_id}/exports/quote-v{version}.xlsx", content
            )
            file = BidProjectFile(
                project_id=project_id,
                file_type=BidFileType.QUOTED_EXPORT.value,
                version_no=version,
                original_filename=filename,
                storage_key=storage_key,
                file_size=len(content),
                sha256=hashlib.sha256(content).hexdigest(),
                created_by=actor_id,
            )
            self.session.add(file)
            if project.status == BidProjectStatus.READY.value:
                self._transition(
                    project, actor_id, BidProjectStatus.EXPORTED, "QUOTE_EXPORTED", None
                )
            else:
                self._add_event(
                    project_id,
                    actor_id,
                    "QUOTE_EXPORTED",
                    BidProjectStatus.EXPORTED.value,
                    BidProjectStatus.EXPORTED.value,
                    f"生成报价文件第{version}版",
                )
            await self.session.flush()
            # MySQL supplies created_at. Load it while async IO is legal and build
            # the DTO before leaving the transaction; Pydantic must not trigger a
            # lazy attribute query after the context has closed.
            await self.session.refresh(file)
            response = self._file_response(file)
        return response

    async def submit(
        self, project_id: uuid.UUID, file_id: uuid.UUID, note: str | None, actor_id: uuid.UUID
    ) -> BidProjectStatusResponse:
        async with transaction_scope(self.session):
            project = await self._project_or_404(project_id, lock=True)
            file = await self.repository.file_by_id(project_id, file_id)
            expected_file_type = {
                BidProjectType.FREE_RECOMMENDATION.value: BidFileType.RECOMMENDATION_EXPORT,
                BidProjectType.PPT_SOLUTION.value: BidFileType.PPT_EXPORT,
            }.get(project.project_type, BidFileType.QUOTED_EXPORT)
            if file is None or file.file_type != expected_file_type.value:
                raise AppError(
                    "BID_SUBMITTED_FILE_INVALID",
                    {
                        BidFileType.RECOMMENDATION_EXPORT: (
                            "提交文件必须是当前项目的自由推品导出文件"
                        ),
                        BidFileType.PPT_EXPORT: "提交文件必须是当前项目生成的 PPT 方案",
                    }.get(expected_file_type, "实际投标文件必须是当前项目的报价导出文件"),
                    409,
                )
            self._transition(
                project, actor_id, BidProjectStatus.SUBMITTED, "PROJECT_SUBMITTED", note
            )
            project.submitted_file_id = file.id
            project.updated_by = actor_id
        return BidProjectStatusResponse(
            id=project.id,
            status=BidProjectStatus(project.status),
            submitted_file_id=project.submitted_file_id,
        )

    async def result(
        self,
        project_id: uuid.UUID,
        target: BidProjectStatus,
        note: str | None,
        actor_id: uuid.UUID,
    ) -> BidProjectStatusResponse:
        async with transaction_scope(self.session):
            project = await self._project_or_404(project_id, lock=True)
            event_type = "PROJECT_WON" if target == BidProjectStatus.WON else "PROJECT_LOST"
            self._transition(project, actor_id, target, event_type, note)
        return BidProjectStatusResponse(
            id=project.id,
            status=BidProjectStatus(project.status),
            submitted_file_id=project.submitted_file_id,
        )

    async def update(
        self, project_id: uuid.UUID, payload: object, actor_id: uuid.UUID
    ) -> BidProjectDetailResponse:
        async with transaction_scope(self.session):
            project = await self._project_or_404(project_id, lock=True)
            if project.status in {
                BidProjectStatus.SUBMITTED.value,
                BidProjectStatus.WON.value,
                BidProjectStatus.LOST.value,
                BidProjectStatus.VOIDED.value,
            }:
                raise AppError(
                    "BID_PROJECT_UPDATE_FORBIDDEN", "当前项目状态不允许修改基本信息", 409
                )
            start_at = getattr(payload, "start_at")
            deadline_at = getattr(payload, "deadline_at")
            if start_at is not None and deadline_at is not None and start_at > deadline_at:
                raise AppError(
                    "BID_PROJECT_INVALID_TIME_RANGE", "项目开始时间不能晚于投标截止时间", 422
                )
            project.project_name = getattr(payload, "project_name").strip()
            project.buyer_name = getattr(payload, "buyer_name").strip()
            project.start_at, project.deadline_at = start_at, deadline_at
            project.remark = (
                getattr(payload, "remark").strip() if getattr(payload, "remark") else None
            )
            project.updated_by = actor_id
            self._add_event(
                project_id,
                actor_id,
                "PROJECT_UPDATED",
                project.status,
                project.status,
                "更新项目基本信息",
            )
        return await self.get(project_id)

    async def void(
        self, project_id: uuid.UUID, reason: str, actor_id: uuid.UUID
    ) -> BidProjectStatusResponse:
        async with transaction_scope(self.session):
            project = await self._project_or_404(project_id, lock=True)
            self._transition(project, actor_id, BidProjectStatus.VOIDED, "PROJECT_VOIDED", reason)
        return BidProjectStatusResponse(
            id=project.id,
            status=BidProjectStatus(project.status),
            submitted_file_id=project.submitted_file_id,
        )

    async def advance_lifecycle(
        self,
        project_id: uuid.UUID,
        target: BidProjectStatus,
        actor_id: uuid.UUID,
        event_type: str,
        note: str | None = None,
    ) -> BidProjectStatusResponse:
        """供匹配服务调用的受控项目状态推进入口。"""
        async with transaction_scope(self.session):
            project = await self._project_or_404(project_id, lock=True)
            if target == BidProjectStatus.READY:
                unfinished_count = await self.session.scalar(
                    select(func.count())
                    .select_from(BidProjectItem)
                    .where(
                        BidProjectItem.project_id == project_id,
                        BidProjectItem.status.not_in(
                            [BidItemStatus.SELECTED.value, BidItemStatus.NO_QUOTE.value]
                        ),
                    )
                )
                if unfinished_count:
                    raise AppError("BID_PROJECT_ITEMS_UNRESOLVED", "仍有需求行未完成选品", 409)
            self._transition(project, actor_id, target, event_type, note)
        return BidProjectStatusResponse(
            id=project.id,
            status=BidProjectStatus(project.status),
            submitted_file_id=project.submitted_file_id,
        )

    async def _recognize_template(
        self,
        file_bytes: bytes,
        recommendation_type: BidRecommendationType | None = None,
    ) -> BidTemplate | None:
        if recommendation_type in {
            BidRecommendationType.TYPE_1_SPECIFICATION,
            BidRecommendationType.TYPE_2_IDENTIFIED_PRODUCT,
        }:
            # 客户表随项目变化。类型 1/2 必须由各自 Agent 解析本次上传，
            # 不能按表头指纹复用其他项目的输入或回填映射。
            return None
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
        except Exception as exc:
            raise AppError("BID_EXCEL_PARSE_FAILED", "Excel 文件无法读取", 422) from exc
        try:
            templates = list(
                (
                    await self.session.scalars(
                        select(BidTemplate).where(BidTemplate.is_active.is_(True))
                    )
                ).all()
            )
            for template in templates:
                if template.sheet_name not in workbook.sheetnames:
                    continue
                sheet = workbook[template.sheet_name]
                headers = [self._text(cell.value) for cell in sheet[template.header_row]]
                if (
                    self._fingerprint(template.sheet_name, template.header_row, headers)
                    == template.fingerprint
                ):
                    return template
        finally:
            workbook.close()
        return None

    def _parse_rows(self, file_bytes: bytes, template: BidTemplate) -> list[dict[str, Any]]:
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
        except Exception as exc:
            raise AppError("BID_EXCEL_PARSE_FAILED", "Excel 文件无法读取", 422) from exc
        try:
            if template.sheet_name not in workbook.sheetnames:
                raise AppError("BID_TEMPLATE_SHEET_MISSING", "模板工作表不存在", 422)
            sheet = workbook[template.sheet_name]
            headers = [self._text(cell.value) for cell in sheet[template.header_row]]
            header_columns = self._header_columns(headers)
            missing = [
                header
                for header in template.import_mapping.values()
                if header not in header_columns
            ]
            if missing:
                raise AppError("BID_TEMPLATE_HEADER_MISSING", "模板表头不完整", 422)
            rows: list[dict[str, Any]] = []
            for row_number, cells in enumerate(
                sheet.iter_rows(min_row=template.data_start_row), template.data_start_row
            ):
                values = [cell.value for cell in cells]
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                source_data = {
                    header: self._json_value(values[index])
                    for header, index in header_columns.items()
                    if index < len(values)
                }
                if (
                    template.template_code.startswith(
                        ("TYPE1_AI_", "TYPE1_ZHONGXIN_SPECIFICATION")
                    )
                    and (
                        "调整后需报价物料" in source_data
                        or "需报价物料" in source_data
                    )
                ):
                    quote_marker = self._text(source_data.get("调整后需报价物料")) or self._text(
                        source_data.get("需报价物料")
                    )
                    if quote_marker != "√":
                        continue
                parsed = {
                    key: self._text(source_data.get(header))
                    for key, header in template.import_mapping.items()
                }
                specification = self._join_nonempty_text(
                    parsed.get("specification"), parsed.get("technical_requirements")
                )
                rows.append(
                    {
                        "sheet_name": template.sheet_name,
                        "source_row_number": row_number,
                        "source_data": source_data,
                        "product_name": parsed.get("product_name"),
                        "brand": parsed.get("brand"),
                        "model": parsed.get("model"),
                        "specification": specification,
                        "category_text": parsed.get("category_text"),
                        "quantity": self._decimal(parsed.get("quantity")),
                        "unit": parsed.get("unit"),
                        "max_price": self._decimal(parsed.get("max_price")),
                        "buyer_item_code": parsed.get("buyer_item_code"),
                        "status": BidItemStatus.PENDING.value,
                    }
                )
                if len(rows) > MAX_PROJECT_ITEMS:
                    raise AppError("BID_EXCEL_ROWS_TOO_MANY", "Excel 需求行超过允许上限", 422)
            return rows
        finally:
            workbook.close()

    @staticmethod
    def _join_nonempty_text(*values: str | None) -> str | None:
        parts = [value.strip() for value in values if value and value.strip()]
        return "；".join(dict.fromkeys(parts)) or None

    @staticmethod
    def _type3_requirement_from_workbook(file_bytes: bytes) -> str:
        """Read explicit category labels only; this is deterministic, not AI header guessing."""
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as exc:
            raise AppError("BID_EXCEL_PARSE_FAILED", "Excel 文件无法读取", 422) from exc
        try:
            categories: list[str] = []
            demand_sheets = [
                sheet for sheet in workbook.worksheets if "参考" not in sheet.title
            ] or list(workbook.worksheets)
            for sheet in demand_sheets:
                for row in sheet.iter_rows(min_col=1, max_col=min(sheet.max_column, 5)):
                    for cell in row:
                        value = BidProjectService._text(cell.value)
                        if value and _TYPE3_CATEGORY_PATTERN.match(value):
                            normalized = re.sub(r"^\s*\d+(?:\.\d+){2}\s+", "", value).strip()
                            if normalized and normalized not in categories:
                                categories.append(normalized)
                        if len(categories) >= 80:
                            break
                    if len(categories) >= 80:
                        break
                if len(categories) >= 80:
                    break
        finally:
            workbook.close()
        if not categories:
            raise AppError(
                "CATEGORY_RECOMMENDATION_CATEGORY_MISSING",
                "类型 3 模板中未识别到明确的三级类目，请在需求说明中填写类目",
                422,
            )
        requirement = "指定三级类目推品：" + "、".join(categories)
        return requirement[:5000]

    async def _output_mapping_context(
        self, project: BidProject
    ) -> tuple[BidRecommendationType, BidTemplate]:
        try:
            if project.recommendation_type is None:
                raise ValueError("missing recommendation type")
            recommendation_type = BidRecommendationType(project.recommendation_type)
        except (TypeError, ValueError) as exc:
            raise AppError(
                "BID_OUTPUT_MAPPING_NOT_SUPPORTED", "只有类型 1、2 支持报价回填映射", 409
            ) from exc
        if recommendation_type not in _OUTPUT_MAPPING_FIELDS:
            raise AppError(
                "BID_OUTPUT_MAPPING_NOT_SUPPORTED", "只有类型 1、2 支持报价回填映射", 409
            )
        if project.template_id is None:
            raise AppError("BID_TEMPLATE_REQUIRED", "项目尚未完成客户表字段识别", 409)
        template = await self.session.get(BidTemplate, project.template_id)
        if template is None:
            raise AppError("BID_TEMPLATE_NOT_FOUND", "项目模板不存在", 404)
        return recommendation_type, template

    @staticmethod
    def _workbook_headers(
        file_bytes: bytes, sheet_name: str, header_row: int
    ) -> list[str | None]:
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
        except Exception as exc:
            raise AppError("BID_EXCEL_PARSE_FAILED", "Excel 文件无法读取", 422) from exc
        try:
            if sheet_name not in workbook.sheetnames:
                raise AppError("BID_TEMPLATE_SHEET_MISSING", "模板工作表不存在", 422)
            sheet = workbook[sheet_name]
            if header_row > sheet.max_row:
                raise AppError("BID_TEMPLATE_HEADER_MISSING", "模板表头行不存在", 422)
            return [BidProjectService._text(cell.value) for cell in sheet[header_row]]
        finally:
            workbook.close()

    @staticmethod
    def _output_mapping_response(
        project: BidProject,
        template: BidTemplate,
        recommendation_type: BidRecommendationType,
        headers: list[str | None],
        mapping: dict[str, str],
        *,
        confirmed: bool,
    ) -> BidOutputMappingResponse:
        available_headers = list(BidProjectService._header_columns(headers))
        return BidOutputMappingResponse(
            project_id=project.id,
            sheet_name=template.sheet_name,
            header_row=template.header_row,
            headers=available_headers,
            suggested_headers=BidProjectService._suggested_output_headers(
                headers,
                mapping,
                template.import_mapping,
            ),
            fields=[
                BidOutputMappingFieldResponse(
                    key=key,
                    label=label,
                    required=required,
                    header=mapping.get(key),
                )
                for key, label, required in _OUTPUT_MAPPING_FIELDS[recommendation_type]
            ],
            mapping=mapping,
            confirmed=confirmed,
            confirmed_by=project.output_mapping_confirmed_by,
            confirmed_at=project.output_mapping_confirmed_at,
        )

    @classmethod
    def _suggested_output_headers(
        cls,
        headers: list[str | None],
        mapping: dict[str, str],
        import_mapping: dict[str, str],
    ) -> list[str]:
        """Show the current workbook's result block instead of every input column."""
        columns = cls._header_columns(headers)
        resolved_outputs = [
            resolved
            for field, configured in mapping.items()
            if (resolved := cls._resolve_export_header(field, configured, headers)) is not None
        ]
        if not resolved_outputs:
            return []
        price_header = cls._resolve_export_header(
            "selected_unit_price",
            mapping.get("selected_unit_price", ""),
            headers,
        )
        anchor_header = price_header or min(resolved_outputs, key=columns.__getitem__)
        anchor_column = columns[anchor_header]
        input_headers = {
            resolved
            for field, configured in import_mapping.items()
            if (resolved := cls._resolve_export_header(field, configured, headers)) is not None
        }
        return [
            header
            for header, column in columns.items()
            if column >= anchor_column and header not in input_headers
        ]

    @classmethod
    def _with_obvious_price_mapping(
        cls,
        headers: list[str | None],
        mapping: dict[str, str],
    ) -> dict[str, str]:
        if mapping.get("selected_unit_price"):
            return mapping
        price_labels = {
            "供应商报价",
            "商品报价",
            "商品报价单价",
            "报价单价",
            "协议价",
        }
        matches = [
            header
            for header in cls._header_columns(headers)
            if cls._layout_insensitive_header(header)
            .replace("（", "")
            .replace("）", "")
            .replace("(", "")
            .replace(")", "")
            in price_labels
        ]
        if len(matches) != 1:
            return mapping
        return {**mapping, "selected_unit_price": matches[0]}

    def _build_export(
        self,
        original_bytes: bytes,
        template: BidTemplate,
        selections: list[tuple[BidProjectItem, Any]],
        *,
        expand_multiple_rows: bool = False,
        export_mapping: dict[str, str] | None = None,
    ) -> bytes:
        configured_mapping = export_mapping or template.export_mapping
        if not configured_mapping.get("selected_unit_price"):
            raise AppError("BID_EXPORT_PRICE_COLUMN_MISSING", "模板未配置报价写入列", 409)
        workbook = load_workbook(BytesIO(original_bytes), data_only=False)
        try:
            sheet = workbook[template.sheet_name]
            headers = [self._text(cell.value) for cell in sheet[template.header_row]]
            header_columns = {
                key: index + 1 for key, index in self._header_columns(headers).items()
            }
            resolved_export_mapping = {
                field: self._resolve_export_header(field, configured_header, headers)
                for field, configured_header in configured_mapping.items()
            }
            if any(header is None for header in resolved_export_mapping.values()):
                raise AppError("BID_EXPORT_COLUMN_MISSING", "原始文件缺少已配置的结果写入列", 409)
            export_rows = (
                self._expand_type1_export_rows(sheet, selections)
                if expand_multiple_rows
                else [(item.source_row_number, item, selection) for item, selection in selections]
            )
            for target_row_number, item, selection in export_rows:
                product = selection.product_snapshot or {}
                supplier = selection.supplier_snapshot or {}
                values = {
                    **(selection.price_snapshot or {}),
                    **product,
                    **supplier,
                    "selected_unit_price": selection.selected_unit_price,
                    "unit": item.unit,
                    "note": selection.note,
                    "is_same_product": (
                        "否"
                        if product.get("type2_match_type") == "TYPE2_AI_SUBSTITUTE"
                        else self._same_selected_product(item, product)
                    ),
                    "discount_compliance": (selection.price_snapshot or {}).get(
                        "discount_compliance"
                    ),
                    "price_valid_until": (selection.price_snapshot or {}).get(
                        "price_valid_until"
                    ),
                }
                for field, resolved_header in resolved_export_mapping.items():
                    assert resolved_header is not None
                    header = resolved_header
                    target = sheet.cell(target_row_number, header_columns[header])
                    if target.value is None or not str(target.value).strip():
                        target.value = values.get(field)
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()

    @staticmethod
    def _expand_type1_export_rows(
        sheet: Any,
        selections: list[tuple[BidProjectItem, Any]],
    ) -> list[tuple[int, BidProjectItem, Any]]:
        grouped: dict[int, list[tuple[BidProjectItem, Any]]] = {}
        for item, selection in selections:
            grouped.setdefault(item.source_row_number, []).append((item, selection))

        for source_row_number in sorted(grouped, reverse=True):
            extra_count = len(grouped[source_row_number]) - 1
            if extra_count <= 0:
                continue
            sheet.insert_rows(source_row_number + 1, amount=extra_count)
            source_dimension = sheet.row_dimensions[source_row_number]
            for offset in range(1, extra_count + 1):
                target_row_number = source_row_number + offset
                target_dimension = sheet.row_dimensions[target_row_number]
                target_dimension.height = source_dimension.height
                target_dimension.hidden = source_dimension.hidden
                target_dimension.outlineLevel = source_dimension.outlineLevel
                for column in range(1, sheet.max_column + 1):
                    source_cell = sheet.cell(source_row_number, column)
                    target_cell = sheet.cell(target_row_number, column)
                    value = source_cell.value
                    if isinstance(value, str) and value.startswith("="):
                        try:
                            value = Translator(
                                value, origin=source_cell.coordinate
                            ).translate_formula(target_cell.coordinate)
                        except ValueError:
                            pass
                    target_cell.value = value
                    if source_cell.has_style:
                        target_cell._style = copy(source_cell._style)
                    if source_cell.hyperlink:
                        target_cell._hyperlink = copy(source_cell.hyperlink)
                    if source_cell.comment:
                        target_cell.comment = copy(source_cell.comment)

        expanded: list[tuple[int, BidProjectItem, Any]] = []
        inserted_before = 0
        for source_row_number in sorted(grouped):
            matches = grouped[source_row_number]
            target_start = source_row_number + inserted_before
            expanded.extend(
                (target_start + offset, item, selection)
                for offset, (item, selection) in enumerate(matches)
            )
            inserted_before += len(matches) - 1
        return expanded

    async def _project_or_404(self, project_id: uuid.UUID, *, lock: bool = False) -> BidProject:
        project = await self.repository.project_by_id(project_id, lock=lock)
        if project is None:
            raise AppError("BID_PROJECT_NOT_FOUND", "投标项目不存在", 404)
        return project

    @staticmethod
    def _ensure_free_project(project: BidProject) -> None:
        if project.project_type != BidProjectType.FREE_RECOMMENDATION.value:
            raise AppError("PROJECT_TYPE_NOT_SUPPORTED", "该项目类型不支持推荐模板", 409)

    async def _next_template_version(self, project_id: uuid.UUID) -> int:
        latest = await self.session.scalar(
            select(func.max(BidProjectFile.version_no)).where(
                BidProjectFile.project_id == project_id,
                BidProjectFile.file_type == BidFileType.RECOMMENDATION_TEMPLATE.value,
            )
        )
        return (latest or 0) + 1

    async def _recommendation_mapping_or_error(
        self, project_id: uuid.UUID, file_id: uuid.UUID, *, lock: bool = False
    ) -> tuple[BidProjectFile, RecommendationTemplateMapping]:
        project = await self._project_or_404(project_id, lock=lock)
        self._ensure_free_project(project)
        statement = (
            select(BidProjectFile, RecommendationTemplateMapping)
            .join(
                RecommendationTemplateMapping,
                RecommendationTemplateMapping.template_file_id == BidProjectFile.id,
            )
            .where(BidProjectFile.id == file_id, BidProjectFile.project_id == project_id)
        )
        if lock:
            statement = statement.with_for_update()
        row = (await self.session.execute(statement)).first()
        if row is None:
            raise AppError("RECOMMENDATION_TEMPLATE_NOT_FOUND", "推荐模板或映射不存在", 404)
        file, mapping = row
        if file.file_type != BidFileType.RECOMMENDATION_TEMPLATE.value:
            raise AppError(
                "RECOMMENDATION_TEMPLATE_PROJECT_MISMATCH", "文件不是该项目的推荐模板", 409
            )
        return file, mapping

    def _mapping_response(
        self, file: BidProjectFile, mapping: RecommendationTemplateMapping
    ) -> RecommendationTemplateMappingResponse:
        return RecommendationTemplateMappingResponse(
            template_file=self._file_response(file),
            sheet_name=mapping.sheet_name,
            header_row=mapping.header_row,
            data_start_row=mapping.data_start_row,
            mapping_json=mapping.mapping_json,
            confirmed_by=mapping.confirmed_by,
            confirmed_at=mapping.confirmed_at,
        )

    def _transition(
        self,
        project: BidProject,
        actor_id: uuid.UUID,
        target: BidProjectStatus,
        event_type: str,
        note: str | None,
    ) -> None:
        ensure_transition(project.status, target)
        before = project.status
        project.status = target.value
        project.updated_by = actor_id
        self._add_event(project.id, actor_id, event_type, before, target.value, note)

    def _add_event(
        self,
        project_id: uuid.UUID,
        actor_id: uuid.UUID,
        event_type: str,
        from_status: str | None,
        to_status: str | None,
        note: str | None,
    ) -> None:
        self.session.add(
            BidProjectEvent(
                project_id=project_id,
                actor_id=actor_id,
                event_type=event_type,
                from_status=from_status,
                to_status=to_status,
                note=note,
            )
        )

    @staticmethod
    def _validate_upload(filename: str, file_bytes: bytes) -> None:
        if Path(filename).suffix.lower() != ".xlsx":
            raise AppError("BID_EXCEL_EXTENSION_UNSUPPORTED", "仅支持 .xlsx 文件", 422)
        if not file_bytes:
            raise AppError("BID_EXCEL_EMPTY", "上传文件不能为空", 422)
        if len(file_bytes) > MAX_FILE_BYTES:
            raise AppError("BID_EXCEL_TOO_LARGE", "上传文件超过大小限制", 422)

    @staticmethod
    def _validate_recommendation_template(filename: str, file_bytes: bytes) -> None:
        if Path(filename).suffix.lower() != ".xlsx" or not file_bytes:
            raise AppError("RECOMMENDATION_TEMPLATE_INVALID", "推荐模板必须是非空 .xlsx 文件", 422)
        if len(file_bytes) > MAX_FILE_BYTES:
            raise AppError("RECOMMENDATION_TEMPLATE_INVALID", "推荐模板超过大小限制", 422)

    @staticmethod
    def _fingerprint(sheet_name: str, header_row: int, headers: list[str | None]) -> str:
        value = json.dumps(
            {"sheet_name": sheet_name, "header_row": header_row, "headers": headers},
            ensure_ascii=False,
            separators=(",", ":"),
        )
        return hashlib.sha256(value.encode("utf-8")).hexdigest()

    @staticmethod
    def _header_columns(headers: list[str | None]) -> dict[str, int]:
        columns: dict[str, int] = {}
        counts: dict[str, int] = {}
        totals: dict[str, int] = {}
        for header in headers:
            if header:
                totals[header] = totals.get(header, 0) + 1
        for index, header in enumerate(headers):
            if not header:
                continue
            counts[header] = counts.get(header, 0) + 1
            columns.setdefault(header, index)
            if totals[header] > 1:
                columns[f"{header}#{counts[header]}"] = index
        return columns

    @classmethod
    def _resolve_export_header(
        cls,
        field: str,
        configured_header: str,
        headers: list[str | None],
    ) -> str | None:
        available_headers = tuple(cls._header_columns(headers))
        if configured_header in available_headers:
            return configured_header

        configured_key = cls._layout_insensitive_header(configured_header)
        layout_matches = [
            header
            for header in available_headers
            if cls._layout_insensitive_header(header) == configured_key
        ]
        if len(layout_matches) == 1:
            return layout_matches[0]

        semantic_prefixes = {
            "discount_compliance": ("是否满足平台承诺折扣率",),
            "price_valid_until": ("价格有效期",),
        }
        prefixes = tuple(
            cls._layout_insensitive_header(prefix)
            for prefix in semantic_prefixes.get(field, ())
        )
        semantic_matches = [
            header
            for header in available_headers
            if prefixes
            and cls._layout_insensitive_header(header).startswith(prefixes)
        ]
        return semantic_matches[0] if len(semantic_matches) == 1 else None

    @staticmethod
    def _layout_insensitive_header(value: str) -> str:
        return re.sub(r"\s+", "", value).casefold()

    @staticmethod
    def _same_selected_product(item: BidProjectItem, product: dict[str, object]) -> str:
        from app.modules.matching.domain.rules import normalize_match_text

        checks = (
            (item.buyer_item_code, product.get("sku")),
            (item.brand, product.get("brand")),
            (item.model, product.get("model")),
        )
        evidence = [
            normalize_match_text(left)
            == normalize_match_text(str(right) if right is not None else None)
            for left, right in checks
            if normalize_match_text(left)
        ]
        return "是" if evidence and all(evidence) else "否"

    @staticmethod
    def _text(value: object) -> str | None:
        if value is None:
            return None
        result = str(value).strip()
        return result or None

    @staticmethod
    def _json_value(value: object) -> str | int | float | None:
        if value is None:
            return None
        if isinstance(value, (str, int, float)):
            return value
        if isinstance(value, Decimal):
            return str(value)
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return str(value)

    @staticmethod
    def _decimal(value: str | None) -> Decimal | None:
        if value is None:
            return None
        try:
            return Decimal(value.replace(",", ""))
        except InvalidOperation as exc:
            raise AppError("BID_EXCEL_NUMBER_INVALID", "Excel 数量或限价格式不正确", 422) from exc

    @staticmethod
    def _list_item(project: BidProject) -> BidProjectListItem:
        return BidProjectListItem.model_validate(project, from_attributes=True)

    @staticmethod
    def _file_response(file: BidProjectFile) -> BidProjectFileResponse:
        return BidProjectFileResponse.model_validate(file, from_attributes=True)
