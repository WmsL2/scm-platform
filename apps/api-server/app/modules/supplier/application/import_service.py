import uuid
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.contracts import AppError
from app.core.transaction import transaction_scope
from app.modules.supplier.domain.rules import ArchiveStatus, CooperationStatus
from app.modules.supplier.infrastructure.models import (
    Supplier,
    SupplierContact,
    SupplierImportBatch,
    SupplierImportRow,
)
from app.modules.supplier.infrastructure.repository import SupplierRepository
from app.modules.supplier.schemas import (
    SupplierImportConfirmResponse,
    SupplierImportPreviewResponse,
    SupplierImportRowResponse,
)
from app.modules.system.service import BusinessSequenceService

IMPORT_HEADERS = ("供应商名称", "主营品牌", "主要优势", "联系人", "联系电话")
MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1_000


class SupplierImportService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.repository = SupplierRepository(session)

    def build_template(self) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        assert worksheet is not None
        worksheet.title = "供应商导入"
        worksheet.append(IMPORT_HEADERS)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = "A1:E1"
        for column, width in zip(("A", "B", "C", "D", "E"), (28, 30, 40, 20, 22), strict=True):
            worksheet.column_dimensions[column].width = width
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        worksheet["E2"].number_format = "@"
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    async def preview(
        self, filename: str, file_bytes: bytes, actor_id: uuid.UUID
    ) -> SupplierImportPreviewResponse:
        if not filename.lower().endswith(".xlsx"):
            raise AppError(
                "SUPPLIER_IMPORT_FILE_TYPE_INVALID", "Only .xlsx files are supported", 422
            )
        if not file_bytes:
            raise AppError("SUPPLIER_IMPORT_FILE_EMPTY", "The import file is empty", 422)
        if len(file_bytes) > MAX_IMPORT_FILE_BYTES:
            raise AppError("SUPPLIER_IMPORT_FILE_TOO_LARGE", "The import file exceeds 5 MB", 422)
        rows = self._parse_rows(file_bytes)
        await self._mark_duplicate_supplier_names(rows)
        valid_rows = sum(row.is_valid for row in rows)
        batch = SupplierImportBatch(
            original_filename=filename[:255],
            status="VALIDATED",
            total_rows=len(rows),
            valid_rows=valid_rows,
            invalid_rows=len(rows) - valid_rows,
            created_by=actor_id,
            rows=rows,
        )
        async with transaction_scope(self.session):
            self.session.add(batch)
            await self.session.flush()
            response = self._preview_response(batch)
        return response

    async def confirm(
        self, batch_id: uuid.UUID, actor_id: uuid.UUID
    ) -> SupplierImportConfirmResponse:
        async with transaction_scope(self.session):
            batch = await self.repository.import_batch_by_id_for_update(batch_id)
            if batch is None:
                raise AppError("SUPPLIER_IMPORT_BATCH_NOT_FOUND", "Import batch not found", 404)
            if batch.created_by != actor_id:
                raise AppError(
                    "SUPPLIER_IMPORT_BATCH_FORBIDDEN",
                    "Only the uploader can confirm this import batch",
                    403,
                )
            if batch.status != "VALIDATED":
                raise AppError(
                    "SUPPLIER_IMPORT_BATCH_ALREADY_CONFIRMED",
                    "Import batch has already been confirmed",
                    409,
                )
            rows = await self.repository.import_rows_by_batch_id_for_update(batch.id)
            if len(rows) != batch.total_rows:
                raise AppError(
                    "SUPPLIER_IMPORT_BATCH_INTEGRITY_ERROR",
                    "Import batch row count does not match its header",
                    409,
                )
            invalid_rows = [row for row in rows if not row.is_valid]
            if invalid_rows or batch.invalid_rows:
                raise AppError(
                    "SUPPLIER_IMPORT_HAS_INVALID_ROWS",
                    "Correct invalid rows and upload a new file before confirmation",
                    409,
                )
            if not rows:
                raise AppError("SUPPLIER_IMPORT_NO_VALID_ROWS", "No valid import rows found", 409)
            if batch.valid_rows != len(rows):
                raise AppError(
                    "SUPPLIER_IMPORT_BATCH_INTEGRITY_ERROR",
                    "Import batch valid row count does not match persisted rows",
                    409,
                )
            processed_count = 0
            for row in rows:
                supplier_name = self._required_value(row.supplier_name)
                existing = await self.repository.by_name_for_update(supplier_name)
                if existing is not None:
                    if not existing.is_deleted:
                        raise AppError("SUPPLIER_NAME_EXISTS", "该供应商已存在", 409)
                    supplier = existing
                    self._restore_deleted_supplier(supplier, row, actor_id)
                else:
                    supplier = Supplier(
                        supplier_code=await BusinessSequenceService(self.session).issue_code(
                            "SUPPLIER"
                        ),
                        supplier_name=supplier_name,
                        main_brands=self._required_value(row.main_brands),
                        advantage=self._required_value(row.advantage),
                        archive_status=ArchiveStatus.ARCHIVED,
                        cooperation_status=CooperationStatus.NORMAL,
                        created_by=actor_id,
                        updated_by=actor_id,
                        archived_by=actor_id,
                        archived_at=datetime.now(),
                        contacts=self._contacts_from_row(row, actor_id),
                    )
                    self.session.add(supplier)
                processed_count += 1

            # Write every supplier before success is returned. A persistence error
            # rolls back both Supplier rows and the batch status together.
            await self.session.flush()
            if processed_count != len(rows):
                raise AppError(
                    "SUPPLIER_IMPORT_BATCH_INTEGRITY_ERROR",
                    "Import batch processing count does not match persisted rows",
                    409,
                )
            batch.status = "CONFIRMED"
            batch.confirmed_by = actor_id
            batch.confirmed_at = datetime.now()
            response = SupplierImportConfirmResponse(
                id=batch.id,
                status=batch.status,
                imported_count=processed_count,
            )
        return response

    def _parse_rows(self, file_bytes: bytes) -> list[SupplierImportRow]:
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=False)
        except Exception as exc:
            raise AppError(
                "SUPPLIER_IMPORT_FILE_INVALID", "The file is not a valid .xlsx workbook", 422
            ) from exc
        worksheet = workbook.active
        assert worksheet is not None
        headers = tuple(
            self._cell_text(cell.value) for cell in next(worksheet.iter_rows(max_row=1))
        )
        if headers != IMPORT_HEADERS:
            expected = "、".join(IMPORT_HEADERS)
            raise AppError(
                "SUPPLIER_IMPORT_TEMPLATE_INVALID",
                f"Template headers must exactly be: {expected}",
                422,
            )
        rows: list[SupplierImportRow] = []
        for source_row_number, cells in enumerate(
            worksheet.iter_rows(min_row=2, max_col=5), start=2
        ):
            if not any(cell.value is not None for cell in cells):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise AppError(
                    "SUPPLIER_IMPORT_TOO_MANY_ROWS",
                    f"An import may contain at most {MAX_IMPORT_ROWS} data rows",
                    422,
                )
            values = [self._cell_text(cell.value) for cell in cells]
            errors = self._row_errors(cells, values)
            rows.append(
                SupplierImportRow(
                    source_row_number=source_row_number,
                    supplier_name=values[0] or None,
                    main_brands=values[1] or None,
                    advantage=values[2] or None,
                    contact_name=values[3] or None,
                    contact_phone=values[4] or None,
                    is_valid=not errors,
                    error_message="；".join(errors) or None,
                )
            )
        if not rows:
            raise AppError("SUPPLIER_IMPORT_NO_DATA_ROWS", "The workbook has no data rows", 422)
        return rows

    async def _mark_duplicate_supplier_names(self, rows: list[SupplierImportRow]) -> None:
        active_names = await self.repository.active_supplier_names(
            {row.supplier_name for row in rows if row.supplier_name is not None}
        )
        first_excel_row_by_name: dict[str, int] = {}
        for row in rows:
            supplier_name = row.supplier_name
            if not supplier_name:
                continue
            errors: list[str] = []
            if supplier_name in active_names:
                errors.append("该供应商已存在")
            first_row = first_excel_row_by_name.get(supplier_name)
            if first_row is None:
                first_excel_row_by_name[supplier_name] = row.source_row_number
            else:
                errors.append(f"与 Excel 第 {first_row} 行供应商名称重复")
            if errors:
                row.is_valid = False
                row.error_message = "；".join(
                    filter(None, [row.error_message, *errors])
                )

    def _restore_deleted_supplier(
        self, supplier: Supplier, row: SupplierImportRow, actor_id: uuid.UUID
    ) -> None:
        supplier.main_brands = self._required_value(row.main_brands)
        supplier.advantage = self._required_value(row.advantage)
        supplier.archive_status = ArchiveStatus.ARCHIVED
        supplier.cooperation_status = CooperationStatus.NORMAL
        supplier.is_deleted = False
        supplier.deleted_by = None
        supplier.deleted_at = None
        supplier.archived_by = actor_id
        supplier.archived_at = datetime.now()
        supplier.updated_by = actor_id
        for contact in supplier.contacts:
            if not contact.is_deleted:
                contact.is_deleted = True
                contact.updated_by = actor_id
        supplier.contacts.extend(self._contacts_from_row(row, actor_id))

    @staticmethod
    def _row_errors(cells: tuple[object, ...], values: list[str]) -> list[str]:
        errors = []
        labels = ("供应商名称", "主营品牌", "主要优势")
        for index, label in enumerate(labels):
            if not values[index]:
                errors.append(f"{label}不能为空")
        for cell in cells:
            if getattr(cell, "data_type", None) == "f":
                errors.append("不支持公式单元格")
                break
        return errors

    @staticmethod
    def _cell_text(value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @staticmethod
    def _required_value(value: str | None) -> str:
        if value is None:
            raise RuntimeError("Validated import row contains a required null field")
        return value

    @staticmethod
    def _contacts_from_row(row: SupplierImportRow, actor_id: uuid.UUID) -> list[SupplierContact]:
        if not row.contact_name and not row.contact_phone:
            return []
        return [
            SupplierContact(
                contact_name=row.contact_name,
                contact_phone=row.contact_phone,
                created_by=actor_id,
                updated_by=actor_id,
            )
        ]

    @staticmethod
    def _preview_response(batch: SupplierImportBatch) -> SupplierImportPreviewResponse:
        return SupplierImportPreviewResponse(
            id=batch.id,
            original_filename=batch.original_filename,
            status=batch.status,
            total_rows=batch.total_rows,
            valid_rows=batch.valid_rows,
            invalid_rows=batch.invalid_rows,
            rows=[
                SupplierImportRowResponse(
                    source_row_number=row.source_row_number,
                    supplier_name=row.supplier_name,
                    main_brands=row.main_brands,
                    advantage=row.advantage,
                    contact_name=row.contact_name,
                    contact_phone=row.contact_phone,
                    is_valid=row.is_valid,
                    error_message=row.error_message,
                )
                for row in batch.rows
            ],
        )
