from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

SUPPLIER_EXCEL_HEADERS = ("供应商名称", "主营品牌", "主要优势", "联系人", "联系电话")
SUPPLIER_EXCEL_SHEET_NAME = "供应商导入"
SUPPLIER_EXCEL_WIDTHS = (28, 30, 40, 20, 22)


def create_supplier_excel_workbook(
    *, include_import_comment: bool = False, include_phone_format_seed: bool = False
) -> tuple[Workbook, Worksheet]:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = SUPPLIER_EXCEL_SHEET_NAME
    worksheet.append(SUPPLIER_EXCEL_HEADERS)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:E1"
    for column, width in zip(("A", "B", "C", "D", "E"), SUPPLIER_EXCEL_WIDTHS, strict=True):
        worksheet.column_dimensions[column].width = width
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    if include_import_comment:
        worksheet["A1"].comment = Comment(
            "必填：供应商名称。主营品牌、主要优势、联系人和联系电话均可留空。", "系统"
        )
    if include_phone_format_seed:
        worksheet["E2"].number_format = "@"
    return workbook, worksheet
