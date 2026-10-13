"""Generate private, fictional JD_DETAIL_RED v5 acceptance decks."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation

from app.modules.recommendation.application import ppt_official_template_v5
from app.modules.recommendation.application.ppt_renderer import PptRenderer


def _png(color: tuple[int, int, int], size: tuple[int, int]) -> bytes:
    output = BytesIO()
    Image.new("RGB", size, color).save(output, format="PNG")
    return output.getvalue()


def _item(index: int, *, partial: bool = False, long: bool = False) -> dict[str, object]:
    product = {
        "id": f"fictional-v5-{index}",
        "product_name": f"虚构办公商品 {index}",
        "brand": "虚构品牌",
        "sku": f"FICT-V5-{index:03d}",
        "product_specification": "虚构规格：环保材质，标准配置",
        "packaging_list": "虚构包装：主件、说明书、配件",
        "selling_points": "虚构卖点一：适用于办公采购\n虚构卖点二：支持受控交付",
        "shipping_courier": "虚构京东物流",
    }
    manual = {"fulfillment_cycle": "3 个工作日"}
    prices = {"market_price": "199.50", "discount_rate": "0.85"}
    if partial:
        product.update(
            {
                "brand": "",
                "product_specification": "",
                "packaging_list": "",
                "selling_points": "",
                "shipping_courier": "",
            }
        )
        manual.clear()
        prices.clear()
    if long:
        product.update(
            {
                "product_name": "虚构超长商品名称用于容量边界验证一二三四五",
                "product_specification": "虚构规格容量边界文本，用于验证模板不缩小字号且仍完整显示",
                "selling_points": (
                    "虚构卖点边界第一行，保持客户模板字体和段落样式\n"
                    "虚构卖点边界第二行，验证完整可读\n"
                    "虚构卖点边界第三行，不与背景或图片重叠"
                ),
            }
        )
    return {"product": product, "prices": prices, "manual": manual}


def _save(
    renderer: PptRenderer, output: Path, source: dict[str, object], images: dict[str, bytes]
) -> None:
    output.write_bytes(
        renderer.render(
            source, product_images=images, template_code="JD_DETAIL_RED", template_version="5"
        )
    )


def _save_customer_price_evidence(
    renderer: PptRenderer, output: Path, source: dict[str, object], images: dict[str, bytes]
) -> None:
    """Generate a renderer-only fictional amount proof without a business source."""
    deck = Presentation(
        BytesIO(
            renderer.render(
                source, product_images=images, template_code="JD_DETAIL_RED", template_version="5"
            )
        )
    )
    prices = next(shape for shape in deck.slides[1].shapes if shape.name == "SCM_FIELD_PRICES")
    # The run's black/red style is part of the private runtime template.  This
    # direct fictional replacement does not create a product price field.
    prices.text_frame.paragraphs[0].runs[1].text = "¥ 321.00"
    deck.save(output)


def _save_supplement_evidence(
    renderer: PptRenderer, output: Path, source: dict[str, object], images: dict[str, bytes]
) -> None:
    """Create a private slot-layout proof with three unrelated fictional images."""
    deck = Presentation(
        BytesIO(
            renderer.render(
                source, product_images=images, template_code="JD_DETAIL_RED", template_version="5"
            )
        )
    )
    page = deck.slides[1]
    for name, image in (
        ("SCM_SLOT_IMAGE_SUPPLEMENT_01", _png((200, 20, 20), (80, 160))),
        ("SCM_SLOT_IMAGE_SUPPLEMENT_02", _png((20, 200, 20), (160, 80))),
        ("SCM_SLOT_IMAGE_SUPPLEMENT_03", _png((20, 20, 200), (100, 100))),
    ):
        slot = next(shape for shape in page.shapes if shape.name == name)
        ppt_official_template_v5._replace_picture(slot, image)
    deck.save(output)


def main() -> None:
    root = Path(__file__).resolve().parents[3]
    runtime = (
        root / ".codex-assets" / "official-ppt" / "runtime" / "JD_DETAIL_RED_OFFICIAL_CLEAN_v5.pptx"
    )
    destination = root / ".codex-assets" / "official-ppt" / "samples" / "v5"
    if not runtime.is_file():
        raise SystemExit("run build_private_official_ppt_template_v5.py first")
    destination.mkdir(parents=True, exist_ok=True)
    renderer = PptRenderer()
    cases = {
        "JD_DETAIL_RED-v5-fictional-complete-one-image.pptx": (
            [_item(1)],
            {"fictional-v5-1": _png((170, 40, 40), (160, 80))},
        ),
        "JD_DETAIL_RED-v5-fictional-partial.pptx": (
            [_item(2, partial=True)],
            {"fictional-v5-2": _png((40, 140, 80), (80, 160))},
        ),
        "JD_DETAIL_RED-v5-fictional-no-image.pptx": ([_item(3)], {}),
        "JD_DETAIL_RED-v5-fictional-five-single-products.pptx": (
            [_item(index) for index in range(4, 9)],
            {
                f"fictional-v5-{index}": _png((20 * index, 30, 180 - 10 * index), (100, 100))
                for index in range(4, 9)
            },
        ),
        "JD_DETAIL_RED-v5-fictional-boundary-safe.pptx": (
            [_item(9, long=True)],
            {"fictional-v5-9": _png((120, 80, 180), (80, 160))},
        ),
    }
    with patch.object(
        ppt_official_template_v5, "resolve_ppt_template_asset", lambda _template: runtime
    ):
        for name, (items, images) in cases.items():
            _save(renderer, destination / name, {"single_products": items}, images)
            print(destination / name)
        _save_customer_price_evidence(
            renderer,
            destination / "JD_DETAIL_RED-v5-fictional-customer-price-evidence.pptx",
            {"single_products": [_item(10)]},
            {"fictional-v5-10": _png((160, 80, 20), (160, 80))},
        )
        _save_supplement_evidence(
            renderer,
            destination / "JD_DETAIL_RED-v5-fictional-supplement-layout-evidence.pptx",
            {"single_products": [_item(11)]},
            {"fictional-v5-11": _png((20, 120, 180), (160, 80))},
        )
        _save(
            renderer,
            destination / "JD_DETAIL_RED-v5-fictional-main-image-fit-evidence.pptx",
            {"single_products": [_item(12), _item(13)]},
            {
                "fictional-v5-12": _png((220, 70, 30), (240, 80)),
                "fictional-v5-13": _png((30, 70, 220), (80, 240)),
            },
        )
        print(destination / "JD_DETAIL_RED-v5-fictional-customer-price-evidence.pptx")
        print(destination / "JD_DETAIL_RED-v5-fictional-supplement-layout-evidence.pptx")
        print(destination / "JD_DETAIL_RED-v5-fictional-main-image-fit-evidence.pptx")


if __name__ == "__main__":
    main()
