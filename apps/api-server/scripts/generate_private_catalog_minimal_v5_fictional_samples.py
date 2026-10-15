"""Generate private, wholly fictional CATALOG_MINIMAL v5 acceptance decks."""

from __future__ import annotations

from argparse import ArgumentParser
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from json import dumps
from os import environ
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation

from app.common.contracts import AppError
from app.core.config import get_settings
from app.modules.recommendation.application import (
    ppt_official_template_catalog_v5 as catalog,
)
from app.modules.recommendation.application.ppt_renderer import PptRenderer
from app.modules.recommendation.application.ppt_template_registry import (
    get_ppt_template,
    resolve_ppt_template_asset,
)
from scripts.build_private_official_ppt_template_v5 import (
    build_catalog_minimal_v5_runtime_asset,
)


def _png(colour: tuple[int, int, int], size: tuple[int, int]) -> bytes:
    output = BytesIO()
    Image.new("RGB", size, colour).save(output, format="PNG")
    return output.getvalue()


def _item(index: int) -> dict[str, object]:
    return {
        "product": {
            "id": f"fictional-catalog-{index}",
            "model": f"FICT-CAT-{index:03d}",
            "product_name": f"完全虚构极简画册商品 {index}",
            "product_specification": "虚构颜色：雾蓝\n虚构尺寸：120×80×60mm",
            "selling_points": "虚构卖点一；虚构卖点二；虚构卖点三",
            "jd_same_product_url": f"https://example.invalid/fictional-jd/{index}",
        },
        "prices": {
            "market_price": "999.00",
            "agreement_price": "239.00",
            "jd_price": "299.00",
        },
    }


def _first_line_boundary_url() -> str:
    base = "https://example.invalid/fictional-jd/"
    suffix = ""
    while (
        catalog._estimated_width(catalog._JD_LINK_LABEL + base + suffix + "a", 9)
        <= catalog._SAFE_TEXT_WIDTH_PT
    ):
        suffix += "a"
    return base + suffix + "a"


def main() -> None:
    parser = ArgumentParser()
    parser.add_argument("--original", type=Path)
    parser.add_argument(
        "--asset-dir",
        type=Path,
        help=(
            "Use an already-reviewed runtime asset directory instead of rebuilding "
            "a private copy."
        ),
    )
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[3]
    original = args.original or root / ".codex-assets/official-ppt/originals/模版3.pptx"
    runtime_root = (
        args.asset_dir
        if args.asset_dir is not None and args.asset_dir.is_absolute()
        else root / (args.asset_dir or Path(".codex-assets/official-ppt/runtime"))
    )
    runtime = runtime_root / "CATALOG_MINIMAL_OFFICIAL_CLEAN_v5.pptx"
    destination = root / ".codex-assets/official-ppt/samples/v5/catalog-minimal"
    destination.mkdir(parents=True, exist_ok=True)
    if args.asset_dir is None:
        build_catalog_minimal_v5_runtime_asset(original, runtime)
    elif not runtime.is_file():
        raise FileNotFoundError(f"reviewed CATALOG_MINIMAL v5 asset is missing: {runtime}")

    complete = _item(1)
    long_safe = deepcopy(_item(2))
    long_safe["product"].update(  # type: ignore[index]
        model="FICT-CATALOG-LONG-MODEL-2026-001",
        product_name="完全虚构的多行商品名称用于模板三画册版安全换行验收",
        product_specification="\n".join(
            f"虚构属性 {index}：" + "参数说明" * 5 for index in range(1, 7)
        ),
        selling_points="\n".join(
            f"虚构卖点 {index}：" + "完整说明" * 5 for index in range(1, 6)
        ),
        jd_same_product_url="https://example.invalid/fictional-jd/" + "continuous-ascii-" * 18,
    )
    first_line_boundary = deepcopy(_item(17))
    first_line_boundary["product"]["jd_same_product_url"] = _first_line_boundary_url()  # type: ignore[index]
    cases: dict[str, tuple[dict[str, object], dict[str, bytes]]] = {
        "complete": (
            {"single_products": [complete]},
            {"fictional-catalog-1": _png((55, 125, 190), (320, 180))},
        ),
        "empty-fields": (
            {"single_products": [{"product": {"id": "fictional-empty"}}]},
            {},
        ),
        "price-mapping": (
            {"single_products": [_item(14)]},
            {"fictional-catalog-14": _png((75, 125, 160), (320, 180))},
        ),
        "price-zero-missing": (
            {
                "single_products": [
                    {
                        **_item(15),
                        "prices": {
                            "market_price": "999.00",
                            "jd_price": 0,
                            "agreement_price": "0.00",
                        },
                    },
                    {
                        **_item(16),
                        "prices": {"market_price": "999.00"},
                    },
                ]
            },
            {
                "fictional-catalog-15": _png((125, 75, 160), (180, 300)),
                "fictional-catalog-16": _png((160, 125, 75), (300, 180)),
            },
        ),
        "no-image": ({"single_products": [_item(3)]}, {}),
        "aspect-ratios": (
            {"single_products": [_item(index) for index in range(4, 8)]},
            {
                "fictional-catalog-4": _png((190, 70, 50), (400, 100)),
                "fictional-catalog-5": _png((60, 160, 90), (100, 400)),
                "fictional-catalog-6": _png((80, 90, 200), (200, 200)),
                "fictional-catalog-7": _png((190, 160, 40), (800, 20)),
            },
        ),
        "long-safe": (
            {"single_products": [long_safe]},
            {"fictional-catalog-2": _png((90, 70, 180), (160, 260))},
        ),
        "url-first-line-boundary": (
            {"single_products": [first_line_boundary]},
            {"fictional-catalog-17": _png((70, 120, 180), (320, 180))},
        ),
        "five-products": (
            {"single_products": [_item(index) for index in range(8, 13)]},
            {
                f"fictional-catalog-{index}": _png(
                    (20 * index, 70, 210 - 10 * index), (140 + index, 100)
                )
                for index in range(8, 13)
            },
        ),
    }
    evidence: dict[str, object] = {
        "original_path": str(original),
        "original_sha256": sha256(original.read_bytes()).hexdigest(),
        "txt_sha256": sha256(original.with_suffix(".txt").read_bytes()).hexdigest(),
        "runtime_path": str(runtime),
        "runtime_sha256": sha256(runtime.read_bytes()).hexdigest(),
        "template_code": "CATALOG_MINIMAL",
        "template_version": "5",
        "data": "wholly fictional; no database access",
    }
    outputs: list[dict[str, object]] = []
    with patch.dict(
        environ,
        {
            "PPT_TEMPLATE_ASSET_DIR": str(runtime_root),
            # Settings validation only; this pure renderer script never opens a DB.
            "DATABASE_URL": "mysql+aiomysql://fictional:fictional@127.0.0.1/unused",
            "AUTH_JWT_SECRET": "fictional-render-only-secret-that-is-long-enough",
        },
    ):
        get_settings.cache_clear()
        try:
            resolved = resolve_ppt_template_asset(get_ppt_template("CATALOG_MINIMAL", "5"))
            assert resolved == runtime
            renderer = PptRenderer(layout_preview=False)
            for name, (source, images) in cases.items():
                data = renderer.render(
                    source,
                    product_images=images,
                    template_code="CATALOG_MINIMAL",
                    template_version="5",
                )
                output = destination / f"CATALOG_MINIMAL-v5-{name}.pptx"
                output.write_bytes(data)
                deck = Presentation(BytesIO(data))
                assert all(
                    "{{" not in shape.text and "自动填入" not in shape.text
                    for slide in deck.slides
                    for shape in slide.shapes
                    if shape.has_text_frame
                )
                outputs.append(
                    {
                        "path": str(output),
                        "sha256": sha256(data).hexdigest(),
                        "slides": len(deck.slides),
                    }
                )
                print(output)
            too_long = deepcopy(_item(13))
            too_long["product"]["selling_points"] = "\n".join(  # type: ignore[index]
                "虚构超限卖点" * 40 for _ in range(30)
            )
            try:
                renderer.render(
                    {"single_products": [too_long]},
                    template_code="CATALOG_MINIMAL",
                    template_version="5",
                )
            except AppError as error:
                assert error.code == "PPT_TEMPLATE_TEXT_CAPACITY_EXCEEDED"
                evidence["rejected_overflow"] = error.code
            else:
                raise AssertionError("genuine text overflow was not rejected")
        finally:
            get_settings.cache_clear()
    evidence["outputs"] = outputs
    report = destination / "CATALOG_MINIMAL-v5-evidence.json"
    report.write_text(dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(report)


if __name__ == "__main__":
    main()
