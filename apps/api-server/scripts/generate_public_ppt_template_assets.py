"""Generate the reviewed, fictional Type-5 PPTX distribution assets.

The resulting files are the only public business-template binaries. They contain
no customer text, branding, product imagery, links, notes, or source slides.
"""

from __future__ import annotations

from base64 import b64decode
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from app.modules.recommendation.application.ppt_template_manifest import (
    PPTX_SLOT_MANIFESTS,
    SLOT_BODY,
    SLOT_IMAGE_PREFIX,
    SLOT_PRICE,
    SLOT_TITLE,
)

_PNG = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9J7VIAAAAASUVORK5CYII="
)
_REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
_OUTPUT_DIR = _REPOSITORY_ROOT / "assets" / "ppt-templates"
_CANVAS_WIDTH = Inches(13.333333)
_CANVAS_HEIGHT = Inches(7.5)


def _rectangle(
    slide: object, left: float, top: float, width: float, height: float, color: RGBColor
) -> None:
    shape = slide.shapes.add_shape(  # type: ignore[attr-defined]
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _slot(slide: object, name: str, left: float, top: float, width: float, height: float) -> None:
    shape = slide.shapes.add_textbox(  # type: ignore[attr-defined]
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.name = name
    shape.text = ""


def _write_template(template_code: str, destination: Path) -> None:
    manifest = PPTX_SLOT_MANIFESTS[template_code]
    deck = Presentation()
    deck.slide_width = _CANVAS_WIDTH
    deck.slide_height = _CANVAS_HEIGHT
    is_red = template_code in {"JD_DETAIL_RED", "JD_FESTIVE_RED"}

    if manifest.cover_slide_index is not None:
        cover = deck.slides.add_slide(deck.slide_layouts[6])
        _rectangle(
            cover,
            0,
            0,
            13.333333,
            7.5,
            RGBColor(168, 15, 32) if template_code == "JD_FESTIVE_RED" else RGBColor(190, 38, 38),
        )
        _slot(cover, "SCM_PUBLIC_COVER_TITLE", 1.0, 3.0, 11.3, 0.75)

    product = deck.slides.add_slide(deck.slide_layouts[6])
    _rectangle(
        product,
        0,
        0,
        13.333333,
        0.25,
        RGBColor(190, 38, 38) if is_red else RGBColor(31, 45, 61),
    )
    if template_code == "JD_FESTIVE_RED":
        # Fictional static background media: retained while named product slots are replaced.
        product.shapes.add_picture(BytesIO(_PNG), 0, 0, _CANVAS_WIDTH, _CANVAS_HEIGHT)
    _slot(product, SLOT_TITLE, 6.1, 0.55, 6.0, 0.7)
    _slot(product, SLOT_BODY, 6.1, 1.4, 6.0, 3.8)
    if SLOT_PRICE in manifest.source_text_slots:
        _slot(product, SLOT_PRICE, 6.1, 5.4, 6.0, 1.1)
    for index in range(len(manifest.source_image_slots)):
        _slot(
            product,
            f"{SLOT_IMAGE_PREFIX}{index + 1}",
            0.6 + (index % 2) * 2.5,
            1.1 + (index // 2) * 2.0,
            2.2,
            1.8,
        )
    deck.save(str(destination))


def main() -> None:
    _OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for template_code, manifest in PPTX_SLOT_MANIFESTS.items():
        _write_template(template_code, _OUTPUT_DIR / manifest.clean_filename)


if __name__ == "__main__":
    main()
